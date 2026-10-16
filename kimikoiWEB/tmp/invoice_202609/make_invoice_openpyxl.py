from pathlib import Path
from openpyxl import load_workbook

src = Path(r'C:\Users\ihogo\Downloads\VA業務\請求書_本番7月末.xlsx')
out = Path(r'C:\Users\ihogo\Downloads\VA業務\請求書_本番9月末.xlsx')
wb = load_workbook(src)
ws = wb['10％、源泉あり']

ws['K6'] = '発行日：2026年09月30日'
ws['D17'] = '2026年10月26日'

ws['B21'] = '8月～9月末'
ws['C21'] = 'ホラゲ本開発　第2フェーズ分'
ws['E21'] = 1
ws['F21'] = '式'
ws['G21'] = 1350000
ws['I21'] = 0.1

ws['B22'] = '8月～9月末'
ws['C22'] = '交通費（8月～9月分）'
ws['E22'] = 1
ws['F22'] = '式'
ws['G22'] = 2136
ws['I22'] = 0

ws['B23'] = '8月～9月末'
ws['C23'] = '固定IPアドレスライセンス代'
ws['E23'] = 1
ws['F23'] = '式'
ws['G23'] = 980
ws['I23'] = 0.1

ws['B40'] = (
    '※上記「ホラゲ本開発 第2フェーズ分」には、対象期間内の成果物\n'
    '作業、週次作業報告書/ビルドファイル/開発資料テンプレート等の作成費用が含まれております。\n'
    '※交通費は出社記録に基づき、2026年8月～9月分を計上しています。\n'
    '※固定IPアドレスライセンス代はロリポップ！固定IPアクセスのお支払い履歴に基づき、2026年8月～9月分を計上しています。'
)

wb.active = wb.sheetnames.index('10％、源泉あり')
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = 'auto'
wb.save(out)
print(out)
