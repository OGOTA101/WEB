from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


template = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260731")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年7月第5週.xlsx"

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年7月第5週"


def set_value(cell_ref: str, value: str) -> None:
    cell = ws[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value


set_value("A2", "対象期間: 2026/07/27〜2026/07/31")
set_value("A4", "2026/07/27〜2026/07/31")
set_value("B4", "7/27, 7/28, 7/29（出社）, 7/30, 7/31（予定）")
set_value("C4", "キミと恋する15日（仮）")
set_value("D4", "Unity本番開発")
set_value(
    "E4",
    "本番開発として、初回死亡後の復帰演出、ADV選択肢と操作制御、LIMEからGoogruへの進行、デスクトップ演出を更新。7/30時点のGit状況を反映し、前週までに完了として報告済みの項目は再掲しない。",
)

implementation_rows = {
    8: {
        "A": "初回死亡後の復帰演出",
        "B": "・初回死亡時の全画面動画を再生する仕組みを追加\n・復活動画の後にタイトル画面へ戻る流れを接続\n・デバッグメニューから演出確認用の切替を追加",
        "D": "SisterHorrorSequenceDirector.cs / フルスクリーン動画 / MasterScenarioController.cs",
        "F": "完了",
        "G": "初回死亡後に演出が途切れず、復活からタイトル復帰までを一連の体験として確認できる状態",
    },
    9: {
        "A": "ADV選択肢/操作感調整",
        "B": "・選択肢が5件以上のとき、表示をコンパクトに切り替える処理を追加\n・選択肢、ログ、名前入力、終了確認へホバーSEを設定\n・画面送り時に不要なSEが鳴らないよう調整",
        "D": "KimikoiAdvChoiceController.cs / KimikoiAdvChoices.prefab / UI SE設定",
        "F": "完了",
        "G": "選択肢が多い場面でも画面を見切れず、操作対象を分かりやすく確認できる状態",
    },
    10: {
        "A": "シナリオ進行中の画面制御",
        "B": "・進行対象外のウィンドウを薄く表示する制御をWindowManagerへ集約\n・LIME二択や主人公入力中に、操作対象が戻らないよう調整\n・キミコイ以外の画面を開いた際、BGMを減衰する処理を追加",
        "D": "DesktopWindowManager.cs / LimeWindowController.cs / DesktopShellController.cs",
        "F": "完了",
        "G": "今操作すべきアプリが明確になり、シナリオ進行中に意図しない画面操作をしにくい状態",
    },
    11: {
        "A": "LIME/Googruシナリオ連携",
        "B": "・LIMEのリンク選択後に、Googruの進行へ正しく移る処理を調整\n・ローカルHTMLと画像素材を使う検索結果画面を追加\n・同じ検索画面を重ねて開かない制御と、入力欄表示を調整中",
        "D": "LimeWindowController.cs / GoogruWindowController.cs / StreamingAssets/GoogruHtml",
        "F": "作業中",
        "G": "LIMEから検索画面へ移るイベントを、プレイヤーが迷わず最後まで進められるよう最終調整中",
    },
    12: {
        "A": "デスクトップ/常駐妹演出",
        "B": "・デスクトップ窓の操作、通知、固定表示を本番確認向けに整理\n・新しい常駐妹画像を取り込み、旧デザインを整理\n・常駐妹表示と各ウィンドウのレイアウトを調整中",
        "D": "SisterResidentWindow.prefab / DesktopWindowManager.cs / sister_base01.png",
        "F": "作業中",
        "G": "常駐妹とデスクトップの見え方を、シナリオ確認時と同じ画面構成で確認できるよう調整中",
    },
    13: {
        "A": "バグエフェクト中央減衰",
        "B": "・バグ演出用のクリーン画面取得タイミングを修正\n・中央減衰の適用順を見直し、表示が崩れにくい状態へ調整\n・検証メニューから演出の見え方を再確認",
        "D": "ProceduralDatamoshGlitchRendererFeature.cs / BugEffectCenterAttenVerify.cs",
        "F": "完了",
        "G": "画面中央の情報を読みやすく保ちながら、周辺の不穏な演出を確認できる状態",
    },
    14: {
        "A": "パラメータ変化SE",
        "B": "・パラメータ上昇用と下降用のSE素材を追加\n・妹パラメータ変化時の再生処理へ反映\n・Unityへの取り込み状態と実機上の聞こえ方を確認中",
        "D": "parameter_up1.wav / parameter_down1.wav / SisterParameterChangeSePlayer.cs",
        "F": "作業中",
        "G": "パラメータが変わったことを音でも把握できるようにし、実際の変化とSEのタイミングを確認中",
    },
    15: {
        "A": "LIME入力欄/フォント表示",
        "B": "・LIME入力欄で長い文言が収まる表示方式へ調整\n・入力中の表示再開位置を調整し、途中で表示が崩れないよう修正\n・TMPフォント再生成後の文字表示を継続確認",
        "D": "LimeWindowController.cs / TextMeshPro SDF assets / LimeWindow.prefab",
        "F": "作業中",
        "G": "LIME会話の文言が途中で読めなくならないよう、文字表示の見え方を最終確認している段階",
    },
}

for row, values in implementation_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

other_rows = {
    19: {
        "A": "シナリオ/マスターデータ更新",
        "B": "・初回導線と共通DAYの演出タイミングを更新\n・怨念、イライラ、ときめきルートの進行データを調整\n・実装側の進行制御に合わせてシナリオCSVを整理",
        "D": "master_ver2_common_day*.csv / master_route_*.csv",
        "F": "完了",
        "G": "実装された画面遷移とシナリオ表の進行が同じ順序で確認できる状態",
    },
    20: {
        "A": "手動テスト/仕様整理",
        "B": "・ウィンドウ操作、Googru、LIME、表示演出の手動テスト手順を作成\n・全データ削除後に残る表示や通知の確認項目を整理\n・本番投入機能とモック差分、シナリオ固有の確認範囲を記録",
        "D": "Docs/ManualTest / Docs/Specifications",
        "F": "作業中",
        "G": "確認担当者が同じ手順で画面・シナリオを確認できるよう、チェック内容を整備中",
    },
    21: {
        "A": "実装状況/確認差分整理",
        "B": "・Git履歴から今週の本番開発更新を確認\n・コミット済みの機能と未コミットの表示調整を分けて整理\n・次回確認が必要なLIME、常駐妹、SE、フォントを記録",
        "D": "Git履歴 / Git作業ツリー / 本番確認対象",
        "F": "完了",
        "G": "今回確認できる完成範囲と、継続して確認する範囲を混在させず共有できる状態",
    },
    22: {
        "A": "報告書整備",
        "B": "・7/27〜7/31分の作業報告をテンプレートに沿って作成\n・7/30時点のGit履歴と未コミット差分を反映\n・出社日と作業中項目を確認しやすい表現に整理",
        "D": "本報告書",
        "F": "完了",
        "G": "今週の実装成果、継続作業、確認時の着眼点を一枚で把握できる状態",
    },
}

for row, values in other_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

set_value(
    "A25",
    "・対象期間は2026/07/27〜2026/07/31。7/29は出社対応日。7/31は予定作業日として記載し、本報告は7/30時点のGit状況を反映。\n"
    "・前週までに完了として報告済みの項目は再掲しない。初回死亡後の復帰演出、ADV操作制御、バグエフェクト中央減衰、シナリオCSV調整を今週の完了項目として整理。\n"
    "・LIME/Googru連携、常駐妹レイアウト、パラメータ変化SE、LIME文字表示とフォント確認は作業中。次回は手動テスト手順に沿った通し確認を継続する。",
)

if ws.row_dimensions[14].height:
    ws.row_dimensions[15].height = ws.row_dimensions[14].height

ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1

wb.save(out)
print(out)
