from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


template = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260625")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年6月第4週.xlsx"

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年6月第4週"


def set_value(cell_ref: str, value):
    cell = ws[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value


set_value("A2", "対象期間: 2026/06/22〜2026/06/25")
set_value("A4", "2026/06/22〜2026/06/25")
set_value("B4", "6/22, 6/23, 6/24, 6/25")
set_value("C4", "キミと恋する15日（仮）")
set_value("D4", "Unity本番開発")
set_value(
    "E4",
    "本番開発として、先週から継続していたLIME/Line基盤、ADV/演出整理、Canvas解像度調整、AI運用規則整理、ビルド番号対応、本番ビルド作成、ドライブ納品まで実施。木曜締めのため作業日は4日分で整理。",
)

implementation_rows = {
    8: {
        "A": "Line/LIME基盤",
        "B": "・Lineトークバブル、スタンプ生成、サイズ調整を別クラスへ分離\n・選択肢、入力欄ドラフト、確定状態を別管理に整理\n・送信待ち状態とパルス演出を分離し、Lineのインスタンス登録も修正",
        "D": "LimeBubblePresenter.cs / LimeScenarioChoiceState.cs / LimeSendButtonWaitingPresenter.cs / LimeWindowController.cs",
        "F": "作業中",
        "G": "LIME会話の見た目と操作を細かく調整できるようになり、企画側の演出確認もしやすくしている段階",
    },
    9: {
        "A": "演出/グリッチ制御",
        "B": "・怨念グリッチ制御クラスを追加\n・怨念系バグ演出のID判定とBGM判定を整理\n・デスクトップ側グリッチ制御とキミコイADV側の同期処理を修正",
        "D": "KimikoiGrudgeGlitchCoordinator.cs / KimikoiAdvBugEffectRunner.cs / SisterPresenceGlitchController.cs",
        "F": "完了",
        "G": "ホラー/バグ演出を条件に応じて出し分けるための管理単位を整理済み",
    },
    10: {
        "A": "Canvas/画面調整",
        "B": "・Canvas基準を2560x1440へ対応\n・各Canvas Prefabのスケール補正を解除\n・キミコイ没入表示中も内部UIを1280x960の論理画面として固定",
        "D": "Main.unity / KimikoiImmersivePresenter.cs / Desktop Prefabs / Rendering Prefabs",
        "F": "完了",
        "G": "画面サイズ差によるUIズレを減らし、確認用ビルドでも同じ見た目で確認しやすくした状態",
    },
    11: {
        "A": "UI/Prefab不具合修正",
        "B": "・実行時に出ていた警告系を修正\n・Inspector指定でMainシーンから外れていた参照を修正\n・Lineトーク欄の表示補正と動画用フォルダMetaの再登録を実施",
        "D": "Main.unity / StartMenuView.prefab / LimeWindow.prefab / Move.meta",
        "F": "完了",
        "G": "プランナー確認時に、参照抜けや表示崩れで確認が止まるリスクを減らした状態",
    },
    12: {
        "A": "MOC/動画素材",
        "B": "・MOCシーンを保存\n・再生テスト用動画をAssetsへ追加\n・動画再生確認に必要なフォルダ構成とMetaを整理",
        "D": "MOC.unity / Assets/Move / 再生テスト用動画.mp4",
        "F": "作業中",
        "G": "映像演出をUnity上で確認するための準備を進めており、今後の演出確認に使える状態へ移行中",
    },
    13: {
        "A": "本番ビルド対応",
        "B": "・ビルド版で4桁のビルド番号を取得できるよう対応\n・ビルド時に番号を自動加算、正規化する機能を追加\n・Infoパネルにビルド番号を表示できるよう更新",
        "D": "KimikoiBuildInfo.cs / KimikoiBuildNumberProcessor.cs / DesktopDebugMenuController.InfoPanel.cs",
        "F": "完了",
        "G": "提出したビルドがどの版か確認しやすくなり、QAや差し戻し時の版管理がしやすい状態",
    },
    14: {
        "A": "本番ビルド/納品",
        "B": "・本番確認用のWindowsビルドを作成\n・EXE_260625_1.zipとして成果物を圧縮\n・作成したビルドをドライブへ納品",
        "D": "D:/KIMIKOI/KIMIKOI/EXE/Horage.exe / EXE_260625_1.zip / ドライブ納品",
        "F": "完了",
        "G": "プランナー/関係者が確認できる形でビルドを用意し、共有先へ提出済み",
    },
    15: {
        "A": "データ/コメント整理",
        "B": "・Data READMEの説明を更新\n・ADV Data内のコメントが不明瞭な箇所を修正\n・Window確認フラグ、フリップ音、Data未設定時の説明を更新",
        "D": "Assets/Data/README.md / KimikoiAdv Data Scripts / WindowDefinition.cs / DesktopEmbeddedCopy.cs",
        "F": "完了",
        "G": "仕様や設定項目の意味が追いやすくなり、企画確認や後続修正時の迷いを減らした状態",
    },
}

for row, values in implementation_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

other_rows = {
    19: {
        "A": "書類/WIKI更新",
        "B": "・Editorツールの開き方をREADMEへ追加\n・デバッグ操作メモとTOP項目を更新\n・開発機能に追加した内容をWIKI/Docsへ追記",
        "D": "README.md / Assets/Docs/README.md / UnityデバッグとEditor操作メモ.md",
        "F": "完了",
        "G": "確認者がEditor操作やデバッグ機能の入口を探しやすいよう、ドキュメント導線を更新済み",
    },
    20: {
        "A": "AI/運用規則整理",
        "B": "・基本的なCursor/AIルールをキミコイ専用ルールへ集約\n・不要になったCursor用Skillsファイルを削除\n・C#とUIの運用規約の参照パスを修正",
        "D": ".cursor/rules/kimikoi-project-rules.mdc / 実装運用規約_CSharp.md / 実装運用規約_UI.md",
        "F": "完了",
        "G": "今後AIや実装担当へ依頼する際のルールが一本化され、手戻りや誤操作を減らしやすい状態",
    },
    21: {
        "A": "先週報告とのすり合わせ",
        "B": "・先週報告書の作業中項目を確認\n・LIME/Line、ADV構成整理、基盤安定化の継続状況を今週分へ反映\n・END/イベントCSVは今週の主作業ではなく、継続課題として扱う形に整理",
        "D": "outputs/work_reports_20260619/エンジニア作業報告書_2026年6月第3週.xlsx",
        "F": "完了",
        "G": "先週からの流れと今週の成果がつながるよう、報告内容を整理済み",
    },
    22: {
        "A": "報告書整備",
        "B": "・6/22〜6/25分の作業報告をテンプレートに沿って作成\n・木曜締めのため作業日を4日分に設定\n・本番ビルド作成とドライブ納品を明記",
        "D": "本報告書",
        "F": "完了",
        "G": "今週の実装、書類更新、ビルド提出状況を業務報告として確認しやすい形に整理済み",
    },
}

for row, values in other_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

set_value(
    "A25",
    "・6/22〜6/25は木曜締めのため、作業日は6/22、6/23、6/24、6/25の4日分として整理。\n"
    "・先週の作業中項目を確認し、今週はLIME/Line分割、Canvas調整、ビルド番号対応、ドキュメント整備へ進行。\n"
    "・本番確認用ビルド EXE_260625_1.zip を作成し、ドライブへ納品済み。",
)

if ws.row_dimensions[14].height:
    ws.row_dimensions[15].height = ws.row_dimensions[14].height

wb.save(out)
print(out)
