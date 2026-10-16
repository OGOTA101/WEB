from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from PIL import ImageFont


template = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260828")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年8月第4週.xlsx"


def set_value(sheet, cell_ref: str, value: str) -> None:
    cell = sheet[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value


def write_values(sheet, rows: dict[int, dict[str, str]]) -> None:
    for row, values in rows.items():
        for column, value in values.items():
            set_value(sheet, f"{column}{row}", value)


def validate(workbook_path: Path, content_rows: list[int]) -> None:
    workbook = load_workbook(workbook_path, data_only=False)
    sheet = workbook.active
    if sheet.print_area != f"'{sheet.title}'!$A$1:$H$28":
        raise ValueError(f"Unexpected print area: {sheet.print_area}")

    font = ImageFont.truetype(r"C:/Windows/Fonts/YuGothR.ttc", 12, index=0)
    for row in content_rows:
        lines = sheet[f"B{row}"].value.splitlines()
        if not 3 <= len(lines) <= 4:
            raise ValueError(f"B{row} must contain 3 to 4 lines: {lines}")
        if sum(line.startswith("・") for line in lines) != 3:
            raise ValueError(f"B{row} must contain three bullet lines: {lines}")
        for line in lines:
            if font.getlength(line) > 248:
                raise ValueError(f"B{row} line is too wide: {line}")

        confirmation_lines = sheet[f"G{row}"].value.splitlines()
        if not 1 <= len(confirmation_lines) <= 5:
            raise ValueError(f"G{row} must contain 1 to 5 lines: {confirmation_lines}")
        for line in confirmation_lines:
            if font.getlength(line) > 258:
                raise ValueError(f"G{row} line is too wide: {line}")

    errors = []
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("#"):
                errors.append(cell.coordinate)
    if errors:
        raise ValueError(f"Formula errors found: {errors}")


# Improve the reusable template with the current reporting rules.
template_wb = load_workbook(template)
template_ws = template_wb.active
set_value(template_ws, "B4", "作業日（出社日は明記）")
set_value(template_ws, "C4", "キミと恋する15日（仮）")
set_value(template_ws, "D4", "Unity本番開発")
set_value(
    template_ws,
    "E4",
    "今週の新規・継続作業と確認状況を記載。完了として報告した項目は再掲しない。",
)
set_value(template_ws, "A6", "1. エンジニアとして実装・調整した作業内容")
set_value(template_ws, "B7", "作業内容")
set_value(template_ws, "G7", "確認内容（プランナー向け）")
set_value(template_ws, "A8", "ゲーム進行/シナリオ")
set_value(template_ws, "A9", "ゲーム内アプリ/仮想OS")
set_value(template_ws, "A10", "UI/Prefab/演出")
set_value(template_ws, "A11", "セーブ/ロード/状態復元")
set_value(template_ws, "A12", "音響/表示/フォント")
set_value(template_ws, "A13", "マスターデータ/ルート")
set_value(template_ws, "A14", "検証/デバッグ")
set_value(template_ws, "A15", "本番確認/調整")
set_value(template_ws, "D8", "例: Scripts / Prefabs / Assets / Master CSV")
set_value(template_ws, "F8", "作業中")
set_value(
    template_ws,
    "B8",
    "・具体的な作業内容を3点以上記載\n・必要時は意味の切れ目で4行目へ\n・行末に一文字だけ残さない",
)
set_value(
    template_ws,
    "G8",
    "・確認する画面・操作\n・期待する表示・進行\n・継続確認の理由を\n　5行以内で記載",
)
set_value(template_ws, "A17", "2. 実装以外の業務・確認／運用")
set_value(template_ws, "B18", "業務内容")
set_value(template_ws, "G18", "確認内容（プランナー向け）")
set_value(template_ws, "A19", "資料/仕様更新")
set_value(template_ws, "A20", "確認環境/運用")
set_value(template_ws, "A21", "実装状況整理")
set_value(template_ws, "A22", "報告書/テンプレート整備")
set_value(template_ws, "D19", "例: Docs / 仕様書 / 手動テスト / 運用資料")
set_value(template_ws, "F19", "作業中")
set_value(template_ws, "G19", "企画側が確認する画面・操作と\n期待する状態を記載")
set_value(template_ws, "A24", "3. 補足 / 継続確認事項")
set_value(
    template_ws,
    "A25",
    "・対象日と出社日を記載。\n"
    "・前週に完了として報告した項目は再掲しない。\n"
    "・次回の確認事項を記載。",
)
template_ws.sheet_properties.pageSetUpPr.fitToPage = True
template_ws.page_setup.orientation = "portrait"
template_ws.page_setup.fitToWidth = 1
template_ws.page_setup.fitToHeight = 1
template_wb.save(template)

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年8月第4週"
set_value(ws, "A2", "対象期間: 2026/08/24〜2026/08/28（8/26出社）")
set_value(ws, "A4", "2026/08/24〜2026/08/28")
set_value(ws, "B4", "8/24〜8/28（8/26出社）")
set_value(ws, "C4", "キミと恋する15日（仮）")
set_value(ws, "D4", "Unity本番開発")
set_value(
    ws,
    "E4",
    "本番開発として、ADV/LIME入力、ルート終盤、音響、デスクトップ表示、マスターデータを調整中。前週までに完了として報告済みの項目は再掲しない。",
)

write_values(
    ws,
    {
        8: {
            "A": "ADV/LIME入力操作",
            "B": "・ADVの入力待ちと選択肢接続を調整\n・選択肢のマウス即時確定を追加\n・LIME入力前のSE抑止を調整",
            "D": "KimikoiAdvChoiceController\nKimikoiAdvPlaybackDirector",
            "F": "作業中",
            "G": "・台詞送りで進行が止まらない\n・選択肢をクリックすると\n　すぐ確定できる\n・入力直前に不要な音が\n　鳴らないことを確認中",
        },
        9: {
            "A": "ルート終盤/ゲームオーバー",
            "B": "・ゲームオーバー後の窓状態を復帰\n・怨念ルートの自動クローズ待ちを保護\n・終盤の窓表示と監視状態を調整",
            "D": "DesktopShellController\nDesktopWindowManager",
            "F": "作業中",
            "G": "・ゲームオーバー後に窓が閉じる\n・怨念ルートの自動閉じが\n　途中で進行を止めない\n・終盤の監視表示が残らない",
        },
        10: {
            "A": "スキンシップ/画面表示",
            "B": "・スキンシップ妹の表示位置を調整\n・クリックできる操作領域を見直し\n・名前入力と公式サイトを調整",
            "D": "SisterSkinshipWindowController\nGoogruWindow / NameInputDialog",
            "F": "作業中",
            "G": "・妹の画像を押せる範囲が\n　見た目と一致している\n・名前入力が正しく表示される\n・公式サイト画面が開ける",
        },
        11: {
            "A": "音響/BGM演出",
            "B": "・BGM/SEの再生ログと停止APIを追加\n・後半ルートのBGMをキューへ接続\n・LIME送受信SEと通知を調整",
            "D": "SoundManager\nKimikoiAdvHeaderAndAudioData",
            "F": "作業中",
            "G": "・台詞と画面切替に合うBGM\n・LIME送受信時のSEと通知\n・後半ルートの曲が切替時に\n　正しく鳴るか確認中",
        },
        12: {
            "A": "デスクトップ操作表示",
            "B": "・デスクトップ窓の開閉フェードを追加\n・クリック用の手カーソルを追加\n・処理中表示とボタン反応を調整",
            "D": "DesktopWindowView\nDesktopCursorStateController",
            "F": "作業中",
            "G": "・窓の開閉が唐突に見えない\n・押せる場所で手カーソルに変わる\n・処理中に誤操作できない",
        },
        13: {
            "A": "セーブ/ロード操作改善",
            "B": "・最新保存枠へNEW表示を追加\n・保存/読込失敗理由の通知を追加\n・復帰時の画面遷移を確認",
            "D": "SnapshotSaveLoadPanelController\nDesktopSnapshotRuntimeService",
            "F": "作業中",
            "G": "・最新の保存枠が見分けられる\n・保存/読込に失敗した場合に\n　理由が画面に表示される\n・復帰後に画面が正しく戻る",
        },
        14: {
            "A": "マスターデータ/演出導線",
            "B": "・共通/初回DAY2の演出導線を調整\n・各ルートのBGMと画面遷移を更新\n・DAY7〜9の旧LIME会話を整理",
            "D": "master_ver2_*\nmaster_route_*",
            "F": "作業中",
            "G": "・DAY2と各ルートで会話から\n　演出へ自然に繋がる\n・旧LIME会話が出ない\n・BGMと画面遷移を確認中",
        },
        15: {
            "A": "イベントCG/フォント/デバッグ",
            "B": "・αイベントCGの仮素材を追加\n・CG表示キーと立ち絵データを設定中\n・日本語TMPフォントの差分を確認",
            "D": "KimikoiAdvPortraitAndCgData\nTextMeshPro SDF assets",
            "F": "作業中",
            "G": "・該当場面でCGが表示される\n・CGと立ち絵が重ならない\n・日本語の文字欠けがない",
        },
        19: {
            "A": "シナリオCSV/資料更新",
            "B": "・Master CSVの書式と演出列を更新\n・終盤シナリオの台詞と効果を見直し\n・演出確認用の参照データを整理",
            "D": "GeneratedVer2 / Routes\nmaster_day_any_ver2.csv",
            "F": "作業中",
            "G": "・台詞、BGM、演出の指定が\n　実装側とずれない\n・終盤の台詞と効果が合う\n・参照データを確認中",
        },
        20: {
            "A": "本番確認環境/版管理",
            "B": "・本番用bundleVersionを0008へ更新\n・Main画面と描画設定の差分を確認\n・検証用ツールと未反映差分を整理",
            "D": "ProjectSettings / URP\nMain.unity / Debug menu",
            "F": "作業中",
            "G": "・本番版の番号で起動できる\n・Main画面の表示が崩れない\n・確認用ツールで変更を追える",
        },
        21: {
            "A": "出社対応/実装状況確認",
            "B": "・8/26に出社し進行状況を確認\n・Git履歴と作業ツリーを照合\n・確認順と継続対応を整理",
            "D": "Git log / working tree\n本番確認リスト",
            "F": "完了",
            "G": "・8/26の出社時に進行を確認\n・Git上の変更と作業中分を照合\n・企画確認の順番を整理済み",
        },
        22: {
            "A": "報告書/テンプレート整備",
            "B": "・8/24〜8/28分の報告書を作成\n・作業内容を3点ずつ具体的に整理\n・4行以内で読める書式へ更新",
            "D": "本報告書\n作業報告書テンプレート",
            "F": "完了",
            "G": "・各項目の確認対象を具体化\n・作業中と完了を見分けやすく\n・5行以内で読める形に更新",
        },
    },
)

set_value(
    ws,
    "A25",
    "・対象：8/24〜8/28。8/26は出社。\n"
    "・本番画面で入力、音響、画面遷移、ルート終盤を確認中。\n"
    "・CG、フォント、シナリオCSV、保存画面の通し確認を継続。",
)

ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
wb.save(out)

validate(template, [8])
validate(out, [*range(8, 16), *range(19, 23)])
print(template)
print(out)
