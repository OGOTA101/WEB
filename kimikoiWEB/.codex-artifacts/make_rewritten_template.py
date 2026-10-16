from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.datavalidation import DataValidation

src = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_2026年5月第1週.xlsx")
out_dir = src.parent
out = out_dir / "エンジニア作業報告書_テンプレート_改訂版.xlsx"

wb = load_workbook(src)
ws = wb.active
ws.title = "作業報告テンプレート"

def set_value(cell, value):
    ws[cell].value = value

def maybe_clear(cell_ref):
    cell = ws[cell_ref]
    if not isinstance(cell, MergedCell):
        cell.value = None

def clear_cells(cells):
    for cell in cells:
        maybe_clear(cell)

set_value("A1", "作業報告書")
set_value("A2", "対象期間: YYYY/MM/DD〜YYYY/MM/DD")

set_value("A4", "YYYY/MM/DD〜YYYY/MM/DD")
set_value("B4", "作業日を記入")
set_value("C4", "プロジェクト名")
set_value("D4", "対象範囲")
set_value("E4", "実施内容・成果・確認内容を1〜2文で記載")

set_value("D7", "主な成果物 / 対象")

implementation_rows = range(8, 14)
other_rows = range(19, 23)

for row in implementation_rows:
    for col in ["B", "C", "D", "E", "F", "G", "H"]:
        maybe_clear(f"{col}{row}")

for row in other_rows:
    for col in ["B", "C", "D", "E", "F", "G", "H"]:
        maybe_clear(f"{col}{row}")

set_value("B8", "・実装/調整した内容を箇条書き\n・対象画面や対象機能を具体的に記載")
set_value("D8", "例: Scripts / Prefabs / Assets / Docs")
set_value("F8", "進行中")
set_value("G8", "確認した内容を記載")

set_value("B19", "・仕様整理、資料作成、調整などを箇条書き")
set_value("D19", "例: 仕様書 / 参考資料 / 報告書")
set_value("F19", "進行中")
set_value("G19", "確認した内容を記載")

set_value("A25", "・工数が不要な場合は時間欄を設けず、実装範囲と確認内容を中心に記載。\n・次回確認事項、残課題、提出前チェックなどを記載。")

for row in [9, 10, 11, 12, 13, 20, 21, 22]:
    maybe_clear(f"F{row}")

for row in range(1, ws.max_row + 1):
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row, col)
        if isinstance(cell, MergedCell):
            continue
        if cell.value == "進行中シンコウチュウ":
            cell.value = "進行中"
        if isinstance(cell.value, str):
            cell.value = cell.value.replace("対象オモ", "対象").replace("オコナ", "")

dv = DataValidation(type="list", formula1='"未着手,進行中,完了,保留"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("F8:F13")
dv.add("F19:F22")

ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.page_margins.left = 0.25
ws.page_margins.right = 0.25
ws.page_margins.top = 0.35
ws.page_margins.bottom = 0.35
ws.print_area = "A1:H28"

wb.save(out)
print(out)
