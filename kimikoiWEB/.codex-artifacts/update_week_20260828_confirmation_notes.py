from pathlib import Path

from openpyxl import load_workbook
from PIL import ImageFont


template = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
report_source = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260828/エンジニア作業報告書_2026年8月第4週.xlsx"
)
report_revision = report_source.with_stem(f"{report_source.stem}_修正版")

template_notes = {
    "G8": "・全体で5行以内\n・操作と結果を書く\n・一文字だけを\n　次行に残さない",
    "G19": "・確認する操作を書く\n・期待する結果を書く\n・全体で5行以内",
}

report_notes = {
    8: "・台詞途中で入力し\n　進行を確認\n・選択肢を押して確定\n・LIME前にSEが出ない",
    9: "・ゲームオーバー後の\n　窓と監視表示を確認\n・怨念ルートの自動終了\n　が止まらないか確認\n・通常操作へ戻る",
    10: "・妹に触れた時の反応\n　と表示位置を確認\n・クリック範囲を確認\n・名前入力とサイトの\n　表示を確認",
    11: "・LIME送受信時の\n　SEを確認\n・後半BGMの切替\n　を確認\n・通知音が残らない",
    12: "・窓の開閉フェード\n　を確認\n・リンク上の手カーソル\n　を確認\n・処理中の誤操作を確認",
    13: "・最新枠のNEW表示\n　を確認\n・読込失敗時の案内\n　を確認\n・読込後も操作できる",
    14: "・各ルートを対象日まで\n　進める\n・会話とBGMの切替\n　を確認\n・旧LIME会話が出ない",
    15: "・対象場面でCGを表示\n・立ち絵とCGの重なり\n　を確認\n・日本語が欠けず読める",
    19: "・該当シナリオを順に\n　進める\n・台詞と演出の内容\n　を確認\n・資料とのずれを確認",
    20: "・本番版で起動と画面\n　を確認\n・描画崩れがないか確認\n・追加差分の確認順\n　を整理",
    21: "・8/26出社時の進行\n　を確認\n・Git差分と現行の動き\n　を照合\n・確認順を整理済み",
    22: "・今週追加した内容\n　を反映\n・確認する場面を\n　項目別に整理\n・次回の継続確認を明記",
}


def update_template() -> None:
    workbook = load_workbook(template)
    sheet = workbook.active
    for cell, value in template_notes.items():
        sheet[cell] = value
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    workbook.save(template)


def update_report(report: Path) -> None:
    workbook = load_workbook(report)
    sheet = workbook.active
    for row, value in report_notes.items():
        sheet[f"G{row}"] = value
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    workbook.save(report)


def validate(path: Path, cells: list[str]) -> None:
    workbook = load_workbook(path, data_only=False)
    sheet = workbook.active
    font = ImageFont.truetype(r"C:/Windows/Fonts/YuGothR.ttc", 12, index=0)
    for cell in cells:
        value = sheet[cell].value
        lines = value.splitlines()
        if not 2 <= len(lines) <= 5:
            raise ValueError(f"{cell} must contain 2 to 5 lines: {lines}")
        if any(font.getlength(line) > 132 for line in lines):
            raise ValueError(f"{cell} exceeds the usable width: {lines}")
    errors = [
        cell.coordinate
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("#")
    ]
    if errors:
        raise ValueError(f"Formula errors found: {errors}")


update_template()
validate(template, list(template_notes))
for report in [report_source, report_revision]:
    if report.exists():
        update_report(report)
        validate(report, [f"G{row}" for row in report_notes])
print(template)
print(report_source)
print(report_revision)
