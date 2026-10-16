from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


template = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260821")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年8月第3週.xlsx"


def set_value(sheet, cell_ref: str, value: str) -> None:
    cell = sheet[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value


# Update the reusable production-development template before producing this week's report.
template_wb = load_workbook(template)
template_ws = template_wb.active
set_value(template_ws, "B4", "作業日（出社日は明記）")
set_value(template_ws, "C4", "キミと恋する15日（仮）")
set_value(template_ws, "D4", "Unity本番開発")
set_value(
    template_ws,
    "E4",
    "今週の新規・継続作業と確認状況を記載。前週に完了として報告した項目は再掲しない。",
)
set_value(template_ws, "A6", "1. エンジニアとして実装・調整した作業内容")
set_value(template_ws, "B7", "作業内容")
set_value(template_ws, "G7", "確認内容（プランナー向け）")
set_value(template_ws, "A8", "ゲーム進行/シナリオ")
set_value(template_ws, "A9", "ゲーム内アプリ/仮想OS")
set_value(template_ws, "A10", "UI/Prefab/演出")
set_value(template_ws, "A11", "セーブ/ロード/状態復元")
set_value(template_ws, "A12", "音響/表示/フォント")
set_value(template_ws, "A13", "マスターデータ/ルート")
set_value(template_ws, "A14", "検証/デバッグ")
set_value(template_ws, "A15", "本番確認/調整")
set_value(template_ws, "D8", "例: Scripts / Prefabs / Assets / Master CSV")
set_value(template_ws, "F8", "作業中")
set_value(
    template_ws,
    "B8",
    "・3〜4行の箇条書きで記載\n・対象画面や機能を具体的に記載\n・行末に一文字だけ残さない",
)
set_value(template_ws, "G8", "確認者が試す内容と\n継続確認事項を記載")
set_value(template_ws, "A17", "2. 実装以外の業務・確認／運用")
set_value(template_ws, "B18", "業務内容")
set_value(template_ws, "G18", "確認内容（プランナー向け）")
set_value(template_ws, "A19", "資料/仕様更新")
set_value(template_ws, "A20", "確認環境/運用")
set_value(template_ws, "A21", "実装状況整理")
set_value(template_ws, "A22", "報告書/テンプレート整備")
set_value(template_ws, "D19", "例: Docs / 仕様書 / 手動テスト / 運用資料")
set_value(template_ws, "F19", "作業中")
set_value(template_ws, "G19", "企画・確認者側で確認できる内容を記載")
set_value(template_ws, "A24", "3. 補足 / 継続確認事項")
set_value(
    template_ws,
    "A25",
    "・対象日と出社日を記載。\n"
    "・前週に完了として報告した項目は再掲しない。\n"
    "・次回の確認事項を記載。",
)
template_ws.sheet_properties.pageSetUpPr.fitToPage = True
template_ws.page_setup.orientation = "portrait"
template_ws.page_setup.fitToWidth = 1
template_ws.page_setup.fitToHeight = 1
template_wb.save(template)

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年8月第3週"

set_value(ws, "A2", "対象期間: 2026/08/10、2026/08/17〜2026/08/21")
set_value(ws, "A4", "2026/08/10、2026/08/17〜2026/08/21")
set_value(ws, "B4", "8/10, 8/17〜8/21")
set_value(ws, "C4", "キミと恋する15日（仮）")
set_value(ws, "D4", "Unity本番開発")
set_value(
    ws,
    "E4",
    "先週の単日稼働（8/10）と今週の本番開発として、妹演出、イライラ進行、仮想OS設定、セーブ／ロードの状態復元を実装・検証中。前週までに完了として報告済みの項目は再掲しない。",
)

implementation_rows = {
    8: {
        "A": "妹カバー/バグ演出204",
        "B": "・妹切替用に強モザイクとカバー演出を追加\n・カバー中のデスクトップ演出と\n　バーストを調整\n・常駐妹、ADV、動画の重なり順を確認",
        "D": "KimikoiAdvBugEffectCatalog\nOuterAdvOverlay / Main・MOC",
        "F": "作業中",
        "G": "表示順と演出を\n実画面で確認中",
    },
    9: {
        "A": "イライラ目/監視イベント",
        "B": "・イライラ目の監視、開閉、リセットを追加\n・閉眼中の表示と共通サイクルを調整\n・DAY7以降とDAY10後の進行を確認",
        "D": "IrairaEyeWindowDirector\nWatchingMock / SisterResident",
        "F": "作業中",
        "G": "イベント進行を\n通し確認中",
    },
    10: {
        "A": "仮想OS設定/システム画面",
        "B": "・壁紙選択とシステムプロパティ画面を追加\n・言語切替の保存処理と\n　フライアウトを追加\n・右クリックから背景設定を開く導線を調整",
        "D": "WallpaperPickerWindowController\nDesktopUiLanguageService / Taskbar",
        "F": "作業中",
        "G": "設定保存と反映を\n再起動後も確認中",
    },
    11: {
        "A": "セーブ/ロード基盤",
        "B": "・デスクトップ状態を採取・復元する\n　基盤を追加\n・セーブ枠のJSONとサムネイルを追加\n・スタートメニューから開く保存画面を実装",
        "D": "DesktopSnapshotRuntimeService\nSaveManager / SaveLoadPanel",
        "F": "作業中",
        "G": "保存・読込後の\n進行を確認中",
    },
    12: {
        "A": "デスクトップ状態の復元",
        "B": "・アプリ配置、最小化、壁紙、\n　ごみ箱状態を保存\n・通知、アプリバッジ、検索ページを復元\n・復元中の操作が進行を起こさないよう調整",
        "D": "DesktopSnapshotRuntimeService\nGoogruWindowController / Shell",
        "F": "作業中",
        "G": "復元後の状態と\nイベント進行を確認",
    },
    13: {
        "A": "ADV/LIMEの状態復元",
        "B": "・ADVの背景、立ち絵、行位置、\n　選択肢を復元\n・LIME履歴、スタンプ、直前台詞を復元\n・外側進行と待機状態の復元を調整",
        "D": "KimikoiAdvWindowController\nLimeWindowController / Snapshot DTO",
        "F": "作業中",
        "G": "会話途中の読込と\n選択肢を確認中",
    },
    14: {
        "A": "ロード時の暗転/操作ガード",
        "B": "・ロード中の暗転オーバーレイと\n　復元APIを追加\n・通知と最小化は画面復元後に戻すよう調整\n・大きい窓とダイアログの復帰を調整",
        "D": "DesktopSnapshotRuntimeService\nSnapshotLoadOverlay / Main・MOC",
        "F": "作業中",
        "G": "読込中の表示と\n誤操作防止を確認",
    },
    15: {
        "A": "セーブ画面/デバッグ操作",
        "B": "・保存枠のサムネイル後読みと削除を追加\n・採取・削除用のデバッグ操作を追加\n・スタートメニューの一覧をスクロール対応",
        "D": "SnapshotSaveLoadPanel\nDebug menu / StartMenuController",
        "F": "作業中",
        "G": "保存枠の操作と\n表示を確認中",
    },
}

for row, values in implementation_rows.items():
    for col, value in values.items():
        set_value(ws, f"{col}{row}", value)

other_rows = {
    19: {
        "A": "マスターデータ/演出接続",
        "B": "・初回、共通、各ルートに演出204を接続\n・妹の見た目が変わらない行の演出を抑制\n・イライラエンド動画と遷移先を更新",
        "D": "master_ver2_* / master_route_*\nBugEffect 204",
        "F": "作業中",
        "G": "全ルートの演出を\n通し確認中",
    },
    20: {
        "A": "セーブ仕様/確認資料",
        "B": "・スナップショットの保存内容を資料化\n・Docs目次から保存内容の参照先を追加\n・実装範囲と確認すべき復元項目を整理",
        "D": "スナップショット保存内容一覧.md\nDocs README",
        "F": "作業中",
        "G": "保存範囲が分かる\n資料を整備中",
    },
    21: {
        "A": "フォント表示/本番確認",
        "B": "・日本語TMPフォントの再生成差分を確認\n・本番画面の文字欠けと表示崩れを確認\n・未コミット差分を本番確認対象として整理",
        "D": "TextMeshPro SDF assets\n本番確認画面",
        "F": "作業中",
        "G": "日本語表示を\n本番画面で確認中",
    },
    22: {
        "A": "報告書/テンプレート整備",
        "B": "・8/10、8/17〜8/21分の作業報告を作成\n・Git履歴と未コミットのフォント差分を反映\n・出社日と継続確認をテンプレートへ反映",
        "D": "本報告書\nエンジニア作業報告書テンプレート",
        "F": "完了",
        "G": "今週の継続確認を\n整理済み",
    },
}

for row, values in other_rows.items():
    for col, value in values.items():
        set_value(ws, f"{col}{row}", value)

set_value(
    ws,
    "A25",
    "・対象：8/10、8/17〜8/21。先週は8/10のみ稼働。\n"
    "・妹演出、イライラ進行、仮想OS設定、セーブ／ロード復元を確認中。\n"
    "・保存→ロード後の通し確認、各ルート、フォント表示を継続確認。",
)

if ws.row_dimensions[14].height:
    ws.row_dimensions[15].height = ws.row_dimensions[14].height

ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1

wb.save(out)
print(template)
print(out)
