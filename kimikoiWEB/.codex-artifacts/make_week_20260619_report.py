from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


template = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260619")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年6月第3週.xlsx"

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年6月第3週"


def set_value(cell_ref: str, value):
    cell = ws[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value


set_value("A2", "対象期間: 2026/06/15〜2026/06/19")
set_value("A4", "2026/06/15〜2026/06/19")
set_value("B4", "6/15, 6/16, 6/17, 6/18, 6/19")
set_value("C4", "キミと恋する15日（仮）")
set_value("D4", "Unity本番開発")
set_value(
    "E4",
    "本番開発として、LIME/Line基盤、ENDシナリオ、イベントCSV、Google Sheets連携、CSV読込整理、ADVフェーズ分離、UI/Prefab修正、書類更新を実施。完了済み項目と継続調整中項目を分けて整理。",
)

implementation_rows = {
    8: {
        "A": "LIME/Line基盤",
        "B": "・Lineスタンプカタログを整備\n・入力欄、相手入力、スタンプ、文字送りの基盤を修正\n・ADV側のLine通知クリックフローを調整",
        "D": "LimeWindowController.cs / LimeStickerCatalog.asset / KimikoiAdvTimelineBridgeController.cs",
        "F": "完了",
        "G": "LIME会話演出で、スタンプや入力中表現を扱うための土台が整った状態",
    },
    9: {
        "A": "ADV表示/選択肢修正",
        "B": "・選択肢時点のテキストメッセージ表示を修正\n・LinePrefabとのInspector紐づけを修正\n・Windowボタン参照外れとUIリスナー二重登録を修正",
        "D": "KimikoiAdvPlaybackDirector.cs / Main.unity / StartMenuView.prefab / KimikoiAdvWindowController.cs",
        "F": "完了",
        "G": "プレイ中に選択肢やウィンドウ操作で止まりにくくなるよう、表示と操作の不具合を修正済み",
    },
    10: {
        "A": "ENDシナリオ対応",
        "B": "・エンディング用CSVを追加\n・END番号以降はEND用CSVから進行するよう修正\n・Google Sheets読込対象にもENDシートを追加",
        "D": "master_scenario_09_END.csv / MasterScenarioController.cs / KimikoiScenarioGoogleSheetsLoader.cs",
        "F": "作業中",
        "G": "エンディング進行をデータで管理できる状態に近づけており、通し確認を進める段階",
    },
    11: {
        "A": "イベントCSV/トリガー",
        "B": "・イベントCSVの定義とトリガーを追加\n・master_day_any.csvとmaster_event_trigger.csvを追加\n・イベントCSVをGoogle Sheets読込に対応",
        "D": "master_day_any.csv / master_event_trigger.csv / KimikoiScenarioRuntimeSource.cs",
        "F": "作業中",
        "G": "日付や条件に応じたイベントを表で管理できるようにし、企画側の調整幅を広げている段階",
    },
    12: {
        "A": "CSV/Google Sheets整理",
        "B": "・CSVの#コメント列をスキップできるように対応\n・CSV/Google SheetsからAdvScriptLineへ変換するロジックを一本化\n・読込入口を整理し、不要なデフォルト指定を削除",
        "D": "KimikoiAdvScenarioCsvLoader.cs / CsvTableCellUtility.cs / KimikoiScenarioSheetSettings.cs",
        "F": "完了",
        "G": "シナリオ表の見通しを保ちながら、Unity側で安定して読み込める形に整理済み",
    },
    13: {
        "A": "ADV構成整理",
        "B": "・ADVのフェーズ管理とビジュアル管理を切り分け\n・CSV/Google Sheetsの入口管理を切り離し\n・移動済みスクリプトを削除し、重複した処理を整理",
        "D": "KimikoiAdvPhaseFlowController.cs / KimikoiAdvWindowController.cs / KimikoiAdvScenarioCsvLoader.cs",
        "F": "作業中",
        "G": "ADV本体を分かりやすく分割し、今後の演出追加や不具合修正を追いやすくしている状態",
    },
    14: {
        "A": "UI/Prefab/TextMeshPro",
        "B": "・TextMeshProフォントを更新\n・StartMenuやWindowボタン参照を修正\n・既存UIが実装変更後も操作できるようにPrefab/Scene側を調整",
        "D": "TextMeshPro assets / StartMenuView.prefab / Main.unity",
        "F": "完了",
        "G": "文字表示と基本操作まわりの崩れを抑え、プランナー確認時に画面を追いやすい状態へ調整済み",
    },
    15: {
        "A": "基盤安定化",
        "B": "・未整理だったコルーチン管理を修正\n・二重登録されていたUIリスナーを整理\n・CSV読込や画面進行の基盤を継続的に軽量化",
        "D": "DesktopShellController.cs / OuterAdvOverlayController.cs / MasterScenarioController.cs",
        "F": "作業中",
        "G": "動作が重複したり想定外に進むリスクを減らすため、内部処理を継続整理中",
    },
}

for row, values in implementation_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

other_rows = {
    19: {
        "A": "書類更新",
        "B": "・KimikoiScenario_Dev1_CurrentCsv.xlsxを更新\n・CSV説明関連のMDを修正\n・Data READMEとmaster_scenario_kimikoi_schema.mdを更新",
        "D": "Docs/Scenario/KimikoiScenario_Dev1_CurrentCsv.xlsx / Assets/Data/README.md / master_scenario_kimikoi_schema.md",
        "F": "完了",
        "G": "シナリオ表やCSV仕様を確認するための資料を、今週の実装内容に合わせて更新済み",
    },
    20: {
        "A": "仕様/データ整理",
        "B": "・CSVデータ説明の文字化けを解消\n・パラメータCSVに解説を追加\n・END、イベント、日付共通データをマスター側へ追加",
        "D": "master_parameter_changes.csv / master_scenario_09_END.csv / master_day_any.csv / master_event_trigger.csv",
        "F": "完了",
        "G": "企画側が表の意味を読み取りやすくなり、シナリオ調整時の確認負荷を下げた状態",
    },
    21: {
        "A": "運用ルール整理",
        "B": "・Cursor/実装ルール系ファイルの更新を確認\n・UI実装運用規約を更新\n・Prefab/Scene優先の作業ルールを継続整理",
        "D": ".cursor/rules / 実装運用規約_UI.md / 実装運用規約_CSharp.md",
        "F": "作業中",
        "G": "実装依頼時のルールを明確にし、手戻りや参照崩れを減らすための運用整理を継続中",
    },
    22: {
        "A": "報告書整備",
        "B": "・6/15〜6/19分の作業報告をテンプレートに沿って作成\n・Git状況、書類更新、実装状況を反映\n・確認内容をプランナー目線で分かる表現に整理",
        "D": "本報告書",
        "F": "完了",
        "G": "今週の実装進捗と継続確認点を、業務報告として読みやすい形に整理済み",
    },
}

for row, values in other_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

set_value(
    "A25",
    "・6/15〜6/19は本番開発として、LIME/Line演出、END進行、イベントCSV、Google Sheets連携、ADV基盤整理を実施。\n"
    "・ENDシナリオ、イベントCSV、ADV構成整理、基盤安定化、運用ルール整理は作業中として継続確認。\n"
    "・書類更新では、シナリオExcel、CSV仕様説明、Data README、マスターCSV関連資料を今週の実装内容に合わせて更新。",
)

if ws.row_dimensions[14].height:
    ws.row_dimensions[15].height = ws.row_dimensions[14].height

wb.save(out)
print(out)
