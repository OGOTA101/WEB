from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

template = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx")
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260523")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年5月第3週.xlsx"

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年5月第3週"

def set_value(cell_ref: str, value):
    cell = ws[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value

set_value("A2", "対象期間: 2026/05/18〜2026/05/23")
set_value("A4", "2026/05/18〜2026/05/23")
set_value("B4", "5/18, 5/19, 5/20, 5/21, 5/22, 5/23")
set_value("C4", "キミと恋する15日（仮）")
set_value("D4", "Unityモック")
set_value(
    "E4",
    "ゲーム本編の実装を通しプレイ前提まで進行。ゲーム内アプリ、シナリオ、UI/Prefab、検証環境は概ね完了扱いとし、妹・ホラー・グリッチ・通知などの演出調整を継続。",
)

implementation_rows = {
    8: {
        "A": "Unity基盤",
        "B": "・TextMeshPro/フォント/Resources構成を整理し、各画面で参照しやすい形に更新\n・不要ファイル除外や素材ロード整理を行い、実装継続に必要な基礎環境を維持",
        "D": "Fonts / TextMeshPro / Resources / .gitignore / ProjectSettings",
        "F": "完了",
        "G": "基盤更新は運用可能な状態まで整理",
    },
    9: {
        "A": "疑似デスクトップ",
        "B": "・Googru検索/STEALページ、通知、タスクマネージャ、フルスクリーン通知を実装\n・LIME通知からWindow展開、通知タップでWindowモード復帰する導線を整備",
        "D": "GoogruWindowController.cs / DesktopShellController.cs / TaskbarController.cs / TaskManagerWindowController.cs",
        "F": "完了",
        "G": "主要デスクトップ操作と通知導線を確認",
    },
    10: {
        "A": "ゲーム内アプリ",
        "B": "・キミコイADV、LIME、Googru、常駐妹、スキンシップ妹を接続\n・選択肢、Skip、メッセージ表示、LIMEリンク/通知、スキンシップWindowを調整",
        "D": "KimikoiAdvWindowController.cs / LimeWindowController.cs / GoogruWindowController.cs / SisterSkinshipWindowController.cs",
        "F": "完了",
        "G": "各アプリの基本導線を実装完了扱いで整理",
    },
    11: {
        "A": "シナリオ連携",
        "B": "・DAY1〜DAY5、共通DAY、選択肢、分岐、パラメータ変動をCSVに反映\n・十日参り、妹びいき選択不可、Day3死亡演出などのシナリオ不具合を修正",
        "D": "master_scenario_*.csv / KimikoiAdvMasterEmbeddedCopy.cs / MasterScenarioController.cs",
        "F": "完了",
        "G": "通し進行に必要なシナリオ反映を完了扱い",
    },
    12: {
        "A": "UI/Prefab調整",
        "B": "・LIMEフォント、REAL ADVメッセージ、常駐妹、スキンシップ、選択肢Prefabを調整\n・Resources配下に素材を再配置し、妹/キミコイ/CG/Referenceの整理を実施",
        "D": "Prefabs/Desktop / Resources / SisterResidentWindow.prefab / LimeWindow.prefab",
        "F": "完了",
        "G": "主要UI/Prefabの配置・参照を整理",
    },
    13: {
        "A": "検証環境",
        "B": "・フルスクリーン、通知、ホラー演出、デバッグパネル、日付進行の確認画像を保存\n・デバッグメニューやローカル設定を追加し、再現確認しやすい状態に整備",
        "D": "Screenshots / DesktopDebugMenuController.cs / DesktopDebugLocalPrefs.cs / UnityデバッグとEditor操作メモ.md",
        "F": "完了",
        "G": "確認用スクリーンショットとデバッグ導線を整備",
    },
    14: {
        "A": "演出調整",
        "B": "・妹初登場、常駐妹リアクション、死亡演出、ブルースクリーン、グリッチ/データモッシュを調整\n・フルスク時のLIME通知アニメーション、ノイズSE/BGM、ホラー演出の見え方を継続調整",
        "D": "SisterHorrorSequenceDirector.cs / KimikoiAdvBugEffectRunner.cs / SisterVisualTransitionDirector.cs / BlueScreenOverlayController.cs",
        "F": "進行中",
        "G": "演出のタイミング・見た目・音を継続調整",
    },
}

for row, values in implementation_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

other_rows = {
    19: {
        "A": "仕様整理",
        "B": "・シナリオ進行整理を更新し、実装済みのDAY進行・演出修正内容を反映\n・仕様書/シナリオ資料とUnity側の実装差分を確認",
        "D": "シナリオ進行整理.md / Resource/specification / Resource/scenario",
        "F": "進行中",
        "G": "実装と資料の差分を同期",
    },
    20: {
        "A": "素材整理",
        "B": "・妹ボイス、LIMEスタンプ、ノイズSE、グリッチBGM、常駐妹素材を追加・分類\n・Resources/Reference/CG/Sister/Kimikoi配下の参照構成を整理",
        "D": "Resource/image / Resource/sound / Resource/voice / Assets/Resources",
        "F": "進行中",
        "G": "素材投入と参照整理を実施",
    },
    21: {
        "A": "進行調整",
        "B": "・B番号単位の修正項目を実装へ反映し、通しプレイ中に見つかった挙動差分を修正\n・フルスク、通知、妹登場、死亡演出などを確認しながら微調整",
        "D": "各種修正コミット / Screenshots / master_scenario csv",
        "F": "進行中",
        "G": "確認と修正を並行",
    },
    22: {
        "A": "報告書整備",
        "B": "・5/18〜5/23分の作業報告を前週テンプレートに沿って整理\n・演出調整を新設し、それ以外の実装項目は完了扱いで整理",
        "D": "本報告書",
        "F": "完了",
        "G": "内容確認済み",
    },
}

for row, values in other_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

set_value(
    "A25",
    "・ゲーム内アプリ、シナリオ連携、UI/Prefab調整、検証環境は概ね完了扱い。\n"
    "・次週は新設した演出調整を中心に、妹・ホラー・グリッチ・通知・音声のタイミングと見え方を詰める。\n"
    "・5/24深夜以降の作業は本報告範囲外とし、次回報告で扱う。",
)

wb.save(out)
print(out)
