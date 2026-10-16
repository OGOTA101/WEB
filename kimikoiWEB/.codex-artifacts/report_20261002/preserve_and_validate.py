import json
import shutil
import sys
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook
from PIL import ImageFont


SUPPORT = Path(__file__).parent
SPEC_PATH = SUPPORT / "authored_values.json"
SPEC = json.loads(SPEC_PATH.read_text(encoding="utf-8"))
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
FONT = ImageFont.truetype(r"C:/Windows/Fonts/YuGothR.ttc", 12, index=0)
ROWS = [*range(8, 16), *range(19, 23)]


def preflight():
    issues = []
    limits = {"A": 82, "B": 244, "D": 228, "G": 128, "H": 112}
    for scope in ["reportCells", "templateCells"]:
        for addr, value in SPEC[scope].items():
            column = addr[0]
            if int(addr[1:]) not in ROWS or column not in limits:
                continue
            lines = value.splitlines()
            for line in lines:
                width = FONT.getlength(line)
                if width > limits[column]:
                    issues.append(f"{scope} {addr}: {width:.1f}px > {limits[column]}px: {line}")
            if column == "B" and value and not 3 <= len(lines) <= 4:
                issues.append(f"{scope} {addr}: {len(lines)} work lines")
            if column == "G" and value and not 3 <= len(lines) <= 5:
                issues.append(f"{scope} {addr}: {len(lines)} confirmation lines")
            if column in ["B", "G"] and value and scope == "reportCells":
                if sum(line.startswith("・") for line in lines) != 3:
                    issues.append(f"{addr}: expected 3 bullets")
            if any(len(line.strip().lstrip("・")) == 1 for line in lines):
                issues.append(f"{scope} {addr}: single-character line")
    for scope in ["reportCells", "templateCells"]:
        for addr, limit in [("A4", 82), ("B4", 101), ("E4", 484), ("A25", 1040)]:
            for line in SPEC[scope].get(addr, "").splitlines():
                if FONT.getlength(line) > limit:
                    issues.append(f"{scope} {addr}: header/remarks width exceeded")
    if issues:
        raise ValueError("\n".join(issues))
    print("PREFLIGHT PASS: report and template line counts/widths")


def cell_values(archive):
    strings = []
    if "xl/sharedStrings.xml" in archive.namelist():
        root = etree.fromstring(archive.read("xl/sharedStrings.xml"))
        strings = ["".join(node.itertext()) for node in root.findall("m:si", NS)]
    root = etree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    values = {}
    for cell in root.findall("m:sheetData/m:row/m:c", NS):
        kind = cell.get("t")
        value = cell.find("m:v", NS)
        if kind == "s" and value is not None:
            text = strings[int(value.text)]
        elif kind == "inlineStr":
            text = "".join(cell.xpath("./m:is//m:t/text()", namespaces=NS))
        elif value is not None:
            text = value.text
        else:
            text = None
        values[cell.get("r")] = text
    return values


def write_native_layout(source, artifact_export, destination, expected, sheet_name):
    with ZipFile(artifact_export) as archive:
        exported_values = cell_values(archive)
    for address, value in expected.items():
        if (exported_values.get(address) or "") != value:
            raise ValueError(f"Artifact content mismatch at {address}")

    with ZipFile(source) as archive:
        entries = [(item, archive.read(item.filename)) for item in archive.infolist()]
    contents = dict((item.filename, data) for item, data in entries)
    worksheet = etree.fromstring(contents["xl/worksheets/sheet1.xml"])
    cells = {cell.get("r"): cell for cell in worksheet.findall("m:sheetData/m:row/m:c", NS)}
    for address, value in expected.items():
        cell = cells[address]
        for child in list(cell):
            if etree.QName(child).localname in ["v", "f", "is"]:
                cell.remove(child)
        cell.attrib.pop("t", None)
        if value:
            cell.set("t", "inlineStr")
            inline = etree.SubElement(cell, f"{{{NS['m']}}}is")
            text_node = etree.SubElement(inline, f"{{{NS['m']}}}t")
            text_node.set(XML_SPACE, "preserve")
            text_node.text = exported_values[address]

    # Artifact Tool has no documented print-layout surface; keep native template styles/merges.
    setup = worksheet.find("m:pageSetup", NS)
    if setup is None:
        raise ValueError("Missing native template pageSetup")
    setup.set("orientation", "portrait")
    setup.set("paperSize", "9")
    setup.set("fitToWidth", "1")
    setup.set("fitToHeight", "1")
    setup_pr = worksheet.find("m:sheetPr/m:pageSetUpPr", NS)
    if setup_pr is None:
        raise ValueError("Missing native template pageSetUpPr")
    setup_pr.set("fitToPage", "1")
    contents["xl/worksheets/sheet1.xml"] = etree.tostring(worksheet, encoding="utf-8")

    workbook = etree.fromstring(contents["xl/workbook.xml"])
    workbook.find("m:sheets/m:sheet", NS).set("name", sheet_name)
    defined_names = workbook.find("m:definedNames", NS)
    if defined_names is None:
        raise ValueError("Missing template print area")
    print_area = defined_names.find("m:definedName[@name='_xlnm.Print_Area']", NS)
    print_area.text = f"'{sheet_name}'!$A$1:$H$28"
    contents["xl/workbook.xml"] = etree.tostring(workbook, encoding="utf-8")

    destination.parent.mkdir(parents=True, exist_ok=True)
    try:
        with ZipFile(destination, "w") as output:
            for item, original in entries:
                output.writestr(item, contents.get(item.filename, original))
    except PermissionError:
        alternate = destination.with_name(f"{destination.stem}_修正版.xlsx")
        with ZipFile(alternate, "w") as output:
            for item, original in entries:
                output.writestr(item, contents.get(item.filename, original))
        return alternate
    return destination


def verify_saved(file_path, expected, source_path):
    sheet = load_workbook(file_path).active
    original = load_workbook(source_path).active
    assert sheet.print_area == f"'{sheet.title}'!$A$1:$H$28"
    assert sheet.page_setup.orientation == "portrait"
    assert (sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight) == (1, 1)
    assert sheet.sheet_properties.pageSetUpPr.fitToPage
    assert set(map(str, sheet.merged_cells.ranges)) == set(map(str, original.merged_cells.ranges))
    for address, value in expected.items():
        assert (sheet[address].value or "") == value, address
        assert sheet[address]._style == original[address]._style, address
    for col in "ABCDEFGH":
        assert sheet.column_dimensions[col].width == original.column_dimensions[col].width
    for row in range(1, 29):
        assert sheet.row_dimensions[row].height == original.row_dimensions[row].height
    errors = [cell.coordinate for row in sheet.iter_rows() for cell in row
              if cell.data_type == "e" or (isinstance(cell.value, str) and cell.value.startswith("#"))]
    assert not errors, errors
    assert not [cell.coordinate for row in sheet.iter_rows() for cell in row if cell.data_type == "f"]
    result = {"path": str(file_path), "sheet": sheet.title, "print_area": sheet.print_area,
              "portrait_fit": "1x1", "formula_errors": errors, "native_styles_preserved": True}
    if "20261002" in str(file_path):
        result["complete_rows"] = [row for row in ROWS if sheet[f"F{row}"].value == "完了"]
        result["working_rows"] = [row for row in ROWS if sheet[f"F{row}"].value == "作業中"]
        result["B_max_lines"] = max(len(sheet[f"B{row}"].value.splitlines()) for row in ROWS)
        result["G_max_lines"] = max(len(sheet[f"G{row}"].value.splitlines()) for row in ROWS)
        result["G_max_pixels"] = max(FONT.getlength(line) for row in ROWS for line in sheet[f"G{row}"].value.splitlines())
    return result


preflight()
if len(sys.argv) > 1 and sys.argv[1] == "save":
    template = Path(SPEC["templatePath"])
    original = SUPPORT / "template_before.xlsx"
    if not original.exists():
        shutil.copy2(template, original)
    report_path = write_native_layout(original, SUPPORT / "report_artifact.xlsx", Path(SPEC["outputPath"]),
                                      SPEC["reportCells"], SPEC["sheetName"])
    template_path = write_native_layout(original, SUPPORT / "template_artifact.xlsx", template,
                                        SPEC["templateCells"], "作業報告テンプレート")
    result = [verify_saved(report_path, SPEC["reportCells"], original),
              verify_saved(template_path, SPEC["templateCells"], original)]
    (SUPPORT / "validation.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
