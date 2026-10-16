from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


template = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260717")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年7月第3週.xlsx"

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年7月第3週"


def set_value(cell_ref: str, value):
    cell = ws[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value


set_value("A2", "対象期間: 2026/07/13〜2026/07/17")
set_value("A4", "2026/07/13〜2026/07/17")
set_value("B4", "7/13, 7/14, 7/15, 7/16, 7/17")
set_value("C4", "キミと恋する15日（仮）")
set_value("D4", "Unity本番開発")
set_value(
    "E4",
    "本番開発として、妹データ削除のゲーム内導線、怨念/イライラ分岐、外側ADVとゲームオーバー、妹ゲージ演出、LIME通知調整、スクリーンショットのSlack投稿、定期自動ビルド構築を進行。前週までに完了済みの項目は再掲せず、今週の新規・継続作業を整理。",
)

implementation_rows = {
    8: {
        "A": "妹データ削除導線",
        "B": "・妹データ削除用の仮想ファイル、Catalog、Window定義を追加\n・ファイル一覧の横並びアイコン表示とゴミ箱連携を実装\n・最終妹データをDAY10までロックし、削除通知の対象を統一",
        "D": "SisterDataDeleteStateCatalog.asset / FileExplorerWindowController.cs / RecycleBinWindowController.cs",
        "F": "作業中",
        "G": "プレイヤーが妹データを削除した結果をゲーム進行と画面表示へ反映するための導線を整備中",
    },
    9: {
        "A": "削除状態/ビジュアル反映",
        "B": "・妹データ削除状態をランタイムで保持する処理を追加\n・常駐妹、スキンシップ妹、ADVビジュアルへvisualLossを反映\n・EVENT30直前の仮命乞い表示を外し、演出表示を調整",
        "D": "SisterDataDeleteRuntimeState.cs / SisterResidentWindowController.cs / KimikoiAdvVisualController.cs",
        "F": "作業中",
        "G": "削除したデータが妹の見た目や反応へ反映されるようにし、選択の重みを画面で伝える準備を進めている段階",
    },
    10: {
        "A": "怨念/イライラ分岐",
        "B": "・怨念、イライラ、ときめきルートのDAY CSVを更新\n・怨念/イライラのDayAny分岐とDAY10分岐を追加\n・しんらい未整備時の怨念ルート退避、イライラ助け成功の記録を追加",
        "D": "master_route_*.csv / master_day_any_ver2.csv / KimikoiSisterParamUtility.cs",
        "F": "作業中",
        "G": "パラメータや選択に応じて、企画意図どおりのルートへ進むかを通し確認する段階",
    },
    11: {
        "A": "イベント/ゲームオーバー",
        "B": "・怨念INTERVAL_TICKの実行、継続加算、開始タイミングを調整\n・助け要請の成功/失敗、外側ADV完了、ゲームオーバーをイベントから接続\n・DayAny終了後にキミコイ本編へ復帰する処理を追加",
        "D": "KimikoiOnnenIntervalTickController.cs / MasterScenarioController.cs / KimikoiAdvSisterGameOverController.cs",
        "F": "作業中",
        "G": "イベント途中や別アプリを挟んだ後も、正しい条件で本編復帰・ゲームオーバーへ進むよう調整中",
    },
    12: {
        "A": "妹ゲージ演出",
        "B": "・妹ゲージごとの枠スケールと順次Fill用Tween基盤を追加\n・常駐妹とスキンシップ妹のゲージを上から順に変化させる処理を実装\n・ゲージ変化の間隔設定を追加",
        "D": "SisterParameterGaugeTween.cs / SisterResidentWindowController.cs / SisterSkinshipWindowController.cs",
        "F": "完了",
        "G": "パラメータ変化を段階的に見せることで、プレイヤーと確認者が変動内容を把握しやすい状態へ改善済み",
    },
    13: {
        "A": "LIME通知/UI調整",
        "B": "・LIME通知行の本文がはみ出す表示を修正\n・通知表示のテキスト配置と収まりを調整\n・通知経由の確認時に内容を読み取りやすくするためのUI修正を反映",
        "D": "DesktopNotificationRow_Lime.prefab / DesktopNotificationRowView.cs / LimeDesktopNotificationRowView.cs",
        "F": "完了",
        "G": "通知本文が読み切れない状態を解消し、シナリオ確認時にメッセージ内容を追いやすくした状態",
    },
    14: {
        "A": "スクリーンショットSlack投稿",
        "B": "・Ctrl+Dからスクリーンショットを投稿する処理を追加\n・送信先設定とスクリーンショット用文言を追加\n・スクリーンショットがSlackに流れるところまで動作確認を完了",
        "D": "DesktopDebugMenuController.cs / DesktopShellSettings.asset / DesktopEmbeddedCopy.asset",
        "F": "完了",
        "G": "確認中の画面をSlackで関係者へ共有できる状態となり、演出確認や不具合共有をすぐに開始できる状態",
    },
    15: {
        "A": "シナリオ/パラメータ調整",
        "B": "・共通DAY、初回DAY、パラメータ変更CSVの調整差分を作業中\n・妹データ削除、助け要請、ルート分岐に合わせてデータ整合を確認\n・未コミット差分を含め、通し確認用のシナリオデータを更新中",
        "D": "GeneratedVer2/master_ver2_*.csv / master_parameter_changes_ver2.csv",
        "F": "作業中",
        "G": "実装した分岐や演出とシナリオ表の内容が一致するよう、企画データの調整と確認を継続中",
    },
}

for row, values in implementation_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

other_rows = {
    19: {
        "A": "マスターデータ更新",
        "B": "・妹リアクションCatalogを更新\n・怨念/イライラ/ときめきのルートCSVを更新\n・削除DayAny、EventTrigger、ADV埋め込み文言を更新",
        "D": "SisterResidentReactionCatalog.asset / SisterSkinshipReactionCatalog.asset / master_event_trigger_ver2.csv",
        "F": "作業中",
        "G": "シナリオ条件と妹の反応データを同期し、企画上の分岐内容を実装へ反映する作業を継続中",
    },
    20: {
        "A": "定期自動ビルド構築",
        "B": "・定期実行する自動ビルドの構成を検討・構築\n・ビルド成果物を確認/共有できる運用を整理\n・本番開発の継続確認に使えるビルド生成フローを準備",
        "D": "定期ビルド設定 / Build Profiles / ビルド成果物運用",
        "F": "作業中",
        "G": "定期的に確認用ビルドを用意できるようにし、関係者が最新状態を試せる運用を構築中",
    },
    21: {
        "A": "実装状況整理",
        "B": "・前週報告の作業中項目から、ルート進行と妹関連の実装を継続\n・完了済みの前週項目は再掲せず、今週の追加分のみを整理\n・未コミットのシナリオ調整差分は次回確認対象として記録",
        "D": "outputs/work_reports_20260710 / Git作業ツリー / GeneratedVer2 CSV",
        "F": "完了",
        "G": "先週からの継続範囲と今回の追加作業を分け、確認対象が混ざらないよう整理済み",
    },
    22: {
        "A": "報告書整備",
        "B": "・7/13〜7/17分の作業報告をテンプレートに沿って作成\n・Git状況と未コミット差分を反映\n・作業中項目をプランナー目線で確認しやすい表現に整理",
        "D": "本報告書",
        "F": "完了",
        "G": "今週の確認対象と継続判断が分かるよう、業務報告として整理済み",
    },
}

for row, values in other_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

set_value(
    "A25",
    "・対象期間は2026/07/13〜2026/07/17。前週までに完了として報告済みの項目は再掲しない。\n"
    "・スクリーンショットのSlack投稿は成功を確認し完了。主な継続作業は、妹データ削除、怨念/イライラ分岐、イベント/ゲームオーバー、定期自動ビルド、シナリオ調整。\n"
    "・次回は、削除状態と各ルートの通し確認、定期自動ビルドの運用確認、シナリオCSVの未コミット差分確認を継続する。",
)

if ws.row_dimensions[14].height:
    ws.row_dimensions[15].height = ws.row_dimensions[14].height

wb.save(out)
print(out)
