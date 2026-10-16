from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


template = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260529")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年5月第4週.xlsx"

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年5月第4週"


def set_value(cell_ref: str, value):
    cell = ws[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value


set_value("A2", "対象期間: 2026/05/24〜2026/05/29")
set_value("A4", "2026/05/24〜2026/05/29")
set_value("B4", "5/24, 5/25, 5/26, 5/27, 5/28, 5/29")
set_value("C4", "キミと恋する15日（仮）")
set_value("D4", "Unityモック")
set_value(
    "E4",
    "ゲーム実装全体を完了扱いに整理。演出調整、QA反映、BGM/SE/ボイス、スキンシップ、通知/アプリ連携、シナリオ反映、工数表更新まで実施。",
)

implementation_rows = {
    8: {
        "A": "Unity基盤",
        "B": "・ProjectSettings、TextMeshPro、AudioOutputSettings、ScriptableObject命名を整理\n・ビルド時のCanvasスケーリング不具合や音量初期値を調整し、実装基盤を最終化",
        "D": "ProjectSettings / TextMeshPro / AudioOutputSettings / DesktopShellSettings",
        "F": "完了",
        "G": "実装運用に必要な基盤設定を整理完了",
    },
    9: {
        "A": "疑似デスクトップ",
        "B": "・ウィンドウ復帰、通知、タスクバー時計、タスクマネージャ、キミコイ終了時の挙動を調整\n・Googru/STEALページやLIME通知導線を更新し、疑似デスクトップ操作を完了扱いに整理",
        "D": "DesktopShellController.cs / DesktopWindowManager.cs / GoogruWindowController.cs / TaskbarView.prefab",
        "F": "完了",
        "G": "主要なデスクトップ操作と通知導線を確認済み",
    },
    10: {
        "A": "ゲーム内アプリ",
        "B": "・キミコイADV、LIME、Googru、常駐妹、スキンシップ妹を接続\n・選択肢、クリック待ち、スキンシップWindow、妹パラメータ、通知/終了挙動を調整",
        "D": "KimikoiAdvWindowController.cs / LimeWindowController.cs / GoogruWindowController.cs / SisterSkinshipWindowController.cs",
        "F": "完了",
        "G": "ゲーム内アプリの基本導線を完了扱いで整理",
    },
    11: {
        "A": "シナリオ連携",
        "B": "・Day2演出の全繋ぎ込み、妹登場後の選択肢前セリフ、クリック待ち箇所を反映\n・master_scenarioと演出発火処理を更新し、BGM/SE/Voice切替とシナリオ進行を同期",
        "D": "master_scenario_*.csv / MasterScenarioController.cs / Resource/scenario",
        "F": "完了",
        "G": "シナリオ進行と演出発火を確認済み",
    },
    12: {
        "A": "UI/Prefab調整",
        "B": "・REAL ADVメッセージ背景、妹立ち絵位置、ゲージ倍率、常駐妹/スキンシップ素材を調整\n・カーソル描画順、Canvas設定、各Prefab参照を見直し、画面表示の崩れを修正",
        "D": "Main.unity / Prefabs/Desktop / Resources/UI / Resources/Sister",
        "F": "完了",
        "G": "主要UI/Prefabの表示と参照を整理済み",
    },
    13: {
        "A": "検証環境",
        "B": "・デバッグパネル、確認用スクリーンショット、フルスク/通知/ゲームオーバー確認を整備\n・キミコイをタスクマネージャで終了した際のエフェクト中断など、再現確認しやすい導線を確認",
        "D": "Screenshots / DesktopDebugMenuController.cs / DesktopDebugMenuController.BugEffectDebug.cs",
        "F": "完了",
        "G": "検証用の確認素材と再現導線を整理済み",
    },
    14: {
        "A": "演出調整",
        "B": "・グリッチ、ノイズ、魚眼、ビネット、白飛び、砂嵐、ホラー演出、マウス/キーボード音を調整\n・妹HORROR、常駐妹ノイズ、スキンシップ、BGM差し替え、音量バランスを最終確認",
        "D": "KimikoiAdvBugEffectRunner.cs / KimikoiAdvBugEffectCatalog.asset / SisterHorrorSequenceDirector.cs / AudioOutputSettings",
        "F": "完了",
        "G": "見た目・音・タイミングを完了扱いで整理",
    },
    15: {
        "A": "QA",
        "B": "・Day2 BGM、STEALページ、エフェクト発生、画面白飛び、Canvasスケーリング、通知/終了挙動を修正\n・確認画像と修正コミットをもとに、QA指摘反映後の再確認を実施",
        "D": "Screenshots / 修正コミット / QA確認項目",
        "F": "完了",
        "G": "指摘反映と再確認を完了扱いで整理",
    },
}

for row, values in implementation_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

other_rows = {
    19: {
        "A": "仕様整理",
        "B": "・関連資料を追加し、実装照合チェックシートと仕様mdを更新\n・Unity実装状況に合わせて、仕様書/シナリオ資料との差分を整理",
        "D": "Resource/specification/md / 実装照合チェックシート.md / Resource/scenario",
        "F": "完了",
        "G": "実装と資料の同期を実施",
    },
    20: {
        "A": "素材整理",
        "B": "・BGM/SE/Voice、妹パラメータ素材、ノイズ素材、UI素材を追加・分類\n・Unity側のResources参照に合わせて素材配置と命名を整理",
        "D": "Resource/image / Resource/sound / Resource/voice / Assets/Resources",
        "F": "完了",
        "G": "素材投入と参照整理を実施",
    },
    21: {
        "A": "工数・MOC対応",
        "B": "・工数表を更新し、MOCを受けた仕様変更に合わせて工数を再計算\n・パラメーター調整概要と関連資料を更新し、資料側の内容を実装状況に合わせて整理",
        "D": "きみこい・工数表_ダウンチャート_自動化完成版.xlsx / パラメーター調整概要.xlsx / パラメーター調整概要.md",
        "F": "完了",
        "G": "MOC反映後の工数再計算を実施",
    },
    22: {
        "A": "報告書整備",
        "B": "・5/24〜5/29分の作業報告を前週テンプレートに沿って整理\n・演出調整とQAを明示し、全体を完了認識の週報としてまとめ直し",
        "D": "本報告書",
        "F": "完了",
        "G": "内容確認済み",
    },
}

for row, values in other_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

set_value(
    "A25",
    "・5/24〜5/29はゲーム実装、演出調整、QA反映まで全体完了扱い。\n"
    "・QA項目を新設し、指摘反映と再確認結果を明示。\n"
    "・工数表更新とMOC後の工数再計算は、実装以外の業務として記録。",
)

if ws.row_dimensions[14].height:
    ws.row_dimensions[15].height = ws.row_dimensions[14].height

wb.save(out)
print(out)
