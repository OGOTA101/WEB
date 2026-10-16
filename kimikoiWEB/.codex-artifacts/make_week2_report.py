from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

template = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx")
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260515")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年5月第2週.xlsx"

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年5月第2週"

def set_value(cell_ref: str, value):
    cell = ws[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value

set_value("A2", "対象期間: 2026/05/09〜2026/05/15")
set_value("A4", "2026/05/09〜2026/05/15")
set_value("B4", "5/9, 5/10, 5/11, 5/12, 5/13, 5/14, 5/15")
set_value("C4", "キミと恋する15日（仮）")
set_value("D4", "Unityモック")
set_value(
    "E4",
    "ゲーム進行基盤を大きく前進。セーブ、DAY進行、シナリオ分割CSV、LIME進行、EVCG/演出、画面調整まで、通しプレイに必要な要素を段階的に整備。",
)

implementation_rows = {
    8: {
        "A": "Unity基盤",
        "B": "・Yu Gothicフォントを追加し、Scene/Prefab全体の表示フォントを統一\n・ビルド設定、起動時ダイアログ、埋め込み文言のScriptableObject化を整備",
        "D": "Fonts / ProjectSettings / DesktopEmbeddedCopy / DesktopShellModalDialogsController.cs",
        "F": "完了",
        "G": "全体表示とビルド周辺の基礎設定を整理",
    },
    9: {
        "A": "疑似デスクトップ",
        "B": "・通知表示、タスクマネージャ起動、デスクトップ色変更を追加\n・2560基準の画面サイズと共有PSDに合わせて配置を再調整",
        "D": "TaskbarController.cs / DesktopShellController.cs / DesktopThemeSettings.cs / Main.unity",
        "F": "完了",
        "G": "主要画面の配置・操作を更新",
    },
    10: {
        "A": "ゲーム内アプリ",
        "B": "・キミコイADV、LIME、Googru、常駐妹まわりを継続実装\n・LIMEクリック進行、送信ボタン表示、選択肢表示、EVCG表示、妹パラメータ管理を追加",
        "D": "KimikoiAdvWindowController.cs / LimeWindowController.cs / GoogruWindowController.cs / SisterStatusWindowController.cs",
        "F": "進行中",
        "G": "DAY進行と各アプリ連携を段階確認",
    },
    11: {
        "A": "シナリオ連携",
        "B": "・マスターCSVを日別に分割し、DAY1〜DAY5の進行とループを整備\n・背景、BGM、CG、立ち絵、分岐、会話内容をシナリオデータへ反映",
        "D": "master_scenario_01〜08.csv / MasterScenarioCsvBundle / SaveManager.cs",
        "F": "進行中",
        "G": "通しプレイ、ループ、日跨ぎの不具合を修正",
    },
    12: {
        "A": "UI/Prefab調整",
        "B": "・UI素材をResources配下へ再整理し、キミコイ/LIME/妹/グリッチ系素材を追加\n・テキストウィンドウ、ログ、背景、通知、EVCG、各種Prefab表示を調整",
        "D": "Resources/UI / Prefabs/Desktop / KimikoiAdvBackgroundData.asset / KimikoiWindow.prefab",
        "F": "進行中",
        "G": "表示差分・演出差分を都度反映",
    },
    13: {
        "A": "検証環境",
        "B": "・DAY5までの通しプレイ、日数経過、選択肢、LIME、CG切替を重点確認\n・デバッグ表示、検証用スクリーンショット、Editor操作メモを整備",
        "D": "Screenshots / DesktopDebugMenuController.cs / UnityデバッグとEditor操作メモ.md",
        "F": "進行中",
        "G": "不具合修正と再確認を継続",
    },
}

for row, values in implementation_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

other_rows = {
    19: {
        "A": "仕様整理",
        "B": "・キミコイ、LIME、仮想デスクトップ、BGM/SE、妹素材などの仕様書を更新\n・シナリオ進行整理とパラメーター調整の資料を追加",
        "D": "Resource/specification/* / シナリオ進行整理.md / パラメーター調整概要.xlsx",
        "F": "進行中",
        "G": "実装と資料の差分を同期",
    },
    20: {
        "A": "素材整理",
        "B": "・妹立ち絵、ホラー演出CG、EVCG仮素材、グリッチノイズ、SE/BGMを追加・分類\n・Resources構成を整理し、実装側から参照しやすい形へ移行",
        "D": "Resource/image / Resource/sound / Assets/Resources",
        "F": "進行中",
        "G": "素材投入と参照整理を実施",
    },
    21: {
        "A": "進行調整",
        "B": "・DAY単位のシナリオ修正、背景/BGM差し替え、選択肢不具合を随時反映\n・ゲーム全体の通し確認に合わせて演出・表示仕様を詰めた",
        "D": "scenario xlsx / master_scenario csv / 各種修正コミット",
        "F": "進行中",
        "G": "実装確認と仕様調整を並行",
    },
    22: {
        "A": "報告書整備",
        "B": "・5/9〜5/15分の作業報告を前週テンプレートに沿って整理\n・完了/進行中の区分を今週時点の認識に合わせて更新",
        "D": "本報告書",
        "F": "完了",
        "G": "内容確認済み",
    },
}

for row, values in other_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

set_value(
    "A25",
    "・Unity基盤と疑似デスクトップは完了扱いとし、以後はゲーム進行・演出・通し確認へ比重を移す。\n"
    "・次週はDAY進行、LIME、EVCG、妹パラメータ、各種演出の整合確認を優先。\n"
    "・シナリオ/仕様資料/実装差分を継続同期し、通しプレイで残る不具合を順次解消する。",
)

wb.save(out)
print(out)
