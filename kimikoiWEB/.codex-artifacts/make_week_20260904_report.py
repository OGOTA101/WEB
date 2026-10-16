from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from PIL import ImageFont


template = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260904")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年9月第1週.xlsx"


def set_value(sheet, cell_ref: str, value: str) -> None:
    cell = sheet[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value


def write_rows(sheet, rows: dict[int, dict[str, str]]) -> None:
    for row, values in rows.items():
        for column, value in values.items():
            set_value(sheet, f"{column}{row}", value)


def apply_page_settings(sheet) -> None:
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1


def validate(path: Path, report_rows: list[int] | None = None) -> None:
    workbook = load_workbook(path, data_only=False)
    sheet = workbook.active
    expected_print_area = f"'{sheet.title}'!$A$1:$H$28"
    if sheet.print_area != expected_print_area:
        raise ValueError(f"Unexpected print area: {sheet.print_area}")

    font = ImageFont.truetype(r"C:/Windows/Fonts/YuGothR.ttc", 12, index=0)
    if report_rows:
        for row in report_rows:
            work_lines = sheet[f"B{row}"].value.splitlines()
            if not 3 <= len(work_lines) <= 4:
                raise ValueError(f"B{row} must contain 3 to 4 lines: {work_lines}")
            if sum(line.startswith("・") for line in work_lines) != 3:
                raise ValueError(f"B{row} must contain three bullet lines: {work_lines}")
            if any(font.getlength(line) > 244 for line in work_lines):
                raise ValueError(f"B{row} exceeds the usable width: {work_lines}")

            confirmation_lines = sheet[f"G{row}"].value.splitlines()
            if not 3 <= len(confirmation_lines) <= 5:
                raise ValueError(f"G{row} must contain 3 to 5 lines: {confirmation_lines}")
            if sum(line.startswith("・") for line in confirmation_lines) != 3:
                raise ValueError(f"G{row} must contain three bullet lines: {confirmation_lines}")
            if any(font.getlength(line) > 128 for line in confirmation_lines):
                raise ValueError(f"G{row} exceeds the usable width: {confirmation_lines}")

            result_lines = sheet[f"D{row}"].value.splitlines()
            if len(result_lines) > 2 or any(font.getlength(line) > 228 for line in result_lines):
                raise ValueError(f"D{row} exceeds the usable width: {result_lines}")

    errors = [
        cell.coordinate
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("#")
    ]
    if errors:
        raise ValueError(f"Formula errors found: {errors}")


# Update the reusable template with the current column-specific writing rules.
template_wb = load_workbook(template)
template_ws = template_wb.active
set_value(template_ws, "B4", "作業日（出社があれば明記）")
set_value(
    template_ws,
    "G8",
    "・確認欄は合計5行以内\n・操作と期待する結果\n　を書く\n・長い内容は改行",
)
set_value(template_ws, "G19", "・操作と結果を書く\n・全体で5行以内")
set_value(
    template_ws,
    "A25",
    "・対象日と出社日を記載（出社がない週も明記）。\n"
    "・作業内容は3〜4行、確認欄は5行以内。\n"
    "・前週に完了として報告した項目は再掲しない。",
)
apply_page_settings(template_ws)
template_wb.save(template)

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年9月第1週"
set_value(ws, "A2", "対象期間: 2026/08/31〜2026/09/04（出社なし）")
set_value(ws, "A4", "2026/08/31〜2026/09/04")
set_value(ws, "B4", "8/31〜9/4\n出社なし")
set_value(ws, "C4", "キミと恋する15日（仮）")
set_value(ws, "D4", "Unity本番開発")
set_value(
    ws,
    "E4",
    "本番開発として、シナリオ／検索導線、LIME・ADV、妹演出、本番素材を調整中。前週完了項目は再掲しない。",
)

write_rows(
    ws,
    {
        8: {
            "A": "シナリオCSV/\nルート・検索接続",
            "B": "・Ver2.0と信頼ルートCSVを反映\n・信頼DAY6〜10の検索語を接続\n・分岐後の数値反映を調整",
            "D": "Master CSV / Route CSV\nGoogruWindowController",
            "F": "作業中",
            "G": "・信頼DAY6〜10を\n　順に進める\n・検索結果とページ\n　を確認\n・分岐先を確認",
        },
        9: {
            "A": "共通DAY2/\n検索・公式サイト",
            "B": "・検索結果と公式サイトをゲーム内へ接続\n・検索結果から公式サイトを開く導線を実装\n・閲覧後にADVへ戻る進行を確認",
            "D": "GoogruWindowController\nGoogruWebBrowserBridge",
            "F": "完了",
            "G": "・共通DAY2を進める\n・検索結果から公式を\n　開く\n・会話の再開を確認",
        },
        10: {
            "A": "LIME/ADV\n表示・入力",
            "B": "・開発者LIME窓と話者トークンを追加\n・LIMEの操作不可状態とSE重複を調整\n・ADVとLIMEの表示遷移を調整",
            "D": "LimeDeveloperWindow\nLimeWindowController",
            "F": "作業中",
            "G": "・LIMEを開き会話を\n　進める\n・操作不可の場面\n　を確認\n・通知音が重ならない",
        },
        11: {
            "A": "イライラ/\n妹演出進行",
            "B": "・イライラ監視の開始条件を調整\n・妹演出のノイズとボイス進行を調整\n・終盤の窓状態と削除待機を調整",
            "D": "IrairaEyeWindowDirector\nKimikoiAdvTimelineBridge",
            "F": "作業中",
            "G": "・監視開始から終盤\n　進める\n・妹のノイズと音声\n　を確認\n・終了時の窓を確認",
        },
        12: {
            "A": "背景/CG/\nフォント",
            "B": "・本番背景の命名とEVCG参照を調整\n・追加イベントCGをADVへ接続\n・TMPフォントのDynamicデータを整理",
            "D": "KimikoiAdvBackgroundData\nPortraitAndCgData",
            "F": "作業中",
            "G": "・各場面で背景を確認\n・CG表示を確認\n・日本語が欠けず\n　読める",
        },
        13: {
            "A": "スキンシップ/\nデスクトップ操作",
            "B": "・スキンシップの回数と表示を調整\n・予告表示とデスクトップカーソルを調整\n・終了後の窓表示と起動画面を調整",
            "D": "SisterSkinshipWindowController\nDesktopShell / TaskManager",
            "F": "作業中",
            "G": "・操作回数を確認\n・予告表示とカーソル\n　を確認\n・終了後の窓を確認",
        },
        14: {
            "A": "終盤ボイス/\nノイズ演出",
            "B": "・怨念終盤の専用ボイス再生を調整\n・LIME起動SEの重複を抑制\n・ノイズと文字化け表示を調整",
            "D": "KimikoiAdvWindowController\nSoundManager",
            "F": "作業中",
            "G": "・怨念終盤まで進める\n・専用ボイスの再生\n　を確認\n・LIME起動音が\n　重ならない",
        },
        15: {
            "A": "本番設定/\n検証導線",
            "B": "・本番用版番号を0011〜0015へ更新\n・起動画面と描画・PlayMode設定を調整\n・ルート分岐とデバッグ導線を調整",
            "D": "ProjectSettings / Main.unity\nDesktopDebugMenu",
            "F": "作業中",
            "G": "・本番版の起動を確認\n・各ルートを選んで\n　進める\n・デバッグ導線を確認",
        },
        19: {
            "A": "シナリオ資料/\nCSV整備",
            "B": "・シナリオExcelをマスターへ反映\n・CSV列定義と資料の対応を整理\n・各ルートの確認対象を一覧化",
            "D": "Scenario Excel Ver2.0\nMaster CSV / Route CSV",
            "F": "作業中",
            "G": "・資料とゲーム内の\n　台詞を見比べる\n・各DAYの確認対象を\n　整理\n・更新漏れを確認",
        },
        20: {
            "A": "本番素材/\nフォント確認",
            "B": "・本番背景・デスクトップ素材を追加中\n・ゲーム用音声・SE素材を追加中\n・TMPフォント差分の表示を確認",
            "D": "Assets/Art/Desktop\nAssets/Audio/Kimikoi",
            "F": "作業中",
            "G": "・本番画面で素材確認\n・BGMとSEの音を確認\n・日本語表示の欠け\n　を確認",
        },
        21: {
            "A": "本番確認/\n差分整理",
            "B": "・日次の本番版更新と差分確認を実施\n・未コミットのフォント・素材差分を整理\n・確認順と継続対応を見直し",
            "D": "Git log / working tree\n本番確認対象",
            "F": "作業中",
            "G": "・最新版で起動する\n・素材と文字を確認\n・未反映の差分を確認",
        },
        22: {
            "A": "報告書/\nテンプレート更新",
            "B": "・8/31〜9/4分の報告書を作成\n・確認欄を5行以内の書式へ更新\n・次回も使えるテンプレートを更新",
            "D": "本報告書\nエンジニア作業報告書テンプレート",
            "F": "完了",
            "G": "・操作と結果を記載\n・5行以内で読める\n・記載例を更新",
        },
    },
)

set_value(
    ws,
    "A25",
    "・対象：8/31〜9/4。出社なし。\n"
    "・シナリオ、検索、LIME、妹演出、本番素材を確認中。\n"
    "・各ルートの通し確認と、本番素材・文字表示を継続確認。",
)
apply_page_settings(ws)
wb.save(out)

rows = [*range(8, 16), *range(19, 23)]
validate(template)
validate(out, rows)
print(template)
print(out)
