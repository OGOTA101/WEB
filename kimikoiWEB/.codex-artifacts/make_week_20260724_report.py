from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


template = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260724")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年7月第4週.xlsx"

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年7月第4週"


def set_value(cell_ref: str, value: str) -> None:
    cell = ws[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value


set_value("A2", "対象期間: 2026/07/20〜2026/07/24")
set_value("A4", "2026/07/20〜2026/07/24")
set_value("B4", "7/20, 7/21, 7/22, 7/23, 7/24")
set_value("C4", "キミと恋する15日（仮）")
set_value("D4", "Unity本番開発")
set_value(
    "E4",
    "本番開発として、ルート分岐と妹データ削除時のパラメータ反映、LIMEスタンプ・通知バッジ、怨念グリッチ演出、検証・開発補助環境を更新。前週までに完了として報告済みの項目は再掲せず、今週の新規・継続作業を整理。",
)

implementation_rows = {
    8: {
        "A": "ルート分岐/削除時パラメータ",
        "B": "・怨念、イライラのDAY7・DAY9・DAY10分岐をシナリオ表に合わせて調整\n・妹データを削除した際、怨念度が正しく反映される処理を追加\n・選択後の遷移先とパラメータ変化が連動する状態まで実装",
        "D": "master_route_*.csv / MasterScenarioController.cs / KimikoiAdvWindowController.cs",
        "F": "完了",
        "G": "選択や妹データ削除の結果に応じて、想定したルートとパラメータでゲームが進む状態を整備済み",
    },
    9: {
        "A": "LIMEスタンプ/通知バッジ",
        "B": "・DAY1のLIME会話へスタンプ表示を接続\n・確認用のスタンプ仮素材と差し替え先を追加\n・デスクトップとタスクバーに未確認通知バッジを実装",
        "D": "LimeStickerCatalog.asset / LimeWindow.prefab / DesktopAppBadgeService.cs",
        "F": "完了",
        "G": "LIMEの新着状況とスタンプ付きの会話内容を、プレイヤーが画面上で把握できる状態を整備済み",
    },
    10: {
        "A": "怨念グリッチ演出調整",
        "B": "・最小化後の復帰時に怨念グリッチが継続する挙動を修正\n・画面中央を読みやすく保ち、周辺ほど強く崩れる表現へ調整\n・中央減衰の見え方を複数パターンで確認中",
        "D": "SisterPresenceGlitchController.cs / ProceduralDatamoshGlitch*.cs / shader",
        "F": "作業中",
        "G": "重要な画面情報を読める状態に保ちながら、不穏さが伝わる演出になるよう最終調整中",
    },
    11: {
        "A": "バグ演出の検証環境",
        "B": "・中央減衰を切り替えて確認する検証メニューを追加\n・デバッグメニューを閉じた際に、表示上書きが残らないよう修正\n・比較用スクリーンショットを出力し、演出差分を確認",
        "D": "BugEffectCenterAttenVerify.cs / DesktopDebugMenuController.BugEffectDebug.cs",
        "F": "完了",
        "G": "演出の強さや見え方を、企画確認時に同じ条件で比較できる確認環境を用意済み",
    },
    12: {
        "A": "妹常駐ウィンドウ調整",
        "B": "・妹常駐ウィンドウの初期表示位置を統一\n・Main/MOCシーンとPrefabの設定を揃えて反映\n・起動時の画面レイアウト差が出ないよう調整",
        "D": "SisterResidentWindow.prefab / Main.unity / MOC.unity",
        "F": "完了",
        "G": "確認環境によって妹ウィンドウの位置が変わらず、画面構成を同じ前提で確認できる状態",
    },
    13: {
        "A": "フォントデータ再生成",
        "B": "・TMPフォントの動的生成データを整理するEditorツールを追加\n・日本語フォントの再生成処理を実行\n・文字表示の確認を継続し、差分を調整中",
        "D": "TmpFontClearDynamicData.cs / TextMeshPro SDF assets",
        "F": "作業中",
        "G": "確認用ビルドでも文字が欠けずに表示されるよう、フォントデータの安定化を継続中",
    },
    14: {
        "A": "Unity開発補助環境",
        "B": "・Unityプロジェクト操作用のMCPパッケージを追加\n・Package ManifestとLock情報を更新\n・Unity内の確認・編集作業を支援する環境を準備",
        "D": "Packages/com.coplaydev.unity-mcp / manifest.json / packages-lock.json",
        "F": "完了",
        "G": "本番開発中のUnity画面・設定確認を効率よく行うための補助環境を追加済み",
    },
    15: {
        "A": "バージョン更新/確認準備",
        "B": "・本番確認用のアプリバージョンを0002へ更新\n・演出、LIME、分岐更新を含む確認対象を整理\n・次回確認で見る画面とシナリオ進行を準備",
        "D": "ProjectSettings.asset / Main.unity / シナリオCSV",
        "F": "完了",
        "G": "今週更新した内容をまとめて確認できる版と、確認時に見る範囲を整理済み",
    },
}

for row, values in implementation_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

other_rows = {
    19: {
        "A": "シナリオ素材リスト更新",
        "B": "・シナリオ素材リストのExcelを更新\n・OS、UI、SEで必要な素材をMDへ整理\n・実装中に必要な素材と未準備の範囲を確認可能に整理",
        "D": "キミコイシナリオVer2.0_素材リスト.xlsx / OS_UI_SE必要素材リスト.md",
        "F": "完了",
        "G": "素材準備の担当・優先度を確認しやすくし、実装待ちになる箇所を把握できる状態を整備済み",
    },
    20: {
        "A": "実装状況/確認差分整理",
        "B": "・前週の継続作業と今週追加した作業を分けて確認\n・Git履歴から分岐、LIME、演出、検証環境の更新を整理\n・未コミットの演出・フォント調整を継続確認対象として記録",
        "D": "Git履歴 / Git作業ツリー / 演出確認スクリーンショット",
        "F": "完了",
        "G": "完了済みの過去作業を混在させず、今回確認する内容と継続対応を分けて共有できる状態",
    },
    21: {
        "A": "演出/フォント確認",
        "B": "・グリッチ中央減衰の比較画像を作成\n・画面中央の読みやすさと周辺演出の強さを確認\n・フォント再生成後の文字表示を確認し、必要な調整を継続",
        "D": "Assets/Screenshots / バグエフェクト設定 / TMPフォント",
        "F": "作業中",
        "G": "演出と文字の見やすさを両立できる状態かを、実画面で最終確認している段階",
    },
    22: {
        "A": "報告書整備",
        "B": "・7/20〜7/24分の作業報告をテンプレートに沿って作成\n・Git履歴と未コミット差分を反映\n・作業中項目をプランナーが確認しやすい表現に整理",
        "D": "本報告書",
        "F": "完了",
        "G": "今週の実装結果、確認対象、継続作業がひと目で分かるように整理済み",
    },
}

for row, values in other_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

set_value(
    "A25",
    "・対象期間は2026/07/20〜2026/07/24。前週までに完了として報告済みの項目は再掲しない。\n"
    "・主な完了項目は、ルート分岐と削除時パラメータ反映、LIMEスタンプ・通知バッジ、検証環境、妹常駐ウィンドウ位置、開発補助環境。\n"
    "・グリッチ中央減衰とフォントデータの表示確認は作業中。次回は実画面での演出・文字表示の最終確認を継続する。",
)

if ws.row_dimensions[14].height:
    ws.row_dimensions[15].height = ws.row_dimensions[14].height

ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1

wb.save(out)
print(out)
