from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


template = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260807")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年8月第1週.xlsx"

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年8月第1週"


def set_value(cell_ref: str, value: str) -> None:
    cell = ws[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value


set_value("A2", "対象期間: 2026/08/03〜2026/08/07")
set_value("A4", "2026/08/03〜2026/08/07")
set_value("B4", "8/3（出社）, 8/4, 8/5, 8/6, 8/7（出社）")
set_value("C4", "キミと恋する15日（仮）")
set_value("D4", "Unity本番開発")
set_value(
    "E4",
    "本番開発として、死亡後の復帰、イライラLIMEイベント、妹・REAL ADVの表示優先、ウィンドウ操作、音響と検証機能を実装・調整中。前週までに完了として報告済みの項目は再掲せず、今週のGit状況を整理。",
)

implementation_rows = {
    8: {
        "A": "死亡後/ROUTE_ENDの進行",
        "B": "・最終削除後やラスト演出で全画面動画を再生する処理を追加\n・ROUTE_ENDとゲームオーバー後、共通DAY1へ戻る流れを修正\n・2回目以降の死亡時も専用動画から共通冒頭へ戻るよう調整",
        "D": "SisterHorrorSequenceDirector.cs / KimikoiAdvWindowController.cs / フルスクリーン動画",
        "F": "作業中",
        "G": "各終了条件から動画・DAY1復帰までが止まらずにつながるか、通しプレイで確認・調整中",
    },
    9: {
        "A": "イライラLIME見られイベント",
        "B": "・LIMEを隠す必要がある場面の見られ判定を追加\n・警戒中の通し会話、イベント分岐、進行状態の保存を接続\n・割り込み時に同じ会話が二重に始まらないよう修正",
        "D": "MasterScenarioController.cs / LimeWindowController.cs / LimeScenarioSessionCoordinator.cs",
        "F": "作業中",
        "G": "LIMEを隠すタイミングに応じたイベントが、会話の重複や進行不整合なく動くかを確認中",
    },
    10: {
        "A": "妹/REAL ADVの表示優先",
        "B": "・常駐妹とスキンシップ妹をバグ演出の影響から分けて表示\n・妹の黒い靄とグリッチパーティクルを追加\n・REAL ADVを妹表示より前面へ出す順序に調整",
        "D": "SisterBugEffectExemptPresenter.cs / OuterAdvOverlayController.cs / SisterResidentWindow.prefab",
        "F": "作業中",
        "G": "不穏な画面演出中でも、妹と重要なシナリオ表示が意図した重なり順になるよう調整中",
    },
    11: {
        "A": "ウィンドウ操作/進行復帰",
        "B": "・ドラッグ時の座標計算を見直し、移動時のずれを修正\n・本文表示を画面外へ出せるウィンドウ設定と範囲制御を追加\n・ログの右クリック閉じ、LIMEのSpace進行、TextViewport未配線を修正",
        "D": "DesktopWindowView.cs / WindowDefinition.cs / KimikoiAdvLogController.cs",
        "F": "作業中",
        "G": "会話中のウィンドウ操作や表示切替で、入力不良や内容の見切れが出ないかを確認中",
    },
    12: {
        "A": "Googru/動画後のシナリオ復帰",
        "B": "・OPEN_GOOGRU後に検索キーワード待ちで止まる挙動を修正\n・動画演出終了後にデスクトップへ戻る流れを修正\n・BLACKOUTのクリック阻害とルート切替後の進行停止を解消",
        "D": "KimikoiAdvTimelineBridgeController.cs / KimikoiAdvWindowController.cs",
        "F": "作業中",
        "G": "検索画面や動画演出を挟んでも、プレイヤーが次のシナリオへ戻れるかを通し確認中",
    },
    13: {
        "A": "音響演出/音量操作",
        "B": "・BGMのハイカット、ローカット効果と試聴用ボタンを追加\n・妹パラメータ上昇・下降時のSEを分けて設定\n・タスクバーから音量を調整するフライアウトを追加",
        "D": "SoundManager.BgmEffects.cs / SisterParameterChangeSePlayer.cs / TaskbarController.cs",
        "F": "作業中",
        "G": "場面の不穏さやパラメータ変化を音でも把握できるよう、音量と再生タイミングを確認中",
    },
    14: {
        "A": "ルート横断デバッグ開始",
        "B": "・デバッグ開始日をルート横断で選べるように変更\n・オールリセットとルート設定を確認しやすい表記へ整理\n・イライラ助け成功回数を直接設定できる検証操作を追加",
        "D": "DesktopDebugMenuController.cs / KimikoiAdvStartDayChoice.cs / SaveManager.cs",
        "F": "作業中",
        "G": "確認したいルートや日付から安全に試せるか、リセットを含む検証手順を確認中",
    },
    15: {
        "A": "TMPフォント再生成/表示確認",
        "B": "・日本語フォントの動的生成データを再生成\n・複数フォントのアセット差分を更新\n・本番確認画面での文字欠けや表示崩れを継続確認",
        "D": "TextMeshPro SDF assets / フォント表示確認",
        "F": "作業中",
        "G": "画面やシナリオ内の日本語が安定して読める状態にするため、最終表示確認を継続中",
    },
}

for row, values in implementation_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

other_rows = {
    19: {
        "A": "シナリオ/マスターデータ更新",
        "B": "・怨念DAY10の告白後に常駐妹を開く導線を追加\n・イライラDAY10のラスト動画と遷移先を接続\n・イライラEVENT20の通し会話向けCSVとトークンを更新",
        "D": "master_route_onnen_day10.csv / master_route_iraira_day10.csv / master_ver2_*.csv",
        "F": "作業中",
        "G": "実装した演出と各ルートのシナリオ進行が、企画した順序でつながるかを確認中",
    },
    20: {
        "A": "本番確認環境整備",
        "B": "・本番確認用のアプリバージョンを0006、0007へ更新\n・Mainシーンへ表示優先とタスクバー音量操作を配線\n・BGM試聴と各ルート確認に必要な設定を整理",
        "D": "ProjectSettings.asset / Main.unity / BgmEffectSettings.asset",
        "F": "作業中",
        "G": "今週更新した画面・音・シナリオを、同じ本番確認環境で確認できるかを調整中",
    },
    21: {
        "A": "実行ログの共有運用",
        "B": "・デバッグメニューにEXE実行ログの表示と送信操作を追加\n・実行ログをDriveへTXT保存する連携を追加\n・保存完了時にSlackへ通知する処理を追加",
        "D": "RuntimeConsoleLogBuffer.cs / KimikoiDebugUploader.gs / Slack通知",
        "F": "作業中",
        "G": "確認中に発生した情報を関係者へ共有し、再現や対応判断へつなげられるかを運用確認中",
    },
    22: {
        "A": "報告書整備",
        "B": "・8/3〜8/7分の作業報告をテンプレートに沿って作成\n・Git履歴と未コミットのフォント差分を反映\n・出社日と作業中項目を確認しやすい表現に整理",
        "D": "本報告書",
        "F": "完了",
        "G": "今週の実装成果、継続作業、確認時の着眼点を一枚で把握できる状態",
    },
}

for row, values in other_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

set_value(
    "A25",
    "・対象期間は2026/08/03〜2026/08/07。8/3および8/7は出社対応日。\n"
    "・前週までに完了として報告済みの項目は再掲しない。今週の開発項目は実装・調整を反映済みだが、通し確認と演出・操作の調整を継続するため作業中として整理。\n"
    "・TMPフォントの日本語表示、各ルート進行、LIMEイベント、音響、実行ログ共有を含め、実機での通し確認を継続する。",
)

if ws.row_dimensions[14].height:
    ws.row_dimensions[15].height = ws.row_dimensions[14].height

ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1

wb.save(out)
print(out)
