from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from PIL import ImageFont


template = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260911")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年9月第2週.xlsx"


def set_value(sheet, cell_ref: str, value: str) -> None:
    cell = sheet[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value


def write_rows(sheet, rows: dict[int, dict[str, str]]) -> None:
    for row, values in rows.items():
        for column, value in values.items():
            set_value(sheet, f"{column}{row}", value)


def apply_page_settings(sheet) -> None:
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1


def validate(path: Path, report_rows: list[int] | None = None) -> None:
    workbook = load_workbook(path, data_only=False)
    sheet = workbook.active
    if sheet.print_area != f"'{sheet.title}'!$A$1:$H$28":
        raise ValueError(f"Unexpected print area: {sheet.print_area}")

    font = ImageFont.truetype(r"C:/Windows/Fonts/YuGothR.ttc", 12, index=0)
    for cell_ref in ["G8", "G19"]:
        lines = sheet[cell_ref].value.splitlines()
        if len(lines) > 5 or any(font.getlength(line) > 128 for line in lines):
            raise ValueError(f"{cell_ref} exceeds the G-column rules: {lines}")

    if report_rows:
        for row in report_rows:
            work_lines = sheet[f"B{row}"].value.splitlines()
            if not 3 <= len(work_lines) <= 4:
                raise ValueError(f"B{row} must contain 3 to 4 lines: {work_lines}")
            if sum(line.startswith("・") for line in work_lines) != 3:
                raise ValueError(f"B{row} must contain three bullet lines: {work_lines}")
            if any(font.getlength(line) > 244 for line in work_lines):
                raise ValueError(f"B{row} exceeds the usable width: {work_lines}")

            confirmation_lines = sheet[f"G{row}"].value.splitlines()
            if not 3 <= len(confirmation_lines) <= 5:
                raise ValueError(f"G{row} must contain 3 to 5 lines: {confirmation_lines}")
            if sum(line.startswith("・") for line in confirmation_lines) != 3:
                raise ValueError(f"G{row} must contain three bullet lines: {confirmation_lines}")
            if any(font.getlength(line) > 128 for line in confirmation_lines):
                raise ValueError(f"G{row} exceeds the usable width: {confirmation_lines}")

            result_lines = sheet[f"D{row}"].value.splitlines()
            if len(result_lines) > 2 or any(font.getlength(line) > 228 for line in result_lines):
                raise ValueError(f"D{row} exceeds the usable width: {result_lines}")

    errors = [
        cell.coordinate
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("#")
    ]
    if errors:
        raise ValueError(f"Formula errors found: {errors}")


# Keep the template instructions aligned with the column widths used in the report.
template_wb = load_workbook(template)
template_ws = template_wb.active
set_value(template_ws, "B4", "作業日（出社日を明記）")
set_value(
    template_ws,
    "G8",
    "・確認欄は合計5行以内\n・操作と結果を書く\n・改行は意味の切れ目\n・一文字だけ残さない",
)
set_value(
    template_ws,
    "G19",
    "・操作と結果を書く\n・確認対象を具体化\n・全体で5行以内",
)
set_value(
    template_ws,
    "A25",
    "・対象日と出社日を記載。\n"
    "・前週に完了として報告した項目は再掲しない。\n"
    "・作業内容は3〜4行、確認欄は5行以内。",
)
apply_page_settings(template_ws)
template_wb.save(template)

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年9月第2週"
set_value(ws, "A2", "対象期間: 2026/09/05〜2026/09/11（9/9出社）")
set_value(ws, "A4", "2026/09/05〜2026/09/11")
set_value(ws, "B4", "9/5〜9/11\n9/9出社")
set_value(ws, "C4", "キミと恋する15日（仮）")
set_value(ws, "D4", "Unity本番開発")
set_value(
    ws,
    "E4",
    "本番開発として、常駐妹、設定画面、Googru、LIME、怨念演出、状態リセットを調整中。前週完了項目は再掲しない。",
)

write_rows(
    ws,
    {
        8: {
            "A": "常駐妹/\nルート遷移",
            "B": "・常駐妹の解禁時点と吹き出し進行を調整\n・ルート遷移中の表示と立ち絵切替を調整\n・怨念DAY6の消去前演出へ妹を接続",
            "D": "KimikoiAdvTimelineBridge\nmaster_route_onnen_day06",
            "F": "作業中",
            "G": "・DAY2解禁後の表示\n・公式サイト後の台詞\n　を確認\n・DAY6消去前を確認",
        },
        9: {
            "A": "ゲーム内設定/\nOS設定",
            "B": "・スタートメニューから設定画面を追加\n・ゲーム設定とOS設定の表示値を同期\n・Windows風画面と既定サイズを調整",
            "D": "KimikoiOsConfigWindow\nSettingsWindow prefab",
            "F": "作業中",
            "G": "・設定画面を開く\n・設定値がそろう\n・閉じても値が残る",
        },
        10: {
            "A": "Googru\nブラウザ操作",
            "B": "・戻る／進むとアドレスバーを常設\n・表示URLと常設ブラウザバーを追加\n・検索ページの再表示操作を調整",
            "D": "GoogruWindowController\nGoogruWindow prefab",
            "F": "作業中",
            "G": "・検索画面を開く\n・戻る／進むを試す\n・URL表示を確認",
        },
        11: {
            "A": "LIME入力/\nスタンプ演出",
            "B": "・LIMEスタンプを本画像へ接続\n・操作前は文字送りとSEを開始しない\n・入力SEの分割再生と驚き演出を追加",
            "D": "LimeWindowController\nLimeStickerCatalog",
            "F": "作業中",
            "G": "・LIMEを開いて待つ\n・入力後に文字と音\n　が出る\n・スタンプと驚き演出\n　を確認",
        },
        12: {
            "A": "怨念/\nエラー演出",
            "B": "・怨念の秒ごとの増加量を調整可能に\n・DAY6乱立中の画面揺れとエラー音を追加\n・増加停止デバッグと前日状態復元を調整",
            "D": "KimikoiOnnenIntervalTick\nDesktopClutterWindow",
            "F": "作業中",
            "G": "・DAY6乱立を進める\n・画面揺れとエラー音\n　を確認\n・増加の停止と再開",
        },
        13: {
            "A": "カーソル/\nクリック操作",
            "B": "・ソフトウェアカーソルを32pxで統一\n・起動後と応答なし解除後の戻りを調整\n・常駐妹のクリック貫通を修正",
            "D": "DesktopCursorStateController\nSisterResidentWindow",
            "F": "作業中",
            "G": "・起動後のカーソル\n・応答なし後に戻る\n・妹の背後をクリック",
        },
        14: {
            "A": "セーブ/\n日次状態リセット",
            "B": "・全データリセット時のシェル初期化を修正\n・外側DAY開始時の前日怨念値を復元\n・リセット後の窓と設定表示を同期",
            "D": "DesktopShellController\nKimikoiAdvPlaybackDirector",
            "F": "作業中",
            "G": "・全データをリセット\n・起動直後の画面\n　に戻る\n・前日の増加値を確認",
        },
        15: {
            "A": "フォント/\n本番版設定",
            "B": "・LIMEとキミコイ窓の\n　タイトルをYu Gothicへ\n・TMPフォント差分の表示を継続確認\n・本番用版番号を0020〜0022へ更新",
            "D": "TextMeshPro SDF assets\nProjectSettings / Prefabs",
            "F": "作業中",
            "G": "・LIMEとADVの見出し\n　を確認\n・日本語文字を表示\n・最新版を起動",
        },
        19: {
            "A": "仕様/実装\n説明整備",
            "B": "・DesktopとADVの処理コメントを更新\n・妹演出とSave/Soundの説明を更新\n・調査時の参照箇所を整理",
            "D": "Desktop / KimikoiAdv\nSave / Sound / Editor",
            "F": "完了",
            "G": "・仕様の参照先\n　を確認\n・調査対象を探す\n・動きと説明を照合",
        },
        20: {
            "A": "本番素材/\n音・スタンプ",
            "B": "・LIMEスタンプを本画像へ接続\n・OSエラー音と効果音を演出へ配線\n・画面タイトルのフォントを変更",
            "D": "LimeStickerCatalog\nSoundCue / Prefabs",
            "F": "作業中",
            "G": "・スタンプを送る\n・エラー音を確認\n・画面タイトルを確認",
        },
        21: {
            "A": "本番確認/\n作業ツリー整理",
            "B": "・本番用版番号を0022まで更新\n・未コミットのTMPフォント差分を確認\n・本番画面での確認対象を整理",
            "D": "ProjectSettings\nTextMeshPro SDF assets",
            "F": "作業中",
            "G": "・最新版で起動する\n・各画面を開く\n・未反映差分を確認",
        },
        22: {
            "A": "報告書/\nテンプレート更新",
            "B": "・9/5〜9/11分の報告書を作成\n・9/9出社と完了／作業中を反映\n・確認欄の5行ルールをテンプレートへ反映",
            "D": "本報告書\n作業報告書テンプレート",
            "F": "完了",
            "G": "・対象日と出社日\n・完了と作業中を区別\n・確認欄を5行以内",
        },
    },
)

set_value(
    ws,
    "A25",
    "・対象：9/5〜9/11。9/9出社。\n"
    "・常駐妹、設定、Googru、LIME、怨念演出を確認中。\n"
    "・ルート通し、本番素材、文字表示、リセット動作を継続確認。",
)
apply_page_settings(ws)
wb.save(out)

rows = [*range(8, 16), *range(19, 23)]
validate(template)
validate(out, rows)
print(template)
print(out)
