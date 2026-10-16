from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


template = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260605")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年6月第1週.xlsx"

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年6月第1週"


def set_value(cell_ref: str, value):
    cell = ws[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value


set_value("A2", "対象期間: 2026/06/01〜2026/06/05")
set_value("A4", "2026/06/01〜2026/06/05")
set_value("B4", "6/1, 6/2, 6/3, 6/4, 6/5")
set_value("C4", "キミと恋する15日（仮）")
set_value("D4", "Unity本番開発")
set_value(
    "E4",
    "本番開発へ移行し、KIMIKOIリポジトリへのリネーム、C#基盤規約対応、asmdef分割、GitHub向けDocs/README整備、工数・MOC・マイルストーン資料更新まで完了扱いで整理。",
)

implementation_rows = {
    8: {
        "A": "本番開発基盤",
        "B": "・Horage系のUnityプロジェクトをKIMIKOI本番開発構成へ移行\n・本番用のフォルダ構成を整理\n・ProjectSettings、Packages、Plugins、Unity Assetsを整理",
        "D": "D:/KIMIKOI / KIMIKOI/Assets / ProjectSettings / Packages",
        "F": "完了",
        "G": "本番開発用リポジトリとして整理完了",
    },
    9: {
        "A": "リポジトリ移行",
        "B": "・Git履歴上でKIMIKOIへのRename対応を実施\n・旧Horage配下のUnity Assetsを移動\n・TextMesh Pro、URP、DOTween、各種設定をKIMIKOI配下へ整理",
        "D": "commit 2eb36f6 / Horage → KIMIKOI rename",
        "F": "完了",
        "G": "本番名義の構成へ移行済み",
    },
    10: {
        "A": "C#実装規約対応",
        "B": "・var廃止、SerializeFieldルール、Find系API廃止を実施\n・AddComponent/GetComponent、Update処理、Coroutineを整理\n・イベントリスナー、nullable、定数化、不要コード整理を段階的に反映",
        "D": "Assets/Scripts / .cursor/rules / 実装運用規約_CSharp.md",
        "F": "完了",
        "G": "本番開発向けの実装ルールへ反映済み",
    },
    11: {
        "A": "Assembly/構成整理",
        "B": "・asmdefを分割\n・Desktop、KimikoiAdv、Managers、Rendering、Editorの責務を整理\n・KimikoiAdv関連スクリプトを専用フォルダへ移動し、参照境界を明確化",
        "D": "Horage.Desktop.asmdef / Horage.Managers.asmdef / Horage.Rendering.asmdef / Horage.Editor.asmdef",
        "F": "完了",
        "G": "コンパイル単位と責務分離を整理",
    },
    12: {
        "A": "ゲーム実装コード整理",
        "B": "・ADV、疑似デスクトップ、演出系のスクリプト配置を見直し\n・マスタCSV、シナリオ進行系の配置を整理\n・既存ゲーム機能を保ったまま、本番開発で追いやすい構成へ調整",
        "D": "KimikoiAdvWindowController.cs / MasterScenarioController.cs / Desktop系Scripts",
        "F": "完了",
        "G": "主要ゲーム実装の配置と保守性を改善",
    },
    13: {
        "A": "データ/仕様参照整理",
        "B": "・Data、Scenario、Specifications、ScreenshotsのREADMEを更新\n・各フォルダの運用資料を整理\n・仕様書Markdown、実装照合チェックシート、パラメーター調整概要をDocs配下へ移動",
        "D": "Docs/Specifications / Assets/Data/README.md / Assets/Screenshots/README.md",
        "F": "完了",
        "G": "仕様・データ・検証資料の参照導線を整理済み",
    },
    14: {
        "A": "GitHub運用整備",
        "B": "・GitHub TOP用READMEを更新\n・Docsハブとブランチ運用ルールを作成/更新\n・master、develop、release、HOTFIX、Git LFS方針をMarkdownで明文化",
        "D": "README.md / Docs/README.md / Docs/BRANCHING.md",
        "F": "完了",
        "G": "GitHub TOPとドキュメント導線を更新済み",
    },
    15: {
        "A": "QA/確認",
        "B": "・本番開発基盤化に伴う参照切れリスクを整理\n・Missing参照、コンパイル単位、Docsリンクの確認観点を整理\n・Unity完了判定ルールとスクリーンショット運用をREADME群へ反映",
        "D": "Assets/Docs/Planning / Assets/Screenshots/README.md / Docs",
        "F": "完了",
        "G": "本番開発移行後の確認観点を整理済み",
    },
}

for row, values in implementation_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

other_rows = {
    19: {
        "A": "仕様整理",
        "B": "・GitHub TOP直下のDocsハブを整備\n・仕様書Markdown一覧とUnity内Docsへの導線を整理\n・シナリオ、仕様書、実装照合、パラメーター調整概要を本番開発側の資料構成へ移行",
        "D": "Docs/README.md / Docs/Specifications / Docs/Scenario / KIMIKOI/Assets/Docs",
        "F": "完了",
        "G": "資料体系を本番開発向けに整理",
    },
    20: {
        "A": "工数・MOC対応",
        "B": "・工数表を更新\n・MOCを受けた仕様変更に合わせて工数を再計算\n・本番制作向けのα版工数表、MOC検証後の見積/ダウンチャート内容を整理",
        "D": "きみこい・工数表_ダウンチャート_α版.xlsx / MOC検証後資料",
        "F": "完了",
        "G": "MOC反映後の工数再計算を実施",
    },
    21: {
        "A": "マイルストーン作成",
        "B": "・6月開始の月次マイルストーン報告書を作成\n・企画向け版のマイルストーン資料を作成\n・本番開発の進行目安と報告単位を整理",
        "D": "キミコイ_月次マイルストーン報告書_2026年6月開始.xlsx / 企画向け.xlsx",
        "F": "完了",
        "G": "月次マイルストーン資料を作成済み",
    },
    22: {
        "A": "報告書整備",
        "B": "・6/1〜6/5分の作業報告をテンプレートに沿って作成\n・本番開発移行、Git状況、TOP README更新を整理\n・Docs整備、工数/MOC、マイルストーン作成を簡潔に整理",
        "D": "本報告書",
        "F": "完了",
        "G": "内容確認済み",
    },
}

for row, values in other_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

set_value(
    "A25",
    "・6/1〜6/5は本番開発移行週として、KIMIKOIリポジトリ整備とC#基盤修正を全体完了扱いで整理。\n"
    "・GitHub TOP README、Docsハブ、ブランチ運用、仕様書Markdown移行を報告対象に含めた。\n"
    "・工数表更新、MOC反映後の工数再計算、6月開始マイルストーン資料作成も実装以外の業務として記録。",
)

if ws.row_dimensions[14].height:
    ws.row_dimensions[15].height = ws.row_dimensions[14].height

wb.save(out)
print(out)
