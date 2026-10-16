from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


template = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260710")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年7月第2週.xlsx"

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年7月第2週"


def set_value(cell_ref: str, value):
    cell = ws[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value


set_value("A2", "対象期間: 2026/07/06〜2026/07/10")
set_value("A4", "2026/07/06〜2026/07/10")
set_value("B4", "7/6, 7/7, 7/8, 7/9, 7/10")
set_value("C4", "キミと恋する15日（仮）")
set_value("D4", "Unity本番開発")
set_value(
    "E4",
    "本番開発として、ルート確定・助け要請・ROUTE_END、ADV/Lineの中断復帰、ゴミ箱/Googru、テスト導線、選択肢Prefab、常駐妹Spine導入準備を進行。前週までに完了済みの項目は再掲せず、今週の継続・新規作業を整理。",
)

implementation_rows = {
    8: {
        "A": "ルート確定/パラメータ",
        "B": "・Day5時点の固定ルートIDから分岐を判断する処理を追加\n・確定後のルートを保存し、再開時にも同じ進行を参照\n・妹パラメータの最大値参照とルート判定の共通化を調整",
        "D": "SaveManager.cs / KimikoiAdvWindowController.cs / KimikoiSisterParamUtility.cs",
        "F": "作業中",
        "G": "プレイヤーの選択結果が以後のルートへ一貫して反映されるよう、分岐条件の安定化を進めている段階",
    },
    9: {
        "A": "助け要請イベント",
        "B": "・助け要請の解禁後にLIMEボタンを表示する処理を追加\n・使用したDayと失敗回数を保存できるように対応\n・Day7/8のイベント発火、失敗分岐、ゲームオーバー接続を実装",
        "D": "LimeWindowController.cs / MasterScenarioEventTriggerCatalog.cs / SaveManager.cs",
        "F": "作業中",
        "G": "助け要請の選択がイベントとゲームオーバーへつながるため、企画意図どおりの条件で発生するか確認を進める状態",
    },
    10: {
        "A": "ROUTE_END/終端演出",
        "B": "・ROUTE_ENDコマンドと演出引数の読込を追加\n・ルート終端処理を外側/内側ADVに接続\n・ルートEND用の全画面暗転を追加",
        "D": "KimikoiAdvPlaybackDirector.cs / KimikoiAdvTimelineBridgeController.cs / BlueScreenOverlayController.cs",
        "F": "作業中",
        "G": "各ルートの終端を演出付きで区切るための実装が入り、通しプレイでの遷移確認を行う段階",
    },
    11: {
        "A": "ADV/Line中断復帰",
        "B": "・キミコイADVが中断された場合も再開できるように調整\n・Lineシナリオセッションの復元処理を追加\n・外側ADV完了命令と復帰時の進行を調整",
        "D": "KimikoiAdvWindowController.cs / LimeScenarioSessionCoordinator.cs / MasterScenarioController.cs",
        "F": "作業中",
        "G": "アプリ切替や中断後も進行が失われにくくなり、確認中に最初からやり直す負担を減らすための対応",
    },
    12: {
        "A": "デスクトップアプリ",
        "B": "・ゴミ箱のスロット表示、ドロップ判定、右クリック復元導線を追加\n・削除時の通知と空状態/ホバー表示を追加\n・GoogruでHTML表示、接続待機、スクロール完了待機を調整",
        "D": "RecycleBinWindowController.cs / RecycleBinSlotView.cs / GoogruWebBrowserBridge.cs",
        "F": "完了",
        "G": "疑似デスクトップ内の操作結果が画面上で分かりやすくなり、アプリ間の確認導線を整備済み",
    },
    13: {
        "A": "テスト/デバッグ導線",
        "B": "・妹パラメータ4種のリセット機能を追加\n・通知テストをアプリテストパネルへ統合\n・キミコイルートテストパネル、仮監視アイコン、Googru参照設定を追加",
        "D": "DesktopDebugMenuController.cs / SisterResidentWatchingMockController.cs / GoogruWindowController.cs",
        "F": "作業中",
        "G": "ルート、通知、妹パラメータを個別に確認しやすい導線を追加し、プランナー確認の準備を進めている段階",
    },
    14: {
        "A": "選択肢Prefab再構成",
        "B": "・選択肢A-Dを明示スロット方式へ移行\n・選択肢CSV解析と表示データの構造を見直し\n・KimikoiAdvChoices Prefabの更新とアップグレード用Editor処理を準備",
        "D": "KimikoiAdvChoices.prefab / AdvChoiceEntry.cs / AdvChoiceSlotUtility.cs / KimikoiAdvChoicesPrefabUpgrader.cs",
        "F": "作業中",
        "G": "選択肢の表示順や対応関係を安定させ、シナリオ表との照合をしやすくするための移行作業中",
    },
    15: {
        "A": "常駐妹Spine導入準備",
        "B": "・常駐妹の仮監視アイコンを実装\n・Spine導入前の技術確認メモと発注書を更新\n・Spine導入後に差し替える表示/Prefab構成を確認",
        "D": "常駐妹_Spine発注前技術確認メモ.xlsx / 常駐妹_Spine発注書.md / SisterResidentWindow.prefab",
        "F": "作業中",
        "G": "常駐妹をSpineアニメーションへ移行する前提を整理し、発注・実装の認識差を減らす準備段階",
    },
}

for row, values in implementation_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

other_rows = {
    19: {
        "A": "書類更新",
        "B": "・常駐妹Spineの発注前技術確認メモを更新\n・常駐妹Spine発注書を追加\n・仕様書READMEに新しい資料への導線を追加",
        "D": "Docs/Specifications/常駐妹_Spine発注前技術確認メモ.xlsx / 常駐妹_Spine発注書.md / README.md",
        "F": "作業中",
        "G": "発注前に必要な表示条件と実装前提を整理し、素材制作側へ確認しやすい資料を準備中",
    },
    20: {
        "A": "Git/Fork資料更新",
        "B": "・Fork/Gitの手順をスクリーンショット付きで更新\n・commit、push、stash、resetの操作説明を追記\n・Docsハブから参照できるよう入口を更新",
        "D": "Docs/GitForkGuide/README.md / Docs/GitForkGuide/images / Docs/README.md",
        "F": "完了",
        "G": "開発参加者がリポジトリ操作を確認しやすくなり、共有作業時の手順迷いを減らす資料を更新済み",
    },
    21: {
        "A": "検証/運用整理",
        "B": "・不具合対応の時期をREADMEで明確化\n・TextMeshPro、Package、ProjectSettingsの更新を反映\n・未コミットの選択肢/Googru/StreamingAssets差分を継続確認",
        "D": "README.md / TextMeshPro Assets / Packages / ProjectSettings / StreamingAssets",
        "F": "作業中",
        "G": "確認対象の版と未反映差分を見分けられるようにし、次回の検証範囲を整理している状態",
    },
    22: {
        "A": "報告書整備",
        "B": "・7/6〜7/10分の作業報告をテンプレートに沿って作成\n・前週までに完了済みの項目は再掲せず、今週の新規/継続作業を記載\n・Git状況と作業中項目をプランナー目線で整理",
        "D": "本報告書",
        "F": "完了",
        "G": "今週の確認対象と継続判断が分かるよう、報告内容を整理済み",
    },
}

for row, values in other_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

set_value(
    "A25",
    "・対象期間は2026/07/06〜2026/07/10。前週までに完了として報告済みの項目は再掲しない。\n"
    "・主な継続作業は、ルート分岐/助け要請/ROUTE_END、ADV/Line中断復帰、テスト導線、選択肢Prefab、常駐妹Spine準備。\n"
    "・次回は、各ルートの通し確認、選択肢Prefab移行、Spine素材導入後の表示確認を継続する。",
)

if ws.row_dimensions[14].height:
    ws.row_dimensions[15].height = ws.row_dimensions[14].height

wb.save(out)
print(out)
