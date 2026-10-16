from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from PIL import ImageFont


template = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260918")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年9月第3週.xlsx"


def set_value(sheet, cell_ref: str, value: str) -> None:
    cell = sheet[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value


def write_rows(sheet, rows: dict[int, dict[str, str]]) -> None:
    for row, values in rows.items():
        for column, value in values.items():
            set_value(sheet, f"{column}{row}", value)


def apply_page_settings(sheet) -> None:
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1


def validate(path: Path, report_rows: list[int] | None = None) -> None:
    workbook = load_workbook(path, data_only=False)
    sheet = workbook.active
    if sheet.print_area != f"'{sheet.title}'!$A$1:$H$28":
        raise ValueError(f"Unexpected print area: {sheet.print_area}")

    font = ImageFont.truetype(r"C:/Windows/Fonts/YuGothR.ttc", 12, index=0)
    for cell_ref in ["G8", "G19"]:
        lines = sheet[cell_ref].value.splitlines()
        if len(lines) > 5 or any(font.getlength(line) > 128 for line in lines):
            raise ValueError(f"{cell_ref} exceeds the G-column rules: {lines}")

    if report_rows:
        for row in report_rows:
            work_lines = sheet[f"B{row}"].value.splitlines()
            if not 3 <= len(work_lines) <= 4:
                raise ValueError(f"B{row} must contain 3 to 4 lines: {work_lines}")
            if sum(line.startswith("・") for line in work_lines) != 3:
                raise ValueError(f"B{row} must contain three bullet lines: {work_lines}")
            if any(font.getlength(line) > 244 for line in work_lines):
                raise ValueError(f"B{row} exceeds the usable width: {work_lines}")

            confirmation_lines = sheet[f"G{row}"].value.splitlines()
            if not 3 <= len(confirmation_lines) <= 5:
                raise ValueError(f"G{row} must contain 3 to 5 lines: {confirmation_lines}")
            if sum(line.startswith("・") for line in confirmation_lines) != 3:
                raise ValueError(f"G{row} must contain three bullet lines: {confirmation_lines}")
            if any(font.getlength(line) > 128 for line in confirmation_lines):
                raise ValueError(f"G{row} exceeds the usable width: {confirmation_lines}")

            result_lines = sheet[f"D{row}"].value.splitlines()
            if len(result_lines) > 2 or any(font.getlength(line) > 228 for line in result_lines):
                raise ValueError(f"D{row} exceeds the usable width: {result_lines}")

    errors = [
        cell.coordinate
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("#")
    ]
    if errors:
        raise ValueError(f"Formula errors found: {errors}")


template_wb = load_workbook(template)
template_ws = template_wb.active
set_value(template_ws, "B4", "作業日（出社日／出社なしを明記）")
set_value(
    template_ws,
    "A25",
    "・対象日と出社日／出社なしを記載。\n"
    "・前週に完了として報告した項目は再掲しない。\n"
    "・作業内容は3〜4行、確認欄は5行以内。",
)
apply_page_settings(template_ws)
template_wb.save(template)

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年9月第3週"
set_value(ws, "A2", "対象期間: 2026/09/12〜2026/09/18（出社なし）")
set_value(ws, "A4", "2026/09/12〜2026/09/18")
set_value(ws, "B4", "9/12〜9/18\n出社なし")
set_value(ws, "C4", "キミと恋する15日（仮）")
set_value(ws, "D4", "Unity本番開発")
set_value(
    ws,
    "E4",
    "本番開発として、DAY戻り、削除演出、既読SKIP、ADV表示、LIME、ルート演出を調整中。前週完了項目は再掲しない。",
)

write_rows(
    ws,
    {
        8: {
            "A": "DAY戻り/\nチェックポイント",
            "B": "・日次チェックポイントの保存と読込を追加\n・死亡後にDAYを選べる\n　全画面メニューを追加\n・デバッグからDAY戻りの導線を追加",
            "D": "KimikoiDayCheckpoint\nKimikoiDayReturnPanel",
            "F": "作業中",
            "G": "・途中のDAYで再開\n・死亡後のDAYを選ぶ\n・選択後の進行を確認",
        },
        9: {
            "A": "妹データ削除/\n怨念エンド",
            "B": "・初回の妹データ削除を操作待ちへ変更\n・乱立した窓を閉じて対象へ誘導\n・全削除から怨念ENDへ接続",
            "D": "DesktopClutterWindow\nKimikoiAdvTimelineBridge",
            "F": "作業中",
            "G": "・大きい窓を閉じる\n・妹データを選ぶ\n・全削除後の遷移\n　を確認",
        },
        10: {
            "A": "既読SKIP/\nLIME文字送り",
            "B": "・既読SKIPをADV、LIME、\n　REAL ADVへ拡張\n・LIME文字送りを一文字ずつのSEへ変更\n・送信プレビューと選択肢の音を調整",
            "D": "KimikoiReadSkipService\nLimeWindowController",
            "F": "作業中",
            "G": "・既読部分をSKIP\n　する\n・LIMEを送信する\n・文字と音の同期\n　を確認",
        },
        11: {
            "A": "ADV表示/\nHUD・背景",
            "B": "・DAY、曜日、場所を文字表示へ変更\n・未登録背景は黒い代替表示へ統一\n・DAY6〜7の背景とEventCG表示を調整",
            "D": "KimikoiAdvVisualController\nKimikoiWindow prefab",
            "F": "作業中",
            "G": "・DAY表記を確認\n・曜日と場所を確認\n・背景とCGを確認",
        },
        12: {
            "A": "ルート演出/\nBGM・動画",
            "B": "・十日参り後に神社BGMへ戻るよう修正\n・DAY10動画の開始位置をシナリオへ反映\n・初回死亡復活用の動画素材を差替え",
            "D": "KimikoiAdvAudioController\nmovie_firstdeath_rebirth",
            "F": "作業中",
            "G": "・十日参りを進める\n・BGMの戻りを確認\n・死亡後の動画を確認",
        },
        13: {
            "A": "ノイズ・怪異/\n画面演出",
            "B": "・ノイズ見た目を台詞後も固定するよう調整\n・目ウィンドウの出現SEを短音へ変更\n・DAY6のLIME隠しを自動操作で演出",
            "D": "SisterPresenceGlitch\nScriptedDesktopCursor",
            "F": "作業中",
            "G": "・ノイズ演出を確認\n・目の出現音を確認\n・LIMEを隠す動き\n　を確認",
        },
        14: {
            "A": "操作チュートリアル/\n画面操作",
            "B": "・対象外の操作を止めるフォーカスを追加\n・対象の窓を閉じる操作へ誘導を追加\n・選択UIの壊れを防ぐ修正を追加",
            "D": "DesktopTutorialFocus\nUiSelectableRegistryRepair",
            "F": "作業中",
            "G": "・誘導対象を操作する\n・他の場所は\n　操作できない\n・窓を閉じて進める",
        },
        15: {
            "A": "本番版設定/\n検証環境",
            "B": "・本番版番号を0023〜0026へ更新\n・MOCシーンと版番号の未反映差分を確認\n・スクリーンショット取得用の導線を追加",
            "D": "ProjectSettings / MOC\nDesktopDebugMenuController",
            "F": "作業中",
            "G": "・最新版を起動する\n・画面を撮影する\n・未反映差分を確認",
        },
        19: {
            "A": "FB修正計画/\n資料更新",
            "B": "・受領したFBの修正計画をドキュメント化\n・シナリオ読込先を最新版シートへ更新\n・信頼ルートCSV出力用の補助ツールを追加",
            "D": "FB修正計画.md\nScenarioSheetSettings",
            "F": "完了",
            "G": "・修正対象を一覧化\n・参照先を最新化\n・CSV対象を確認",
        },
        20: {
            "A": "マスタ/シナリオ\n分岐調整",
            "B": "・DAY6の閉じる捨てを\n　プレイヤー待ちへ変更\n・DayAnyのパラメータ反映処理を追加\n・隠す分岐の話者と怨念開始を調整",
            "D": "master_day_any_ver2.csv\nMasterScenarioController",
            "F": "作業中",
            "G": "・選択待ちになる\n・内心の話者を確認\n・怨念開始位置を確認",
        },
        21: {
            "A": "本番確認/\nフィードバック対応",
            "B": "・画面遷移と削除演出の確認対象を整理\n・DAY戻り、既読SKIPの通し確認を継続\n・MOCシーンの未反映差分を確認中",
            "D": "Main / MOC scene\nFeedback correction plan",
            "F": "作業中",
            "G": "・各DAYへ戻る\n・既読をSKIPする\n・演出と画面を確認",
        },
        22: {
            "A": "報告書/\nテンプレート更新",
            "B": "・9/12〜9/18分の報告書を作成\n・出社なしと完了／作業中を反映\n・確認欄は5行以内で\n　テンプレートへ反映",
            "D": "本報告書\n作業報告書テンプレート",
            "F": "完了",
            "G": "・対象期間を確認\n・出社なしを明記\n・状態を分けて記載",
        },
    },
)

set_value(
    ws,
    "A25",
    "・対象：9/12〜9/18。出社なし。\n"
    "・DAY戻り、削除演出、既読SKIP、ルート演出を確認中。\n"
    "・本番用の表示、音、素材、進行を継続確認。",
)
apply_page_settings(ws)
wb.save(out)

rows = [*range(8, 16), *range(19, 23)]
validate(template)
validate(out, rows)
print(template)
print(out)
