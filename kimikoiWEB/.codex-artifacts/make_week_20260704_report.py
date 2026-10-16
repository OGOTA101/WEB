from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


template = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260704")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_2026年7月第1週.xlsx"

wb = load_workbook(template)
ws = wb.active
ws.title = "2026年7月第1週"


def set_value(cell_ref: str, value):
    cell = ws[cell_ref]
    if isinstance(cell, MergedCell):
        raise ValueError(f"{cell_ref} is a merged child cell")
    cell.value = value


set_value("A2", "対象期間: 2026/06/29〜2026/07/04（7/2木曜分除外）")
set_value("A4", "2026/06/29〜2026/07/04")
set_value("B4", "6/29, 6/30, 7/1, 7/3, 7/4（7/2除外）")
set_value("C4", "キミと恋する15日（仮）")
set_value("D4", "Unity本番開発")
set_value(
    "E4",
    "本番開発として、Ver2シナリオ変換、Day5後ルートCSV、妹リアクションカタログ、Ver2再生/ジャンプ、テストパネル、TextMeshPro復旧、Git/Fork資料整備を実施。木曜日分は除外して整理。",
)

implementation_rows = {
    8: {
        "A": "Ver2シナリオ変換",
        "B": "・Ver2変換後CSVをGeneratedVer2配下へ追加\n・共通DAY、初回DAY、intro系の変換データを登録\n・変換レポートを追加し、変換結果を追えるように整理",
        "D": "GeneratedVer2/*.csv / convert_report.json / convert_kimikoi_ver2_xlsx.py",
        "F": "完了",
        "G": "Ver2シナリオをUnity側で読み込むためのデータ土台を用意済み",
    },
    9: {
        "A": "ルートCSV/ジャンプ",
        "B": "・Day5後ルート用CSVを追加\n・ときめき、イライラ、怨念のDay6〜Day10ルートを登録\n・Ver2シナリオ読込とジャンプ動作の基盤を追加",
        "D": "Routes/master_route_*_day06-10.csv / ScenarioJumpRequest.cs / MasterScenarioEventTriggerCatalog.cs",
        "F": "作業中",
        "G": "ルート分岐のデータは入ったため、今後は通し再生で分岐先が正しいか確認する段階",
    },
    10: {
        "A": "Ver2再生/選択肢",
        "B": "・Ver2再生、選択肢、ループ処理を調整\n・Day8 END時にルート遷移する処理を追加\n・不要になったルートジャンプ処理とタイトル戻り処理を廃止",
        "D": "KimikoiAdvPlaybackDirector.cs / KimikoiAdvChoiceController.cs / KimikoiAdvTimelineBridgeController.cs",
        "F": "作業中",
        "G": "企画上の分岐やEND到達をゲーム内進行へ反映するための実装を進めている状態",
    },
    11: {
        "A": "妹カタログ/素材",
        "B": "・妹常駐、スキンシップ、データ削除状態のカタログを追加\n・妹リアクションJSONとScriptableObjectを登録\n・ルートモック用の妹ボイス素材を追加",
        "D": "SisterResidentReactionCatalog / SisterSkinshipReactionCatalog / Resource/voice/sister",
        "F": "作業中",
        "G": "妹ルート演出を表データと素材で管理し、後続の演出差し替えをしやすくしている段階",
    },
    12: {
        "A": "Desktop/ADV連携",
        "B": "・Desktop関連とキミコイADVの連携を修正\n・Googru、LIME、Save、Window周りのつながりを調整\n・妹カタログと常駐/スキンシップ画面の連携を追加",
        "D": "DesktopWindowManager.cs / LimeWindowController.cs / GoogruWindowController.cs / SaveManager.cs",
        "F": "完了",
        "G": "デスクトップ側アプリとADV本編のつながりを確認しやすい状態へ整理済み",
    },
    13: {
        "A": "テストパネル/デバッグ",
        "B": "・通知テスト類をアプリテストパネルへ統合\n・キミコイルートテストパネルを追加\n・デバッグ文言とジャンプ行数を追加",
        "D": "DesktopDebugMenuController.cs / DesktopEmbeddedCopy.asset / KimikoiAdvMasterTokens.cs",
        "F": "完了",
        "G": "プランナー確認時に、通知やルート進行を個別に試しやすい確認導線を追加済み",
    },
    14: {
        "A": "表示/フォント復旧",
        "B": "・TextMeshProとフォントアセットの破損箇所を退避/修正\n・TextMeshPro更新を実施\n・Mainシーン側の表示崩れリスクを抑えるための調整を実施",
        "D": "TextMeshPro Font Assets / Main.unity / Font Materials",
        "F": "完了",
        "G": "文字表示が崩れるリスクを下げ、確認用画面で読みやすさを保つための対応済み",
    },
    15: {
        "A": "演出/パラメータ補完",
        "B": "・パラメータ調整データを更新\n・共通DAY4/5、初回DAY3の演出補完を反映\n・不具合修正時期の資料表記を調整",
        "D": "master_parameter_changes_ver2.csv / master_ver2_common_day4-5.csv / 読み解きMD",
        "F": "作業中",
        "G": "演出とパラメータの整合を詰めており、シナリオ確認と合わせて継続調整する状態",
    },
}

for row, values in implementation_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

other_rows = {
    19: {
        "A": "書類更新",
        "B": "・Ver2シナリオ読み解きMDを追加\n・Docs READMEへ読み解き資料の導線を追加\n・不具合ランクと不具合修正時期の表記を調整",
        "D": "Docs/Scenario/キミコイシナリオVer2.0_読み解き.md / README.md / Docs/README.md",
        "F": "完了",
        "G": "企画・確認者がVer2シナリオの構造や確認優先度を追いやすいよう資料導線を整理済み",
    },
    20: {
        "A": "Git/Fork資料",
        "B": "・GitHub/Fork解説用の資料を追加\n・画像付きでclone手順を確認できるように整理\n・DocsハブへGitForkGuideの入口を追加",
        "D": "Docs/GitForkGuide/README.md / images / Docs/README.md",
        "F": "完了",
        "G": "メンバーがリポジトリ取得やFork作業を迷わず進めやすい資料を整備済み",
    },
    21: {
        "A": "素材/フォルダ整理",
        "B": "・Moc用素材フォルダへ画像、音声、BGM、SEを整理\n・ルートモック用妹ボイスを追加\n・シナリオ素材リスト更新は未コミット差分として継続確認",
        "D": "Resource/Moc / Resource/voice/sister / キミコイシナリオVer2.0_素材リスト.xlsx",
        "F": "作業中",
        "G": "素材の所在は整理が進んでいるが、素材リスト側は確認しながら更新中",
    },
    22: {
        "A": "報告書整備",
        "B": "・7/4締めの作業報告をテンプレートに沿って作成\n・7/2木曜分の作業は報告対象から除外\n・Git状況、書類更新、作業中項目を整理",
        "D": "本報告書",
        "F": "完了",
        "G": "今週の成果と継続確認点を、確認者が追いやすい形で整理済み",
    },
}

for row, values in other_rows.items():
    for col, value in values.items():
        set_value(f"{col}{row}", value)

set_value(
    "A25",
    "・対象は2026/06/29〜2026/07/04の今週分。ただし2026/07/02(木)の作業分は除外。\n"
    "・主な進捗はVer2シナリオCSV、ルート分岐、妹カタログ/ボイス、テストパネル、Git/Fork資料整備。\n"
    "・ルートCSV/ジャンプ、Ver2再生、妹素材、演出/パラメータ補完、素材リスト更新は作業中として継続確認。",
)

if ws.row_dimensions[14].height:
    ws.row_dimensions[15].height = ws.row_dimensions[14].height

wb.save(out)
print(out)
