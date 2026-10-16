from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


base = Path(
    r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/work_reports_20260508/エンジニア作業報告書_テンプレート_改訂版.xlsx"
)
out_dir = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/engineer_report_template")
out_dir.mkdir(parents=True, exist_ok=True)
out = out_dir / "エンジニア作業報告書_汎用テンプレート.xlsx"

wb = load_workbook(base)
ws = wb.active
ws.title = "作業報告書"


def clear_cell_comments(sheet):
    for row in sheet.iter_rows():
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.comment = None


def set_page(sheet, print_area: str, orientation: str):
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = orientation
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = print_area


def clear_report_values(sheet):
    clear_cell_comments(sheet)
    sheet["A1"].value = "作業報告書"
    sheet["A2"].value = "対象期間:"
    for cell in ["A4", "B4", "C4", "D4", "E4"]:
        sheet[cell].value = None

    for row in range(8, 14):
        for cell in [f"A{row}", f"B{row}", f"D{row}", f"F{row}", f"G{row}", f"H{row}"]:
            sheet[cell].value = None

    for row in range(19, 23):
        for cell in [f"A{row}", f"B{row}", f"D{row}", f"F{row}", f"G{row}", f"H{row}"]:
            sheet[cell].value = None

    sheet["A25"].value = None


def refresh_validation(sheet):
    sheet.data_validations.dataValidation = []
    dv = DataValidation(
        type="list",
        formula1='"未着手,進行中,完了,保留,要確認"',
        allow_blank=True,
        showDropDown=False,
    )
    dv.errorTitle = "入力形式の確認"
    dv.error = "状態はリストから選択してください。"
    dv.promptTitle = "状態"
    dv.prompt = "未着手 / 進行中 / 完了 / 保留 / 要確認 から選択できます。"
    sheet.add_data_validation(dv)
    dv.add("F8:F13")
    dv.add("F19:F22")


def copy_row_style(src_sheet, dst_sheet, src_row: int, dst_row: int, max_col: int):
    dst_sheet.row_dimensions[dst_row].height = src_sheet.row_dimensions[src_row].height
    for col in range(1, max_col + 1):
        src_cell = src_sheet.cell(row=src_row, column=col)
        dst_cell = dst_sheet.cell(row=dst_row, column=col)
        if src_cell.has_style:
            dst_cell._style = copy(src_cell._style)
        dst_cell.number_format = src_cell.number_format
        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.protection = copy(src_cell.protection)


def extend_example_columns(sheet):
    for col in ["I", "J", "K", "L"]:
        sheet.column_dimensions[col].width = 22
    for row in range(1, 29):
        for col in range(9, 13):
            cell = sheet.cell(row=row, column=col)
            cell.font = Font(name="Yu Gothic", size=9, color="404040")
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = Border()


def merge_write(sheet, cell_range: str, value, fill_color: str, font_color: str = "1F1F1F", bold: bool = False):
    sheet.merge_cells(cell_range)
    start = cell_range.split(":")[0]
    cell = sheet[start]
    cell.value = value
    cell.fill = PatternFill("solid", fgColor=fill_color)
    cell.font = Font(name="Yu Gothic", size=9, color=font_color, bold=bold)
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    side = Side(style="thin", color="B7C3CF")
    border = Border(left=side, right=side, top=side, bottom=side)
    for row in sheet[cell_range]:
        for merged_cell in row:
            merged_cell.fill = PatternFill("solid", fgColor=fill_color)
            merged_cell.border = border


def add_example_guides(sheet):
    extend_example_columns(sheet)
    merge_write(
        sheet,
        "I1:L1",
        "記入ガイド",
        "1F4E78",
        "FFFFFF",
        True,
    )
    merge_write(
        sheet,
        "I2:L2",
        "上部項目: 期間、作業日、プロジェクト、対象、要約を記載。",
        "DDEBF7",
        "1F4E78",
        True,
    )
    guides = [
        (
            "I3:L5",
            "要約欄\n今週の成果を1〜2文で記載。",
        ),
        (
            "I7:L10",
            "具体内容\n「・」ごとに改行して記載。",
        ),
        (
            "I11:L13",
            "対象欄\nファイル名、機能名、資料名などを記載。",
        ),
        (
            "I14:L16",
            "状態・確認内容\n状態と確認結果を短く記載。",
        ),
        (
            "I18:L22",
            "実装以外の業務\n資料、工数、連絡、調整などを記載。",
        ),
        (
            "I24:L28",
            "補足 / 次回確認事項\n次週予定や未確定事項を記載。",
        ),
    ]
    for cell_range, value in guides:
        merge_write(sheet, cell_range, value, "FFF2CC")

    comments = {
        "A2": "対象期間を記載します。",
        "A4": "報告対象の期間を記載します。",
        "B4": "実際に作業した日を記載します。",
        "C4": "プロジェクト名を記載します。",
        "D4": "開発区分や対象範囲を記載します。",
        "E4": "今週の成果を短く記載します。",
        "B8": "1行1内容で箇条書きにします。",
        "D8": "代表的な対象を記載します。",
        "F8": "状態はプルダウンから選択します。",
        "B19": "コード以外の作業を記載します。",
        "A25": "次週予定や補足を記載します。",
    }
    for cell, text in comments.items():
        sheet[cell].comment = Comment(text, "エンジニア")


def fill_example(sheet):
    sheet.title = "記入例と説明"
    sheet["A1"].value = "作業報告書"
    sheet["A2"].value = "対象期間: 2026/06/01〜2026/06/05"
    sheet["A4"].value = "2026/06/01〜2026/06/05"
    sheet["B4"].value = "6/1, 6/2, 6/3, 6/4, 6/5"
    sheet["C4"].value = "星降る迷宮（仮）"
    sheet["D4"].value = "Unity本番開発"
    sheet["E4"].value = "本番開発開始に向けて、Unity基盤、ゲームシステム、UI、データ連携、QA確認を完了扱いで整理。資料更新、工数再計算、マイルストーン作成も実施。"

    implementation = {
        8: ("開発基盤", "・Unityプロジェクトを本番開発用構成へ整理\n・Assembly Definitionを分割\n・共通設定とパッケージ参照を確認", "ProjectSettings / Packages / Scripts/*.asmdef", "完了", "コンパイル単位と参照境界を整理済み", ""),
        9: ("プレイヤー制御", "・移動、ジャンプ、ダッシュ処理を実装\n・入力設定をゲームパッド/キーボード両対応へ調整\n・壁接触時の挙動を修正", "PlayerController.cs / InputActions", "完了", "基本操作を確認済み", ""),
        10: ("ステージ実装", "・チュートリアルステージを作成\n・床、壁、扉、チェックポイントを配置\n・ステージ遷移時の初期位置を調整", "Stage01.unity / StageGate.cs", "完了", "開始からクリアまで通し確認済み", ""),
        11: ("敵AI", "・巡回、発見、追跡、帰還の状態遷移を実装\n・視野角と追跡距離を調整\n・デバッグ表示で検知範囲を確認", "EnemyStateMachine.cs / EnemyViewSensor.cs", "完了", "通常プレイで破綻しない状態まで確認", ""),
        12: ("UI/Prefab調整", "・HPゲージ、アイテム欄、ポーズ画面を作成\n・Prefab参照切れを修正\n・解像度変更時のレイアウト崩れを調整", "HUD.prefab / PauseMenu.prefab / CanvasScaler", "完了", "主要解像度で表示確認済み", ""),
        13: ("検証環境", "・通しプレイで停止系不具合がないことを確認\n・Missing参照とConsole Errorを確認\n・確認用スクリーンショットを整理", "QAチェックリスト / Screenshots / Console Log", "完了", "重大な停止系不具合なし", ""),
    }
    for row, values in implementation.items():
        for col, value in zip(["A", "B", "D", "F", "G", "H"], values):
            sheet[f"{col}{row}"].value = value

    other = {
        19: ("仕様整理", "・仕様書を実装状況に合わせて更新\n・画面仕様とマスタ仕様の差分を整理\n・READMEから関連資料へ移動できる導線を追加", "Docs/README.md / 仕様書 / 画面仕様", "完了", "資料と実装の差分を整理", ""),
        20: ("資料作成", "・工数表を更新\n・MOCで増えた仕様変更分を再計算\n・月次マイルストーン資料を作成", "工数表.xlsx / MOC議事メモ / マイルストーン表.xlsx", "完了", "共有用資料へ反映済み", ""),
        21: ("運用整理", "・Git運用ルールを整理\n・レビュー対象を一覧化\n・次週着手分を分離", "README.md / BRANCHING.md / タスク一覧", "完了", "次週の着手順を整理済み", ""),
        22: ("報告書整備", "・今週のGit差分を確認\n・実装作業と実装以外の業務を分けて整理\n・1ページで読める粒度に調整", "本報告書 / Git log", "完了", "内容確認済み", ""),
    }
    for row, values in other.items():
        for col, value in zip(["A", "B", "D", "F", "G", "H"], values):
            sheet[f"{col}{row}"].value = value

    sheet["A25"].value = "・今週は本番開発開始に必要な基盤と主要機能を完了扱いで整理。\n・次週はステージ追加、敵AIのパターン追加、QA指摘の軽微修正を予定。\n・工数再計算とマイルストーンは、仕様変更が出た場合に再更新する。"


clear_report_values(ws)
refresh_validation(ws)
set_page(ws, "A1:H28", "portrait")

example_ws = wb.copy_worksheet(ws)
fill_example(example_ws)
refresh_validation(example_ws)
add_example_guides(example_ws)
set_page(example_ws, "A1:L28", "landscape")

# Keep the blank report as the first sheet and the annotated example as the second sheet.
wb._sheets = [ws, example_ws]

wb.save(out)
print(out)
