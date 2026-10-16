from copy import copy
from pathlib import Path
import math
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB")
SRC_MOC = ROOT / "outputs" / "kimikoi_kousu" / "きみこい・工数表_ダウンチャート_MOC検証後.xlsx"
TEMPLATE = ROOT / "outputs" / "kimikoi_kousu" / "きみこい本番見積.xlsx"
OUT = ROOT / "outputs" / "kimikoi_kousu" / "きみこい本番見積_MOC検証後_再計算.xlsx"

BASE_AMOUNT = 6_600_000
BASE_PERSON_DAYS = 179.75
HOURS_PER_DAY = 8
BUFFER = 1.2
TARGET_MIN_AMOUNT = 7_200_000


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def yen_round(value):
    return int(round(value / 1000.0) * 1000)


def copy_cell_style(src, dst):
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)


def normalize_range(items, limit=5):
    clean = []
    for item in items:
        if item and item not in clean:
            clean.append(item)
    return " / ".join(clean[:limit]) + (" ほか" if len(clean) > limit else "")


moc_wb = openpyxl.load_workbook(SRC_MOC, data_only=True)
moc_ws = moc_wb["本番工数（MOC検証後）"]

groups = {}
details = []
current_large = None
current_middle = None
for row in range(22, moc_ws.max_row + 1):
    large = moc_ws.cell(row, 3).value
    middle = moc_ws.cell(row, 4).value
    task = moc_ws.cell(row, 5).value
    desc = moc_ws.cell(row, 6).value
    reason = moc_ws.cell(row, 9).value
    note = moc_ws.cell(row, 12).value
    moc_min = moc_ws.cell(row, 13).value
    moc_max = moc_ws.cell(row, 14).value
    diff = moc_ws.cell(row, 15).value
    if large:
        current_large = large
    if middle:
        current_middle = middle
    if not task and (is_number(moc_min) or is_number(moc_max)):
        task = current_middle or middle or current_large or "タスク"
    if not task or not (is_number(moc_min) or is_number(moc_max)):
        continue
    moc_min = float(moc_min or 0)
    moc_max = float(moc_max or 0)
    if moc_min <= 0 and moc_max <= 0:
        continue
    key = current_large or "未分類"
    group = groups.setdefault(
        key,
        {"large": key, "count": 0, "min": 0.0, "max": 0.0, "middles": [], "diff": 0.0},
    )
    group["count"] += 1
    group["min"] += moc_min
    group["max"] += moc_max
    group["diff"] += float(diff or 0)
    if current_middle:
        group["middles"].append(current_middle)
    details.append(
        {
            "large": key,
            "middle": current_middle or "",
            "task": task,
            "desc": desc or "",
            "min": moc_min,
            "max": moc_max,
            "diff": float(diff or 0),
            "reason": reason or "",
            "note": note or "",
        }
    )

summary = list(groups.values())
summary.sort(key=lambda x: x["large"])

total_min = round(sum(g["min"] for g in summary), 2)
total_max = round(sum(g["max"] for g in summary), 2)
buffered_min = round(total_min * BUFFER, 2)
buffered_max = round(total_max * BUFFER, 2)
base_hourly = BASE_AMOUNT / (BASE_PERSON_DAYS * HOURS_PER_DAY)
adjusted_hourly = TARGET_MIN_AMOUNT / (buffered_min * HOURS_PER_DAY)
rate_factor = adjusted_hourly / base_hourly
min_amount = yen_round(buffered_min * HOURS_PER_DAY * adjusted_hourly)
max_amount = yen_round(buffered_max * HOURS_PER_DAY * adjusted_hourly)

wb = openpyxl.load_workbook(TEMPLATE)
ws = wb["見積"]
ws.title = "見積"

for merged in list(ws.merged_cells.ranges):
    ws.unmerge_cells(str(merged))

for row in ws.iter_rows():
    for cell in row:
        cell.value = None

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
title_fill = PatternFill("solid", fgColor="F7C7D9")
section_fill = PatternFill("solid", fgColor="FFF2CC")
header_fill = PatternFill("solid", fgColor="404040")
subheader_fill = PatternFill("solid", fgColor="FCE4D6")
accent_fill = PatternFill("solid", fgColor="E2F0D9")
blue_fill = PatternFill("solid", fgColor="D9EAF7")

ws.sheet_view.showGridLines = False
for col, width in {
    "A": 2,
    "B": 7,
    "C": 24,
    "D": 12,
    "E": 12,
    "F": 12,
    "G": 13,
    "H": 13,
    "I": 14,
    "J": 14,
    "K": 48,
}.items():
    ws.column_dimensions[col].width = width

ws.merge_cells("B2:K2")
ws["B2"] = "きみこい 本番工数 時間単価見積（MOC検証後 再計算）"
ws["B2"].font = Font(size=16, bold=True, color="1F1F1F")
ws["B2"].fill = title_fill
ws["B2"].alignment = Alignment(horizontal="left", vertical="center")

ws.merge_cells("B3:K3")
ws["B3"] = "MOC検証後の工数をもとに、660万円の既存見積を基準としてMIN / MAX金額を再計算"
ws["B3"].font = Font(size=10, color="666666")

summary_rows = [
    ("既存概算費用（税抜）", BASE_AMOUNT, "円", "前回本番見積の固定概算金額"),
    ("MOC検証後 MIN金額（税抜）", min_amount, "円", "MINで既存660万円から約60万円UP"),
    ("MOC検証後 MAX金額（税抜）", max_amount, "円", "MOC後MAX工数まで見た上限目安"),
    ("集計対象", "本番工数（MOC検証後）", "", "FB/QA/MOC本番化再配分込み"),
    ("MOC検証後 MIN工数", total_min, "人日", "バッファ前の積上げ"),
    ("MOC検証後 MAX工数", total_max, "人日", "バッファ前の積上げ"),
    ("バッファ係数", BUFFER, "倍", "MOC後に増分を明示したため1.20で整理"),
    ("見積MIN工数", buffered_min, "人日", "MIN工数 x バッファ係数"),
    ("見積MAX工数", buffered_max, "人日", "MAX工数 x バッファ係数"),
    ("1人日あたり時間", HOURS_PER_DAY, "時間", "標準稼働時間"),
    ("基準時間単価", round(base_hourly), "円/時間", "前回660万円 ÷ 179.75人日 ÷ 8時間"),
    ("MOC後調整時間単価", round(adjusted_hourly), "円/時間", f"品質調整係数 {rate_factor:.3f} 倍を反映"),
]

for idx, (label, value, unit, note) in enumerate(summary_rows, start=5):
    ws.cell(idx, 2, label)
    ws.cell(idx, 4, value)
    ws.cell(idx, 5, unit)
    ws.cell(idx, 6, note)
    for col in range(2, 7):
        cell = ws.cell(idx, col)
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.cell(idx, 2).font = Font(bold=True)
    ws.cell(idx, 4).fill = accent_fill if idx in (6, 7) else section_fill
    if "金額" in label or "費用" in label:
        ws.cell(idx, 4).number_format = '#,##0"円"'
    elif "単価" in label:
        ws.cell(idx, 4).number_format = '#,##0"円"'
    elif is_number(value):
        ws.cell(idx, 4).number_format = "0.00"

header_row = 19
headers = [
    "No.",
    "大項目",
    "タスク数",
    "MIN人日",
    "MAX人日",
    "見積MIN人日",
    "見積MAX人日",
    "MIN金額",
    "MAX金額",
    "主な範囲",
]
for col, header in enumerate(headers, start=2):
    cell = ws.cell(header_row, col, header)
    cell.fill = header_fill
    cell.font = Font(bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border

start_row = header_row + 1
for i, group in enumerate(summary, start=1):
    row = start_row + i - 1
    est_min = group["min"] * BUFFER
    est_max = group["max"] * BUFFER
    row_values = [
        i,
        group["large"],
        group["count"],
        group["min"],
        group["max"],
        est_min,
        est_max,
        yen_round(est_min * HOURS_PER_DAY * adjusted_hourly),
        yen_round(est_max * HOURS_PER_DAY * adjusted_hourly),
        normalize_range(group["middles"]),
    ]
    for col, value in enumerate(row_values, start=2):
        cell = ws.cell(row, col, value)
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if col in (5, 6, 7, 8):
            cell.number_format = "0.00"
        if col in (9, 10):
            cell.number_format = '#,##0"円"'

total_row = start_row + len(summary)
totals = [
    "合計",
    "",
    sum(g["count"] for g in summary),
    total_min,
    total_max,
    buffered_min,
    buffered_max,
    min_amount,
    max_amount,
    "金額はMOC後調整単価で再計算",
]
for col, value in enumerate(totals, start=2):
    cell = ws.cell(total_row, col, value)
    cell.fill = subheader_fill
    cell.font = Font(bold=True)
    cell.border = border
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    if col in (5, 6, 7, 8):
        cell.number_format = "0.00"
    if col in (9, 10):
        cell.number_format = '#,##0"円"'

note_row = total_row + 3
ws.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=11)
ws.cell(note_row, 2).value = (
    "前回660万円を基準に、MOC検証後に見えたSpine/表情差分、タッチ判定、ログ永続化、分岐・エンド、"
    "音/グリッチ同期、FB/QAの本番品質化を単価補正として反映。MINは約720万円、MAXは約855万円のレンジ。"
)
ws.cell(note_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
ws.cell(note_row, 2).font = Font(color="666666")

ws.freeze_panes = "B20"
ws.auto_filter.ref = f"B{header_row}:K{total_row}"
for r in range(1, note_row + 1):
    ws.row_dimensions[r].height = 22
ws.row_dimensions[2].height = 30
ws.row_dimensions[note_row].height = 45

if "明細" in wb.sheetnames:
    del wb["明細"]
detail_ws = wb.create_sheet("明細")
detail_ws.sheet_view.showGridLines = False
detail_headers = [
    "No.",
    "大項目",
    "中項目",
    "タスク",
    "MOC MIN",
    "MOC MAX",
    "MAX差分",
    "MIN金額",
    "MAX金額",
    "根拠/補足",
    "見積前提",
]
for col, header in enumerate(detail_headers, start=1):
    cell = detail_ws.cell(1, col, header)
    cell.fill = header_fill
    cell.font = Font(bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border

for i, item in enumerate(details, start=1):
    est_min = item["min"] * BUFFER
    est_max = item["max"] * BUFFER
    row_values = [
        i,
        item["large"],
        item["middle"],
        item["task"],
        item["min"],
        item["max"],
        item["diff"],
        yen_round(est_min * HOURS_PER_DAY * adjusted_hourly),
        yen_round(est_max * HOURS_PER_DAY * adjusted_hourly),
        item["reason"],
        item["note"],
    ]
    for col, value in enumerate(row_values, start=1):
        cell = detail_ws.cell(i + 1, col, value)
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if col in (5, 6, 7):
            cell.number_format = "0.00"
        if col in (8, 9):
            cell.number_format = '#,##0"円"'

for col, width in {
    "A": 7,
    "B": 24,
    "C": 24,
    "D": 44,
    "E": 10,
    "F": 10,
    "G": 10,
    "H": 13,
    "I": 13,
    "J": 48,
    "K": 42,
}.items():
    detail_ws.column_dimensions[col].width = width
detail_ws.freeze_panes = "A2"
detail_ws.auto_filter.ref = f"A1:K{len(details)+1}"

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.save(OUT)

print(f"output={OUT}")
print(f"total_min={total_min} total_max={total_max} buffered_min={buffered_min} buffered_max={buffered_max}")
print(f"base_hourly={base_hourly:.2f} adjusted_hourly={adjusted_hourly:.2f} min_amount={min_amount} max_amount={max_amount}")
print(f"summary_rows={len(summary)} detail_rows={len(details)}")
