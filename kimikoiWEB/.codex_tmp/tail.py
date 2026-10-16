from openpyxl import load_workbook
path=r'C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/kimikoi_kousu/きみこい・工数表_ダウンチャート_自動化完成版.xlsx'
wb=load_workbook(path, data_only=False)
for s in ['ダウンチャート']:
 ws=wb[s]
 for row in range(250, ws.max_row+1):
  vals=[ws.cell(row,col).value for col in range(1,15)]
  if any(v is not None for v in vals): print(row, vals)
