from openpyxl import load_workbook
orig=r'C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/kimikoi_kousu/きみこい・工数表_ダウンチャート_自動化完成版.xlsx'
out=r'C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/kimikoi_kousu/きみこい・工数表_ダウンチャート_α版11月末調整.xlsx'
wb1=load_workbook(orig, data_only=False)
wb2=load_workbook(out, data_only=False)
for name in ['本番工数','ダウンチャート','祝日設定']:
    ws1=wb1[name]; ws2=wb2[name]
    diffs=[]
    if (ws1.max_row, ws1.max_column)!=(ws2.max_row, ws2.max_column):
        diffs.append(('dimension',(ws1.max_row, ws1.max_column),(ws2.max_row, ws2.max_column)))
    for r in range(1, ws1.max_row+1):
        for c in range(1, ws1.max_column+1):
            if ws1.cell(r,c).value != ws2.cell(r,c).value:
                diffs.append((ws1.cell(r,c).coordinate, ws1.cell(r,c).value, ws2.cell(r,c).value))
                break
        if len(diffs)>5: break
    print(name, 'diffs=', len(diffs), diffs[:5])
