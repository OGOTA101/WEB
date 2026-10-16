from pathlib import Path
import openpyxl

path = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_ダウンチャート_MOC検証後.xlsx")
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb["本番工数（MOC検証後）"]

rows = []
for row in range(22, ws.max_row + 1):
    task = ws.cell(row, 5).value
    mn = ws.cell(row, 13).value
    mx = ws.cell(row, 14).value
    if isinstance(mn, (int, float)) or isinstance(mx, (int, float)):
        if (mn or 0) > 0 or (mx or 0) > 0:
            rows.append((row, ws.cell(row, 3).value, ws.cell(row, 4).value, task, mn, mx))
print(len(rows), sum((r[4] or 0) for r in rows), sum((r[5] or 0) for r in rows))
for r in rows:
    if not r[3]:
        print("blank task", r)
