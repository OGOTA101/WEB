from openpyxl import load_workbook
from pathlib import Path
path=Path(r'C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/kimikoi_kousu/きみこい・工数表_α版11月末調整.xlsx')
wb=load_workbook(path, data_only=False)
print('sheets=', wb.sheetnames)
ws=wb['α版ダウンチャート']
alpha_ends=[]; post_starts=[]; total=0; alpha_total=0; post_total=0
for row in range(6, 256):
    phase=ws.cell(row,2).value
    if not phase: continue
    planned=float(ws.cell(row,7).value or 0)
    total += planned
    if phase.startswith('α版'):
        alpha_total += planned
        alpha_ends.append(ws.cell(row,9).value.date())
    else:
        post_total += planned
        post_starts.append(ws.cell(row,8).value.date())
print('total planned=', round(total,3), 'alpha=', round(alpha_total,3), 'post=', round(post_total,3))
print('alpha max end=', max(alpha_ends), 'post min start=', min(post_starts), 'final end=', max(ws.cell(row,9).value.date() for row in range(6,256) if ws.cell(row,9).value))
errors=[]
for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            v=cell.value
            if isinstance(v,str) and any(e in v for e in ['#REF!','#DIV/0!','#VALUE!','#NAME?','#N/A']):
                errors.append((sheet.title, cell.coordinate, v))
print('formula error strings=', len(errors))
for e in errors[:20]: print(e)
# basic circular-like self reference check for formulas in new sheets only
for name in ['α版工数_11月末','α版ダウンチャート']:
    sheet=wb[name]
    circ=[]
    for row in sheet.iter_rows():
        for cell in row:
            v=cell.value
            if isinstance(v,str) and v.startswith('=') and cell.coordinate in v:
                circ.append((cell.coordinate,v))
    print(name, 'self refs=', circ[:10], 'count=', len(circ))
