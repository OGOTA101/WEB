from openpyxl import load_workbook
from pathlib import Path
path = Path(r'C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/kimikoi_kousu/きみこい・工数表_ダウンチャート_自動化完成版.xlsx')
wb = load_workbook(path, data_only=False)
print(wb.sheetnames)
for ws in wb.worksheets:
    print('\n---', ws.title, ws.max_row, ws.max_column)
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row,40), min_col=1, max_col=min(ws.max_column,20), values_only=False):
        vals=[]
        for c in r:
            v=c.value
            if v is None: vals.append('')
            else: vals.append(str(v)[:80].replace('\n',' '))
        if any(vals): print(c.row, vals)
