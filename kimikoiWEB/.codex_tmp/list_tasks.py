from openpyxl import load_workbook
from pathlib import Path
path=Path(r'C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/kimikoi_kousu/きみこい・工数表_ダウンチャート_自動化完成版.xlsx')
wb=load_workbook(path, data_only=True)
ws=wb['ダウンチャート']
for row in range(6, ws.max_row+1):
    no=ws.cell(row,1).value
    if not no: continue
    large=ws.cell(row,2).value or ''
    mid=ws.cell(row,3).value or ''
    task=ws.cell(row,4).value or ''
    effort=ws.cell(row,6).value or 0
    start=ws.cell(row,7).value
    end=ws.cell(row,8).value
    print(f'{no}\t{large}\t{mid}\t{task}\t{effort}\t{start:%Y-%m-%d}\t{end:%Y-%m-%d}')
