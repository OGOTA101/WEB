from openpyxl import load_workbook
wb=load_workbook(r'C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/kimikoi_kousu/きみこい・工数表_ダウンチャート_自動化完成版.xlsx', data_only=True)
ws=wb['本番工数']
for row in range(340, 420):
 vals=[ws.cell(row,col).value for col in range(2,13)]
 if any(v is not None for v in vals): print(row, vals)
