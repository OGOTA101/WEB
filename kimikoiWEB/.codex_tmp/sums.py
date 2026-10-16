from openpyxl import load_workbook
from datetime import datetime
wb=load_workbook(r'C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/kimikoi_kousu/きみこい・工数表_ダウンチャート_自動化完成版.xlsx', data_only=True)
ws=wb['ダウンチャート']
for cutoff in [150,155,157,158]:
    total=sum(ws.cell(r,6).value or 0 for r in range(6, 6+cutoff))
    print(cutoff,total)
