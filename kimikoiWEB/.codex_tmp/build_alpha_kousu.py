from copy import copy
from datetime import date, datetime, timedelta
from math import ceil, floor
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


INPUT = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/kimikoi_kousu/きみこい・工数表_ダウンチャート_自動化完成版.xlsx")
OUTPUT = Path(r"C:/Users/ihogo/Desktop/WEB/kimikoiWEB/outputs/kimikoi_kousu/きみこい・工数表_ダウンチャート_α版11月末調整.xlsx")

ALPHA_CUTOFF_NO = 155
START_DATE = date(2026, 6, 1)
ALPHA_TARGET = date(2026, 11, 30)
POST_ALPHA_START = date(2026, 12, 1)


def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return value


def is_workday(day, holidays):
    return day.weekday() < 5 and day not in holidays


def workdays_from(start_day, count, holidays):
    days = []
    current = start_day
    while len(days) < count:
        if is_workday(current, holidays):
            days.append(current)
        current += timedelta(days=1)
    return days


def schedule_phase(tasks, start_day, holidays, daily_capacity=1.0):
    total_effort = sum(task["planned"] for task in tasks)
    workday_count = max(1, ceil(total_effort / daily_capacity))
    workdays = workdays_from(start_day, workday_count, holidays)
    cumulative = 0.0
    for task in tasks:
        before = cumulative
        after = cumulative + task["planned"]
        start_index = floor(before / daily_capacity) + 1
        end_index = max(start_index, ceil(after / daily_capacity))
        task["phase_cum_before"] = before
        task["phase_cum_after"] = after
        task["start_index"] = start_index
        task["end_index"] = end_index
        task["start_date"] = workdays[start_index - 1]
        task["end_date"] = workdays[end_index - 1]
        cumulative = after
    return tasks


def style_header(ws, row, min_col, max_col, fill="1F4E78"):
    for col in range(min_col, max_col + 1):
        cell = ws.cell(row, col)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color="B7B7B7"))


def clear_sheet_if_exists(wb, name):
    if name in wb.sheetnames:
        del wb[name]


def copy_column_dimensions(source_ws, target_ws, max_col):
    for col in range(1, max_col + 1):
        letter = get_column_letter(col)
        if letter in source_ws.column_dimensions:
            target_ws.column_dimensions[letter].width = source_ws.column_dimensions[letter].width


def build():
    wb = load_workbook(INPUT, data_only=False)

    clear_sheet_if_exists(wb, "α版工数_11月末")
    clear_sheet_if_exists(wb, "α版工数")
    clear_sheet_if_exists(wb, "α版ダウンチャート")

    source_chart = wb["ダウンチャート"]
    source_effort = wb["本番工数"]
    holidays_ws = wb["祝日設定"]
    holidays = {
        as_date(holidays_ws.cell(row, 1).value)
        for row in range(2, holidays_ws.max_row + 1)
        if holidays_ws.cell(row, 1).value
    }

    tasks = []
    for row in range(6, source_chart.max_row + 1):
        no = source_chart.cell(row, 1).value
        if not isinstance(no, int):
            continue
        task = {
            "no": no,
            "large": source_chart.cell(row, 2).value,
            "mid": source_chart.cell(row, 3).value,
            "task": source_chart.cell(row, 4).value,
            "effort": float(source_chart.cell(row, 5).value or 0),
            "planned": float(source_chart.cell(row, 6).value or 0),
            "reason": source_chart.cell(row, 9).value,
            "source_row": source_chart.cell(row, 10).value,
        }
        task["phase"] = "α版（全ルート通しプレイ可能）" if no <= ALPHA_CUTOFF_NO else "11月末以降（仕上げ・演出・QA）"
        tasks.append(task)

    alpha_tasks = [task for task in tasks if task["no"] <= ALPHA_CUTOFF_NO]
    post_tasks = [task for task in tasks if task["no"] > ALPHA_CUTOFF_NO]
    schedule_phase(alpha_tasks, START_DATE, holidays)
    schedule_phase(post_tasks, POST_ALPHA_START, holidays)
    scheduled = alpha_tasks + post_tasks

    effort_ws = wb.create_sheet("α版工数")
    effort_ws.sheet_view.showGridLines = False
    effort_ws.freeze_panes = "A8"
    effort_ws["A1"] = "きみこい 工数表（α版11月末調整）"
    effort_ws["A1"].font = Font(size=16, bold=True, color="1F4E78")
    effort_ws["A2"] = "前提"
    effort_ws["B2"] = "全ルートを遊べるところまで（パラメータや不具合は許容）を2026年11月末までに完了。残タスクは2026年12月1日以降に実行。"
    effort_ws["B2"].alignment = Alignment(wrap_text=True)
    total_planned = sum(task["planned"] for task in scheduled)
    alpha_planned = sum(task["planned"] for task in alpha_tasks)
    post_planned = sum(task["planned"] for task in post_tasks)
    alpha_complete = max(task["end_date"] for task in alpha_tasks)

    effort_ws["A3"] = "合計予定工数"
    effort_ws["B3"] = total_planned
    effort_ws["C3"] = "人日（元の工数から変更なし）"
    effort_ws["A4"] = "α版予定工数"
    effort_ws["B4"] = alpha_planned
    effort_ws["A5"] = "11月末以降予定工数"
    effort_ws["B5"] = post_planned
    effort_ws["A6"] = "α版完了目安"
    effort_ws["B6"] = alpha_complete
    effort_ws["C6"] = "目標: 2026-11-30 以内"

    headers = ["No", "フェーズ", "大項目", "中項目", "タスク", "工数", "予定工数", "根拠/備考", "元行", "開始日", "終了日"]
    for col, header in enumerate(headers, 1):
        effort_ws.cell(7, col).value = header
    style_header(effort_ws, 7, 1, len(headers))

    for idx, task in enumerate(scheduled, 8):
        values = [
            task["no"],
            task["phase"],
            task["large"],
            task["mid"],
            task["task"],
            task["effort"],
            task["planned"],
            task["reason"],
            task["source_row"],
            task["start_date"],
            task["end_date"],
        ]
        for col, value in enumerate(values, 1):
            effort_ws.cell(idx, col).value = value
            effort_ws.cell(idx, col).alignment = Alignment(vertical="top", wrap_text=(col in [2, 5, 8]))
            effort_ws.cell(idx, col).border = Border(bottom=Side(style="hair", color="D9E2F3"))

    widths = [8, 26, 22, 24, 42, 10, 12, 38, 8, 12, 12]
    for col, width in enumerate(widths, 1):
        effort_ws.column_dimensions[get_column_letter(col)].width = width
    for row in range(8, 258):
        effort_ws.row_dimensions[row].height = 34
    for row in [3, 4, 5, 6]:
        effort_ws.cell(row, 2).number_format = "0.000"
    effort_ws["B6"].number_format = "yyyy-mm-dd"
    effort_ws.auto_filter.ref = "A7:K257"
    effort_ws.conditional_formatting.add(
        "B8:B257",
        CellIsRule(operator="equal", formula=['"α版（全ルート通しプレイ可能）"'], fill=PatternFill("solid", fgColor="E2F0D9")),
    )
    effort_ws.conditional_formatting.add(
        "B8:B257",
        CellIsRule(operator="equal", formula=['"11月末以降（仕上げ・演出・QA）"'], fill=PatternFill("solid", fgColor="FFF2CC")),
    )

    chart_ws = wb.create_sheet("α版ダウンチャート")
    chart_ws.sheet_view.showGridLines = False
    chart_ws.freeze_panes = "O6"

    chart_ws["A1"] = "きみこい α版11月末 ダウンチャート"
    chart_ws["A1"].font = Font(size=16, bold=True, color="1F4E78")
    top_labels = [
        ("A2", "開始日", START_DATE),
        ("B2", "α版目標", ALPHA_TARGET),
        ("C2", "α版完了目安", max(task["end_date"] for task in alpha_tasks)),
        ("D2", "終了目安", max(task["end_date"] for task in scheduled)),
        ("E2", "タスク数", len(scheduled)),
        ("F2", "予定工数合計", sum(task["planned"] for task in scheduled)),
        ("G2", "α版予定工数", sum(task["planned"] for task in alpha_tasks)),
        ("H2", "11月末以降予定工数", sum(task["planned"] for task in post_tasks)),
        ("I2", "日次消化力", 1),
        ("J2", "休日前提", "土日祝休み"),
    ]
    for cell_ref, label, value in top_labels:
        col = chart_ws[cell_ref].column
        chart_ws.cell(2, col).value = label
        chart_ws.cell(3, col).value = value
        chart_ws.cell(2, col).font = Font(bold=True, color="FFFFFF")
        chart_ws.cell(2, col).fill = PatternFill("solid", fgColor="1F4E78")
        chart_ws.cell(2, col).alignment = Alignment(horizontal="center", wrap_text=True)
        chart_ws.cell(3, col).alignment = Alignment(horizontal="center", wrap_text=True)
        if isinstance(value, date):
            chart_ws.cell(3, col).number_format = "yyyy-mm-dd"
        if isinstance(value, float):
            chart_ws.cell(3, col).number_format = "0.000"

    chart_headers = ["No", "フェーズ", "大項目", "中項目", "タスク", "工数", "予定工数", "開始日", "終了日", "根拠/備考", "元行", "累計前", "累計後", "開始Index", "終了Index"]
    for col, header in enumerate(chart_headers, 1):
        chart_ws.cell(5, col).value = header
    style_header(chart_ws, 5, 1, len(chart_headers))

    for idx, task in enumerate(scheduled, 6):
        values = [
            task["no"],
            task["phase"],
            task["large"],
            task["mid"],
            task["task"],
            task["effort"],
            task["planned"],
            task["start_date"],
            task["end_date"],
            task["reason"],
            task["source_row"],
            task["phase_cum_before"],
            task["phase_cum_after"],
            task["start_index"],
            task["end_index"],
        ]
        for col, value in enumerate(values, 1):
            cell = chart_ws.cell(idx, col)
            cell.value = value
            cell.alignment = Alignment(vertical="top", wrap_text=(col in [2, 5, 10]))
            cell.border = Border(bottom=Side(style="hair", color="D9E2F3"))
            if col in [8, 9]:
                cell.number_format = "yyyy-mm-dd"
            if col in [6, 7, 12, 13]:
                cell.number_format = "0.000"

    chart_widths = [7, 28, 20, 22, 40, 9, 11, 12, 12, 34, 8, 10, 10, 10, 10]
    for col, width in enumerate(chart_widths, 1):
        chart_ws.column_dimensions[get_column_letter(col)].width = width
    for row in range(6, 256):
        chart_ws.row_dimensions[row].height = 28
    chart_ws.auto_filter.ref = "A5:O255"

    min_timeline = START_DATE
    max_timeline = max(task["end_date"] for task in scheduled)
    days = []
    current = min_timeline
    while current <= max_timeline:
        days.append(current)
        current += timedelta(days=1)

    timeline_start_col = 16
    weekday_labels = ["月", "火", "水", "木", "金", "土", "日"]
    for offset, day in enumerate(days):
        col = timeline_start_col + offset
        month_label = f"{day.year}/{day.month}" if day.day == 1 else ""
        chart_ws.cell(3, col).value = month_label
        chart_ws.cell(4, col).value = day
        chart_ws.cell(5, col).value = weekday_labels[day.weekday()]
        chart_ws.cell(4, col).number_format = "m/d"
        chart_ws.column_dimensions[get_column_letter(col)].width = 3.2
        for row in [3, 4, 5]:
            chart_ws.cell(row, col).alignment = Alignment(horizontal="center")
            chart_ws.cell(row, col).font = Font(size=8, bold=(row == 3))
        if day.weekday() >= 5:
            for row in range(4, 256):
                chart_ws.cell(row, col).fill = PatternFill("solid", fgColor="E7E6E6")
        if day in holidays:
            for row in range(4, 256):
                chart_ws.cell(row, col).fill = PatternFill("solid", fgColor="FCE4D6")
        if day == ALPHA_TARGET:
            for row in range(3, 256):
                chart_ws.cell(row, col).fill = PatternFill("solid", fgColor="F4B183")

    alpha_fill = PatternFill("solid", fgColor="70AD47")
    post_fill = PatternFill("solid", fgColor="5B9BD5")
    for idx, task in enumerate(scheduled, 6):
        fill = alpha_fill if task["no"] <= ALPHA_CUTOFF_NO else post_fill
        for offset, day in enumerate(days):
            if task["start_date"] <= day <= task["end_date"]:
                cell = chart_ws.cell(idx, timeline_start_col + offset)
                cell.fill = fill
                cell.value = "■"
                cell.font = Font(color="FFFFFF", size=6)
                cell.alignment = Alignment(horizontal="center", vertical="center")

    legend_row = len(scheduled) + 8
    chart_ws.cell(legend_row, 1).value = "凡例"
    chart_ws.cell(legend_row, 2).value = "α版（全ルート通しプレイ可能）"
    chart_ws.cell(legend_row, 2).fill = alpha_fill
    chart_ws.cell(legend_row, 3).value = "11月末以降（仕上げ・演出・QA）"
    chart_ws.cell(legend_row, 3).fill = post_fill
    chart_ws.cell(legend_row, 4).value = "土日"
    chart_ws.cell(legend_row, 4).fill = PatternFill("solid", fgColor="E7E6E6")
    chart_ws.cell(legend_row, 5).value = "祝日"
    chart_ws.cell(legend_row, 5).fill = PatternFill("solid", fgColor="FCE4D6")
    chart_ws.cell(legend_row, 6).value = "α版目標日"
    chart_ws.cell(legend_row, 6).fill = PatternFill("solid", fgColor="F4B183")

    for ws in [effort_ws, chart_ws]:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cell.alignment = copy(cell.alignment)
                    cell.alignment = Alignment(
                        horizontal=cell.alignment.horizontal,
                        vertical=cell.alignment.vertical or "top",
                        wrap_text=cell.alignment.wrap_text,
                    )

    # Ensure the original sheets remain first and untouched in order; new sheets are appended after them.
    wb.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
