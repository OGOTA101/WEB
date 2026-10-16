from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile
import re
import xml.etree.ElementTree as ET

from openpyxl import load_workbook


PATH = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_本番割増_詳細版.xlsx")
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
ET.register_namespace("", NS_MAIN)
ET.register_namespace("r", NS_REL)


def q(ns, tag):
    return f"{{{ns}}}{tag}"


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def cache_for_sheet(ws):
    g_sum = sum(cell.value for cell in ws["G"] if number(cell.value))
    h_sum = sum(cell.value for cell in ws["H"] if number(cell.value))
    risk = ws["D6"].value or 0
    people = ws["D15"].value or 1
    days = ws["D16"].value or 20
    unit_price = ws["D11"].value or 0
    d8 = g_sum * risk
    d9 = h_sum * risk
    return {
        "D4": g_sum,
        "D5": h_sum,
        "D8": d8,
        "D9": d9,
        "D12": d8 * unit_price,
        "D13": d9 * unit_price,
        "D17": d8 / days / people if days and people else 0,
        "D18": d9 / days / people if days and people else 0,
    }


wb_values = load_workbook(PATH, data_only=False)
target_sheets = {
    name: cache_for_sheet(wb_values[name])
    for name in ["工数表ベース", "工数表ベース原本", "本番割増_詳細"]
}

with ZipFile(PATH, "r") as zin:
    workbook_xml = ET.fromstring(zin.read("xl/workbook.xml"))
    rels_xml = ET.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
    rid_to_target = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in rels_xml.findall(q(NS_PKG_REL, "Relationship"))
    }
    sheet_to_part = {}
    for sheet in workbook_xml.find(q(NS_MAIN, "sheets")).findall(q(NS_MAIN, "sheet")):
        name = sheet.attrib["name"]
        if name not in target_sheets:
            continue
        rid = sheet.attrib[q(NS_REL, "id")]
        target = rid_to_target[rid].lstrip("/")
        if not target.startswith("xl/"):
            target = f"xl/{target}"
        sheet_to_part[name] = target

    tmp = NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp_path = Path(tmp.name)
    tmp.close()
    with ZipFile(tmp_path, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            sheet_name = next((name for name, part in sheet_to_part.items() if part == item.filename), None)
            if sheet_name:
                root = ET.fromstring(data)
                wanted = target_sheets[sheet_name]
                sheet_data = root.find(q(NS_MAIN, "sheetData"))
                for row in sheet_data.findall(q(NS_MAIN, "row")):
                    for cell in row.findall(q(NS_MAIN, "c")):
                        ref = re.sub(r"\$|'.*'!", "", cell.attrib.get("r", ""))
                        if ref not in wanted:
                            continue
                        v = cell.find(q(NS_MAIN, "v"))
                        if v is None:
                            v = ET.SubElement(cell, q(NS_MAIN, "v"))
                        v.text = f"{wanted[ref]:.10g}"
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            zout.writestr(item, data)

tmp_path.replace(PATH)
for sheet, values in target_sheets.items():
    print(sheet, values)
print(PATH)
