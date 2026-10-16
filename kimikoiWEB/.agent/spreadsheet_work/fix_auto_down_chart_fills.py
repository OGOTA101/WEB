from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Border, PatternFill, Side


WORKBOOK = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_ダウンチャート_自動化完成版.xlsx")
SHEET = "ダウンチャート"
TASK_START_ROW = 6
DATE_START_COL = 15
DATE_DAYS = 730

BAR_FILL = PatternFill("solid", fgColor="9DC3E6")
WEEKEND_FILL = PatternFill("solid", fgColor="E7E6E6")
HOLIDAY_FILL = PatternFill("solid", fgColor="D9D2E9")
BLANK_FILL = PatternFill(fill_type=None)
HAIR = Side(style="hair", color="D9D9D9")
HAIR_BORDER = Border(left=HAIR, right=HAIR, top=HAIR, bottom=HAIR)


def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    return None


data_wb = load_workbook(WORKBOOK, data_only=True)
data_ws = data_wb[SHEET]
holiday_ws = data_wb["祝日設定"]
holidays = {
    value.date()
    for (value,) in holiday_ws.iter_rows(min_row=2, max_col=1, values_only=True)
    if isinstance(value, datetime)
}

end_row = TASK_START_ROW
for row in range(TASK_START_ROW, data_ws.max_row + 1):
    if data_ws.cell(row, 1).value:
        end_row = row

date_values = []
for offset in range(DATE_DAYS):
    value = data_ws.cell(4, DATE_START_COL + offset).value
    date_values.append(as_date(value))

tasks = []
for row in range(TASK_START_ROW, end_row + 1):
    start = as_date(data_ws.cell(row, 7).value)
    end = as_date(data_ws.cell(row, 8).value)
    tasks.append((row, start, end))

wb = load_workbook(WORKBOOK)
ws = wb[SHEET]

for offset, current in enumerate(date_values):
    col = DATE_START_COL + offset
    if current is None:
        continue

    is_holiday = current in holidays
    is_weekend = current.weekday() >= 5
    date_fill = HOLIDAY_FILL if is_holiday else WEEKEND_FILL if is_weekend else BLANK_FILL

    for row in (4, 5):
        ws.cell(row, col).fill = date_fill

    for row, start, end in tasks:
        cell = ws.cell(row, col)
        cell.border = HAIR_BORDER
        if start is None or end is None:
            cell.fill = BLANK_FILL
            continue
        if is_holiday:
            cell.fill = HOLIDAY_FILL
        elif is_weekend:
            cell.fill = WEEKEND_FILL
        elif start <= current <= end:
            cell.fill = BAR_FILL
        else:
            cell.fill = BLANK_FILL

wb.save(WORKBOOK)

print(WORKBOOK)
print("painted_rows", len(tasks))
print("date_columns", sum(1 for value in date_values if value is not None))
