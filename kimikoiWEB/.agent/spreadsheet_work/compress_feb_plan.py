from __future__ import annotations

from collections import defaultdict
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side


SOURCE = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_本番割増_再調整_単独シート.xlsx")
TEMP = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\.agent\spreadsheet_work\feb_compressed_source.xlsx")
SHEET = "本番割増_再調整"
TARGET_RAW_MAX = 145.8
BUFFER_RATE = 1.25


def as_number(value):
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value)
        except ValueError:
            return None
    return None


def base_factor(text: str, original_max: float) -> tuple[float, str, bool]:
    """Return factor, reason, protected flag.

    Protected tasks are intentionally kept heavier because they are core runtime,
    new specification work, or debug/QA needed to finish, not merely to demo.
    """
    reason = "共通化・AI補助・実装範囲固定で圧縮"
    protected = False
    factor = 0.50

    if original_max <= 0.2:
        factor = 0.50
        reason = "軽作業として0.1人日単位へ圧縮"
    elif original_max <= 0.5:
        factor = 0.40
        reason = "定型・短時間作業として軽量化"

    if re.search(r"手動|エディタ|インポート|設定作業|定型|コマンド|ボタン|短い|ごく少量|Prefab作成", text):
        factor = min(factor, 0.30)
        reason = "エディタ/定型作業のため軽量化"

    if re.search(r"Steam|Steamworks|DRM|実績API|ストア", text):
        factor = min(factor, 0.10)
        reason = "2月末版では後送/最小接続にして圧縮"

    if re.search(r"Localization|Localize|ローカライズ|多言語", text):
        factor = min(factor, 0.20)
        reason = "初回は日本語本線に絞り多言語対応を後送"

    if re.search(r"ゲームパッド|コントローラー|仮想カーソル|MoveCursor|デバイス検知|マグネット|カーソル加速", text):
        factor = min(factor, 0.30)
        reason = "PCマウス前提へ寄せ、操作拡張を後送"

    if re.search(r"リサイズ|Resize|サイズ更新|最小サイズ", text):
        factor = min(factor, 0.15)
        reason = "固定/限定サイズ運用へ寄せて圧縮"

    if re.search(r"Open\(\)|Close\(\)|ヘッダー色|アニメーション|DOTween演出|演出寄り|視認性|フェード|クロスフェード", text):
        factor = min(factor, 0.30)
        reason = "演出調整は最低限に限定して圧縮"

    if re.search(r"Googru|メール|ゴミ箱|推し活|チートツール|便利ツール|オカルトサイト|検索履歴", text):
        factor = min(factor, 0.35)
        reason = "周辺アプリはシナリオ必須範囲に絞って圧縮"

    if re.search(r"WindowBase|疑似OS|WindowManager|ドラッグ|フォーカス|デスクトップ|タスクバー|ウィンドウ", text):
        factor = max(factor, 0.45)
        reason = "疑似OS中核は簡略化しつつ保持"

    if re.search(r"シナリオ|ADV|分岐|既読|バックログ|セーブ|ロード|変数|フラグ|ステート|コマンド", text):
        factor = max(factor, 0.55)
        reason = "進行/ADV中核は重さを残して圧縮"
        protected = True

    if re.search(r"タッチ|Spine|表情|Expression|スキンシップ|Collider|コライダー|Reaction", text):
        factor = max(factor, 0.58)
        reason = "新仕様の中核なので削りすぎず圧縮"
        protected = True

    if re.search(r"Debug|デバッグ|Exe|ビルド|検証|QA|統合|クラッシュ|最適化|パフォーマンス|Profiler|リグレッション", text):
        factor = max(factor, 0.70)
        reason = "デバッグ込み完了のため枠を保持"
        protected = True

    return factor, reason, protected


def rounded_effort(value: float) -> float:
    rounded = round(value, 1)
    if value > 0 and rounded == 0:
        rounded = 0.1
    return rounded


wb = load_workbook(SOURCE)
ws = wb[SHEET]

tasks = []
major = ""
mid = ""
for row in range(22, ws.max_row + 1):
    major_value = ws.cell(row, 3).value
    mid_value = ws.cell(row, 4).value
    if isinstance(major_value, str) and major_value.strip():
        major = major_value.strip()
    if isinstance(mid_value, str) and mid_value.strip():
        mid = mid_value.strip()

    original_max = as_number(ws.cell(row, 8).value)
    original_min = as_number(ws.cell(row, 7).value)
    if original_max is None or original_max <= 0:
        continue

    text = " ".join(
        str(value)
        for value in [
            major,
            mid,
            ws.cell(row, 5).value,
            ws.cell(row, 6).value,
            ws.cell(row, 9).value,
            ws.cell(row, 10).value,
            ws.cell(row, 11).value,
            ws.cell(row, 12).value,
        ]
        if value
    )
    factor, reason, protected = base_factor(text, original_max)
    new_max = rounded_effort(original_max * factor)
    tasks.append(
        {
            "row": row,
            "major": major,
            "original_min": original_min or original_max,
            "original_max": original_max,
            "new_max": new_max,
            "reason": reason,
            "protected": protected,
        }
    )

current_total = sum(task["new_max"] for task in tasks)
eligible = [task for task in tasks if not task["protected"] and task["new_max"] > 0.1]
protected_total = sum(task["new_max"] for task in tasks if task["protected"])
eligible_total = sum(task["new_max"] for task in eligible)
scale = (TARGET_RAW_MAX - protected_total) / eligible_total
scale = max(0.1, min(1.0, scale))

for task in eligible:
    task["new_max"] = rounded_effort(task["new_max"] * scale)

total = sum(task["new_max"] for task in tasks)
if total > TARGET_RAW_MAX:
    # Shave 0.1 from the least risky unprotected tasks until the plan fits.
    for task in sorted(
        [task for task in tasks if not task["protected"] and task["new_max"] > 0.1],
        key=lambda x: (x["new_max"], x["original_max"]),
        reverse=True,
    ):
        if total <= TARGET_RAW_MAX:
            break
        task["new_max"] = round(task["new_max"] - 0.1, 1)
        total = round(total - 0.1, 1)

reduction_fill = PatternFill("solid", fgColor="E2F0D9")
note_fill = PatternFill("solid", fgColor="DDEBF7")
thin = Side(style="thin", color="7F7F7F")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for task in tasks:
    row = task["row"]
    new_max = task["new_max"]
    new_min = rounded_effort(max(0.1, new_max * 0.85))
    if new_min > new_max:
        new_min = new_max
    ws.cell(row, 7, new_min)
    ws.cell(row, 8, new_max)

    old_reason = ws.cell(row, 12).value
    reason = f"2月末圧縮: {task['original_max']:.1f}->{new_max:.1f}。{task['reason']}"
    if old_reason and "2月末圧縮" not in str(old_reason):
        reason = f"{old_reason} / {reason}"
    ws.cell(row, 12, reason)

    for col in (7, 8, 12):
        ws.cell(row, col).fill = reduction_fill if col in (7, 8) else note_fill
        ws.cell(row, col).border = border
    ws.cell(row, 12).font = Font(size=9)

raw_min = sum(as_number(ws.cell(task["row"], 7).value) or 0 for task in tasks)
raw_max = sum(as_number(ws.cell(task["row"], 8).value) or 0 for task in tasks)
buffered_min = raw_min * BUFFER_RATE
buffered_max = raw_max * BUFFER_RATE

ws["D6"] = BUFFER_RATE
ws["D14"] = f"2月末圧縮版: 原本MAX 265.2→{raw_max:.1f} / バッファ込MAX {buffered_max:.1f}人日"
ws["F14"] = "2027/2末完了想定。軽作業/後送可能範囲を圧縮し、デバッグ・統合枠は保持"
ws["D19"] = "※2月末圧縮版: 薄緑セルは工数削減済み。右端L列にエンジニア視点の削減理由を記載"

for cell in ("D14", "F14", "D19"):
    ws[cell].fill = note_fill
    ws[cell].font = Font(bold=True if cell == "D19" else False, size=9)

for row in range(22, ws.max_row + 1):
    for col in range(2, 13):
        ws.cell(row, col).border = border

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
TEMP.parent.mkdir(parents=True, exist_ok=True)
wb.save(TEMP)

by_major = defaultdict(float)
for task in tasks:
    by_major[task["major"]] += task["new_max"]

print(TEMP)
print("raw_min", round(raw_min, 1))
print("raw_max", round(raw_max, 1))
print("buffered_min", round(buffered_min, 1))
print("buffered_max", round(buffered_max, 1))
print("base_rule_total", round(current_total, 1))
for major, value in by_major.items():
    print(major, round(value, 1), round(value * BUFFER_RATE, 1))
