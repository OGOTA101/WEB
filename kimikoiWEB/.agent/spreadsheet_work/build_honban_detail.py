from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


INPUT = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\きみこい・工数表.xlsx")
OUTPUT_DIR = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu")
OUTPUT = OUTPUT_DIR / "きみこい・工数表_本番割増_詳細版.xlsx"
SOURCE_SHEET = "工数表ベース原本"
TARGET_SHEET = "本番割増_詳細"


def as_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


detail_rows = [
    ("企画/仕様棚卸し", "最新企画書の本番スコープ再整理", "Sweet Sister Nightmareで本番化する範囲と除外範囲を決める", 0.4, 0.5, "企画書の読み替えと実装粒度の整理", "仕様棚卸し", "元PDF差分を見ながら整理"),
    ("企画/仕様棚卸し", "受入条件リスト作成", "各機能が完成したと言える条件を箇条書き化する", 0.3, 0.5, "後工程のQA観点に直結するため", "Definition of Done", None),
    ("企画/仕様棚卸し", "10日間進行の分岐表更新", "日付・イベント・好感度/恐怖値の関係を表に落とす", 0.5, 0.8, "ルート漏れ防止のため", "10日目到達条件", None),
    ("企画/仕様棚卸し", "既存工数表との重複確認", "原本の既存項目と今回の本番割増項目の重複を避ける", 0.3, 0.4, "不自然な二重計上を防ぐ", "見積調整", None),
    ("企画/仕様棚卸し", "演出優先度の剪定", "必須演出、できれば演出、見送り候補を分類する", 0.2, 0.3, "技術検証前の仮説整理", "優先度付け", None),
    ("妹 表情差分/Spine", "妹ちゃん表情差分リスト作成", "通常、照れ、恐怖、怒り、虚ろ、笑顔崩れ等の必要表情を洗う", 0.4, 0.6, "素材発注/実装受入の前提", "表情差分", None),
    ("妹 表情差分/Spine", "表情ID/命名規則整備", "シナリオから表情を呼び出すためのID名を決める", 0.2, 0.3, "データ駆動にするため", "Expression ID", None),
    ("妹 表情差分/Spine", "Spineスロット/アタッチメント設計", "目、眉、口、頬、影など差分パーツの切替設計を行う", 0.8, 1.2, "素材構造に依存して調整が増える", "Spine構成", "素材PSD/Spineデータの粒度次第"),
    ("妹 表情差分/Spine", "表情ブレンド制御", "笑顔から不穏、怯えから狂気など中間表情を制御する", 0.8, 1.4, "かわいい/怖いの境界調整が難しい", "Expression Blend", None),
    ("妹 表情差分/Spine", "瞬き/呼吸/待機モーション調整", "妹ちゃんの常時表示時に破綻しない待機演出を入れる", 0.5, 0.8, "Spineタイムライン調整", "Idle Motion", None),
    ("妹 表情差分/Spine", "恐怖進行に応じた見え方変化", "パラメータや日数で表情・影・ノイズ感を段階的に変える", 0.7, 1.3, "演出品質の確認が必要", "怖いのにかわいい", None),
    ("妹 表情差分/Spine", "差分素材の受入/差し替え確認", "表情差分素材を実装し、ズレや破綻を確認する", 0.5, 0.9, "素材戻しが出やすい", "素材受入", None),
    ("妹 表情差分/Spine", "Unity上のSpine表示統合", "既存UI/疑似デスクトップ上にSpineキャラを安定表示する", 0.7, 1.2, "Canvas/Render順の調整が必要", "Unity Spine", None),
    ("妹 表情差分/Spine", "静止画フォールバック対応", "Spine未対応箇所や軽量表示用に静止画差分でも出せるようにする", 0.3, 0.5, "安全策として最低限用意", "Fallback", None),
    ("妹 表情差分/Spine", "Atlas/メモリ使用量確認", "表情差分増加によるテクスチャ容量とロード負荷を確認する", 0.4, 0.7, "ビルド後の負荷確認が必要", "Performance", None),
    ("妹 表情差分/Spine", "シナリオイベントからの表情呼び出し", "台詞やイベントに合わせてSpine表情を切り替える口を作る", 0.6, 1.0, "データ形式との接続が必要", "Scenario Hook", None),
    ("技術剪定/演出検証", "グリッチ演出プロトタイプ", "画面ノイズ、色収差、ズレ表現の成立ラインを試す", 0.7, 1.2, "本実装前の見た目検証", "Shader検証", "重い演出は削る前提"),
    ("技術剪定/演出検証", "疑似デスクトップ重ね合わせ検証", "LIMEやボード、妹表示を同時に出した時の破綻を確認する", 0.5, 0.8, "Canvas階層と入力優先度の確認", "Desktop検証", None),
    ("技術剪定/演出検証", "偽カーソル/入力乗っ取り検証", "ビルド状態でカーソル演出が安全に見せられるか試す", 0.4, 0.7, "OS依存リスクの確認", "Input検証", None),
    ("技術剪定/演出検証", "動画/RenderTexture検証", "実写風・異物感演出をゲーム画面内に自然に混ぜる", 0.4, 0.7, "再生負荷と解像度の確認", "VideoPlayer", None),
    ("技術剪定/演出検証", "セーブ/ロード中の演出割り込み検証", "異変演出とセーブデータ更新が衝突しないか確認する", 0.5, 0.9, "進行不能リスクが高い", "Save Safety", None),
    ("技術剪定/演出検証", "低スペック想定の負荷予算確認", "演出を盛った時のFPS/メモリ目安を確認する", 0.4, 0.6, "製品版での最低限の安定性", "負荷確認", None),
    ("技術剪定/演出検証", "見送り候補の技術判断", "重すぎる/危ない演出を本番から落とす判断材料を作る", 0.3, 0.5, "剪定のための整理", "実装判断", None),
    ("技術剪定/演出検証", "検証結果メモ整理", "採用/保留/廃止の理由を後で追えるように残す", 0.3, 0.4, "チーム共有用", "検証ログ", None),
    ("タッチアクション", "妹ちゃんタッチ判定エリア設計", "頭、顔、肩、手元など反応箇所を定義する", 0.6, 0.9, "トリッカル系の触り心地を再現", "Touch Area", None),
    ("タッチアクション", "タッチ反応テーブル作成", "場所/日数/パラメータ別の反応台詞と表情を整理する", 0.4, 0.7, "シナリオと表情の接続が必要", "Reaction Table", None),
    ("タッチアクション", "パラメータ変動連携", "タッチ結果で好感度、イライラ度、恐怖値等が変わるようにする", 0.5, 0.9, "既存パラメータ管理との接続", "Parameter Hook", None),
    ("タッチアクション", "Spineリアクション再生", "タッチ箇所に応じて表情/モーション/揺れを再生する", 0.5, 0.9, "素材と演出調整が絡む", "Spine Reaction", None),
    ("タッチアクション", "連打/長押し/禁止操作の分岐", "押し方によって反応が変わる処理を追加する", 0.4, 0.7, "嫌がる/照れる/壊れる反応差分", "Input Pattern", None),
    ("タッチアクション", "クールタイム/暴発防止", "連続入力で演出が重ならないよう制御する", 0.3, 0.4, "演出破綻防止", "Cooldown", None),
    ("タッチアクション", "ボイス/SE連携", "タッチ時のSE、短い声、ノイズ音を同期する", 0.4, 0.7, "音量/タイミング調整込み", "Audio Hook", None),
    ("タッチアクション", "初回チュートリアル/誘導表示", "触れることが自然に伝わる最低限の導線を作る", 0.2, 0.3, "UI文言と表示条件のみ", "Guide", None),
    ("タッチアクション", "タッチ不能状態の例外処理", "イベント中や暗転中はタッチできないようにする", 0.3, 0.5, "進行不能防止", "Guard", None),
    ("Exeシナリオ読込/デバッグ", "ビルド用シナリオ配置設計", "StreamingAssets等に置いたシナリオをExeから読む構成にする", 0.6, 1.1, "Editor専用パスを避ける", "Build Path", None),
    ("Exeシナリオ読込/デバッグ", "ファイルパス/文字コード検証", "Windowsビルドで日本語ファイル名や改行が壊れないか確認する", 0.5, 0.8, "環境差分の確認", "Encoding", None),
    ("Exeシナリオ読込/デバッグ", "シナリオ再読込デバッグ機能", "ビルド状態でもシナリオを再読込して確認できる導線を作る", 0.7, 1.2, "調整速度を上げるため", "Hot Reload", None),
    ("Exeシナリオ読込/デバッグ", "シナリオ構文チェック", "ID欠け、表情ID不一致、ジャンプ先欠けを検出する", 0.8, 1.3, "ビルド後の原因特定を早くする", "Validator", None),
    ("Exeシナリオ読込/デバッグ", "日付/イベントジャンプ機能", "10日間進行の任意地点へビルド中に飛べるようにする", 0.6, 1.0, "通し確認時間を短縮", "Scene Jump", None),
    ("Exeシナリオ読込/デバッグ", "パラメータ編集パネル", "好感度や恐怖値などをビルド中に変更して分岐確認する", 0.5, 0.8, "既存デバッグ機能の本番寄せ", "Param Debug", None),
    ("Exeシナリオ読込/デバッグ", "フラグビューア", "現在ONになっているイベント/ED/回想フラグを一覧表示する", 0.5, 0.9, "分岐不具合の調査用", "Flag Viewer", None),
    ("Exeシナリオ読込/デバッグ", "EDルート簡易シミュレータ", "現在値から到達しうるEDを確認する簡易機能を作る", 0.7, 1.2, "複数EDの確認負荷を下げる", "ED Debug", None),
    ("Exeシナリオ読込/デバッグ", "ログ出力/コピー機能", "ビルド中の進行ログを外部へ出せるようにする", 0.4, 0.6, "不具合報告をしやすくする", "Log Export", None),
    ("Exeシナリオ読込/デバッグ", "デバッグメニュー表示制御", "製品版で誤ってデバッグ機能が出ないようにする", 0.4, 0.6, "ビルド設定との連動", "Debug Gate", None),
    ("Exeシナリオ読込/デバッグ", "セーブデータ初期化/差し替え", "ビルド中にテスト用セーブへ戻せるようにする", 0.3, 0.5, "テスト反復用", "Save Reset", None),
    ("Exeシナリオ読込/デバッグ", "読込失敗時のエラー表示", "シナリオ欠損や壊れたデータを画面/ログで分かるようにする", 0.4, 0.7, "原因不明の停止を防ぐ", "Error Handling", None),
    ("統合/QA", "1周通しスモークテスト", "開始から10日目/EDまで最低1ルートを通す", 0.7, 1.2, "本番化項目の総合確認", "Smoke Test", None),
    ("統合/QA", "表情×タッチ反応マトリクス確認", "主要表情とタッチ反応の組み合わせ破綻を確認する", 0.5, 0.9, "組み合わせ確認", "Matrix QA", None),
    ("統合/QA", "Windows Exeビルド確認", "EditorではなくExeでシナリオ/Spine/デバッグを確認する", 0.5, 0.9, "ビルド差分の確認", "Exe QA", None),
    ("統合/QA", "回帰修正/微調整", "追加実装により既存LIME/疑似OS/EDが壊れていないか直す", 0.4, 0.6, "小さい修正の受け皿", "Regression", None),
    ("統合/QA", "引き継ぎメモ/残課題整理", "未確定演出、素材待ち、次工程の注意点をまとめる", 0.3, 0.4, "次工程共有用", "Handoff", None),
]

assert round(sum(row[3] for row in detail_rows), 1) == 23.5
assert round(sum(row[4] for row in detail_rows), 1) == 38.9


OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
wb = load_workbook(INPUT)
if TARGET_SHEET in wb.sheetnames:
    del wb[TARGET_SHEET]

source = wb[SOURCE_SHEET]
ws = wb.copy_worksheet(source)
ws.title = TARGET_SHEET
wb._sheets.remove(ws)
source_index = wb.sheetnames.index(SOURCE_SHEET)
wb._sheets.insert(source_index + 1, ws)

for sheet in wb.worksheets:
    sheet.sheet_view.tabSelected = False
ws.sheet_view.tabSelected = True
wb.active = wb.sheetnames.index(TARGET_SHEET)

labels = {
    "C4": "合計工数 MIN（生数字）",
    "C5": "合計工数 MAX（生数字）",
    "C8": "総合見積工数 MIN",
    "C9": "総合見積工数 MAX",
    "C12": "概算見積額 MIN (バッファ)(税抜)",
    "C13": "概算見積額 MAX (バッファ)(税抜)",
    "C14": "本番割増（生数字）",
    "C17": "想定開発期間 MIN",
    "C18": "想定開発期間 MAX",
    "G21": "MIN工数",
    "H21": "MAX工数",
    "I21": "根拠/補足",
}
for cell, value in labels.items():
    ws[cell] = value

ws["F4"] = "原本タスク＋詳細本番割増MIN"
ws["F5"] = "原本タスク＋詳細本番割増MAX"
ws["D14"] = "MIN 23.5 / MAX 38.9"
ws["E14"] = "人日"
ws["F14"] = "0.1刻みの積み上げ。月営業日20日換算で約1.2〜1.9ヶ月相当"
ws["D19"] = "※最新企画書＋追加要望（妹表情/Spine、技術剪定、タッチ、Exeデバッグ）を詳細化"

for row in range(22, 293):
    g = ws.cell(row=row, column=7).value
    h = ws.cell(row=row, column=8).value
    if as_number(g):
        ws.cell(row=row, column=8).value = g
    elif as_number(h):
        ws.cell(row=row, column=8).value = None

section_fill = PatternFill(fill_type="solid", fgColor="FCE4D6")
task_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
section_font = Font(bold=True)
wrap = Alignment(wrap_text=True, vertical="center")

start_row = 293
current_section = None
for index, (section, task, desc, min_days, max_days, reason, note, comment) in enumerate(detail_rows, start=start_row):
    if section != current_section:
        ws.cell(row=index, column=3).value = section if current_section is None else ws.cell(row=index, column=3).value
    current_section = section
    ws.cell(row=index, column=2).value = f'=IF(ISBLANK(E{index}), "", COUNTIF($E$25:E{index}, "<>"))'
    ws.cell(row=index, column=3).value = section if index == start_row or detail_rows[index - start_row - 1][0] != section else None
    ws.cell(row=index, column=4).value = section
    ws.cell(row=index, column=5).value = task
    ws.cell(row=index, column=6).value = desc
    ws.cell(row=index, column=7).value = min_days
    ws.cell(row=index, column=8).value = max_days
    ws.cell(row=index, column=9).value = reason
    ws.cell(row=index, column=10).value = note
    ws.cell(row=index, column=11).value = comment
    for col in range(3, 12):
        ws.cell(row=index, column=col).fill = task_fill
        ws.cell(row=index, column=col).alignment = wrap
    ws.row_dimensions[index].height = 46
    if ws.cell(row=index, column=3).value:
        ws.cell(row=index, column=3).fill = section_fill
        ws.cell(row=index, column=3).font = section_font

for col_letter, width in {
    "C": 32,
    "D": 30,
    "E": 44,
    "F": 56,
    "G": 11,
    "H": 11,
    "I": 40,
    "J": 26,
    "K": 34,
}.items():
    ws.column_dimensions[col_letter].width = width

ws["D4"] = "=SUM(G:G)"
ws["D5"] = "=SUM(H:H)"
ws["D8"] = "=D4*D6"
ws["D9"] = "=D5*D6"
ws["D12"] = "=D8*D11"
ws["D13"] = "=D9*D11"
ws["D17"] = "=D8/D16/D15"
ws["D18"] = "=D9/D16/D15"

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUTPUT)
print(OUTPUT)
