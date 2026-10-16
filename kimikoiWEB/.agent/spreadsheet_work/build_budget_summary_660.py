from __future__ import annotations

from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB")
SOURCE = ROOT / "outputs" / "kimikoi_kousu" / "きみこい・工数表_ダウンチャート.xlsx"
OUTPUT = ROOT / "outputs" / "kimikoi_kousu" / "きみこい・概算予算表_660万円.xlsx"

BUDGET = 6_600_000
BUFFER_RATE = 1.25
TAX_NOTE = "税別想定"


def as_number(value) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    return 0.0


def short_text(value, limit=70) -> str:
    text = "" if value is None else str(value).strip()
    return text[:limit]


def clean_client_note(value) -> str:
    text = short_text(value, 70)
    replacements = {
        "既存項目との差分のみ追加": "既存実装との差分対応",
        "減額": "対応範囲を明確化",
        "削減": "対応範囲を明確化",
        "圧縮": "見積範囲を整理",
    }
    for before, after in replacements.items():
        text = text.replace(before, after)
    return text.strip()


wb_src = load_workbook(SOURCE, data_only=True)
ws_src = wb_src["本番工数"]

groups: OrderedDict[str, dict] = OrderedDict()
current_major = ""
for row in range(22, ws_src.max_row + 1):
    major_value = ws_src.cell(row, 3).value
    if isinstance(major_value, str) and major_value.strip():
        current_major = major_value.strip()

    min_effort = as_number(ws_src.cell(row, 7).value)
    max_effort = as_number(ws_src.cell(row, 8).value)
    if max_effort <= 0:
        continue

    group = groups.setdefault(
        current_major or "未分類",
        {
            "min": 0.0,
            "max": 0.0,
            "count": 0,
            "tasks": [],
            "notes": [],
        },
    )
    group["min"] += min_effort
    group["max"] += max_effort
    group["count"] += 1
    task = ws_src.cell(row, 5).value or ws_src.cell(row, 4).value
    if task and len(group["tasks"]) < 4:
        group["tasks"].append(str(task).strip())
    note = clean_client_note(ws_src.cell(row, 12).value or ws_src.cell(row, 9).value)
    if note and note not in group["notes"] and len(group["notes"]) < 2:
        group["notes"].append(note)

total_min = round(sum(g["min"] for g in groups.values()), 2)
total_max = round(sum(g["max"] for g in groups.values()), 2)
total_scheduled = round(total_max * BUFFER_RATE, 2)
unit_price = BUDGET / total_scheduled

rows = []
running_amount = 0
for index, (major, group) in enumerate(groups.items(), start=1):
    scheduled = round(group["max"] * BUFFER_RATE, 2)
    share = scheduled / total_scheduled if total_scheduled else 0
    if index < len(groups):
        amount = round(BUDGET * share / 1000) * 1000
        running_amount += amount
    else:
        amount = BUDGET - running_amount
    rows.append(
        {
            "no": index,
            "major": major,
            "scope": " / ".join(group["tasks"]),
            "min": round(group["min"], 1),
            "max": round(group["max"], 1),
            "scheduled": scheduled,
            "share": share,
            "amount": amount,
            "note": " / ".join(group["notes"]) or "本番実装範囲として見積",
            "count": group["count"],
        }
    )

wb = Workbook()
ws = wb.active
ws.title = "概算予算表"
ws.sheet_view.showGridLines = False

dark = "1F4E5F"
dark2 = "173A47"
blue = "D9EAF7"
green = "E2F0D9"
yellow = "FFF2CC"
orange = "FCE4D6"
white = "FFFFFF"
gray = "E7E6E6"
thin = Side(style="thin", color="B7B7B7")
medium = Side(style="medium", color="666666")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
header_border = Border(left=thin, right=thin, top=medium, bottom=medium)

ws.merge_cells("B2:J2")
ws["B2"] = "きみこい 本番開発 概算予算表"
ws["B2"].fill = PatternFill("solid", fgColor=dark)
ws["B2"].font = Font(bold=True, color=white, size=16)
ws["B2"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[2].height = 28

ws.merge_cells("B3:I3")
ws["B3"] = "工数表ベースの項目別集計。提示額はMIN/MAXにかかわらず 6,600,000円 とする前提。"
ws["B3"].fill = PatternFill("solid", fgColor=dark)
ws["B3"].font = Font(color=white, size=10)
ws["B3"].alignment = Alignment(horizontal="left", vertical="center")
ws["J3"] = "作成日"
ws["K3"] = "2026/5/7"
for cell in ("J3", "K3"):
    ws[cell].fill = PatternFill("solid", fgColor=dark)
    ws[cell].font = Font(color=white, size=10, bold=cell == "J3")
    ws[cell].alignment = Alignment(horizontal="center", vertical="center")

summary = [
    ("提示額", BUDGET, "円", TAX_NOTE),
    ("MIN工数", total_min, "人日", "本番工数シート集計"),
    ("MAX工数", total_max, "人日", "本番工数シート集計"),
    ("バッファ込MAX", total_scheduled, "人日", f"MAX x {BUFFER_RATE}"),
    ("逆算人日単価", unit_price, "円/人日", "660万円÷バッファ込MAX"),
]
for offset, (label, value, unit, note) in enumerate(summary):
    col = 2 + offset * 2
    ws.cell(5, col, label)
    ws.cell(5, col + 1, value)
    ws.cell(6, col, unit)
    ws.cell(6, col + 1, note)
    for r in (5, 6):
        for c in (col, col + 1):
            cell = ws.cell(r, c)
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor=dark2 if r == 5 else blue)
            cell.font = Font(bold=(r == 5), color=white if r == 5 else "000000", size=9)
    ws.cell(5, col + 1).number_format = '#,##0'
ws["C5"].number_format = '"¥"#,##0'
ws["J5"].number_format = '"¥"#,##0'

headers = [
    "No",
    "大項目",
    "主な対象タスク",
    "MIN人日",
    "MAX人日",
    "バッファ込MAX",
    "工数比率",
    "予算配分",
    "見積メモ",
    "タスク数",
]
start_row = 9
for col, header in enumerate(headers, start=2):
    cell = ws.cell(start_row, col, header)
    cell.fill = PatternFill("solid", fgColor=dark)
    cell.font = Font(bold=True, color=white)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = header_border

for row_index, row_data in enumerate(rows, start=start_row + 1):
    values = [
        row_data["no"],
        row_data["major"],
        row_data["scope"],
        row_data["min"],
        row_data["max"],
        row_data["scheduled"],
        row_data["share"],
        row_data["amount"],
        row_data["note"],
        row_data["count"],
    ]
    for col, value in enumerate(values, start=2):
        cell = ws.cell(row_index, col, value)
        cell.border = border
        cell.alignment = Alignment(
            horizontal="left" if col in {3, 4, 10} else "center",
            vertical="center",
            wrap_text=col in {4, 10},
        )
        cell.font = Font(size=9)
        if col in {5, 6, 7}:
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.number_format = "0.0"
        elif col == 8:
            cell.fill = PatternFill("solid", fgColor=green)
            cell.number_format = "0.0%"
        elif col == 9:
            cell.fill = PatternFill("solid", fgColor=yellow)
            cell.number_format = '"¥"#,##0'
        elif col == 11:
            cell.fill = PatternFill("solid", fgColor=gray)
    ws.row_dimensions[row_index].height = 36

total_row = start_row + 1 + len(rows)
ws.cell(total_row, 3, "合計")
ws.cell(total_row, 5, total_min)
ws.cell(total_row, 6, total_max)
ws.cell(total_row, 7, total_scheduled)
ws.cell(total_row, 8, 1.0)
ws.cell(total_row, 9, BUDGET)
ws.cell(total_row, 10, "固定概算見積。仕様変更・追加演出・素材追加は別途調整。")
for col in range(2, 12):
    cell = ws.cell(total_row, col)
    cell.fill = PatternFill("solid", fgColor=orange)
    cell.font = Font(bold=True)
    cell.border = header_border
    cell.alignment = Alignment(horizontal="center" if col not in {3, 10} else "left", vertical="center", wrap_text=True)
for col in (5, 6, 7):
    ws.cell(total_row, col).number_format = "0.0"
ws.cell(total_row, 8).number_format = "0.0%"
ws.cell(total_row, 9).number_format = '"¥"#,##0'
ws.row_dimensions[total_row].height = 34

note_row = total_row + 3
ws.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=11)
ws.cell(note_row, 2, "前提: 本資料は概算予算用。工数表の本番実装・統合QA・デバッグ対応を対象とし、仕様追加・演出追加・素材差し替え・検収後の大幅修正は別途見積。")
ws.cell(note_row, 2).fill = PatternFill("solid", fgColor=blue)
ws.cell(note_row, 2).border = border
ws.cell(note_row, 2).alignment = Alignment(wrap_text=True, vertical="center")
ws.cell(note_row, 2).font = Font(size=9)
ws.row_dimensions[note_row].height = 32

chart = BarChart()
chart.title = "項目別 予算配分"
chart.y_axis.title = "円"
chart.x_axis.title = "大項目"
data = Reference(ws, min_col=9, min_row=start_row, max_row=total_row - 1)
cats = Reference(ws, min_col=3, min_row=start_row + 1, max_row=total_row - 1)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.height = 8
chart.width = 18
ws.add_chart(chart, f"B{note_row + 2}")

widths = {
    "A": 3,
    "B": 6,
    "C": 22,
    "D": 48,
    "E": 11,
    "F": 11,
    "G": 14,
    "H": 11,
    "I": 14,
    "J": 44,
    "K": 16,
}
for col, width in widths.items():
    ws.column_dimensions[col].width = width
for row in range(1, total_row + 1):
    if ws.row_dimensions[row].height is None:
        ws.row_dimensions[row].height = 20
ws.row_dimensions[6].height = 34

ws.freeze_panes = "B10"
ws.auto_filter.ref = f"B{start_row}:K{total_row}"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.print_area = f"B2:K{note_row + 20}"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
if OUTPUT.exists():
    OUTPUT.unlink()
wb.save(OUTPUT)

print(OUTPUT)
print("groups", len(rows))
print("min", total_min)
print("max", total_max)
print("scheduled", total_scheduled)
print("budget", BUDGET)
