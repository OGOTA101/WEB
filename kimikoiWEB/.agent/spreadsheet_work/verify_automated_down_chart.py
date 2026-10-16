from __future__ import annotations

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer


WORKBOOK = r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_ダウンチャート_自動化版.xlsx"

wb = load_workbook(WORKBOOK, data_only=False)
ws = wb["ダウンチャート"]

formula_count = 0
bad = []
for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                formula_count += 1
                try:
                    Tokenizer(value).items
                except Exception as exc:
                    bad.append((sheet.title, cell.coordinate, str(exc), value))

banned = ["圧縮", "原本MAX", "本番割増", "割増", "→", "削減"]
banned_counts = {
    word: sum(
        1
        for sheet in wb.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and word in cell.value
    )
    for word in banned
}

print("sheets", wb.sheetnames)
print("summary", [ws.cell(3, col).value for col in range(1, 10)])
print("row6", [ws.cell(6, col).value for col in range(1, 14)])
print("formula_count", formula_count)
print("bad_formula_count", len(bad))
for item in bad[:10]:
    print("bad", item)
print("banned_counts", banned_counts)
print("conditional_ranges", len(ws.conditional_formatting))
print("hidden_helpers", [ws.column_dimensions[col].hidden for col in ["J", "K", "L", "M"]])
