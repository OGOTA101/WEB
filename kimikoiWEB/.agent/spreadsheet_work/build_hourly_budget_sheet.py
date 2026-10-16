from __future__ import annotations

from collections import OrderedDict
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile
import shutil
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB")
SOURCE = ROOT / "outputs" / "kimikoi_kousu" / "きみこい・工数表_ダウンチャート.xlsx"
OUTPUT = ROOT / "outputs" / "kimikoi_kousu" / "きみこい・本番工数_時間単価見積_660万円.xlsx"

TOTAL_AMOUNT = 6_600_000
HOURS_PER_DAY = 8
BUFFER_RATE = 1.25


def as_number(value) -> float:
    return float(value) if isinstance(value, (int, float)) else 0.0


source_wb = load_workbook(SOURCE, data_only=True)
source_ws = source_wb["本番工数"]

groups: OrderedDict[str, dict] = OrderedDict()
current_major = ""

for row in range(22, source_ws.max_row + 1):
    major_value = source_ws.cell(row, 3).value
    mid_value = source_ws.cell(row, 4).value
    min_days = as_number(source_ws.cell(row, 7).value)
    max_days = as_number(source_ws.cell(row, 8).value)

    if isinstance(major_value, str) and major_value.strip():
        current_major = major_value.strip()
    if not current_major or max_days <= 0:
        continue

    if current_major not in groups:
        groups[current_major] = {
            "tasks": 0,
            "min_days": 0.0,
            "max_days": 0.0,
            "mids": OrderedDict(),
        }

    group = groups[current_major]
    group["tasks"] += 1
    group["min_days"] += min_days
    group["max_days"] += max_days
    if isinstance(mid_value, str) and mid_value.strip():
        group["mids"][mid_value.strip()] = True

wb = Workbook()
ws = wb.active
ws.title = "本番工数_時間見積"
ws.sheet_view.showGridLines = False

dark = "1F4E5F"
dark2 = "173A47"
blue = "DDEBF7"
green = "E2F0D9"
yellow = "FFF2CC"
white = "FFFFFF"
gray = "F2F2F2"
thin = Side(style="thin", color="B7B7B7")
medium = Side(style="medium", color="666666")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
strong_border = Border(left=thin, right=thin, top=medium, bottom=medium)

ws.merge_cells("B2:I2")
ws["B2"] = "きみこい 本番工数 時間単価見積"
ws["B2"].fill = PatternFill("solid", fgColor=dark)
ws["B2"].font = Font(bold=True, color=white, size=16)
ws["B2"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[2].height = 30

ws.merge_cells("B3:I3")
ws["B3"] = "大項目ごとの工数時間を集計し、全体工数に対して概算費用 660万円を置いた資料"
ws["B3"].fill = PatternFill("solid", fgColor=dark)
ws["B3"].font = Font(color=white, size=10)
ws["B3"].alignment = Alignment(horizontal="left", vertical="center")

summary_rows = [
    ("概算費用（税抜）", TOTAL_AMOUNT, "円", "固定の概算金額"),
    ("集計対象", "本番工数 MAX（バッファ込み）", "", "デバッグ・統合QA込み"),
    ("見積工数", "", "人日", "大項目別工数の合計"),
    ("1人日あたり時間", HOURS_PER_DAY, "時間", "標準稼働時間"),
    ("見積総時間", "=D7*D8", "時間", "工数を時間へ換算"),
    ("概算時間単価", "=D5/D9", "円/時間", "総額を総時間で除算"),
]

for offset, (label, value, unit, note) in enumerate(summary_rows, start=5):
    ws.cell(offset, 2, label)
    ws.cell(offset, 3, "")
    ws.cell(offset, 4, value)
    ws.cell(offset, 5, unit)
    ws.cell(offset, 6, note)
    for col in range(2, 7):
        cell = ws.cell(offset, col)
        cell.border = border
        cell.alignment = Alignment(horizontal="left" if col in {2, 6} else "center", vertical="center")
        if col == 2:
            cell.fill = PatternFill("solid", fgColor=dark2)
            cell.font = Font(bold=True, color=white)
        elif col == 4:
            cell.fill = PatternFill("solid", fgColor=yellow if offset in {5, 10} else blue)
            cell.font = Font(bold=True)
        else:
            cell.fill = PatternFill("solid", fgColor=gray)

for row in range(5, 11):
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
    ws.cell(row, 2).alignment = Alignment(horizontal="left", vertical="center")

ws["D5"].number_format = "#,##0"
ws["D7"].number_format = "0.0"
ws["D8"].number_format = "0.0"
ws["D9"].number_format = "#,##0.0"
ws["D10"].number_format = "#,##0"

headers = ["No.", "大項目", "タスク数", "MIN人日", "MAX人日", "見積人日", "見積時間", "主な範囲"]
header_row = 12
for col, header in enumerate(headers, start=2):
    cell = ws.cell(header_row, col, header)
    cell.fill = PatternFill("solid", fgColor=dark)
    cell.font = Font(bold=True, color=white)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = strong_border

start_row = header_row + 1
for idx, (major, data) in enumerate(groups.items(), start=1):
    row = start_row + idx - 1
    scheduled_days = data["max_days"] * BUFFER_RATE
    mids = list(data["mids"].keys())
    scope = " / ".join(mids[:5])
    if len(mids) > 5:
        scope += " ほか"
    values = [
        idx,
        major,
        data["tasks"],
        data["min_days"],
        data["max_days"],
        scheduled_days,
        scheduled_days * HOURS_PER_DAY,
        scope,
    ]
    for col, value in enumerate(values, start=2):
        cell = ws.cell(row, col, value)
        cell.border = border
        cell.alignment = Alignment(
            horizontal="left" if col in {3, 9} else "center",
            vertical="center",
            wrap_text=col == 9,
        )
        cell.font = Font(size=10)
        if col in {6, 7, 8}:
            cell.fill = PatternFill("solid", fgColor=green)
        elif row % 2 == 0:
            cell.fill = PatternFill("solid", fgColor="F8FBFD")
    ws.row_dimensions[row].height = 34

total_row = start_row + len(groups)
total_tasks = sum(data["tasks"] for data in groups.values())
total_min = sum(data["min_days"] for data in groups.values())
total_max = sum(data["max_days"] for data in groups.values())
total_days = sum(data["max_days"] * BUFFER_RATE for data in groups.values())
total_hours = total_days * HOURS_PER_DAY
hourly_rate = TOTAL_AMOUNT / total_hours

ws["D5"] = f"{TOTAL_AMOUNT:,}"
ws["D7"] = total_days
ws["D9"] = total_hours
ws["D10"] = round(hourly_rate)
ws.cell(total_row, 2, "合計")
ws.cell(total_row, 3, "")
ws.cell(total_row, 4, total_tasks)
ws.cell(total_row, 5, total_min)
ws.cell(total_row, 6, total_max)
ws.cell(total_row, 7, total_days)
ws.cell(total_row, 8, total_hours)
ws.cell(total_row, 9, "金額は項目別配分せず、総時間から一括算出")
for col in range(2, 10):
    cell = ws.cell(total_row, col)
    cell.fill = PatternFill("solid", fgColor=dark2)
    cell.font = Font(bold=True, color=white)
    cell.border = strong_border
    cell.alignment = Alignment(horizontal="center" if col != 9 else "left", vertical="center")

note_row = total_row + 3
ws.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=9)
ws.cell(note_row, 2, "前提: 本資料は項目別の費用配分表ではなく、実装工数を時間換算し、総額660万円の根拠を示すための概算資料です。")
ws.cell(note_row, 2).fill = PatternFill("solid", fgColor=yellow)
ws.cell(note_row, 2).border = border
ws.cell(note_row, 2).alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[note_row].height = 34

for row in range(start_row, total_row + 1):
    ws.cell(row, 5).number_format = "0.0"
    ws.cell(row, 6).number_format = "0.0"
    ws.cell(row, 7).number_format = "0.0"
    ws.cell(row, 8).number_format = "#,##0.0"
for row in range(5, 11):
    ws.row_dimensions[row].height = 22
ws.row_dimensions[12].height = 24

widths = {
    "B": 16,
    "C": 24,
    "D": 16,
    "E": 11,
    "F": 11,
    "G": 12,
    "H": 13,
    "I": 56,
}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

chart = BarChart()
chart.title = "大項目別 見積時間"
chart.y_axis.title = "時間"
chart.x_axis.title = "大項目"
chart.height = 8
chart.width = 20
data = Reference(ws, min_col=8, min_row=12, max_row=total_row - 1)
cats = Reference(ws, min_col=3, min_row=start_row, max_row=total_row - 1)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.legend = None
ws.add_chart(chart, "K5")

ws.freeze_panes = "B13"
ws.auto_filter.ref = f"B12:I{total_row}"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0

for col in range(2, 10):
    ws.cell(4, col).border = Border(bottom=medium)

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)

ns_uri = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ns = {"main": ns_uri}
ET.register_namespace("", ns_uri)

with TemporaryDirectory() as tmp:
    tmp_root = Path(tmp)
    with ZipFile(OUTPUT, "r") as zin:
        zin.extractall(tmp_root)

    workbook_tree = ET.parse(tmp_root / "xl" / "workbook.xml")
    sheet_el = workbook_tree.getroot().find(".//main:sheet[@name='本番工数_時間見積']", ns)
    rel_id = sheet_el.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
    rels_tree = ET.parse(tmp_root / "xl" / "_rels" / "workbook.xml.rels")
    target = None
    for rel in rels_tree.getroot():
        if rel.attrib.get("Id") == rel_id:
            target = rel.attrib["Target"]
            break
    if target.startswith("/"):
        sheet_path = tmp_root / target.lstrip("/")
    elif target.startswith("worksheets/"):
        sheet_path = tmp_root / "xl" / target
    else:
        sheet_path = tmp_root / "xl" / target
    tree = ET.parse(sheet_path)
    root = tree.getroot()
    cells = {cell.attrib["r"]: cell for cell in root.findall(".//main:c", ns) if "r" in cell.attrib}
    caches = {
        "D7": total_days,
        "D9": total_hours,
        "D10": round(hourly_rate),
        f"D{total_row}": sum(data["tasks"] for data in groups.values()),
        f"E{total_row}": total_min,
        f"F{total_row}": total_max,
        f"G{total_row}": total_days,
        f"H{total_row}": total_hours,
    }
    for ref, value in caches.items():
        cell = cells.get(ref)
        if cell is None:
            continue
        for value_el in cell.findall("main:v", ns):
            cell.remove(value_el)
        v = ET.SubElement(cell, f"{{{ns_uri}}}v")
        cell.attrib.pop("t", None)
        v.text = str(round(value, 8))
    tree.write(sheet_path, encoding="utf-8", xml_declaration=True)

    patched = OUTPUT.with_suffix(".patched.xlsx")
    if patched.exists():
        patched.unlink()
    with ZipFile(patched, "w", ZIP_DEFLATED) as zout:
        for file in tmp_root.rglob("*"):
            if file.is_file():
                zout.write(file, file.relative_to(tmp_root).as_posix())
    shutil.move(patched, OUTPUT)

print(OUTPUT)
print("groups", len(groups))
print("scheduled_days", round(total_days, 2))
print("hours", round(total_hours, 1))
