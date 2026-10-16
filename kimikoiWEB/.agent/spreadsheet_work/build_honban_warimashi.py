from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.styles import PatternFill


INPUT = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\きみこい・工数表.xlsx")
OUTPUT_DIR = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu")
OUTPUT = OUTPUT_DIR / "きみこい・工数表_本番割増_MIN20_MAX40_最終.xlsx"
SOURCE_SHEET = "工数表ベース原本"
TARGET_SHEET = "本番割増_MIN20_MAX40"


def as_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
wb = load_workbook(INPUT)

if TARGET_SHEET in wb.sheetnames:
    del wb[TARGET_SHEET]

source = wb[SOURCE_SHEET]
ws = wb.copy_worksheet(source)
ws.title = TARGET_SHEET
wb._sheets.remove(ws)
source_index = wb.sheetnames.index(SOURCE_SHEET)
wb._sheets.insert(source_index + 1, ws)

ws.sheet_view.tabSelected = True
wb.active = wb.sheetnames.index(TARGET_SHEET)
for sheet in wb.worksheets:
    if sheet.title != TARGET_SHEET:
        sheet.sheet_view.tabSelected = False

# Repurpose the copied sheet as a MIN/MAX production estimate.
summary_labels = {
    "C4": "合計工数 MIN（生数字）",
    "C5": "合計工数 MAX（生数字）",
    "C8": "総合見積工数 MIN",
    "C9": "総合見積工数 MAX",
    "C12": "概算見積額 MIN (バッファ)(税抜)",
    "C13": "概算見積額 MAX (バッファ)(税抜)",
    "C17": "想定開発期間 MIN",
    "C18": "想定開発期間 MAX",
    "G21": "MIN工数",
    "H21": "MAX工数",
    "I21": "根拠/補足",
}
for cell, value in summary_labels.items():
    ws[cell] = value

ws["F4"] = "原本タスク＋本番割増MIN"
ws["F5"] = "原本タスク＋本番割増MAX"
ws["F8"] = "既存リスク係数を含む安全値"
ws["F9"] = "既存リスク係数を含む安全値"
ws["C14"] = "本番割増（生数字）"
ws["D14"] = "MIN 20 / MAX 40"
ws["E14"] = "人日"
ws["F14"] = "月営業日20日換算で1〜2ヶ月相当"
ws["D19"] = "※本番割増は最新企画書（Sweet Sister Nightmare）の追加範囲を反映"

# Existing tasks become the common baseline. The H column was AI effort in the source,
# but this scenario sheet uses it as MAX effort.
for row in range(22, 293):
    g = ws.cell(row=row, column=7).value
    h = ws.cell(row=row, column=8).value
    if as_number(g):
        ws.cell(row=row, column=8).value = g
    elif as_number(h):
        ws.cell(row=row, column=8).value = None

production_rows = [
    {
        "row": 293,
        "c": "08. 本番割増（最新企画）",
        "h": "-",
        "k": "PDF【最新】企画書.pdfをもとに、本番化で増える調整・実装・QAを上乗せ。",
    },
    {
        "row": 294,
        "d": "本番スコープ整理",
        "e": "企画書ベースの仕様棚卸し・受入条件整理",
        "f": "美少女ホラーADV、10日間進行、疑似デスクトップ、LIME、複数EDの本番範囲を確定する",
        "g": 2.0,
        "h": 4.0,
        "i": "不確定要素の洗い出し、実装順序、除外条件の確認",
        "j": "Sweet Sister Nightmare",
    },
    {
        "row": 295,
        "d": "パラメータ体験",
        "e": "ニケ風パラメータ調整・失敗時死亡演出の本番調整",
        "f": "かわいい反応とホラー反転が成立するバランスへ調整する",
        "g": 3.0,
        "h": 6.0,
        "i": "選択肢、反応演出、即死・進行不能条件の調整幅",
        "j": "好感度/イライラ度/ときめき度/しんらい度",
    },
    {
        "row": 296,
        "d": "シナリオ/ED",
        "e": "10日間進行・ルート分岐・多彩なエンディングの接続",
        "f": "10日目到達、通常/隠し/グッド/バッド系EDの発火確認を行う",
        "g": 4.0,
        "h": 8.0,
        "i": "条件組み合わせ、回想導線、分岐漏れ確認",
        "j": "企画書p16-p18相当",
    },
    {
        "row": 297,
        "d": "疑似デスクトップ演出",
        "e": "デスクトップ画面・LIME・スキンシップボードの本番統合",
        "f": "通常UIに異変を混ぜ、怖いのにかわいい体験として見せる",
        "g": 4.0,
        "h": 8.0,
        "i": "複数アプリ表示、異変割り込み、演出崩れの確認",
        "j": "画面イメージ/ゲームコンセプト反映",
    },
    {
        "row": 298,
        "d": "ホラー演出品質",
        "e": "グリッチ・ストーカー/見え方変化・画像演出の仕上げ",
        "f": "だんだん可愛く見える変化と、恐怖の鮮度を両立させる",
        "g": 3.0,
        "h": 6.0,
        "i": "見た目調整と素材差し替えでブレやすい",
        "j": "怖いのにカワイイ",
    },
    {
        "row": 299,
        "d": "本番QA/リリース調整",
        "e": "通しプレイ・クラッシュ/セーブ/Steam導線確認",
        "f": "本番配布前に、1周目からED到達までの破綻を潰す",
        "g": 4.0,
        "h": 8.0,
        "i": "分岐、セーブ、実績、ビルド差分の総合確認",
        "j": "本番用割増",
    },
]

for item in production_rows:
    row = item["row"]
    mapping = {
        "c": 3,
        "d": 4,
        "e": 5,
        "f": 6,
        "g": 7,
        "h": 8,
        "i": 9,
        "j": 10,
        "k": 11,
    }
    for key, col in mapping.items():
        if key in item:
            ws.cell(row=row, column=col).value = item[key]

# Make the added block visually stand out while keeping the source sheet style.
section_fill = PatternFill(fill_type="solid", fgColor="FCE4D6")
task_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
for col in range(3, 12):
    ws.cell(row=293, column=col).fill = section_fill
for row in range(294, 300):
    for col in range(4, 11):
        ws.cell(row=row, column=col).fill = task_fill

for col_letter, width in {
    "C": 28,
    "D": 24,
    "E": 48,
    "F": 62,
    "G": 11,
    "H": 11,
    "I": 42,
    "J": 36,
    "K": 42,
}.items():
    ws.column_dimensions[col_letter].width = max(ws.column_dimensions[col_letter].width or 0, width)

wrap_alignment = Alignment(wrap_text=True, vertical="center")
for row in range(293, 300):
    ws.row_dimensions[row].height = 38 if row == 293 else 58
    for col in range(3, 12):
        ws.cell(row=row, column=col).alignment = wrap_alignment

# Ensure formulas use the scenario meaning.
ws["D4"] = "=SUM(G:G)"
ws["D5"] = "=SUM(H:H)"
ws["D8"] = "=D4*D6"
ws["D9"] = "=D5*D6"
ws["D12"] = "=D8*D11"
ws["D13"] = "=D9*D11"
ws["D17"] = "=D8/D16/D15"
ws["D18"] = "=D9/D16/D15"

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUTPUT)
print(OUTPUT)
