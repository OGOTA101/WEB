from __future__ import annotations

from copy import copy
from datetime import date, timedelta
import os
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SOURCE = Path(os.environ.get("KIMIKOI_GANTT_SOURCE", r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_本番割増_再調整_単独シート.xlsx"))
OUTPUT = Path(os.environ.get("KIMIKOI_GANTT_OUTPUT", r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_本番割増_横長工程表.xlsx"))
BASE_SHEET = os.environ.get("KIMIKOI_GANTT_BASE_SHEET", "本番割増_再調整")
SHEET = "横長工程表"

START_DATE = date.fromisoformat(os.environ.get("KIMIKOI_GANTT_START", "2026-06-01"))
DAILY_CAPACITY = 1.0
BUFFER_RATE = 1.25


HOLIDAYS = {
    date(2026, 1, 1),
    date(2026, 1, 12),
    date(2026, 2, 11),
    date(2026, 2, 23),
    date(2026, 3, 20),
    date(2026, 4, 29),
    date(2026, 5, 3),
    date(2026, 5, 4),
    date(2026, 5, 5),
    date(2026, 5, 6),
    date(2026, 7, 20),
    date(2026, 8, 11),
    date(2026, 9, 21),
    date(2026, 9, 22),
    date(2026, 9, 23),
    date(2026, 10, 12),
    date(2026, 11, 3),
    date(2026, 11, 23),
    date(2027, 1, 1),
    date(2027, 1, 11),
    date(2027, 2, 11),
    date(2027, 2, 23),
    date(2027, 3, 21),
    date(2027, 3, 22),
    date(2027, 4, 29),
    date(2027, 5, 3),
    date(2027, 5, 4),
    date(2027, 5, 5),
    date(2027, 7, 19),
    date(2027, 8, 11),
    date(2027, 9, 20),
    date(2027, 9, 23),
    date(2027, 10, 11),
    date(2027, 11, 3),
    date(2027, 11, 23),
}


MAJOR_COLORS = [
    "9DC3E6",
    "A9D18E",
    "F4B183",
    "FFD966",
    "B4C6E7",
    "C6E0B4",
    "D9EAD3",
    "F8CBAD",
    "D6DCE4",
    "EADCF8",
    "FCE4D6",
    "DDEBF7",
    "E2F0D9",
]


def as_number(value) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value)
        except ValueError:
            return None
    return None


def is_workday(day: date) -> bool:
    return day.weekday() < 5 and day not in HOLIDAYS


def extract_tasks(ws):
    major = ""
    mid = ""
    tasks = []
    for row in range(22, ws.max_row + 1):
        major_value = ws.cell(row, 3).value
        mid_value = ws.cell(row, 4).value
        detail_value = ws.cell(row, 5).value
        description = ws.cell(row, 6).value
        min_effort = as_number(ws.cell(row, 7).value)
        max_effort = as_number(ws.cell(row, 8).value)
        basis = ws.cell(row, 9).value
        note = ws.cell(row, 10).value
        comment = ws.cell(row, 11).value
        reason = ws.cell(row, 12).value

        if isinstance(major_value, str) and major_value.strip():
            major = major_value.strip()
        if isinstance(mid_value, str) and mid_value.strip():
            mid = mid_value.strip()

        if max_effort is None or max_effort <= 0:
            continue

        title = detail_value if isinstance(detail_value, str) and detail_value.strip() else mid
        tasks.append(
            {
                "source_row": row,
                "major": major,
                "mid": mid,
                "title": str(title).strip(),
                "description": description or "",
                "min": min_effort or 0,
                "max": max_effort,
                "scheduled": round(max_effort * BUFFER_RATE, 3),
                "basis": basis or "",
                "note": note or "",
                "comment": comment or "",
                "reason": reason or "",
                "allocations": {},
            }
        )
    return tasks


def schedule_tasks(tasks):
    current = START_DATE
    remaining_capacity = DAILY_CAPACITY
    for task in tasks:
        remaining_effort = task["scheduled"]
        while remaining_effort > 0.0001:
            if not is_workday(current):
                current += timedelta(days=1)
                remaining_capacity = DAILY_CAPACITY
                continue
            used = min(remaining_effort, remaining_capacity)
            task["allocations"][current] = round(used, 3)
            remaining_effort = round(remaining_effort - used, 3)
            remaining_capacity = round(remaining_capacity - used, 3)
            if remaining_capacity <= 0.0001:
                current += timedelta(days=1)
                remaining_capacity = DAILY_CAPACITY

    last_day = max(day for task in tasks for day in task["allocations"])
    dates = []
    day = START_DATE
    while day <= last_day:
        dates.append(day)
        day += timedelta(days=1)
    return dates, last_day


def month_ranges(dates):
    ranges = []
    start_index = 0
    current_key = (dates[0].year, dates[0].month)
    for index, day in enumerate(dates):
        key = (day.year, day.month)
        if key != current_key:
            ranges.append((start_index, index - 1, current_key))
            start_index = index
            current_key = key
    ranges.append((start_index, len(dates) - 1, current_key))
    return ranges


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb = load_workbook(SOURCE)
if SHEET in wb.sheetnames:
    del wb[SHEET]
ws_base = wb[BASE_SHEET]
tasks = extract_tasks(ws_base)
dates, end_date = schedule_tasks(tasks)

ws = wb.create_sheet(SHEET, index=1)
ws.sheet_view.showGridLines = False
ws.freeze_panes = "J6"

for sheet in wb.worksheets:
    sheet.sheet_view.tabSelected = False
ws.sheet_view.tabSelected = True
wb.active = wb.sheetnames.index(SHEET)

dark = "1F4E5F"
dark2 = "173A47"
white = "FFFFFF"
light_blue = "D9EAF7"
light_yellow = "FFF2CC"
light_green = "E2F0D9"
weekend_fill = PatternFill("solid", fgColor="E7E6E6")
holiday_fill = PatternFill("solid", fgColor="D9D2E9")
thin = Side(style="thin", color="B7B7B7")
hair = Side(style="hair", color="D9D9D9")
medium = Side(style="medium", color="666666")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
date_border = Border(left=hair, right=hair, top=thin, bottom=thin)

title_fill = PatternFill("solid", fgColor=dark)
header_fill = PatternFill("solid", fgColor=dark)
subheader_fill = PatternFill("solid", fgColor=dark2)
linked_fill = PatternFill("solid", fgColor=light_blue)
assumption_fill = PatternFill("solid", fgColor=light_yellow)
summary_fill = PatternFill("solid", fgColor=light_green)

ws.merge_cells("A1:I1")
ws["A1"] = "きみこい 本番割増 横長工程表"
ws["A1"].fill = title_fill
ws["A1"].font = Font(bold=True, color=white, size=16)
ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 28

summary = [
    ("開始日", START_DATE),
    ("終了目安", end_date),
    ("タスク数", len(tasks)),
    ("原本MAX合計", round(sum(t["max"] for t in tasks), 2)),
    ("バッファ係数", BUFFER_RATE),
    ("予定工数合計", round(sum(t["scheduled"] for t in tasks), 2)),
    ("日次消化力", DAILY_CAPACITY),
    ("前提", "土日祝休み"),
]
for col, (label, value) in enumerate(summary, start=1):
    cell = ws.cell(2, col, label)
    cell.fill = subheader_fill
    cell.font = Font(bold=True, color=white, size=9)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border
    value_cell = ws.cell(3, col, value)
    value_cell.fill = linked_fill if col in {2, 3, 4, 6} else assumption_fill
    value_cell.font = Font(bold=True, size=9)
    value_cell.alignment = Alignment(horizontal="center", vertical="center")
    value_cell.border = border
for cell in ("A3", "B3"):
    ws[cell].number_format = "yyyy/m/d"
for cell in ("D3", "E3", "F3", "G3"):
    ws[cell].number_format = "0.0"

left_headers = ["No", "大項目", "中項目", "タスク", "原本MAX", "予定工数", "開始日", "終了日", "根拠/備考"]
for col, header in enumerate(left_headers, start=1):
    cell = ws.cell(5, col, header)
    cell.fill = header_fill
    cell.font = Font(bold=True, color=white)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(left=thin, right=thin, top=medium, bottom=medium)

widths = {"A": 6, "B": 18, "C": 20, "D": 44, "E": 10, "F": 10, "G": 12, "H": 12, "I": 34}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

date_start_col = 10
for start_idx, end_idx, (year, month) in month_ranges(dates):
    start_col = date_start_col + start_idx
    end_col = date_start_col + end_idx
    ws.merge_cells(start_row=3, start_column=start_col, end_row=3, end_column=end_col)
    cell = ws.cell(3, start_col, f"{year}/{month}")
    cell.fill = header_fill
    cell.font = Font(bold=True, color=white, size=9)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for col in range(start_col, end_col + 1):
        ws.cell(3, col).border = border

weekday_names = ["月", "火", "水", "木", "金", "土", "日"]
for index, day in enumerate(dates):
    col = date_start_col + index
    col_letter = get_column_letter(col)
    ws.column_dimensions[col_letter].width = 3

    is_holiday = day in HOLIDAYS
    non_work = not is_workday(day)
    fill = holiday_fill if is_holiday else (weekend_fill if non_work else PatternFill("solid", fgColor="FFFFFF"))

    day_cell = ws.cell(4, col, day.day)
    day_cell.fill = fill
    day_cell.font = Font(size=8, bold=is_holiday or day.weekday() >= 5)
    day_cell.alignment = Alignment(horizontal="center", vertical="center")
    day_cell.border = date_border

    week_cell = ws.cell(5, col, weekday_names[day.weekday()])
    week_cell.fill = fill
    week_cell.font = Font(size=8, color="C00000" if non_work else "000000")
    week_cell.alignment = Alignment(horizontal="center", vertical="center")
    week_cell.border = date_border

color_by_major = {}
for task in tasks:
    if task["major"] not in color_by_major:
        color_by_major[task["major"]] = MAJOR_COLORS[len(color_by_major) % len(MAJOR_COLORS)]

first_task_row = 6
for row_offset, task in enumerate(tasks):
    row = first_task_row + row_offset
    color = color_by_major[task["major"]]
    bar_fill = PatternFill("solid", fgColor=color)
    major_fill = PatternFill("solid", fgColor=color)
    first_day = min(task["allocations"])
    last_day = max(task["allocations"])

    left_values = [
        row_offset + 1,
        task["major"],
        task["mid"],
        task["title"],
        task["max"],
        task["scheduled"],
        first_day,
        last_day,
        " / ".join(str(x) for x in [task["basis"], task["note"], task["reason"]] if x),
    ]
    for col, value in enumerate(left_values, start=1):
        cell = ws.cell(row, col, value)
        cell.border = border
        cell.alignment = Alignment(horizontal="left" if col in {2, 3, 4, 9} else "center", vertical="center", wrap_text=col in {4, 9})
        cell.font = Font(size=8)
        if col == 2:
            cell.fill = major_fill
        elif col in {5, 6}:
            cell.fill = linked_fill
            cell.number_format = "0.0"
        elif col in {7, 8}:
            cell.number_format = "yyyy/m/d"
    ws.row_dimensions[row].height = 18

    for index, day in enumerate(dates):
        col = date_start_col + index
        cell = ws.cell(row, col)
        cell.border = Border(left=hair, right=hair, top=hair, bottom=hair)
        if day in task["allocations"]:
            cell.fill = bar_fill
        elif day in HOLIDAYS:
            cell.fill = holiday_fill
        elif day.weekday() >= 5:
            cell.fill = weekend_fill

last_task_row = first_task_row + len(tasks) - 1
last_date_col = date_start_col + len(dates) - 1

for col in range(1, last_date_col + 1):
    ws.cell(last_task_row + 2, col).border = Border(top=medium)

legend_row = last_task_row + 3
ws.cell(legend_row, 1, "凡例")
ws.cell(legend_row, 1).font = Font(bold=True)
legend_items = [("予定作業", MAJOR_COLORS[0]), ("土日", "E7E6E6"), ("祝日", "D9D2E9")]
for i, (label, color) in enumerate(legend_items, start=2):
    ws.cell(legend_row, i, label)
    ws.cell(legend_row, i).fill = PatternFill("solid", fgColor=color)
    ws.cell(legend_row, i).border = border
    ws.cell(legend_row, i).alignment = Alignment(horizontal="center")

ws.auto_filter.ref = f"A5:I{last_task_row}"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0

for row in range(1, 6):
    ws.row_dimensions[row].height = 20
ws.row_dimensions[3].height = 18
ws.row_dimensions[4].height = 18
ws.row_dimensions[5].height = 22

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUTPUT)

print(OUTPUT)
print("tasks", len(tasks))
print("date_columns", len(dates))
print("raw_max", round(sum(t["max"] for t in tasks), 2))
print("scheduled", round(sum(t["scheduled"] for t in tasks), 2))
print("end", end_date.isoformat())
