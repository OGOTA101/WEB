from __future__ import annotations

from pathlib import Path
import os
import re
import subprocess

from openpyxl import load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side


ROOT = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB")
PYTHON = Path(r"C:\Users\ihogo\.cache\codex-runtimes\codex-primary-runtime\dependencies\python\python.exe")
WORK = ROOT / ".agent" / "spreadsheet_work"
SOURCE = WORK / "feb_compressed_source.xlsx"
CLEAN_SOURCE = WORK / "down_chart_clean_source.xlsx"
OUTPUT = ROOT / "outputs" / "kimikoi_kousu" / "きみこい・工数表_ダウンチャート.xlsx"


def clean_note(value):
    if value is None:
        return None
    text = str(value)
    if "2月末圧縮" in text:
        # Drop the numeric before/after wording and keep only a client-facing premise.
        parts = re.split(r"2月末圧縮:\s*[^。]*。", text)
        text = " / ".join(part.strip(" /") for part in parts if part.strip(" /"))
    replacements = {
        "エディタ/定型作業のため軽量化": "定型作業として見積",
        "定型・短時間作業として軽量化": "短時間作業として見積",
        "軽作業として0.1人日単位へ圧縮": "軽作業として見積",
        "初回は日本語本線に絞り多言語対応を後送": "初回実装は日本語本線を前提",
        "PCマウス前提へ寄せ、操作拡張を後送": "PCマウス操作前提",
        "固定/限定サイズ運用へ寄せて圧縮": "固定/限定サイズ運用を前提",
        "演出調整は最低限に限定して圧縮": "演出調整は本番要件範囲を前提",
        "周辺アプリはシナリオ必須範囲に絞って圧縮": "シナリオ必須範囲を前提",
        "2月末版では後送/最小接続にして圧縮": "初回版では最小接続を前提",
        "疑似OS中核は簡略化しつつ保持": "疑似OS中核として実装枠を確保",
        "進行/ADV中核は重さを残して圧縮": "進行/ADV中核として実装枠を確保",
        "新仕様の中核なので削りすぎず圧縮": "新仕様の中核として実装枠を確保",
        "デバッグ込み完了のため枠を保持": "デバッグ/統合確認込み",
        "共通化・AI補助・実装範囲固定で圧縮": "共通化・実装範囲固定を前提",
    }
    for before, after in replacements.items():
        text = text.replace(before, after)
    text = re.sub(r"\s*/\s*/\s*", " / ", text).strip(" /")
    banned = ["圧縮", "原本MAX", "原本", "→", "削減"]
    for word in banned:
        text = text.replace(word, "")
    return text.strip(" /") or None


if not SOURCE.exists():
    subprocess.run([str(PYTHON), str(WORK / "compress_feb_plan.py")], check=True, cwd=str(ROOT))

wb = load_workbook(SOURCE)
ws = wb["本番割増_再調整"]
ws.title = "本番工数"

ws["C14"] = "本番工数サマリ"
ws["D14"] = "MIN 152.3 / MAX 179.8"
ws["E14"] = "人日"
ws["F14"] = "本番想定工数。デバッグ・統合QA込み"
ws["D19"] = "※右端L列に見積前提・備考を記載"
ws["F4"] = "本番想定工数"
ws["F5"] = "本番想定工数"
ws["L21"] = "見積前提"

note_fill = PatternFill("solid", fgColor="DDEBF7")
green_fill = PatternFill("solid", fgColor="E2F0D9")
thin = Side(style="thin", color="7F7F7F")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for row in range(22, ws.max_row + 1):
    ws.cell(row, 12, clean_note(ws.cell(row, 12).value))
    if ws.cell(row, 12).value:
        ws.cell(row, 12).fill = note_fill
        ws.cell(row, 12).font = Font(size=9)
    for col in range(2, 13):
        ws.cell(row, col).border = border
    if isinstance(ws.cell(row, 7).value, (int, float)) or isinstance(ws.cell(row, 8).value, (int, float)):
        ws.cell(row, 7).fill = green_fill
        ws.cell(row, 8).fill = green_fill

for row in range(1, ws.max_row + 1):
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row, col).value
        if isinstance(value, str):
            cleaned = clean_note(value)
            if cleaned != value:
                ws.cell(row, col, cleaned)

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
CLEAN_SOURCE.parent.mkdir(parents=True, exist_ok=True)
wb.save(CLEAN_SOURCE)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
if OUTPUT.exists():
    OUTPUT.unlink()

env = os.environ.copy()
env["KIMIKOI_GANTT_SOURCE"] = str(CLEAN_SOURCE)
env["KIMIKOI_GANTT_OUTPUT"] = str(OUTPUT)
env["KIMIKOI_GANTT_BASE_SHEET"] = "本番工数"
subprocess.run([str(PYTHON), str(WORK / "build_gantt_schedule.py")], check=True, cwd=str(ROOT), env=env)

wb = load_workbook(OUTPUT)
if "本番工数" not in wb.sheetnames and "本番割増_再調整" in wb.sheetnames:
    wb["本番割増_再調整"].title = "本番工数"
if "横長工程表" in wb.sheetnames:
    wb["横長工程表"].title = "ダウンチャート"

ws = wb["ダウンチャート"]
ws["A1"] = "きみこい 本番工数 ダウンチャート"
ws["D2"] = "MAX工数合計"
ws["E2"] = "バッファ係数"
ws["F2"] = "予定工数合計"
ws["I5"] = "根拠/備考"
for row in range(1, ws.max_row + 1):
    for col in range(1, min(ws.max_column, 12) + 1):
        value = ws.cell(row, col).value
        if isinstance(value, str):
            ws.cell(row, col, clean_note(value))

for sheet in wb.worksheets:
    sheet.sheet_view.tabSelected = False
wb["本番工数"].sheet_view.tabSelected = True
wb.active = wb.sheetnames.index("本番工数")

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUTPUT)

print(OUTPUT)
print(wb.sheetnames)
