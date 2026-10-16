from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB")
SOURCE = ROOT / "outputs" / "kimikoi_kousu" / "きみこい・工数表_ダウンチャート.xlsx"
OUTPUT = ROOT / "outputs" / "kimikoi_kousu" / "kimikoi_downchart_google_sheets_minimal.xlsx"

DATE_DAYS = 330
DATE_START_COL = 10
TASK_START_ROW = 6

HOLIDAYS = [
    date(2026, 1, 1), date(2026, 1, 12), date(2026, 2, 11), date(2026, 2, 23),
    date(2026, 3, 20), date(2026, 4, 29), date(2026, 5, 3), date(2026, 5, 4),
    date(2026, 5, 5), date(2026, 5, 6), date(2026, 7, 20), date(2026, 8, 11),
    date(2026, 9, 21), date(2026, 9, 22), date(2026, 9, 23), date(2026, 10, 12),
    date(2026, 11, 3), date(2026, 11, 23), date(2027, 1, 1), date(2027, 1, 11),
    date(2027, 2, 11), date(2027, 2, 23), date(2027, 3, 21), date(2027, 3, 22),
    date(2027, 4, 29), date(2027, 5, 3), date(2027, 5, 4), date(2027, 5, 5),
]


def as_number(value) -> float:
    return float(value) if isinstance(value, (int, float)) else 0.0


src_wb = load_workbook(SOURCE, data_only=True)
src = src_wb["本番工数"]

tasks = []
major = ""
mid = ""
for row in range(22, src.max_row + 1):
    major_value = src.cell(row, 3).value
    mid_value = src.cell(row, 4).value
    task_value = src.cell(row, 5).value
    desc_value = src.cell(row, 6).value
    min_effort = as_number(src.cell(row, 7).value)
    max_effort = as_number(src.cell(row, 8).value)
    basis = src.cell(row, 9).value or ""
    premise = src.cell(row, 12).value or ""

    if isinstance(major_value, str) and major_value.strip():
        major = major_value.strip()
    if isinstance(mid_value, str) and mid_value.strip():
        mid = mid_value.strip()
    if max_effort <= 0:
        continue

    title = task_value if isinstance(task_value, str) and task_value.strip() else mid
    note = " / ".join(str(x) for x in [basis, premise] if x)
    tasks.append([len(tasks) + 1, major, mid, title, desc_value or "", min_effort, max_effort, note])

wb = Workbook()
settings = wb.active
settings.title = "Settings"
task_ws = wb.create_sheet("Tasks")
chart = wb.create_sheet("DownChart")

dark = "1F4E5F"
dark2 = "173A47"
white = "FFFFFF"
input_fill = "FFF2CC"
linked_fill = "DDEBF7"
bar_color = "2F75B5"
thin = Side(style="thin", color="B7B7B7")
hair = Side(style="hair", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hair_border = Border(left=hair, right=hair, top=hair, bottom=hair)

# Settings sheet: visible and simple.
settings.sheet_view.showGridLines = False
settings["A1"] = "Kimikoi DownChart Settings"
settings["A1"].font = Font(bold=True, size=15, color=white)
settings["A1"].fill = PatternFill("solid", fgColor=dark)
settings.merge_cells("A1:D1")
settings_values = [
    ("Start Date", date(2026, 6, 1), "Change this date to shift the schedule"),
    ("Buffer Rate", 1.25, "Scheduled days = MAX days x buffer"),
    ("Daily Capacity", 1.0, "Person-days burned per workday"),
    ("Holiday Mode", "土日祝休み", "土日祝休み / 土日休み / 休みなし"),
]
for row, (label, value, note) in enumerate(settings_values, start=3):
    settings.cell(row, 1, label)
    settings.cell(row, 2, value)
    settings.cell(row, 3, note)
    for col in range(1, 4):
        cell = settings.cell(row, col)
        cell.border = border
        cell.alignment = Alignment(vertical="center")
    settings.cell(row, 1).fill = PatternFill("solid", fgColor=dark2)
    settings.cell(row, 1).font = Font(bold=True, color=white)
    settings.cell(row, 2).fill = PatternFill("solid", fgColor=input_fill)
settings["B3"].number_format = "yyyy/m/d"
settings["B4"].number_format = "0.00"
settings["B5"].number_format = "0.00"

settings["E1"] = "Holidays"
settings["E1"].font = Font(bold=True)
for idx, holiday in enumerate(HOLIDAYS, start=2):
    settings.cell(idx, 5, holiday)
    settings.cell(idx, 5).number_format = "yyyy/m/d"

for col, width in {"A": 18, "B": 16, "C": 48, "E": 14}.items():
    settings.column_dimensions[col].width = width

# Tasks sheet: plain values, easy for Google Sheets to import/edit.
task_headers = ["No", "Major", "Middle", "Task", "Description", "MIN Days", "MAX Days", "Basis / Premise"]
for col, header in enumerate(task_headers, start=1):
    cell = task_ws.cell(1, col, header)
    cell.fill = PatternFill("solid", fgColor=dark)
    cell.font = Font(bold=True, color=white)
    cell.border = border
    cell.alignment = Alignment(horizontal="center")
for row_idx, row_values in enumerate(tasks, start=2):
    for col, value in enumerate(row_values, start=1):
        cell = task_ws.cell(row_idx, col, value)
        cell.border = border
        cell.alignment = Alignment(wrap_text=col in {4, 5, 8}, vertical="center")
        if col in {6, 7}:
            cell.number_format = "0.0"
for col, width in {"A": 6, "B": 20, "C": 22, "D": 42, "E": 46, "F": 10, "G": 10, "H": 46}.items():
    task_ws.column_dimensions[col].width = width
task_ws.freeze_panes = "A2"
task_ws.auto_filter.ref = f"A1:H{len(tasks)+1}"

# DownChart sheet.
chart.sheet_view.showGridLines = False
chart.freeze_panes = f"{get_column_letter(DATE_START_COL)}{TASK_START_ROW}"
chart["A1"] = "Kimikoi DownChart for Google Sheets"
chart["A1"].fill = PatternFill("solid", fgColor=dark)
chart["A1"].font = Font(bold=True, color=white, size=15)
chart.merge_cells("A1:I1")

summary = [
    ("Start Date", "=Settings!B3"),
    ("End Date", f'=IFERROR(MAX($H${TASK_START_ROW}:$H${TASK_START_ROW + len(tasks) - 1}),"")'),
    ("Task Count", f'=COUNTIF(Tasks!$G$2:$G${len(tasks)+1},">0")'),
    ("MAX Days", f"=SUM(Tasks!$G$2:$G${len(tasks)+1})"),
    ("Buffer", "=Settings!B4"),
    ("Scheduled Days", "=D3*E3"),
    ("Daily Capacity", "=Settings!B5"),
    ("Holiday Mode", "=Settings!B6"),
]
for col, (label, value) in enumerate(summary, start=1):
    chart.cell(2, col, label)
    chart.cell(2, col).fill = PatternFill("solid", fgColor=dark2)
    chart.cell(2, col).font = Font(bold=True, color=white)
    chart.cell(2, col).border = border
    chart.cell(2, col).alignment = Alignment(horizontal="center")
    chart.cell(3, col, value)
    chart.cell(3, col).fill = PatternFill("solid", fgColor=linked_fill)
    chart.cell(3, col).font = Font(bold=True)
    chart.cell(3, col).border = border
    chart.cell(3, col).alignment = Alignment(horizontal="center")
chart["A3"].number_format = "yyyy/m/d"
chart["B3"].number_format = "yyyy/m/d"
for ref in ["D3", "E3", "F3", "G3"]:
    chart[ref].number_format = "0.0"

chart_headers = ["No", "Major", "Middle", "Task", "MAX", "Scheduled", "Start", "End", "Basis"]
for col, header in enumerate(chart_headers, start=1):
    cell = chart.cell(5, col, header)
    cell.fill = PatternFill("solid", fgColor=dark)
    cell.font = Font(bold=True, color=white)
    cell.border = border
    cell.alignment = Alignment(horizontal="center")
for col, width in {"A": 6, "B": 18, "C": 20, "D": 42, "E": 9, "F": 10, "G": 12, "H": 12, "I": 34}.items():
    chart.column_dimensions[col].width = width

for idx in range(1, len(tasks) + 1):
    task_row = idx + 1
    row = TASK_START_ROW + idx - 1
    chart.cell(row, 1, f'=IF(Tasks!G{task_row}>0,Tasks!A{task_row},"")')
    chart.cell(row, 2, f'=IF($A{row}="","",Tasks!B{task_row})')
    chart.cell(row, 3, f'=IF($A{row}="","",Tasks!C{task_row})')
    chart.cell(row, 4, f'=IF($A{row}="","",Tasks!D{task_row})')
    chart.cell(row, 5, f'=IF($A{row}="","",Tasks!G{task_row})')
    chart.cell(row, 6, f'=IF($A{row}="","",$E{row}*$E$3)')
    previous_sum = f"SUM($F${TASK_START_ROW}:F{row - 1})" if row > TASK_START_ROW else "0"
    current_sum = f"SUM($F${TASK_START_ROW}:F{row})"
    start_index = f"INT(({previous_sum})/$G$3)+1"
    end_index = f"ROUNDUP(({current_sum})/$G$3,0)"
    chart.cell(row, 7, f'=IF($A{row}="","",IF($H$3="休みなし",$A$3+{start_index}-1,IF($H$3="土日祝休み",WORKDAY.INTL($A$3-1,{start_index},"0000011",Settings!$E$2:$E$80),WORKDAY.INTL($A$3-1,{start_index},"0000011"))))')
    chart.cell(row, 8, f'=IF($A{row}="","",IF($H$3="休みなし",$A$3+{end_index}-1,IF($H$3="土日祝休み",WORKDAY.INTL($A$3-1,{end_index},"0000011",Settings!$E$2:$E$80),WORKDAY.INTL($A$3-1,{end_index},"0000011"))))')
    chart.cell(row, 9, f'=IF($A{row}="","",Tasks!H{task_row})')
    for col in range(1, 10):
        cell = chart.cell(row, col)
        cell.border = border
        cell.font = Font(size=8)
        cell.alignment = Alignment(horizontal="left" if col in {2, 3, 4, 9} else "center", vertical="center", wrap_text=col in {4, 9})
        if col in {5, 6}:
            cell.number_format = "0.0"
            cell.fill = PatternFill("solid", fgColor=linked_fill)
        if col in {7, 8}:
            cell.number_format = "yyyy/m/d"

last_task_row = TASK_START_ROW + len(tasks) - 1
for date_idx in range(DATE_DAYS):
    col = DATE_START_COL + date_idx
    letter = get_column_letter(col)
    chart.column_dimensions[letter].width = 3
    chart.cell(3, col, f'=IF({letter}$4="","",IF(DAY({letter}$4)=1,TEXT({letter}$4,"yyyy/m"),""))')
    chart.cell(4, col, f'=IF(OR($A$3="",$B$3="",$A$3+{date_idx}>$B$3),"",$A$3+{date_idx})')
    chart.cell(5, col, f'=IF({letter}$4="","",CHOOSE(WEEKDAY({letter}$4,2),"月","火","水","木","金","土","日"))')
    for r in (3, 4, 5):
        chart.cell(r, col).border = hair_border
        chart.cell(r, col).alignment = Alignment(horizontal="center")
        chart.cell(r, col).font = Font(size=8)
    chart.cell(4, col).number_format = "d"

for row in range(TASK_START_ROW, last_task_row + 1):
    for col in range(DATE_START_COL, DATE_START_COL + DATE_DAYS):
        letter = get_column_letter(col)
        cell = chart.cell(row, col)
        cell.value = f'=IF(AND({letter}$4<>"",$A{row}<>"",{letter}$4>=$G{row},{letter}$4<=$H{row},OR($H$3="休みなし",WEEKDAY({letter}$4,2)<=5),OR($H$3<>"土日祝休み",COUNTIF(Settings!$E$2:$E$80,{letter}$4)=0)),"■","")'
        cell.font = Font(size=8, color=bar_color, bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = hair_border

for sheet in wb.worksheets:
    sheet.sheet_view.tabSelected = False
chart.sheet_view.tabSelected = True
wb.active = wb.sheetnames.index("DownChart")

# Keep workbook intentionally simple for Google Sheets import.
wb.calculation.calcMode = "auto"
wb.calculation.fullCalcOnLoad = False
wb.calculation.forceFullCalc = False

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)
print(OUTPUT)
print("tasks", len(tasks))
print("rows", chart.max_row, "cols", chart.max_column)
