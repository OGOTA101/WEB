from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


INPUT = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_本番割増_詳細版.xlsx")
OUTPUT = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_本番割増_詳細調整版.xlsx")
SHEET = "本番割増_詳細"

rows = [
    ("実装影響整理", "企画PDF差分の実装影響メモ", "企画書から実装に影響する差分だけを短く洗い出す", 0.1, 0.2, "仕様整理そのものは加算しないため最小限", "差分確認", "仕様作業寄りのため減額"),
    ("実装影響整理", "受入条件リスト作成", "完成条件の文章化は基本的に仕様側で持つ", 0.0, 0.0, "実装作業ではないため今回の増分から除外", "除外", "工数0"),
    ("実装影響整理", "10日進行データ構造差分確認", "10日制とパラメータ分岐が既存データ構造で吸収できるか確認する", 0.2, 0.4, "実装影響の確認に限定", "Data Model", None),
    ("実装影響整理", "既存工数表との重複確認", "既存項目に含まれる作業の重複を避ける確認", 0.0, 0.0, "見積整理作業であり実装増分ではない", "除外", "工数0"),
    ("実装影響整理", "演出リスクの技術剪定", "実装危険度の高い演出をプロトタイプ対象/見送り対象に分ける", 0.4, 0.8, "技術判断が必要なため残す", "Tech Triage", None),
    ("妹 表情差分/Spine", "表情差分データ定義/インポート設定", "表情IDをSpine/静止画切替テーブルとして実装側で扱える形にする", 0.6, 0.9, "仕様リストではなくデータ定義と取り込み設定", "Expression Data", None),
    ("妹 表情差分/Spine", "表情IDマスタ/Enum生成", "シナリオ指定と実装参照がズレないよう表情IDをコード側に固定する", 0.3, 0.5, "参照ミス防止のため", "Enum/Master", None),
    ("妹 表情差分/Spine", "Spineスロット/アタッチメント設計確認", "目、眉、口、頬、影などの差分パーツをRuntimeで切り替え可能か確認する", 0.7, 1.2, "素材構造に依存して調整が増える", "Spine構成", None),
    ("妹 表情差分/Spine", "表情ブレンド制御", "笑顔から不穏、怯えから狂気など中間表情を制御する", 0.9, 1.6, "かわいい/怖いの境界調整が難しい", "Expression Blend", None),
    ("妹 表情差分/Spine", "瞬き/呼吸/待機モーション実装", "常時表示時に破綻しない瞬き、呼吸、待機モーションを再生する", 0.6, 1.0, "SpineタイムラインとUnity制御が絡む", "Idle Motion", None),
    ("妹 表情差分/Spine", "恐怖進行パラメータ連動", "日数やパラメータに応じて表情、影、ノイズ感を段階的に切り替える", 0.7, 1.2, "演出条件と表示制御の接続", "State Hook", None),
    ("妹 表情差分/Spine", "差分素材差し替えパイプライン", "差分素材更新時にAtlas/参照を壊さず差し替えられる導線を作る", 0.4, 0.8, "素材戻し対応を効率化", "Asset Pipeline", None),
    ("妹 表情差分/Spine", "Unity上のSpine表示統合", "既存UI/疑似デスクトップ上にSpineキャラを安定表示する", 0.8, 1.3, "Canvas/Render順の調整が必要", "Unity Spine", None),
    ("妹 表情差分/Spine", "静止画フォールバック対応", "Spine未対応箇所や軽量表示用に静止画差分でも出せるようにする", 0.3, 0.5, "安全策として最低限用意", "Fallback", None),
    ("妹 表情差分/Spine", "Atlas/メモリ使用量チェック", "表情差分増加によるテクスチャ容量とロード負荷を確認する", 0.3, 0.6, "ビルド後の負荷確認が必要", "Performance", None),
    ("妹 表情差分/Spine", "シナリオイベントからの表情呼び出し", "台詞やイベントに合わせてSpine表情を切り替えるAPIを作る", 0.6, 1.0, "データ形式との接続が必要", "Scenario Hook", None),
    ("技術剪定/演出検証", "グリッチ演出プロトタイプ", "画面ノイズ、色収差、ズレ表現の成立ラインを試す", 0.6, 1.0, "本実装前の見た目検証", "Shader検証", None),
    ("技術剪定/演出検証", "疑似デスクトップ同時表示検証", "LIME、ボード、妹表示を同時に出した時のCanvas破綻を確認する", 0.4, 0.7, "既存疑似OS項目と重複するため検証に限定", "Desktop検証", "既存項目分を減額"),
    ("技術剪定/演出検証", "偽カーソルOS依存検証", "ビルド状態でカーソル演出が安全に見せられるか試す", 0.4, 0.8, "OS依存リスクの確認", "Input検証", None),
    ("技術剪定/演出検証", "動画/RenderTexture検証", "実写風・異物感演出をゲーム画面内に自然に混ぜる", 0.3, 0.6, "既存動画連携項目があるため検証差分のみ", "VideoPlayer", "既存項目分を減額"),
    ("技術剪定/演出検証", "セーブ/ロード中の演出割り込み安全確認", "異変演出とセーブデータ更新が衝突しないか確認する", 0.5, 0.9, "進行不能リスクが高い", "Save Safety", None),
    ("技術剪定/演出検証", "負荷予算確認", "演出を盛った時のFPS/メモリ目安を確認する", 0.4, 0.7, "製品版での最低限の安定性", "Performance", None),
    ("技術剪定/演出検証", "見送り候補の技術判断メモ", "重すぎる/危ない演出を本番から落とす判断材料を残す", 0.1, 0.2, "整理作業寄りのため最小限", "Cut List", "減額"),
    ("技術剪定/演出検証", "検証結果メモ整理", "採用/保留/廃止の理由整理は仕様共有作業として扱う", 0.0, 0.0, "実装増分ではないため除外", "除外", "工数0"),
    ("タッチアクション", "妹ちゃんタッチ判定エリア実装", "頭、顔、肩、手元など反応箇所のHitAreaを実装する", 0.7, 1.1, "トリッカル系の触り心地を再現", "Touch Area", None),
    ("タッチアクション", "タッチ反応データテーブル接続", "場所/日数/パラメータ別の反応台詞と表情を実装データへ接続する", 0.4, 0.8, "シナリオと表情の接続が必要", "Reaction Table", None),
    ("タッチアクション", "パラメータ変動連携", "タッチ結果で好感度、イライラ度、恐怖値等が変わるようにする", 0.5, 0.9, "既存パラメータ管理との接続", "Parameter Hook", None),
    ("タッチアクション", "Spineリアクション再生", "タッチ箇所に応じて表情/モーション/揺れを再生する", 0.5, 0.9, "素材と演出調整が絡む", "Spine Reaction", None),
    ("タッチアクション", "連打/長押し/禁止操作の分岐", "押し方によって反応が変わる処理を追加する", 0.4, 0.7, "嫌がる/照れる/壊れる反応差分", "Input Pattern", None),
    ("タッチアクション", "クールタイム/暴発防止", "連続入力で演出が重ならないよう制御する", 0.3, 0.5, "演出破綻防止", "Cooldown", None),
    ("タッチアクション", "ボイス/SE連携", "タッチ時のSE、短い声、ノイズ音を同期する", 0.3, 0.6, "音量/タイミング調整込み", "Audio Hook", None),
    ("タッチアクション", "初回チュートリアル/誘導表示", "触れることの誘導表示はUI/仕様寄りのため今回の増分から外す", 0.0, 0.0, "実装差分としては優先度低", "除外", "工数0"),
    ("タッチアクション", "タッチ不能状態ガード", "イベント中や暗転中はタッチできないようにする", 0.3, 0.5, "進行不能防止", "Guard", None),
    ("Exeシナリオ読込/デバッグ", "ビルド用シナリオパス解決", "StreamingAssets等に置いたシナリオをExeから読む構成にする", 0.6, 1.1, "Editor専用パスを避ける", "Build Path", None),
    ("Exeシナリオ読込/デバッグ", "文字コード/ファイルパス検証", "Windowsビルドで日本語ファイル名や改行が壊れないか確認する", 0.4, 0.7, "環境差分の確認", "Encoding", None),
    ("Exeシナリオ読込/デバッグ", "シナリオ再読込デバッグ機能", "ビルド状態でもシナリオを再読込して確認できる導線を作る", 0.7, 1.3, "調整速度を上げるため", "Hot Reload", None),
    ("Exeシナリオ読込/デバッグ", "シナリオ構文チェック", "ID欠け、表情ID不一致、ジャンプ先欠けを検出する", 0.8, 1.4, "ビルド後の原因特定を早くする", "Validator", None),
    ("Exeシナリオ読込/デバッグ", "日付/イベントジャンプ機能", "10日間進行の任意地点へビルド中に飛べるようにする", 0.5, 0.9, "既存デバッグ項目と重複するため差分のみ", "Scene Jump", "既存項目分を減額"),
    ("Exeシナリオ読込/デバッグ", "パラメータ編集パネル", "好感度や恐怖値などをビルド中に変更して分岐確認する", 0.4, 0.8, "既存デバッグ項目と重複するため差分のみ", "Param Debug", "既存項目分を減額"),
    ("Exeシナリオ読込/デバッグ", "フラグビューア", "現在ONのイベント/ED/回想フラグを一覧表示する", 0.4, 0.8, "分岐不具合の調査用", "Flag Viewer", None),
    ("Exeシナリオ読込/デバッグ", "EDルート簡易シミュレータ", "現在値から到達しうるEDを確認する簡易機能を作る", 0.4, 0.8, "複数EDの確認負荷を下げるが既存ED判定と重複するため減額", "ED Debug", "減額"),
    ("Exeシナリオ読込/デバッグ", "ログ出力/コピー機能", "ビルド中の進行ログを外部へ出せるようにする", 0.3, 0.6, "不具合報告をしやすくする", "Log Export", None),
    ("Exeシナリオ読込/デバッグ", "デバッグメニュー表示制御", "製品版で誤ってデバッグ機能が出ないようにする", 0.3, 0.5, "ビルド設定との連動", "Debug Gate", None),
    ("Exeシナリオ読込/デバッグ", "セーブデータ初期化/差し替え", "ビルド中にテスト用セーブへ戻せるようにする", 0.2, 0.4, "既存デバッグ項目と重複するため差分のみ", "Save Reset", "既存項目分を減額"),
    ("Exeシナリオ読込/デバッグ", "読込失敗時のエラー表示", "シナリオ欠損や壊れたデータを画面/ログで分かるようにする", 0.5, 0.8, "原因不明の停止を防ぐ", "Error Handling", None),
    ("統合/QA", "追加差分スモークテスト", "表情/タッチ/Exe読込の追加差分だけを重点的に通す", 0.5, 0.9, "全体QAではなく追加仕様差分に限定", "Smoke Test", "減額"),
    ("統合/QA", "表情×タッチ反応マトリクス確認", "主要表情とタッチ反応の組み合わせ破綻を確認する", 0.5, 0.9, "組み合わせ確認", "Matrix QA", None),
    ("統合/QA", "Windows Exe差分確認", "EditorではなくExeでシナリオ/Spine/デバッグ差分を確認する", 0.4, 0.8, "ビルド差分の確認", "Exe QA", None),
    ("統合/QA", "追加差分の回帰修正", "追加実装により既存LIME/疑似OS/EDが壊れていないか直す", 0.3, 0.6, "既存全体回帰ではなく追加差分のみ", "Regression", "減額"),
    ("統合/QA", "引き継ぎメモ/残課題整理", "残課題整理は共有作業として扱う", 0.0, 0.0, "実装増分ではないため除外", "除外", "工数0"),
]

assert round(sum(r[3] for r in rows), 1) == 20.2
assert round(sum(r[4] for r in rows), 1) == 36.2

wb = load_workbook(INPUT)
ws = wb[SHEET]

ws["D14"] = "MIN 20.2 / MAX 36.2"
ws["F14"] = "仕様寄り/既存重複を減額。追加仕様から増える実装差分のみ"
ws["D19"] = "※仕様整理や既存表と重複する作業は0または減額し、実装/接続/検証差分へ調整"

section_fill = PatternFill(fill_type="solid", fgColor="FCE4D6")
task_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
zero_fill = PatternFill(fill_type="solid", fgColor="E7E6E6")
section_font = Font(bold=True)
wrap = Alignment(wrap_text=True, vertical="center")

start = 293
for offset, row_data in enumerate(rows):
    row = start + offset
    section, task, desc, min_days, max_days, reason, note, comment = row_data
    previous_section = rows[offset - 1][0] if offset else None
    ws.cell(row=row, column=2).value = f'=IF(ISBLANK(E{row}), "", COUNTIF($E$25:E{row}, "<>"))'
    ws.cell(row=row, column=3).value = section if section != previous_section else None
    ws.cell(row=row, column=4).value = section
    ws.cell(row=row, column=5).value = task
    ws.cell(row=row, column=6).value = desc
    ws.cell(row=row, column=7).value = min_days
    ws.cell(row=row, column=8).value = max_days
    ws.cell(row=row, column=9).value = reason
    ws.cell(row=row, column=10).value = note
    ws.cell(row=row, column=11).value = comment
    fill = zero_fill if min_days == 0 and max_days == 0 else task_fill
    for col in range(3, 12):
        ws.cell(row=row, column=col).fill = fill
        ws.cell(row=row, column=col).alignment = wrap
    if ws.cell(row=row, column=3).value:
        ws.cell(row=row, column=3).fill = section_fill
        ws.cell(row=row, column=3).font = section_font
    ws.row_dimensions[row].height = 46

for col_letter, width in {
    "C": 32,
    "D": 30,
    "E": 44,
    "F": 58,
    "G": 11,
    "H": 11,
    "I": 42,
    "J": 26,
    "K": 34,
}.items():
    ws.column_dimensions[col_letter].width = width

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUTPUT)
print(OUTPUT)
