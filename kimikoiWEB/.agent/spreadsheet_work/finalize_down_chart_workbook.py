from pathlib import Path

from openpyxl import load_workbook


INPUT = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_ダウンチャート.xlsx")
OUTPUT = INPUT

wb = load_workbook(INPUT)

if "本番割増_再調整" in wb.sheetnames:
    wb["本番割増_再調整"].title = "本番工数"
if "横長工程表" in wb.sheetnames:
    wb["横長工程表"].title = "ダウンチャート"

if "ダウンチャート" in wb.sheetnames:
    ws = wb["ダウンチャート"]
    ws["A1"] = "きみこい 本番工数 ダウンチャート"
    ws["I5"] = "根拠/備考"

for sheet in wb.worksheets:
    sheet.sheet_view.tabSelected = False

if "本番工数" in wb.sheetnames:
    wb["本番工数"].sheet_view.tabSelected = True
    wb.active = wb.sheetnames.index("本番工数")

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUTPUT)

print(OUTPUT)
print(wb.sheetnames)
