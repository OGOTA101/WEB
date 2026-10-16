from __future__ import annotations

from datetime import date, timedelta
from math import ceil
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile
import shutil
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel


WORKBOOK = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_ダウンチャート_自動化完成版.xlsx")
SHEET = "ダウンチャート"
START_DATE = date(2026, 6, 1)
BUFFER_RATE = 1.25
DAILY_CAPACITY = 1.0
DATE_DAYS = 730
DATE_START_COL = 15
TASK_START_ROW = 6

HOLIDAYS = {
    date(2026, 1, 1),
    date(2026, 1, 12),
    date(2026, 2, 11),
    date(2026, 2, 23),
    date(2026, 3, 20),
    date(2026, 4, 29),
    date(2026, 5, 3),
    date(2026, 5, 4),
    date(2026, 5, 5),
    date(2026, 5, 6),
    date(2026, 7, 20),
    date(2026, 8, 11),
    date(2026, 9, 21),
    date(2026, 9, 22),
    date(2026, 9, 23),
    date(2026, 10, 12),
    date(2026, 11, 3),
    date(2026, 11, 23),
    date(2027, 1, 1),
    date(2027, 1, 11),
    date(2027, 2, 11),
    date(2027, 2, 23),
}


def is_workday(day: date) -> bool:
    return day.weekday() < 5 and day not in HOLIDAYS


def nth_workday(start: date, index: int) -> date:
    count = 0
    day = start
    while True:
        if is_workday(day):
            count += 1
            if count == index:
                return day
        day += timedelta(days=1)


def excel_date(value: date) -> float:
    return float(to_excel(value))


def col_letter(col: int) -> str:
    result = ""
    while col:
        col, rem = divmod(col - 1, 26)
        result = chr(65 + rem) + result
    return result


def sheet_xml_path(root: Path, sheet_name: str) -> Path:
    ns = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pkgrel": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    workbook_tree = ET.parse(root / "xl" / "workbook.xml")
    sheet_el = None
    for candidate in workbook_tree.getroot().findall(".//main:sheet", ns):
        if candidate.attrib.get("name") == sheet_name:
            sheet_el = candidate
            break
    if sheet_el is None:
        raise ValueError(sheet_name)
    rel_id = sheet_el.attrib[f"{{{ns['rel']}}}id"]
    rels_tree = ET.parse(root / "xl" / "_rels" / "workbook.xml.rels")
    for rel in rels_tree.getroot().findall("pkgrel:Relationship", ns):
        if rel.attrib.get("Id") == rel_id:
            target = rel.attrib["Target"]
            if target.startswith("/"):
                return root / target.lstrip("/")
            if target.startswith("worksheets/"):
                return root / "xl" / target
            return root / "xl" / target
    raise ValueError(rel_id)


source_wb = load_workbook(WORKBOOK, data_only=True)
source_ws = source_wb["本番工数"]

major = ""
mid = ""
tasks = []
for row in range(22, 1501):
    major_value = source_ws.cell(row, 3).value
    mid_value = source_ws.cell(row, 4).value
    detail_value = source_ws.cell(row, 5).value
    max_effort = source_ws.cell(row, 8).value
    basis = source_ws.cell(row, 9).value

    if isinstance(major_value, str) and major_value.strip():
        major = major_value.strip()
    if isinstance(mid_value, str) and mid_value.strip():
        mid = mid_value.strip()
    if not isinstance(max_effort, (int, float)) or max_effort <= 0:
        continue

    title = detail_value if isinstance(detail_value, str) and detail_value.strip() else mid
    scheduled = round(float(max_effort) * BUFFER_RATE, 10)
    tasks.append(
        {
            "source_row": row,
            "major": major,
            "mid": mid,
            "title": str(title),
            "max": float(max_effort),
            "scheduled": scheduled,
            "basis": basis or "",
        }
    )

cum = 0.0
for task in tasks:
    start_index = int(cum / DAILY_CAPACITY) + 1
    cum += task["scheduled"]
    end_index = ceil(cum / DAILY_CAPACITY)
    task["cum_before"] = round(cum - task["scheduled"], 10)
    task["cum_end"] = round(cum, 10)
    task["start_index"] = start_index
    task["end_index"] = end_index
    task["start_date"] = nth_workday(START_DATE, start_index)
    task["end_date"] = nth_workday(START_DATE, end_index)

end_date = max(task["end_date"] for task in tasks)
raw_max = sum(task["max"] for task in tasks)
scheduled_total = sum(task["scheduled"] for task in tasks)

caches: dict[str, object] = {
    "A3": excel_date(START_DATE),
    "B3": excel_date(end_date),
    "C3": len(tasks),
    "D3": round(raw_max, 10),
    "F3": round(scheduled_total, 10),
}

weekday_names = ["月", "火", "水", "木", "金", "土", "日"]
for offset in range(DATE_DAYS):
    current = START_DATE + timedelta(days=offset)
    col = col_letter(DATE_START_COL + offset)
    if current <= end_date:
        caches[f"{col}3"] = f"{current.year}/{current.month}" if current.day == 1 else ""
        caches[f"{col}4"] = excel_date(current)
        caches[f"{col}5"] = weekday_names[current.weekday()]
    else:
        caches[f"{col}3"] = ""
        caches[f"{col}4"] = ""
        caches[f"{col}5"] = ""

for idx, task in enumerate(tasks, start=0):
    row = TASK_START_ROW + idx
    caches[f"A{row}"] = idx + 1
    caches[f"B{row}"] = task["major"]
    caches[f"C{row}"] = task["mid"]
    caches[f"D{row}"] = task["title"]
    caches[f"E{row}"] = round(task["max"], 10)
    caches[f"F{row}"] = round(task["scheduled"], 10)
    caches[f"G{row}"] = excel_date(task["start_date"])
    caches[f"H{row}"] = excel_date(task["end_date"])
    caches[f"I{row}"] = str(task["basis"])
    caches[f"J{row}"] = task["source_row"]
    caches[f"K{row}"] = task["cum_before"]
    caches[f"L{row}"] = task["cum_end"]
    caches[f"M{row}"] = task["start_index"]
    caches[f"N{row}"] = task["end_index"]

ns_uri = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ns = {"main": ns_uri}
ET.register_namespace("", ns_uri)

with TemporaryDirectory() as tmp:
    tmp_root = Path(tmp)
    with ZipFile(WORKBOOK, "r") as zin:
        zin.extractall(tmp_root)

    xml_path = sheet_xml_path(tmp_root, SHEET)
    tree = ET.parse(xml_path)
    root = tree.getroot()
    cells = {cell.attrib["r"]: cell for cell in root.findall(".//main:c", ns) if "r" in cell.attrib}

    for ref, value in caches.items():
        cell = cells.get(ref)
        if cell is None:
            continue
        for value_el in cell.findall("main:v", ns):
            cell.remove(value_el)
        if value == "":
            cell.attrib.pop("t", None)
            continue
        v = ET.SubElement(cell, f"{{{ns_uri}}}v")
        if isinstance(value, str):
            cell.attrib["t"] = "str"
            v.text = value
        else:
            cell.attrib.pop("t", None)
            v.text = str(value)

    tree.write(xml_path, encoding="utf-8", xml_declaration=True)

    patched = WORKBOOK.with_suffix(".patched.xlsx")
    if patched.exists():
        patched.unlink()
    with ZipFile(patched, "w", ZIP_DEFLATED) as zout:
        for file in tmp_root.rglob("*"):
            if file.is_file():
                zout.write(file, file.relative_to(tmp_root).as_posix())
    shutil.move(patched, WORKBOOK)

print(WORKBOOK)
print("tasks", len(tasks))
print("raw_max", round(raw_max, 2))
print("scheduled", round(scheduled_total, 2))
print("end", end_date.isoformat())
