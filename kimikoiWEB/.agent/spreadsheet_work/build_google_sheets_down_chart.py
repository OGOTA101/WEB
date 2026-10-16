from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB")
SOURCE = ROOT / "outputs" / "kimikoi_kousu" / "きみこい・工数表_ダウンチャート.xlsx"
OUTPUT = ROOT / "outputs" / "kimikoi_kousu" / "きみこい・工数表_ダウンチャート_Googleスプレッドシート対応版.xlsx"

TASK_START_ROW = 6
DATE_START_COL = 15  # O
DATE_DAYS = 340

HOLIDAYS = [
    date(2026, 1, 1), date(2026, 1, 12), date(2026, 2, 11), date(2026, 2, 23),
    date(2026, 3, 20), date(2026, 4, 29), date(2026, 5, 3), date(2026, 5, 4),
    date(2026, 5, 5), date(2026, 5, 6), date(2026, 7, 20), date(2026, 8, 11),
    date(2026, 9, 21), date(2026, 9, 22), date(2026, 9, 23), date(2026, 10, 12),
    date(2026, 11, 3), date(2026, 11, 23), date(2027, 1, 1), date(2027, 1, 11),
    date(2027, 2, 11), date(2027, 2, 23), date(2027, 3, 21), date(2027, 3, 22),
    date(2027, 4, 29), date(2027, 5, 3), date(2027, 5, 4), date(2027, 5, 5),
    date(2027, 7, 19), date(2027, 8, 11), date(2027, 9, 20), date(2027, 9, 23),
    date(2027, 10, 11), date(2027, 11, 3), date(2027, 11, 23),
]


def q(name: str) -> str:
    return f"'{name}'"


def as_number(value) -> float:
    return float(value) if isinstance(value, (int, float)) else 0.0


wb = load_workbook(SOURCE)
for sheet_name in ["ダウンチャート", "タスク計算", "祝日設定"]:
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]

base = wb["本番工数"]

tasks = []
major = ""
mid = ""
for row in range(22, base.max_row + 1):
    major_value = base.cell(row, 3).value
    mid_value = base.cell(row, 4).value
    detail_value = base.cell(row, 5).value
    max_effort = as_number(base.cell(row, 8).value)
    basis = base.cell(row, 9).value or ""
    if isinstance(major_value, str) and major_value.strip():
        major = major_value.strip()
    if isinstance(mid_value, str) and mid_value.strip():
        mid = mid_value.strip()
    if max_effort <= 0:
        continue
    title = detail_value if isinstance(detail_value, str) and detail_value.strip() else mid
    tasks.append((row, major, mid, title, basis))

ws = wb.create_sheet("ダウンチャート", index=1)
calc = wb.create_sheet("タスク計算", index=2)
holiday_ws = wb.create_sheet("祝日設定", index=3)

holiday_ws["A1"] = "祝日"
for idx, holiday in enumerate(HOLIDAYS, start=2):
    holiday_ws.cell(idx, 1, holiday)
    holiday_ws.cell(idx, 1).number_format = "yyyy/m/d"
holiday_ws.sheet_state = "hidden"
calc.sheet_state = "hidden"

dark = "1F4E5F"
dark2 = "173A47"
white = "FFFFFF"
input_fill = "FFF2CC"
linked_fill = "DDEBF7"
summary_fill = "E2F0D9"
bar_color = "2F75B5"
thin = Side(style="thin", color="B7B7B7")
hair = Side(style="hair", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hair_border = Border(left=hair, right=hair, top=hair, bottom=hair)

# Calculation sheet: simple formulas only, with current task source rows fixed.
calc_headers = ["No", "元行", "大項目", "中項目", "タスク", "根拠/備考", "MAX工数", "予定工数", "累計前", "累計後", "開始Index", "終了Index", "開始日", "終了日"]
for col, header in enumerate(calc_headers, start=1):
    calc.cell(1, col, header)

for idx, (source_row, major, mid, title, basis) in enumerate(tasks, start=2):
    calc.cell(idx, 1, idx - 1)
    calc.cell(idx, 2, source_row)
    calc.cell(idx, 3, major)
    calc.cell(idx, 4, mid)
    calc.cell(idx, 5, f'=IF(INDEX({q("本番工数")}!$E:$E,$B{idx})<>"",INDEX({q("本番工数")}!$E:$E,$B{idx}),D{idx})')
    calc.cell(idx, 6, basis)
    calc.cell(idx, 7, f'=INDEX({q("本番工数")}!$H:$H,$B{idx})')
    calc.cell(idx, 8, f'=IF(G{idx}>0,G{idx}*{q("ダウンチャート")}!$E$3,0)')
    calc.cell(idx, 9, f'=SUM($H$2:H{idx - 1})')
    calc.cell(idx, 10, f'=SUM($H$2:H{idx})')
    calc.cell(idx, 11, f'=IF(H{idx}<=0,"",INT(I{idx}/{q("ダウンチャート")}!$G$3)+1)')
    calc.cell(idx, 12, f'=IF(H{idx}<=0,"",ROUNDUP(J{idx}/{q("ダウンチャート")}!$G$3,0))')
    calc.cell(idx, 13, f'=IF(K{idx}="","",IF({q("ダウンチャート")}!$H$3="休みなし",{q("ダウンチャート")}!$A$3+K{idx}-1,IF({q("ダウンチャート")}!$H$3="土日祝休み",WORKDAY.INTL({q("ダウンチャート")}!$A$3-1,K{idx},"0000011",{q("祝日設定")}!$A$2:$A$80),WORKDAY.INTL({q("ダウンチャート")}!$A$3-1,K{idx},"0000011"))))')
    calc.cell(idx, 14, f'=IF(L{idx}="","",IF({q("ダウンチャート")}!$H$3="休みなし",{q("ダウンチャート")}!$A$3+L{idx}-1,IF({q("ダウンチャート")}!$H$3="土日祝休み",WORKDAY.INTL({q("ダウンチャート")}!$A$3-1,L{idx},"0000011",{q("祝日設定")}!$A$2:$A$80),WORKDAY.INTL({q("ダウンチャート")}!$A$3-1,L{idx},"0000011"))))')

last_calc_row = len(tasks) + 1

# Dashboard / chart sheet.
ws.sheet_view.showGridLines = False
ws.freeze_panes = f"{get_column_letter(DATE_START_COL)}{TASK_START_ROW}"
ws.merge_cells("A1:I1")
ws["A1"] = "きみこい 本番工数 ダウンチャート（Googleスプレッドシート対応版）"
ws["A1"].fill = PatternFill("solid", fgColor=dark)
ws["A1"].font = Font(bold=True, color=white, size=15)
ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

labels = ["開始日", "終了目安", "タスク数", "MAX工数合計", "バッファ係数", "予定工数合計", "日次消化力", "休日前提"]
values = [
    date(2026, 6, 1),
    f'=IFERROR(MAX({q("タスク計算")}!$N$2:$N${last_calc_row}),"")',
    f'=COUNTIF({q("タスク計算")}!$G$2:$G${last_calc_row},">0")',
    f'=SUM({q("タスク計算")}!$G$2:$G${last_calc_row})',
    1.25,
    "=D3*E3",
    1,
    "土日祝休み",
]
for col, (label, value) in enumerate(zip(labels, values), start=1):
    ws.cell(2, col, label)
    ws.cell(2, col).fill = PatternFill("solid", fgColor=dark2)
    ws.cell(2, col).font = Font(bold=True, color=white)
    ws.cell(2, col).alignment = Alignment(horizontal="center")
    ws.cell(2, col).border = border
    ws.cell(3, col, value)
    ws.cell(3, col).fill = PatternFill("solid", fgColor=input_fill if col in {1, 5, 7, 8} else linked_fill)
    ws.cell(3, col).font = Font(bold=True)
    ws.cell(3, col).alignment = Alignment(horizontal="center")
    ws.cell(3, col).border = border

ws["A3"].number_format = "yyyy/m/d"
ws["B3"].number_format = "yyyy/m/d"
for ref in ["D3", "E3", "F3", "G3"]:
    ws[ref].number_format = "0.0"

ws.merge_cells("I2:N3")
ws["I2"] = "黄色セルを変更すると、Googleスプレッドシート上で再計算されます。互換性優先のため、バーは条件付き書式ではなく記号で表示しています。"
ws["I2"].fill = PatternFill("solid", fgColor=summary_fill)
ws["I2"].alignment = Alignment(wrap_text=True, vertical="center")
ws["I2"].border = border

dv_premise = DataValidation(type="list", formula1='"土日祝休み,土日休み,休みなし"', allow_blank=False)
ws.add_data_validation(dv_premise)
dv_premise.add(ws["H3"])

headers = ["No", "大項目", "中項目", "タスク", "MAX工数", "予定工数", "開始日", "終了日", "根拠/備考"]
for col, header in enumerate(headers, start=1):
    cell = ws.cell(5, col, header)
    cell.value = header
    cell.fill = PatternFill("solid", fgColor=dark)
    cell.font = Font(bold=True, color=white)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border

widths = {"A": 6, "B": 18, "C": 20, "D": 44, "E": 10, "F": 10, "G": 12, "H": 12, "I": 34}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

for idx in range(2, last_calc_row + 1):
    row = TASK_START_ROW + idx - 2
    ws.cell(row, 1, f'=IF({q("タスク計算")}!$G{idx}>0,{q("タスク計算")}!$A{idx},"")')
    ws.cell(row, 2, f'=IF($A{row}="","",{q("タスク計算")}!$C{idx})')
    ws.cell(row, 3, f'=IF($A{row}="","",{q("タスク計算")}!$D{idx})')
    ws.cell(row, 4, f'=IF($A{row}="","",{q("タスク計算")}!$E{idx})')
    ws.cell(row, 5, f'=IF($A{row}="","",{q("タスク計算")}!$G{idx})')
    ws.cell(row, 6, f'=IF($A{row}="","",{q("タスク計算")}!$H{idx})')
    ws.cell(row, 7, f'=IF($A{row}="","",{q("タスク計算")}!$M{idx})')
    ws.cell(row, 8, f'=IF($A{row}="","",{q("タスク計算")}!$N{idx})')
    ws.cell(row, 9, f'=IF($A{row}="","",{q("タスク計算")}!$F{idx})')
    for col in range(1, 10):
        cell = ws.cell(row, col)
        cell.border = border
        cell.font = Font(size=8)
        cell.alignment = Alignment(horizontal="left" if col in {2, 3, 4, 9} else "center", vertical="center", wrap_text=col in {4, 9})
        if col in {5, 6}:
            cell.fill = PatternFill("solid", fgColor=linked_fill)
            cell.number_format = "0.0"
        if col in {7, 8}:
            cell.number_format = "yyyy/m/d"
    ws.row_dimensions[row].height = 18

last_task_row = TASK_START_ROW + len(tasks) - 1
last_date_col = DATE_START_COL + DATE_DAYS - 1

for col in range(DATE_START_COL, last_date_col + 1):
    letter = get_column_letter(col)
    offset = col - DATE_START_COL
    ws.column_dimensions[letter].width = 3
    ws.cell(3, col, f'=IF({letter}$4="","",IF(DAY({letter}$4)=1,TEXT({letter}$4,"yyyy/m"),""))')
    ws.cell(4, col, f'=IF(OR($A$3="",$B$3="",$A$3+{offset}>$B$3),"",$A$3+{offset})')
    ws.cell(5, col, f'=IF({letter}$4="","",CHOOSE(WEEKDAY({letter}$4,2),"月","火","水","木","金","土","日"))')
    for r in (3, 4, 5):
        ws.cell(r, col).border = hair_border
        ws.cell(r, col).alignment = Alignment(horizontal="center")
        ws.cell(r, col).font = Font(size=8, bold=r == 3)
    ws.cell(4, col).number_format = "d"

for row in range(TASK_START_ROW, last_task_row + 1):
    for col in range(DATE_START_COL, last_date_col + 1):
        letter = get_column_letter(col)
        cell = ws.cell(row, col)
        cell.value = f'=IF(AND({letter}$4<>"",$A{row}<>"",{letter}$4>=$G{row},{letter}$4<=$H{row},OR($H$3="休みなし",WEEKDAY({letter}$4,2)<=5),OR($H$3<>"土日祝休み",COUNTIF({q("祝日設定")}!$A$2:$A$80,{letter}$4)=0)),"■","")'
        cell.font = Font(size=8, color=bar_color, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = hair_border

for row in range(1, 6):
    ws.row_dimensions[row].height = 20
ws.row_dimensions[1].height = 30
ws.row_dimensions[5].height = 24

ws.auto_filter.ref = f"A5:I{last_task_row}"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0

for sheet in wb.worksheets:
    sheet.sheet_view.tabSelected = False
ws.sheet_view.tabSelected = True
wb.active = wb.sheetnames.index("ダウンチャート")

wb.calculation.calcMode = "auto"
wb.calculation.fullCalcOnLoad = False
wb.calculation.forceFullCalc = False

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)
print(OUTPUT)
print("tasks", len(tasks))
print("date_days", DATE_DAYS)
