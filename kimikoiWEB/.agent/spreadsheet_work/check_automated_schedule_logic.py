from __future__ import annotations

from datetime import date, timedelta
from openpyxl import load_workbook


WORKBOOK = r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_ダウンチャート_自動化版.xlsx"
HOLIDAYS = {
    date(2026, 1, 1),
    date(2026, 1, 12),
    date(2026, 2, 11),
    date(2026, 2, 23),
    date(2026, 3, 20),
    date(2026, 4, 29),
    date(2026, 5, 3),
    date(2026, 5, 4),
    date(2026, 5, 5),
    date(2026, 5, 6),
    date(2026, 7, 20),
    date(2026, 8, 11),
    date(2026, 9, 21),
    date(2026, 9, 22),
    date(2026, 9, 23),
    date(2026, 10, 12),
    date(2026, 11, 3),
    date(2026, 11, 23),
    date(2027, 1, 1),
    date(2027, 1, 11),
    date(2027, 2, 11),
    date(2027, 2, 23),
    date(2027, 3, 21),
    date(2027, 3, 22),
    date(2027, 4, 29),
    date(2027, 5, 3),
    date(2027, 5, 4),
    date(2027, 5, 5),
}


def is_workday(day: date) -> bool:
    return day.weekday() < 5 and day not in HOLIDAYS


def add_workdays(start: date, index: int) -> date:
    day = start
    count = 0
    while True:
        if is_workday(day):
            count += 1
            if count == index:
                return day
        day += timedelta(days=1)


wb = load_workbook(WORKBOOK, data_only=True)
ws = wb["本番工数"]
tasks = []
major = ""
for row in range(22, 1020):
    if isinstance(ws.cell(row, 3).value, str) and ws.cell(row, 3).value.strip():
        major = ws.cell(row, 3).value.strip()
    max_value = ws.cell(row, 8).value
    if isinstance(max_value, (int, float)) and max_value > 0:
        tasks.append((major, float(max_value)))

scheduled = sum(value * 1.25 for _, value in tasks)
end_index = int(-(-scheduled // 1))
for start in [date(2026, 6, 1), date(2026, 9, 1)]:
    print(start.isoformat(), "tasks", len(tasks), "scheduled", round(scheduled, 2), "end", add_workdays(start, end_index).isoformat())
