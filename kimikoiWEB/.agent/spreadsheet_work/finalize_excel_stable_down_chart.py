from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


WORKBOOK = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_ダウンチャート_自動化完成版.xlsx")
SHEET = "ダウンチャート"
TASK_START_ROW = 6
DATE_START_COL = 15
DATE_DAYS = 730
START_DATE = date(2026, 6, 1)
BUFFER_RATE = 1.25
DAILY_CAPACITY = 1.0

HOLIDAYS = {
    date(2026, 1, 1), date(2026, 1, 12), date(2026, 2, 11), date(2026, 2, 23),
    date(2026, 3, 20), date(2026, 4, 29), date(2026, 5, 3), date(2026, 5, 4),
    date(2026, 5, 5), date(2026, 5, 6), date(2026, 7, 20), date(2026, 8, 11),
    date(2026, 9, 21), date(2026, 9, 22), date(2026, 9, 23), date(2026, 10, 12),
    date(2026, 11, 3), date(2026, 11, 23), date(2027, 1, 1), date(2027, 1, 11),
    date(2027, 2, 11), date(2027, 2, 23), date(2027, 3, 21), date(2027, 3, 22),
    date(2027, 4, 29), date(2027, 5, 3), date(2027, 5, 4), date(2027, 5, 5),
}


def is_workday(day: date) -> bool:
    return day.weekday() < 5 and day not in HOLIDAYS


def nth_workday(start: date, index: int) -> date:
    count = 0
    day = start
    while True:
        if is_workday(day):
            count += 1
            if count == index:
                return day
        day += timedelta(days=1)


def as_number(value) -> float:
    return float(value) if isinstance(value, (int, float)) else 0.0


def col_letter(col: int) -> str:
    result = ""
    while col:
        col, rem = divmod(col - 1, 26)
        result = chr(65 + rem) + result
    return result


data_wb = load_workbook(WORKBOOK, data_only=True)
base = data_wb["本番工数"]

tasks = []
major = ""
mid = ""
for source_row in range(22, base.max_row + 1):
    major_value = base.cell(source_row, 3).value
    mid_value = base.cell(source_row, 4).value
    detail_value = base.cell(source_row, 5).value
    max_effort = as_number(base.cell(source_row, 8).value)
    basis = base.cell(source_row, 9).value or ""
    if isinstance(major_value, str) and major_value.strip():
        major = major_value.strip()
    if isinstance(mid_value, str) and mid_value.strip():
        mid = mid_value.strip()
    if max_effort <= 0:
        continue
    title = detail_value if isinstance(detail_value, str) and detail_value.strip() else mid
    scheduled = round(max_effort * BUFFER_RATE, 10)
    tasks.append({
        "no": len(tasks) + 1,
        "source_row": source_row,
        "major": major,
        "mid": mid,
        "title": str(title),
        "max": max_effort,
        "scheduled": scheduled,
        "basis": basis,
    })

cum = 0.0
for task in tasks:
    start_idx = int(cum / DAILY_CAPACITY) + 1
    cum += task["scheduled"]
    end_idx = int(-(-cum // DAILY_CAPACITY))
    task["start"] = nth_workday(START_DATE, start_idx)
    task["end"] = nth_workday(START_DATE, end_idx)
    task["cum_before"] = round(cum - task["scheduled"], 10)
    task["cum_end"] = round(cum, 10)
    task["start_idx"] = start_idx
    task["end_idx"] = end_idx

end_date = max(task["end"] for task in tasks)
date_values = [START_DATE + timedelta(days=i) for i in range((end_date - START_DATE).days + 1)]

wb = load_workbook(WORKBOOK)
ws = wb[SHEET]

thin = Side(style="thin", color="B7B7B7")
hair = Side(style="hair", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hair_border = Border(left=hair, right=hair, top=hair, bottom=hair)
linked_fill = PatternFill("solid", fgColor="DDEBF7")
bar_fill = PatternFill("solid", fgColor="9DC3E6")
weekend_fill = PatternFill("solid", fgColor="E7E6E6")
holiday_fill = PatternFill("solid", fgColor="D9D2E9")
blank_fill = PatternFill(fill_type=None)

ws["A3"] = START_DATE
ws["B3"] = end_date
ws["C3"] = len(tasks)
ws["D3"] = round(sum(t["max"] for t in tasks), 2)
ws["E3"] = BUFFER_RATE
ws["F3"] = round(sum(t["scheduled"] for t in tasks), 2)
ws["G3"] = DAILY_CAPACITY
ws["H3"] = "土日祝休み"
ws["A3"].number_format = "yyyy/m/d"
ws["B3"].number_format = "yyyy/m/d"
for ref in ["D3", "E3", "F3", "G3"]:
    ws[ref].number_format = "0.0"

last_task_row = TASK_START_ROW + len(tasks) - 1
max_row_to_clear = 655

for idx, task in enumerate(tasks):
    row = TASK_START_ROW + idx
    values = [
        task["no"], task["major"], task["mid"], task["title"], task["max"], task["scheduled"],
        task["start"], task["end"], task["basis"], task["source_row"], task["cum_before"],
        task["cum_end"], task["start_idx"], task["end_idx"],
    ]
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row, col, value)
        cell.border = border
        cell.font = Font(size=8)
        cell.alignment = Alignment(
            horizontal="left" if col in {2, 3, 4, 9} else "center",
            vertical="center",
            wrap_text=col in {4, 9},
        )
        if col in {5, 6}:
            cell.fill = linked_fill
            cell.number_format = "0.0"
        elif col in {7, 8}:
            cell.number_format = "yyyy/m/d"
    ws.row_dimensions[row].height = 20

for row in range(last_task_row + 1, max_row_to_clear + 1):
    for col in range(1, 15):
        ws.cell(row, col).value = None
        ws.cell(row, col).fill = blank_fill
    ws.row_dimensions[row].height = 18

for offset in range(DATE_DAYS):
    col = DATE_START_COL + offset
    current = START_DATE + timedelta(days=offset)
    if current <= end_date:
        ws.cell(3, col, f"{current.year}/{current.month}" if current.day == 1 else None)
        ws.cell(4, col, current)
        ws.cell(5, col, ["月", "火", "水", "木", "金", "土", "日"][current.weekday()])
        ws.cell(4, col).number_format = "d"
    else:
        for r in (3, 4, 5):
            ws.cell(r, col).value = None
            ws.cell(r, col).fill = blank_fill
    for r in (3, 4, 5):
        ws.cell(r, col).border = hair_border
        ws.cell(r, col).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, col).font = Font(size=8, bold=r == 3)

for offset in range(DATE_DAYS):
    col = DATE_START_COL + offset
    current = START_DATE + timedelta(days=offset)
    is_holiday = current in HOLIDAYS
    is_weekend = current.weekday() >= 5
    header_fill = holiday_fill if is_holiday else weekend_fill if is_weekend else blank_fill
    if current <= end_date:
        ws.cell(4, col).fill = header_fill
        ws.cell(5, col).fill = header_fill
    for task in tasks:
        row = TASK_START_ROW + task["no"] - 1
        cell = ws.cell(row, col)
        cell.value = None
        cell.border = hair_border
        if current > end_date:
            cell.fill = blank_fill
        elif is_holiday:
            cell.fill = holiday_fill
        elif is_weekend:
            cell.fill = weekend_fill
        elif task["start"] <= current <= task["end"]:
            cell.fill = bar_fill
        else:
            cell.fill = blank_fill
    for row in range(last_task_row + 1, max_row_to_clear + 1):
        ws.cell(row, col).value = None
        ws.cell(row, col).fill = blank_fill
        ws.cell(row, col).border = hair_border

for col, width in {
    "A": 7, "B": 22, "C": 24, "D": 48, "E": 10, "F": 12, "G": 12, "H": 12, "I": 40,
}.items():
    ws.column_dimensions[col].width = width
for col in range(10, 15):
    ws.column_dimensions[col_letter(col)].hidden = True

ws.auto_filter.ref = f"A5:I{last_task_row}"
ws.freeze_panes = "O6"

wb.save(WORKBOOK)

print(WORKBOOK)
print("tasks", len(tasks))
print("end", end_date.isoformat())
