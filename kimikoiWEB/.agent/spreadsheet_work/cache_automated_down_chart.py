from __future__ import annotations

from datetime import date, timedelta
from math import ceil, floor
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile
import shutil
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel


WORKBOOK = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_ダウンチャート_自動化版.xlsx")
BASE_SHEET = "本番工数"
CHART_SHEET = "ダウンチャート"
HOLIDAY_SHEET = "祝日設定"
SOURCE_FIRST_ROW = 22
SOURCE_LAST_ROW = 1019
FIRST_TASK_ROW = 6
FIRST_DATE_COL = 14
DATE_DAYS = 730
BUFFER_RATE = 1.25
DAILY_CAPACITY = 1.0
START_DATE = date(2026, 6, 1)

HOLIDAYS = {
    date(2026, 1, 1),
    date(2026, 1, 12),
    date(2026, 2, 11),
    date(2026, 2, 23),
    date(2026, 3, 20),
    date(2026, 4, 29),
    date(2026, 5, 3),
    date(2026, 5, 4),
    date(2026, 5, 5),
    date(2026, 5, 6),
    date(2026, 7, 20),
    date(2026, 8, 11),
    date(2026, 9, 21),
    date(2026, 9, 22),
    date(2026, 9, 23),
    date(2026, 10, 12),
    date(2026, 11, 3),
    date(2026, 11, 23),
    date(2027, 1, 1),
    date(2027, 1, 11),
    date(2027, 2, 11),
    date(2027, 2, 23),
    date(2027, 3, 21),
    date(2027, 3, 22),
    date(2027, 4, 29),
    date(2027, 5, 3),
    date(2027, 5, 4),
    date(2027, 5, 5),
    date(2027, 7, 19),
    date(2027, 8, 11),
    date(2027, 9, 20),
    date(2027, 9, 23),
    date(2027, 10, 11),
    date(2027, 11, 3),
    date(2027, 11, 23),
}


def is_workday(day: date) -> bool:
    return day.weekday() < 5 and day not in HOLIDAYS


def workday_by_index(start: date, index: int) -> date:
    day = start
    count = 0
    while True:
        if is_workday(day):
            count += 1
            if count == index:
                return day
        day += timedelta(days=1)


def sheet_xml_path(root: Path, sheet_name: str) -> Path:
    ns = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pkgrel": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    workbook_tree = ET.parse(root / "xl" / "workbook.xml")
    sheet_el = None
    for candidate in workbook_tree.getroot().findall(".//main:sheet", ns):
        if candidate.attrib.get("name") == sheet_name:
            sheet_el = candidate
            break
    if sheet_el is None:
        raise ValueError(sheet_name)
    rel_id = sheet_el.attrib[f"{{{ns['rel']}}}id"]
    rels_tree = ET.parse(root / "xl" / "_rels" / "workbook.xml.rels")
    for rel in rels_tree.getroot().findall("pkgrel:Relationship", ns):
        if rel.attrib.get("Id") == rel_id:
            target = rel.attrib["Target"]
            if target.startswith("/"):
                return root / target.lstrip("/")
            if target.startswith("worksheets/"):
                return root / "xl" / target
            return root / "xl" / target
    raise ValueError(rel_id)


wb = load_workbook(WORKBOOK, data_only=True)
ws = wb[BASE_SHEET]

tasks = []
major = ""
mid = ""
helper_values = {}
for source_row in range(SOURCE_FIRST_ROW, SOURCE_LAST_ROW + 1):
    major_value = ws.cell(source_row, 3).value
    mid_value = ws.cell(source_row, 4).value
    if isinstance(major_value, str) and major_value.strip():
        major = major_value.strip()
    if isinstance(mid_value, str) and mid_value.strip():
        mid = mid_value.strip()
    helper_values[source_row] = (major, mid)

    max_value = ws.cell(source_row, 8).value
    if not isinstance(max_value, (int, float)) or max_value <= 0:
        continue

    task_value = ws.cell(source_row, 5).value
    title = task_value if isinstance(task_value, str) and task_value.strip() else mid
    note_parts = []
    for col in (9, 10, 12):
        value = ws.cell(source_row, col).value
        if value:
            note_parts.append(str(value))

    scheduled = float(max_value) * BUFFER_RATE
    prev = sum(task["scheduled"] for task in tasks)
    start_idx = floor(prev / DAILY_CAPACITY) + 1
    end_idx = ceil((prev + scheduled) / DAILY_CAPACITY)
    tasks.append(
        {
            "source_row": source_row,
            "major": major,
            "mid": mid,
            "title": str(title),
            "max": float(max_value),
            "scheduled": scheduled,
            "prev": prev,
            "start_idx": start_idx,
            "end_idx": end_idx,
            "start_date": workday_by_index(START_DATE, start_idx),
            "end_date": workday_by_index(START_DATE, end_idx),
            "note": " / ".join(note_parts),
        }
    )

raw_max = sum(task["max"] for task in tasks)
scheduled_total = sum(task["scheduled"] for task in tasks)
end_date = tasks[-1]["end_date"]
display_days = min(730, max(365, (end_date - START_DATE).days + 30))

ns_uri = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ns = {"main": ns_uri}
ET.register_namespace("", ns_uri)

with TemporaryDirectory() as tmp:
    tmp_root = Path(tmp)
    with ZipFile(WORKBOOK, "r") as zin:
        zin.extractall(tmp_root)

    def patch_sheet(sheet_name: str):
        xml_path = sheet_xml_path(tmp_root, sheet_name)
        tree = ET.parse(xml_path)
        root = tree.getroot()
        cells = {cell.attrib["r"]: cell for cell in root.findall(".//main:c", ns) if "r" in cell.attrib}
        return xml_path, tree, cells

    def set_cache(cells, ref: str, value):
        cell = cells.get(ref)
        if cell is None:
            return
        for value_el in cell.findall("main:v", ns):
            cell.remove(value_el)
        if value is None:
            return
        v = ET.SubElement(cell, f"{{{ns_uri}}}v")
        if isinstance(value, str):
            cell.attrib["t"] = "str"
            v.text = value
        elif isinstance(value, date):
            cell.attrib.pop("t", None)
            v.text = str(to_excel(value))
        else:
            cell.attrib.pop("t", None)
            v.text = str(round(float(value), 10))

    base_xml_path, base_tree, base_cells = patch_sheet(BASE_SHEET)
    for source_row, (cached_major, cached_mid) in helper_values.items():
        set_cache(base_cells, f"AF{source_row}", cached_major)
        set_cache(base_cells, f"AG{source_row}", cached_mid)
    base_tree.write(base_xml_path, encoding="utf-8", xml_declaration=True)

    chart_xml_path, chart_tree, chart_cells = patch_sheet(CHART_SHEET)

    set_cache(chart_cells, "B3", end_date)
    set_cache(chart_cells, "C3", len(tasks))
    set_cache(chart_cells, "D3", raw_max)
    set_cache(chart_cells, "F3", scheduled_total)
    set_cache(chart_cells, "I3", display_days)

    for offset in range(DATE_DAYS):
        col_letter = get_column_letter(FIRST_DATE_COL + offset)
        day = START_DATE + timedelta(days=offset)
        if offset < display_days:
            month_label = f"{day.year}/{day.month}" if offset == 0 or day.day == 1 else ""
            set_cache(chart_cells, f"{col_letter}3", month_label)
            set_cache(chart_cells, f"{col_letter}4", day)
            set_cache(chart_cells, f"{col_letter}5", ["月", "火", "水", "木", "金", "土", "日"][day.weekday()])

    for idx, task in enumerate(tasks, start=0):
        row = FIRST_TASK_ROW + idx
        set_cache(chart_cells, f"A{row}", idx + 1)
        set_cache(chart_cells, f"B{row}", task["major"])
        set_cache(chart_cells, f"C{row}", task["mid"])
        set_cache(chart_cells, f"D{row}", task["title"])
        set_cache(chart_cells, f"E{row}", task["max"])
        set_cache(chart_cells, f"F{row}", task["scheduled"])
        set_cache(chart_cells, f"G{row}", task["start_date"])
        set_cache(chart_cells, f"H{row}", task["end_date"])
        set_cache(chart_cells, f"I{row}", task["note"])
        set_cache(chart_cells, f"J{row}", task["source_row"])
        set_cache(chart_cells, f"K{row}", task["prev"])
        set_cache(chart_cells, f"L{row}", task["start_idx"])
        set_cache(chart_cells, f"M{row}", task["end_idx"])

    chart_tree.write(chart_xml_path, encoding="utf-8", xml_declaration=True)

    patched = WORKBOOK.with_suffix(".patched.xlsx")
    if patched.exists():
        patched.unlink()
    with ZipFile(patched, "w", ZIP_DEFLATED) as zout:
        for file in tmp_root.rglob("*"):
            if file.is_file():
                zout.write(file, file.relative_to(tmp_root).as_posix())
    shutil.move(patched, WORKBOOK)

print("cached", len(tasks), round(raw_max, 2), round(scheduled_total, 2), end_date.isoformat())
