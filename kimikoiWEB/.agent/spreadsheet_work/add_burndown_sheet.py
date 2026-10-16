from copy import copy
from datetime import date, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile
import math
import shutil
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


SOURCE = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_本番割増_再調整_単独シート.xlsx")
OUTPUT = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_本番割増_バーンダウン付き_実値入り.xlsx")
BASE_SHEET = "本番割増_再調整"
SHEET = "バーンダウン"


HOLIDAYS = [
    (date(2026, 1, 1), "元日"),
    (date(2026, 1, 12), "成人の日"),
    (date(2026, 2, 11), "建国記念の日"),
    (date(2026, 2, 23), "天皇誕生日"),
    (date(2026, 3, 20), "春分の日"),
    (date(2026, 4, 29), "昭和の日"),
    (date(2026, 5, 3), "憲法記念日"),
    (date(2026, 5, 4), "みどりの日"),
    (date(2026, 5, 5), "こどもの日"),
    (date(2026, 5, 6), "休日"),
    (date(2026, 7, 20), "海の日"),
    (date(2026, 8, 11), "山の日"),
    (date(2026, 9, 21), "敬老の日"),
    (date(2026, 9, 22), "休日"),
    (date(2026, 9, 23), "秋分の日"),
    (date(2026, 10, 12), "スポーツの日"),
    (date(2026, 11, 3), "文化の日"),
    (date(2026, 11, 23), "勤労感謝の日"),
    (date(2027, 1, 1), "元日"),
    (date(2027, 1, 11), "成人の日"),
    (date(2027, 2, 11), "建国記念の日"),
    (date(2027, 2, 23), "天皇誕生日"),
    (date(2027, 3, 21), "春分の日"),
    (date(2027, 3, 22), "休日"),
    (date(2027, 4, 29), "昭和の日"),
    (date(2027, 5, 3), "憲法記念日"),
    (date(2027, 5, 4), "みどりの日"),
    (date(2027, 5, 5), "こどもの日"),
    (date(2027, 7, 19), "海の日"),
    (date(2027, 8, 11), "山の日"),
    (date(2027, 9, 20), "敬老の日"),
    (date(2027, 9, 23), "秋分の日"),
    (date(2027, 10, 11), "スポーツの日"),
    (date(2027, 11, 3), "文化の日"),
    (date(2027, 11, 23), "勤労感謝の日"),
]


def is_workday(d: date, holidays: set[date]) -> bool:
    return d.weekday() < 5 and d not in holidays


def add_workdays(start: date, workdays: int, holidays: set[date]) -> date:
    d = start - timedelta(days=1)
    remaining = workdays
    while remaining > 0:
        d += timedelta(days=1)
        if is_workday(d, holidays):
            remaining -= 1
    return d


def excel_serial(d: date) -> int:
    return (d - date(1899, 12, 30)).days


def month_add(d: date, months: int) -> date:
    month_index = d.month - 1 + months
    return date(d.year + month_index // 12, month_index % 12 + 1, 1)


def forecast_values(min_effort: float, max_effort: float, selected_effort: float) -> dict[str, object]:
    holidays = {d for d, _ in HOLIDAYS}
    start = date(2026, 6, 1)
    daily_capacity = 1.0
    needed_workdays = math.ceil(selected_effort / daily_capacity)
    rows = 520

    daily = []
    cumulative = 0.0
    cumulative_workdays = 0
    min_end = None
    max_end = None
    selected_end = None
    for idx in range(1, rows + 1):
        current = start + timedelta(days=idx - 1)
        workday = 1 if is_workday(current, holidays) else 0
        burn = workday * daily_capacity
        cumulative += burn
        cumulative_workdays += workday
        remaining_min = max(0.0, min_effort - cumulative)
        remaining_max = max(0.0, max_effort - cumulative)
        remaining_selected = max(0.0, selected_effort - cumulative)
        ideal_remaining = max(
            0.0,
            selected_effort - (selected_effort / needed_workdays) * min(cumulative_workdays, needed_workdays),
        )
        if min_end is None and remaining_min == 0:
            min_end = current
        if max_end is None and remaining_max == 0:
            max_end = current
        if selected_end is None and remaining_selected == 0:
            selected_end = current
        daily.append(
            {
                "no": idx,
                "date": current,
                "weekday": "月火水木金土日"[current.weekday()],
                "workday": workday,
                "burn": burn,
                "cumulative": cumulative,
                "remaining_min": remaining_min,
                "remaining_max": remaining_max,
                "remaining_selected": remaining_selected,
                "ideal_remaining": ideal_remaining,
            }
        )

    monthly = []
    planned_cumulative = 0.0
    for offset in range(24):
        month_start = month_add(start, offset)
        next_month = month_add(start, offset + 1)
        month_rows = [r for r in daily if month_start <= r["date"] < next_month]
        workdays = sum(r["workday"] for r in month_rows)
        planned = sum(r["burn"] for r in month_rows)
        planned_cumulative += planned
        monthly.append(
            {
                "date": month_start,
                "workdays": workdays,
                "planned": planned,
                "remaining_min": max(0.0, min_effort - planned_cumulative),
                "remaining_max": max(0.0, max_effort - planned_cumulative),
                "remaining_selected": max(0.0, selected_effort - planned_cumulative),
            }
        )

    return {
        "start": start,
        "min_effort": min_effort,
        "max_effort": max_effort,
        "selected_effort": selected_effort,
        "daily_capacity": daily_capacity,
        "needed_workdays": needed_workdays,
        "min_end": min_end,
        "max_end": max_end,
        "selected_end": selected_end,
        "calendar_days": (selected_end - start).days + 1,
        "months": round(needed_workdays / 20, 2),
        "daily": daily,
        "monthly": monthly,
    }


def get_sheet_xml_path(root: Path, sheet_name: str) -> Path:
    ns = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pkgrel": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    workbook = ET.parse(root / "xl" / "workbook.xml")
    sheet_el = None
    for candidate in workbook.getroot().findall(".//main:sheet", ns):
        if candidate.attrib.get("name") == sheet_name:
            sheet_el = candidate
            break
    if sheet_el is None:
        raise ValueError(f"Sheet not found: {sheet_name}")

    rel_id = sheet_el.attrib[f"{{{ns['rel']}}}id"]
    rels = ET.parse(root / "xl" / "_rels" / "workbook.xml.rels")
    target = None
    for rel in rels.getroot().findall("pkgrel:Relationship", ns):
        if rel.attrib.get("Id") == rel_id:
            target = rel.attrib["Target"]
            break
    if target is None:
        raise ValueError(f"Relationship not found: {sheet_name}")
    if target.startswith("/"):
        return root / target.lstrip("/")
    if target.startswith("worksheets/"):
        return root / "xl" / target
    return root / "xl" / target


def patch_formula_caches(path: Path, caches_by_sheet: dict[str, dict[str, object]]) -> None:
    ns_uri = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ns = {"main": ns_uri}
    ET.register_namespace("", ns_uri)

    with TemporaryDirectory() as tmp:
        tmp_root = Path(tmp)
        with ZipFile(path, "r") as zin:
            zin.extractall(tmp_root)

        for sheet_name, caches in caches_by_sheet.items():
            sheet_path = get_sheet_xml_path(tmp_root, sheet_name)
            tree = ET.parse(sheet_path)
            root = tree.getroot()
            cells = {
                cell.attrib["r"]: cell
                for cell in root.findall(".//main:c", ns)
                if "r" in cell.attrib
            }
            for ref, value in caches.items():
                cell = cells.get(ref)
                if cell is None:
                    continue
                for value_el in cell.findall("main:v", ns):
                    cell.remove(value_el)
                for inline_el in cell.findall("main:is", ns):
                    cell.remove(inline_el)
                if value is None:
                    continue
                v = ET.SubElement(cell, f"{{{ns_uri}}}v")
                if isinstance(value, date):
                    cell.attrib.pop("t", None)
                    v.text = str(excel_serial(value))
                elif isinstance(value, str):
                    cell.attrib["t"] = "str"
                    v.text = value
                else:
                    cell.attrib.pop("t", None)
                    v.text = f"{float(value):.10f}".rstrip("0").rstrip(".")
            tree.write(sheet_path, encoding="utf-8", xml_declaration=True)

        patched = path.with_suffix(".patched.xlsx")
        with ZipFile(patched, "w", ZIP_DEFLATED) as zout:
            for file in tmp_root.rglob("*"):
                if file.is_file():
                    zout.write(file, file.relative_to(tmp_root).as_posix())
        shutil.move(patched, path)


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb = load_workbook(SOURCE)
if SHEET in wb.sheetnames:
    del wb[SHEET]

ws = wb.create_sheet(SHEET, index=1)
ws.sheet_view.showGridLines = False
ws.freeze_panes = "A45"

for sheet in wb.worksheets:
    sheet.sheet_view.tabSelected = False
ws.sheet_view.tabSelected = True
wb.active = wb.sheetnames.index(SHEET)

dark = "1F4E5F"
blue = "D9EAF7"
green = "E2F0D9"
yellow = "FFF2CC"
orange = "FCE4D6"
gray = "F2F2F2"
white = "FFFFFF"
red_text = "9C0006"
green_text = "006100"

thin = Side(style="thin", color="7F7F7F")
medium = Side(style="medium", color="000000")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
section_border = Border(left=medium, right=medium, top=medium, bottom=medium)

title_fill = PatternFill("solid", fgColor=dark)
input_fill = PatternFill("solid", fgColor=yellow)
linked_fill = PatternFill("solid", fgColor=blue)
summary_fill = PatternFill("solid", fgColor=green)
holiday_fill = PatternFill("solid", fgColor=orange)
header_fill = PatternFill("solid", fgColor=dark)
gray_fill = PatternFill("solid", fgColor=gray)

title_font = Font(bold=True, color=white, size=16)
header_font = Font(bold=True, color=white)
bold = Font(bold=True)
small = Font(size=9)
center = Alignment(horizontal="center", vertical="center")
left = Alignment(horizontal="left", vertical="center")
right = Alignment(horizontal="right", vertical="center")

for col, width in {
    "A": 18,
    "B": 13,
    "C": 12,
    "D": 10,
    "E": 11,
    "F": 12,
    "G": 13,
    "H": 13,
    "I": 13,
    "J": 13,
    "K": 11,
    "L": 12,
    "M": 13,
    "N": 26,
    "P": 12,
    "Q": 11,
    "R": 12,
    "S": 13,
    "T": 13,
    "U": 13,
    "V": 12,
    "W": 13,
    "X": 8,
}.items():
    ws.column_dimensions[col].width = width

ws.merge_cells("A1:N1")
ws["A1"] = "きみこい 本番割増スケジュール / バーンダウン"
ws["A1"].fill = title_fill
ws["A1"].font = title_font
ws["A1"].alignment = left
ws.row_dimensions[1].height = 28

ws["A3"] = "設定"
ws["A3"].fill = header_fill
ws["A3"].font = header_font
ws["A3"].alignment = center
ws.merge_cells("A3:C3")

settings = [
    ("作業開始日", date(2026, 6, 1), "変更可"),
    ("採用シナリオ", "MAX", "MIN/MAXを選択"),
    ("MIN工数(バッファ込)", f"=IFERROR('{BASE_SHEET}'!D8,307.675)", "見積シート連動"),
    ("MAX工数(バッファ込)", f"=IFERROR('{BASE_SHEET}'!D9,331.55)", "見積シート連動"),
    ("採用工数", '=IF($B$5="MIN",$B$6,$B$7)', "チャート基準"),
    ("担当人数", 1, "変更可"),
    ("1人あたり日次消化", 1, "人日/日"),
    ("日次消化力", "=$B$9*$B$10", "自動計算"),
    ("週末パターン", 1, "1=土日休み"),
    ("必要営業日", "=ROUNDUP($B$8/$B$11,0)", "採用工数÷日次消化"),
    ("予定終了日", '=IFERROR(INDEX($B$46:$B$565,MATCH(0,$I$46:$I$565,0)),DATE(2027,10,6))', "土日祝除外"),
    ("暦日数", "=$B$14-$B$4+1", "開始日含む"),
    ("想定月数", "=ROUND($B$13/20,2)", "20営業日/月換算"),
]
for i, (label, value, note) in enumerate(settings, start=4):
    ws.cell(i, 1, label)
    ws.cell(i, 2, value)
    ws.cell(i, 3, note)
    for col in range(1, 4):
        cell = ws.cell(i, col)
        cell.border = border
        cell.alignment = left if col != 2 else right
    ws.cell(i, 1).font = bold
    ws.cell(i, 2).fill = input_fill if label in {"作業開始日", "採用シナリオ", "担当人数", "1人あたり日次消化", "週末パターン"} else linked_fill
    ws.cell(i, 3).font = small

for cell in ["B4", "B14"]:
    ws[cell].number_format = "yyyy/m/d"
for cell in ["B6", "B7", "B8", "B11", "B13", "B15", "B16"]:
    ws[cell].number_format = "0.0"

dv = DataValidation(type="list", formula1='"MIN,MAX"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(ws["B5"])

ws["D3"] = "祝日一覧（変更可）"
ws["D3"].fill = header_fill
ws["D3"].font = header_font
ws["D3"].alignment = center
ws.merge_cells("D3:F3")
for col, label in enumerate(["日付", "名称", "備考"], start=4):
    ws.cell(4, col, label)
    ws.cell(4, col).fill = holiday_fill
    ws.cell(4, col).font = bold
    ws.cell(4, col).alignment = center
    ws.cell(4, col).border = border
for i, (holiday, name) in enumerate(HOLIDAYS, start=5):
    ws.cell(i, 4, holiday)
    ws.cell(i, 5, name)
    ws.cell(i, 6, "内閣府公表")
    for col in range(4, 7):
        ws.cell(i, col).border = border
        ws.cell(i, col).alignment = left
    ws.cell(i, 4).number_format = "yyyy/m/d"

ws["H3"] = "概算スケジュール"
ws["H3"].fill = header_fill
ws["H3"].font = header_font
ws["H3"].alignment = center
ws.merge_cells("H3:N3")
kpis = [
    ("MIN終了目安", '=IFERROR(INDEX($B$46:$B$565,MATCH(0,$G$46:$G$565,0)),DATE(2027,8,31))', "軽めに進んだ場合"),
    ("MAX終了目安", '=IFERROR(INDEX($B$46:$B$565,MATCH(0,$H$46:$H$565,0)),DATE(2027,10,6))', "安全側"),
    ("採用シナリオ終了", "=$B$14", "現在はMAX基準"),
    ("開始時点の残工数", "=$B$8", "見積シート連動"),
    ("月平均消化", "=ROUND($B$11*20,1)", "20営業日換算"),
    ("モック後の目安", '="約"&TEXT($B$16,"0.0")&"ヶ月"', "実際の仕様確定で変動"),
]
for r, (label, value, note) in enumerate(kpis, start=4):
    ws.cell(r, 8, label)
    ws.cell(r, 9, value)
    ws.cell(r, 10, note)
    ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=14)
    for col in range(8, 15):
        ws.cell(r, col).border = border
        ws.cell(r, col).alignment = left if col != 9 else right
    ws.cell(r, 8).font = bold
    ws.cell(r, 9).fill = summary_fill
    ws.cell(r, 10).font = small
for cell in ["I4", "I5", "I6"]:
    ws[cell].number_format = "yyyy/m/d"
for cell in ["I7", "I8"]:
    ws[cell].number_format = "0.0"

ws["H12"] = "読み方"
ws["H12"].fill = gray_fill
ws["H12"].font = bold
ws["H13"] = "黄色セルは変更可、青セルは見積シート連動。実績消化をK列に入れると実績残もグラフ化されます。"
ws.merge_cells("H13:N15")
ws["H13"].alignment = Alignment(wrap_text=True, vertical="top")
for row in range(12, 16):
    for col in range(8, 15):
        ws.cell(row, col).border = border

daily_start = 45
headers = [
    "No",
    "日付",
    "曜日",
    "稼働日",
    "日次消化",
    "累計消化",
    "残MIN",
    "残MAX",
    "残採用",
    "理想残",
    "実績消化",
    "実績累計",
    "実績残",
    "メモ",
]
for col, label in enumerate(headers, start=1):
    c = ws.cell(daily_start, col, label)
    c.fill = header_fill
    c.font = header_font
    c.alignment = center
    c.border = section_border

rows = 520
for r in range(daily_start + 1, daily_start + rows + 1):
    idx = r - daily_start
    ws.cell(r, 1, f"=ROW()-{daily_start}")
    ws.cell(r, 2, f"=$B$4+A{r}-1")
    ws.cell(r, 3, f'=CHOOSE(WEEKDAY(B{r},1),"日","月","火","水","木","金","土")')
    ws.cell(r, 4, f"=--AND(WEEKDAY(B{r},2)<=5,COUNTIF($D$5:$D$40,B{r})=0)")
    ws.cell(r, 5, f"=D{r}*$B$11")
    ws.cell(r, 6, f"=SUM($E${daily_start + 1}:E{r})")
    ws.cell(r, 7, f"=MAX(0,$B$6-F{r})")
    ws.cell(r, 8, f"=MAX(0,$B$7-F{r})")
    ws.cell(r, 9, f"=MAX(0,$B$8-F{r})")
    ws.cell(r, 10, f"=MAX(0,$B$8-($B$8/$B$13)*MIN(SUM($D${daily_start + 1}:D{r}),$B$13))")
    ws.cell(r, 12, f"=SUM($K${daily_start + 1}:K{r})")
    ws.cell(r, 13, f'=IF(COUNT($K${daily_start + 1}:K{r})=0,"",MAX(0,$B$8-L{r}))')
    for col in range(1, 15):
        cell = ws.cell(r, col)
        cell.border = border
        cell.alignment = center if col <= 4 or col in {7, 8, 9, 10, 11, 12, 13} else left
    ws.cell(r, 2).number_format = "yyyy/m/d"
    for col in range(5, 14):
        ws.cell(r, col).number_format = "0.0"
    ws.cell(r, 11).fill = input_fill
    if idx % 2 == 0:
        for col in range(1, 15):
            if col != 11:
                ws.cell(r, col).fill = PatternFill("solid", fgColor="FAFAFA")

ws.conditional_formatting.add(
    f"A{daily_start + 1}:N{daily_start + rows}",
    FormulaRule(formula=[f"$D{daily_start + 1}=0"], fill=gray_fill),
)
ws.conditional_formatting.add(
    f"I{daily_start + 1}:I{daily_start + rows}",
    FormulaRule(formula=[f"$I{daily_start + 1}=0"], font=Font(color=green_text, bold=True)),
)
ws.conditional_formatting.add(
    f"M{daily_start + 1}:M{daily_start + rows}",
    FormulaRule(formula=[f"AND($M{daily_start + 1}<>\"\",$M{daily_start + 1}>$I{daily_start + 1})"], font=Font(color=red_text, bold=True)),
)

monthly_start = 45
monthly_col = 16
monthly_headers = ["月", "営業日", "予定消化", "月末残MIN", "月末残MAX", "月末残採用", "実績消化", "実績残", "表示月"]
for col, label in enumerate(monthly_headers, start=monthly_col):
    c = ws.cell(monthly_start, col, label)
    c.fill = header_fill
    c.font = header_font
    c.alignment = center
    c.border = section_border

for r in range(monthly_start + 1, monthly_start + 25):
    offset = r - monthly_start - 1
    month_date = f"DATE(YEAR($B$4),MONTH($B$4)+{offset},1)"
    next_month = f"DATE(YEAR($B$4),MONTH($B$4)+{offset + 1},1)"
    ws.cell(r, monthly_col, f"={month_date}")
    ws.cell(r, monthly_col + 1, f'=SUMIFS($D${daily_start + 1}:$D${daily_start + rows},$B${daily_start + 1}:$B${daily_start + rows},">="&{month_date},$B${daily_start + 1}:$B${daily_start + rows},"<"&{next_month})')
    ws.cell(r, monthly_col + 2, f'=SUMIFS($E${daily_start + 1}:$E${daily_start + rows},$B${daily_start + 1}:$B${daily_start + rows},">="&{month_date},$B${daily_start + 1}:$B${daily_start + rows},"<"&{next_month})')
    ws.cell(r, monthly_col + 3, f"=MAX(0,$B$6-SUM($R${monthly_start + 1}:R{r}))")
    ws.cell(r, monthly_col + 4, f"=MAX(0,$B$7-SUM($R${monthly_start + 1}:R{r}))")
    ws.cell(r, monthly_col + 5, f"=MAX(0,$B$8-SUM($R${monthly_start + 1}:R{r}))")
    ws.cell(r, monthly_col + 6, f'=SUMIFS($K${daily_start + 1}:$K${daily_start + rows},$B${daily_start + 1}:$B${daily_start + rows},">="&{month_date},$B${daily_start + 1}:$B${daily_start + rows},"<"&{next_month})')
    ws.cell(r, monthly_col + 7, f'=IF(SUM($V${monthly_start + 1}:V{r})=0,"",MAX(0,$B$8-SUM($V${monthly_start + 1}:V{r})))')
    ws.cell(r, monthly_col + 8, f'=TEXT(P{r},"m月")')
    for col in range(monthly_col, monthly_col + 9):
        ws.cell(r, col).border = border
        ws.cell(r, col).alignment = center
    for col in range(monthly_col + 1, monthly_col + 8):
        ws.cell(r, col).number_format = "0.0"
    ws.cell(r, monthly_col).number_format = "yyyy/m"

line = LineChart()
line.title = "バーンダウン（残工数）"
line.style = 13
line.y_axis.title = "残工数（人日）"
line.x_axis.title = "月"
line.x_axis.numFmt = "yyyy/m"
line.x_axis.tickLblPos = "none"
line.height = 11
line.width = 22
cats = Reference(ws, min_col=monthly_col, min_row=monthly_start + 1, max_row=monthly_start + 24)
for col in [monthly_col + 3, monthly_col + 4, monthly_col + 5, monthly_col + 7]:
    data = Reference(ws, min_col=col, min_row=monthly_start, max_row=monthly_start + 24)
    line.add_data(data, titles_from_data=True)
line.set_categories(cats)
for i, color in enumerate(["5B9BD5", "ED7D31", "70AD47", "C00000"]):
    line.series[i].graphicalProperties.line.solidFill = color
    line.series[i].graphicalProperties.line.width = 22000
for series, title in zip(line.series, ["MIN残", "MAX残", "採用残", "実績残"]):
    series.tx = SeriesLabel(v=title)
ws.add_chart(line, "H18")

bar = BarChart()
bar.title = "月別予定消化"
bar.y_axis.title = "人日"
bar.x_axis.title = "月"
bar.x_axis.numFmt = "yyyy/m"
bar.x_axis.tickLblPos = "none"
bar.height = 11
bar.width = 14
bar_data = Reference(ws, min_col=monthly_col + 2, min_row=monthly_start, max_row=monthly_start + 24)
bar_cats = Reference(ws, min_col=monthly_col, min_row=monthly_start + 1, max_row=monthly_start + 24)
bar.add_data(bar_data, titles_from_data=True)
bar.set_categories(bar_cats)
bar.series[0].graphicalProperties.solidFill = "5B9BD5"
bar.series[0].tx = SeriesLabel(v="予定消化")
ws.add_chart(bar, "P18")

for row in range(1, daily_start + rows + 1):
    ws.row_dimensions[row].height = 18
for row in [1, 3, 4, daily_start]:
    ws.row_dimensions[row].height = 24

for row in range(1, daily_start + rows + 1):
    for col in range(1, 21):
        ws.cell(row, col).font = copy(ws.cell(row, col).font)
        if row > daily_start:
            ws.cell(row, col).font = Font(size=9)

for row in range(4, 18):
    ws.cell(row, 2).font = Font(bold=True)

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUTPUT)

min_effort = 307.675
max_effort = 331.55
forecast = forecast_values(min_effort, max_effort, max_effort)
burndown_cache = {
    "B6": min_effort,
    "B7": max_effort,
    "B8": max_effort,
    "B11": forecast["daily_capacity"],
    "B13": forecast["needed_workdays"],
    "B14": forecast["selected_end"],
    "B15": forecast["calendar_days"],
    "B16": forecast["months"],
    "I4": forecast["min_end"],
    "I5": forecast["max_end"],
    "I6": forecast["selected_end"],
    "I7": max_effort,
    "I8": 20.0,
    "I9": f"約{forecast['months']:.1f}ヶ月",
}

for i, row in enumerate(forecast["daily"], start=daily_start + 1):
    burndown_cache[f"A{i}"] = row["no"]
    burndown_cache[f"B{i}"] = row["date"]
    burndown_cache[f"C{i}"] = row["weekday"]
    burndown_cache[f"D{i}"] = row["workday"]
    burndown_cache[f"E{i}"] = row["burn"]
    burndown_cache[f"F{i}"] = row["cumulative"]
    burndown_cache[f"G{i}"] = row["remaining_min"]
    burndown_cache[f"H{i}"] = row["remaining_max"]
    burndown_cache[f"I{i}"] = row["remaining_selected"]
    burndown_cache[f"J{i}"] = row["ideal_remaining"]
    burndown_cache[f"L{i}"] = 0.0
    burndown_cache[f"M{i}"] = None

for i, row in enumerate(forecast["monthly"], start=monthly_start + 1):
    burndown_cache[f"P{i}"] = row["date"]
    burndown_cache[f"Q{i}"] = row["workdays"]
    burndown_cache[f"R{i}"] = row["planned"]
    burndown_cache[f"S{i}"] = row["remaining_min"]
    burndown_cache[f"T{i}"] = row["remaining_max"]
    burndown_cache[f"U{i}"] = row["remaining_selected"]
    burndown_cache[f"V{i}"] = 0.0
    burndown_cache[f"W{i}"] = None
    burndown_cache[f"X{i}"] = f"{row['date'].month}月"

patch_formula_caches(
    OUTPUT,
    {
        BASE_SHEET: {
            "D4": 246.14,
            "D5": 265.24,
            "D8": 307.675,
            "D9": 331.55,
            "D14": "追加30.7/49.8 削減10.4/10.4 差分+20.3/+39.4",
            "D17": 15.38375,
            "D18": 16.5775,
        },
        SHEET: burndown_cache,
    },
)

holiday_set = {d for d, _ in HOLIDAYS}
start = date(2026, 6, 1)
min_end = add_workdays(start, 308, holiday_set)
max_end = add_workdays(start, 332, holiday_set)
print(OUTPUT)
print("sheets", wb.sheetnames)
print("estimated_end_min", min_end.isoformat(), "estimated_end_max", max_end.isoformat())
