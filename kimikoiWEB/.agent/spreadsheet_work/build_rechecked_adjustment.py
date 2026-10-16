from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile
import re
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


SOURCE = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\きみこい・工数表.xlsx")
OUTPUT = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_本番割増_再調整版.xlsx")
BASE_SHEET = "工数表ベース原本"
TARGET_SHEET = "本番割増_再調整"


tasks = [
    ("妹 表情差分/Spine", "表情差分データ定義/インポート設定", "表情IDをSpine/静止画切替テーブルとして実装側で扱える形にする", 0.6, 0.9, "仕様リスト作成ではなく、実装データ化の差分", "Expression Data", None),
    ("妹 表情差分/Spine", "表情IDマスタ/Enum生成", "シナリオ指定と実装参照がズレないよう表情IDをコード側に固定する", 0.3, 0.5, "参照ミス防止のため", "Enum/Master", None),
    ("妹 表情差分/Spine", "Spine差分パーツRuntime切替確認", "目、眉、口、頬、影などの差分パーツをRuntimeで切替できるか確認する", 0.6, 1.0, "Spine導入/プレハブ作成は既存項目にあるため切替差分のみ", "Spine Runtime", "既存Spine項目と重複しない範囲"),
    ("妹 表情差分/Spine", "表情ブレンド制御", "笑顔から不穏、怯えから狂気など中間表情を制御する", 0.8, 1.2, "かわいい/怖いの境界調整が難しい", "Expression Blend", None),
    ("妹 表情差分/Spine", "ホラー用Material/Attachment制御", "影、赤み、ノイズ等を表情状態に合わせて切り替える", 0.6, 1.0, "表情差分だけではなく見え方の実装差分がある", "Material State", None),
    ("妹 表情差分/Spine", "表情ステート優先度Resolver", "シナリオ指定、タッチ反応、恐怖進行の表情指定が衝突した時の優先順位を解決する", 0.6, 0.8, "追加仕様で表情指定元が増えるため", "Expression Resolver", None),
    ("妹 表情差分/Spine", "表情差分キャッシュ/遅延ロード", "増えた表情差分を必要なタイミングで読み込み、初回表示の引っかかりを抑える", 0.5, 0.6, "表情差分増加に伴う実装差分", "Expression Cache", None),
    ("妹 表情差分/Spine", "Spineイベント/表情タグ同期", "Spine側イベントとシナリオ表情タグの発火タイミングを同期する", 0.6, 0.8, "演出ズレを防ぐため", "Spine Event Sync", None),
    ("妹 表情差分/Spine", "Spine/静止画フォールバックResolver", "Spine未対応箇所や軽量表示用に静止画差分へ切り替えるResolverを作る", 0.3, 0.5, "安全策として最低限用意", "Fallback", None),
    ("妹 表情差分/Spine", "Atlas/メモリ使用量チェック", "表情差分増加によるテクスチャ容量とロード負荷を確認する", 0.3, 0.6, "ビルド後の負荷確認が必要", "Performance", None),
    ("妹 表情差分/Spine", "表情タグ検証", "シナリオ内の表情タグが未定義IDを参照していないか検出する", 0.3, 0.6, "既存の表情変更コマンドに対する検証差分", "Tag Validator", None),
    ("タッチアクション", "妹ちゃんタッチ判定エリア実装", "頭、顔、肩、手元など反応箇所のHitAreaを実装する", 0.7, 1.1, "トリッカル系の触り心地を再現", "Touch Area", None),
    ("タッチアクション", "疑似デスクトップ上の入力優先度制御", "LIMEやウィンドウ表示中にタッチ判定が暴発しないよう制御する", 0.6, 1.0, "既存疑似OS上に新しい入力層が乗るため", "Input Priority", None),
    ("タッチアクション", "タッチ反応データテーブル接続", "場所/日数/パラメータ別の反応台詞と表情を実装データへ接続する", 0.4, 0.8, "シナリオと表情の接続が必要", "Reaction Table", None),
    ("タッチアクション", "パラメータ変動連携", "タッチ結果で好感度、イライラ度、恐怖値等が変わるようにする", 0.5, 0.9, "既存パラメータ管理との接続差分", "Parameter Hook", None),
    ("タッチアクション", "Spineリアクション再生", "タッチ箇所に応じて表情/モーション/揺れを再生する", 0.5, 0.9, "素材と演出調整が絡む", "Spine Reaction", None),
    ("タッチアクション", "タッチ反応キュー/キャンセル制御", "連続タッチや会話開始時に古い反応が残らないよう反応キューを整理する", 0.6, 0.8, "反応演出の衝突防止", "Reaction Queue", None),
    ("タッチアクション", "UI遮蔽/透過レイヤー検証", "ウィンドウ裏や透明UI上でタッチ判定が通る/通らない条件を確認する", 0.5, 0.7, "疑似OSとタッチの境界調整", "Raycast Layer", None),
    ("タッチアクション", "反応ログ/チューニング出力", "どの部位を触ってどの反応に分岐したかを調整用ログとして出す", 0.5, 0.6, "実装後の調整効率を上げる", "Touch Log", None),
    ("タッチアクション", "連打/長押し/禁止操作の分岐", "押し方によって反応が変わる処理を追加する", 0.4, 0.7, "嫌がる/照れる/壊れる反応差分", "Input Pattern", None),
    ("タッチアクション", "クールタイム/暴発防止", "連続入力で演出が重ならないよう制御する", 0.3, 0.5, "演出破綻防止", "Cooldown", None),
    ("タッチアクション", "ボイス/SE同期", "タッチ時のSE、短い声、ノイズ音を表情切替と同期する", 0.3, 0.6, "音量/タイミング調整込み", "Audio Hook", None),
    ("タッチアクション", "タッチ履歴フラグ保存", "触った場所や拒否反応を後続イベント条件に使えるよう保存する", 0.3, 0.6, "新タッチ仕様の状態管理差分", "Touch Flags", None),
    ("技術剪定/演出検証", "グリッチ演出採用ライン検証", "既存グリッチ実装を前提に、追加仕様で必要な強度/発火条件だけ検証する", 0.4, 0.7, "シェーダー作成自体は既存項目にあるため減額", "Shader検証", "既存グリッチ項目と重複しない範囲"),
    ("技術剪定/演出検証", "疑似デスクトップ同時表示検証", "LIME、ボード、妹表示、タッチ判定を同時に出した時のCanvas破綻を確認する", 0.4, 0.7, "既存疑似OS項目と重複するため検証差分のみ", "Desktop検証", None),
    ("技術剪定/演出検証", "偽カーソルOS依存検証", "ビルド状態でカーソル演出が安全に見せられるか試す", 0.4, 0.8, "OS依存リスクの確認", "Input検証", None),
    ("技術剪定/演出検証", "セーブ/ロード中の演出割り込み安全確認", "異変演出とセーブデータ更新が衝突しないか確認する", 0.5, 0.9, "進行不能リスクが高い", "Save Safety", None),
    ("技術剪定/演出検証", "負荷予算確認", "Spine、タッチ、グリッチ同時発火時のFPS/メモリ目安を確認する", 0.4, 0.7, "製品版での最低限の安定性", "Performance", None),
    ("技術剪定/演出検証", "演出設定のScriptableObject化", "追加演出の強度や発火条件をデータ側で調整できるようにする", 0.5, 0.9, "不確定演出を剪定しやすくするため", "Effect Config", None),
    ("技術剪定/演出検証", "ビルド別演出ON/OFFフラグ", "危険な演出をビルド設定で無効化できるようにする", 0.3, 0.5, "採用/見送りを切り替えやすくする", "Feature Flag", None),
    ("技術剪定/演出検証", "演出プリセット比較Scene", "グリッチ、暗転、カーソル演出を同条件で見比べる検証Sceneを作る", 0.7, 0.8, "採用判断を早くするため", "Effect Sandbox", None),
    ("Exeシナリオ読込/デバッグ", "ビルド用シナリオパス解決", "StreamingAssets等に置いたシナリオをExeから読む構成にする", 0.6, 1.1, "Editor専用パスを避ける", "Build Path", None),
    ("Exeシナリオ読込/デバッグ", "文字コード/ファイルパス検証", "Windowsビルドで日本語ファイル名や改行が壊れないか確認する", 0.4, 0.7, "環境差分の確認", "Encoding", None),
    ("Exeシナリオ読込/デバッグ", "シナリオBundleバージョン/Checksum", "Exeに同梱したシナリオとデバッグ読み込みデータの版ズレを検出する", 0.4, 0.8, "ビルド状態の事故防止", "Scenario Version", None),
    ("Exeシナリオ読込/デバッグ", "シナリオ再読込デバッグ機能", "ビルド状態でもシナリオを再読込して確認できる導線を作る", 0.7, 1.2, "調整速度を上げるため", "Hot Reload", None),
    ("Exeシナリオ読込/デバッグ", "シナリオ構文チェック拡張", "ジャンプ先、表情ID、タッチ反応ID、演出タグ欠けを検出する", 0.8, 1.4, "ビルド後の原因特定を早くする", "Validator", None),
    ("Exeシナリオ読込/デバッグ", "行番号/再生中コマンド表示", "Exe上で現在行、イベントID、実行中コマンドを確認できるようにする", 0.6, 0.7, "シナリオ調整時の原因特定用", "Line Debug", None),
    ("Exeシナリオ読込/デバッグ", "読込失敗時のエラー表示", "シナリオ欠損や壊れたデータを画面/ログで分かるようにする", 0.5, 0.8, "原因不明の停止を防ぐ", "Error Handling", None),
    ("Exeシナリオ読込/デバッグ", "ビルド用デバッグメニューGate", "製品版で誤ってデバッグ機能が出ないようにする", 0.3, 0.5, "既存デバッグUIへのビルド差分", "Debug Gate", None),
    ("Exeシナリオ読込/デバッグ", "ログ出力/コピー機能", "ビルド中の進行ログを外部へ出せるようにする", 0.3, 0.6, "不具合報告をしやすくする", "Log Export", None),
    ("Exeシナリオ読込/デバッグ", "追加仕様フラグビューア", "表情、タッチ履歴、ED候補など追加仕様の状態を一覧表示する", 0.3, 0.6, "既存ログ表示に対する追加差分", "Flag Viewer", None),
    ("Exeシナリオ読込/デバッグ", "デバッグHUDの追加仕様タブ", "表情、タッチ、シナリオ読込だけを確認するビルド用タブを追加する", 0.4, 0.7, "既存デバッグメニューの拡張差分", "Debug HUD", None),
    ("Exeシナリオ読込/デバッグ", "追加仕様Debug Overlay詳細化", "表情状態、タッチ入力、恐怖進行、ED候補を同時に確認するOverlayを作る", 0.6, 0.7, "追加仕様に絞ったデバッグ表示", "Debug Overlay", None),
    ("Exeシナリオ読込/デバッグ", "テスト用セーブプリセット", "新パラメータ/タッチ履歴込みのテスト用セーブを切り替える", 0.5, 0.9, "既存セーブ削除では足りない追加差分", "Save Preset", None),
    ("Exeシナリオ読込/デバッグ", "デバッグコマンドライン引数", "Exe起動時にデバッグモードやシナリオパスを指定できるようにする", 0.4, 0.8, "ビルド状態の検証速度を上げる", "Debug Args", None),
    ("Exeシナリオ読込/デバッグ", "Release/Debug Build Profile分離", "外部シナリオ読込やデバッグOverlayを製品ビルドから確実に切り離す", 0.6, 0.7, "製品事故防止", "Build Profile", None),
    ("Exeシナリオ読込/デバッグ", "外部シナリオ差し替え保護", "開発用差し替えデータが製品版に混ざらないようガードする", 0.3, 0.6, "事故防止", "Build Guard", None),
    ("Exeシナリオ読込/デバッグ", "内蔵シナリオFallback", "外部シナリオ読み込みに失敗した時、同梱データへ戻す処理を入れる", 0.4, 0.6, "Exeでの停止事故を防ぐ", "Scenario Fallback", None),
    ("統合/QA", "表情×タッチ反応マトリクス確認", "主要表情とタッチ反応の組み合わせ破綻を確認する", 0.5, 0.9, "組み合わせ確認", "Matrix QA", None),
    ("統合/QA", "表情/タッチデータ検証テスト", "未定義ID、空反応、禁止状態の漏れを自動チェックする", 0.4, 0.7, "追加データの事故防止", "Data Test", None),
    ("統合/QA", "Windows Exe差分確認", "EditorではなくExeでシナリオ/Spine/デバッグ差分を確認する", 0.4, 0.8, "ビルド差分の確認", "Exe QA", None),
    ("統合/QA", "追加差分の回帰修正", "追加実装により既存LIME/疑似OS/EDが壊れていないか直す", 0.3, 0.6, "既存全体回帰ではなく追加差分のみ", "Regression", None),
    ("統合/QA", "新パラメータ用テストデータ作成", "表情、タッチ、恐怖進行を短時間で確認できるテストデータを用意する", 0.3, 0.5, "追加仕様の検証効率化", "Test Data", None),
    ("統合/QA", "分岐到達ログの確認", "10日進行で追加表情/タッチ条件が期待通り発火したかログ確認する", 0.4, 0.8, "追加仕様差分の通し確認", "Route Log", None),
    ("統合/QA", "自動シナリオSmoke Runner", "主要ルートのシナリオを自動で流し、未定義タグや停止を検出する", 0.8, 0.9, "Exe読込追加に伴う最低限の自動検査", "Scenario Smoke", None),
    ("統合/QA", "10日進行×追加仕様通し導線", "10日分の表情、タッチ、演出発火を短縮確認する通し導線を作る", 0.7, 0.8, "企画の進行軸に合わせた確認用", "10Days QA", None),
    ("統合/QA", "追加状態のセーブSchema Versioning", "表情、タッチ、恐怖進行、ED候補の保存形式にバージョンを持たせる", 0.6, 0.7, "旧セーブ破損を防ぐ", "Save Schema", None),
    ("統合/QA", "追加アセット参照のビルド前検査", "表情差分、Spine、SE、外部シナリオの参照漏れをビルド前に検査する", 0.6, 0.7, "ビルド後の欠損事故防止", "Prebuild Check", None),
    ("統合/QA", "削減/見送り演出の未参照確認", "見送りにした演出タグや素材参照が残っていないか確認する", 0.4, 0.7, "剪定後の壊れ防止", "Dead Ref Check", None),
    ("統合/QA", "Steam/製品ビルド設定との干渉確認", "デバッグGateや外部シナリオ読込が製品ビルド設定と干渉しないか確認する", 0.5, 0.9, "Exe状態の安全確認", "Build QA", None),
    ("統合/QA", "セーブマイグレーション確認", "追加した表情/タッチ/恐怖進行データで旧セーブが壊れないか確認する", 0.4, 0.8, "製品版事故防止", "Save Migration", None),
    ("統合/QA", "不確定演出の最終ON/OFF確認", "剪定した演出が設定通り有効/無効になるか確認する", 0.4, 0.8, "演出剪定の最終確認", "Feature Flag QA", None),
    ("統合/QA", "表情差分ロード順確認", "初回表示、日付変更、再読込時に表情差分のロード順が崩れないか確認する", 0.5, 0.8, "Spine/静止画混在の確認", "Load Order", None),
    ("統合/QA", "追加ED条件のDry Run", "新しい表情/タッチ履歴がED条件に影響する場合の到達確認を行う", 0.4, 0.7, "既存ED判定への追加差分", "ED Dry Run", None),
    ("統合/QA", "タッチ入力の端末差分確認", "マウス/仮想カーソル/パッド操作時のタッチ判定差を確認する", 0.3, 0.6, "入力方式差分の確認", "Input QA", None),
]

reductions = {
    129: (5.0, 5.0, 4.0, 4.0, "既存CSV/シナリオ形式を流用し、新仕様では表情/タッチ/演出タグの拡張に限定するため減額"),
    220: (1.5, 1.5, 0.7, 0.7, "チートツール単体ではなくExe用Debug Overlayへ寄せるため減額"),
    242: (2.5, 2.5, 2.0, 2.0, "10日進行と追加状態のSchema化に寄せ、リセット対象の精査を少し圧縮"),
    244: (2.0, 2.0, 1.0, 1.0, "強くてニューゲームは重い周回強化ではなく最低限の引き継ぎに留めるため減額"),
    249: (2.0, 2.0, 1.2, 1.2, "性格診断画面は新企画では主軸外のため、ED後の簡易結果表示に圧縮"),
    251: (1.0, 1.0, 0.7, 0.7, "ED条件の追加は行うが、全面互換ではなく追加状態のVersioningで吸収するため減額"),
    268: (1.0, 1.0, 0.5, 0.5, "BSOD系は採用候補として扱い、静的表示の最低限に圧縮"),
    269: (1.0, 1.0, 0.4, 0.4, "偽タスクマネージャーは製品リスクが高く、演出剪定対象として簡易化"),
    270: (1.0, 1.0, 0.5, 0.5, "偽アプリ終了は危険演出をON/OFF管理する方針にして実装範囲を圧縮"),
    272: (0.5, 0.5, 0.2, 0.2, "実写動画演出は新仕様の主軸外のため、検証・保留前提に圧縮"),
    273: (1.0, 1.0, 0.4, 0.4, "VideoPlayer制御は必須演出ではなく、採用時のみ最低限対応に圧縮"),
    274: (1.5, 1.5, 0.6, 0.6, "動画終了連携は演出採用時の簡易戻しに留めるため減額"),
    277: (2.0, 2.0, 1.5, 1.5, "SteamDeck等の広い対応は後段に回し、まずExe本体検証を優先するため減額"),
    282: (3.0, 3.0, 2.0, 2.0, "汎用デバッグボタン量産ではなく、Exe用の追加差分デバッグに絞るため減額"),
    285: (1.0, 1.0, 0.6, 0.6, "全開放は追加仕様確認に必要な範囲へ絞り、汎用アンロックは圧縮"),
    286: (1.0, 1.0, 0.6, 0.6, "セーブ削除単体ではなくテスト用セーブプリセットへ寄せるため減額"),
    287: (0.5, 0.5, 0.2, 0.2, "汎用ログ表示は追加仕様Debug Overlay/ログ出力へ統合するため減額"),
}


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def cache_for_sheet(ws):
    g_sum = sum(cell.value for cell in ws["G"] if is_number(cell.value))
    h_sum = sum(cell.value for cell in ws["H"] if is_number(cell.value))
    risk = ws["D6"].value or 0
    people = ws["D15"].value or 1
    days = ws["D16"].value or 20
    unit_price = ws["D11"].value or 0
    d8 = g_sum * risk
    d9 = h_sum * risk
    return {
        "D4": g_sum,
        "D5": h_sum,
        "D8": d8,
        "D9": d9,
        "D12": d8 * unit_price,
        "D13": d9 * unit_price,
        "D17": d8 / days / people if days and people else 0,
        "D18": d9 / days / people if days and people else 0,
    }


added_min = round(sum(task[3] for task in tasks), 1)
added_max = round(sum(task[4] for task in tasks), 1)
reduced_min = round(sum(old_min - new_min for old_min, _, new_min, _, _ in reductions.values()), 1)
reduced_max = round(sum(old_max - new_max for _, old_max, _, new_max, _ in reductions.values()), 1)
net_min = round(added_min - reduced_min, 1)
net_max = round(added_max - reduced_max, 1)
assert (added_min, added_max, reduced_min, reduced_max, net_min, net_max) == (30.7, 49.8, 10.4, 10.4, 20.3, 39.4)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb = load_workbook(SOURCE)
if TARGET_SHEET in wb.sheetnames:
    del wb[TARGET_SHEET]

source = wb[BASE_SHEET]
ws = wb.copy_worksheet(source)
ws.title = TARGET_SHEET
wb._sheets.remove(ws)
source_index = wb.sheetnames.index(BASE_SHEET)
wb._sheets.insert(source_index + 1, ws)
for sheet in wb.worksheets:
    sheet.sheet_view.tabSelected = False
ws.sheet_view.tabSelected = True
wb.active = wb.sheetnames.index(TARGET_SHEET)

labels = {
    "C4": "合計工数 MIN（生数字）",
    "C5": "合計工数 MAX（生数字）",
    "C8": "総合見積工数 MIN",
    "C9": "総合見積工数 MAX",
    "C12": "概算見積額 MIN (バッファ)(税抜)",
    "C13": "概算見積額 MAX (バッファ)(税抜)",
    "C14": "追加/削減差分（生数字）",
    "C17": "想定開発期間 MIN",
    "C18": "想定開発期間 MAX",
    "G21": "MIN工数",
    "H21": "MAX工数",
    "I21": "根拠/補足",
    "L21": "調整理由",
}
for cell, value in labels.items():
    ws[cell] = value
ws["F4"] = "原本＋追加仕様差分−削減候補"
ws["F5"] = "原本＋追加仕様差分−削減候補"
ws["D14"] = f"追加{added_min}/{added_max} 削減{reduced_min}/{reduced_max} 差分+{net_min}/+{net_max}"
ws["E14"] = "人日"
ws["F14"] = "仕様/既存重複を外し、追加仕様から増える実装差分のみを計上"
ws["D19"] = "※追加タスクは薄黄、削減した既存タスクは薄緑で表示。右端L列に調整理由を記載"

# Existing tasks become common MIN/MAX baseline.
for row in range(22, 293):
    g = ws.cell(row=row, column=7).value
    h = ws.cell(row=row, column=8).value
    if is_number(g):
        ws.cell(row=row, column=8).value = g
    elif is_number(h):
        ws.cell(row=row, column=8).value = None

add_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
section_fill = PatternFill(fill_type="solid", fgColor="FCE4D6")
reduction_fill = PatternFill(fill_type="solid", fgColor="E2F0D9")
header_fill = PatternFill(fill_type="solid", fgColor="D9EAD3")
section_font = Font(bold=True)
wrap = Alignment(wrap_text=True, vertical="center")
center = Alignment(vertical="center")
detail_row_height = source.row_dimensions[22].height or source.sheet_format.defaultRowHeight or 15.75
thin_black = Side(style="thin", color="000000")
table_border = Border(left=thin_black, right=thin_black, top=thin_black, bottom=thin_black)

for row, (_, _, new_min, new_max, reason) in reductions.items():
    ws.cell(row=row, column=7).value = new_min
    ws.cell(row=row, column=8).value = new_max
    ws.cell(row=row, column=12).value = reason
    for col in range(2, 13):
        ws.cell(row=row, column=col).border = table_border
        ws.cell(row=row, column=col).alignment = center
        if col >= 3:
            ws.cell(row=row, column=col).fill = reduction_fill
    ws.row_dimensions[row].height = detail_row_height

start = 293
for offset, task_data in enumerate(tasks):
    row = start + offset
    section, task, desc, min_days, max_days, reason, note, comment = task_data
    prev_section = tasks[offset - 1][0] if offset else None
    ws.cell(row=row, column=2).value = f'=IF(ISBLANK(E{row}), "", COUNTIF($E$25:E{row}, "<>"))'
    ws.cell(row=row, column=3).value = section if section != prev_section else None
    ws.cell(row=row, column=4).value = section
    ws.cell(row=row, column=5).value = task
    ws.cell(row=row, column=6).value = desc
    ws.cell(row=row, column=7).value = min_days
    ws.cell(row=row, column=8).value = max_days
    ws.cell(row=row, column=9).value = reason
    ws.cell(row=row, column=10).value = note
    ws.cell(row=row, column=11).value = comment
    ws.cell(row=row, column=12).value = "既存項目との差分のみ追加"
    for col in range(2, 13):
        ws.cell(row=row, column=col).border = table_border
        ws.cell(row=row, column=col).alignment = center
        if col >= 3:
            ws.cell(row=row, column=col).fill = add_fill
    if ws.cell(row=row, column=3).value:
        ws.cell(row=row, column=3).fill = section_fill
        ws.cell(row=row, column=3).font = section_font
    ws.row_dimensions[row].height = detail_row_height

for col in range(3, 13):
    ws.cell(row=21, column=col).fill = header_fill
    ws.cell(row=21, column=col).font = Font(bold=True)
    ws.cell(row=21, column=col).alignment = center
    ws.cell(row=21, column=col).border = table_border

for col_letter, width in {
    "B": 4.5,
    "C": 28,
    "D": 34,
    "E": 38,
    "F": 56,
    "G": 10,
    "H": 10,
    "I": 36,
    "J": 22,
    "K": 26,
    "L": 42,
}.items():
    ws.column_dimensions[col_letter].width = width

last_detail_row = start + len(tasks) - 1
for row in range(21, last_detail_row + 1):
    for col in range(2, 13):
        ws.cell(row=row, column=col).border = table_border

for row in range(293, start + len(tasks)):
    ws.row_dimensions[row].height = detail_row_height
for row in range(3, 22):
    ws.row_dimensions[row].height = 18
ws.row_dimensions[14].height = 24
ws["D14"].alignment = wrap

ws["D4"] = "=SUM(G:G)"
ws["D5"] = "=SUM(H:H)"
ws["D8"] = "=D4*D6"
ws["D9"] = "=D5*D6"
ws["D12"] = "=D8*D11"
ws["D13"] = "=D9*D11"
ws["D17"] = "=D8/D16/D15"
ws["D18"] = "=D9/D16/D15"

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUTPUT)

# Cache summary values so previews show numbers before Excel recalculates.
wb_values = load_workbook(OUTPUT, data_only=False)
target_sheets = {name: cache_for_sheet(wb_values[name]) for name in ["工数表ベース", "工数表ベース原本", TARGET_SHEET]}

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
ET.register_namespace("", NS_MAIN)
ET.register_namespace("r", NS_REL)

def q(ns, tag):
    return f"{{{ns}}}{tag}"

with ZipFile(OUTPUT, "r") as zin:
    workbook_xml = ET.fromstring(zin.read("xl/workbook.xml"))
    rels_xml = ET.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
    rid_to_target = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in rels_xml.findall(q(NS_PKG_REL, "Relationship"))
    }
    sheet_to_part = {}
    for sheet in workbook_xml.find(q(NS_MAIN, "sheets")).findall(q(NS_MAIN, "sheet")):
        name = sheet.attrib["name"]
        if name not in target_sheets:
            continue
        rid = sheet.attrib[q(NS_REL, "id")]
        target = rid_to_target[rid].lstrip("/")
        if not target.startswith("xl/"):
            target = f"xl/{target}"
        sheet_to_part[name] = target

    tmp = NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp_path = Path(tmp.name)
    tmp.close()
    with ZipFile(tmp_path, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            sheet_name = next((name for name, part in sheet_to_part.items() if part == item.filename), None)
            if sheet_name:
                root = ET.fromstring(data)
                wanted = target_sheets[sheet_name]
                sheet_data = root.find(q(NS_MAIN, "sheetData"))
                for xrow in sheet_data.findall(q(NS_MAIN, "row")):
                    for cell in xrow.findall(q(NS_MAIN, "c")):
                        ref = re.sub(r"\$|'.*'!", "", cell.attrib.get("r", ""))
                        if ref not in wanted:
                            continue
                        v = cell.find(q(NS_MAIN, "v"))
                        if v is None:
                            v = ET.SubElement(cell, q(NS_MAIN, "v"))
                        v.text = f"{wanted[ref]:.10g}"
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            zout.writestr(item, data)
tmp_path.replace(OUTPUT)

print(OUTPUT)
print("added", added_min, added_max, "reduced", reduced_min, reduced_max, "net", net_min, net_max)
print(target_sheets[TARGET_SHEET])
