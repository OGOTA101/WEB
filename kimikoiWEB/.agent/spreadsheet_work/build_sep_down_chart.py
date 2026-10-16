from __future__ import annotations

import os
import shutil
import subprocess
import xml.etree.ElementTree as ET
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook


ROOT = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB")
PYTHON = Path(r"C:\Users\ihogo\.cache\codex-runtimes\codex-primary-runtime\dependencies\python\python.exe")
WORK = ROOT / ".agent" / "spreadsheet_work"
SOURCE = WORK / "down_chart_clean_source.xlsx"
OUTPUT = ROOT / "outputs" / "kimikoi_kousu" / "きみこい・工数表_ダウンチャート_9月開始.xlsx"
SHEET = "本番工数"
CACHES = {
    "D4": 121.8,
    "D5": 143.8,
    "D8": 152.25,
    "D9": 179.75,
    "D17": 7.6125,
    "D18": 8.9875,
}


def sheet_xml_path(root: Path, sheet_name: str) -> Path:
    ns = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pkgrel": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    workbook_tree = ET.parse(root / "xl" / "workbook.xml")
    sheet_el = None
    for candidate in workbook_tree.getroot().findall(".//main:sheet", ns):
        if candidate.attrib.get("name") == sheet_name:
            sheet_el = candidate
            break
    if sheet_el is None:
        raise ValueError(sheet_name)
    rel_id = sheet_el.attrib[f"{{{ns['rel']}}}id"]
    rels_tree = ET.parse(root / "xl" / "_rels" / "workbook.xml.rels")
    for rel in rels_tree.getroot().findall("pkgrel:Relationship", ns):
        if rel.attrib.get("Id") == rel_id:
            target = rel.attrib["Target"]
            if target.startswith("/"):
                return root / target.lstrip("/")
            if target.startswith("worksheets/"):
                return root / "xl" / target
            return root / "xl" / target
    raise ValueError(rel_id)


if not SOURCE.exists():
    subprocess.run([str(PYTHON), str(WORK / "build_clean_down_chart.py")], check=True, cwd=str(ROOT))

if OUTPUT.exists():
    OUTPUT.unlink()

env = os.environ.copy()
env["KIMIKOI_GANTT_SOURCE"] = str(SOURCE)
env["KIMIKOI_GANTT_OUTPUT"] = str(OUTPUT)
env["KIMIKOI_GANTT_BASE_SHEET"] = "本番工数"
env["KIMIKOI_GANTT_START"] = "2026-09-01"
subprocess.run([str(PYTHON), str(WORK / "build_gantt_schedule.py")], check=True, cwd=str(ROOT), env=env)

wb = load_workbook(OUTPUT)
if "横長工程表" in wb.sheetnames:
    wb["横長工程表"].title = "ダウンチャート"

ws = wb["ダウンチャート"]
ws["A1"] = "きみこい 本番工数 ダウンチャート"
ws["D2"] = "MAX工数合計"
ws["E2"] = "バッファ係数"
ws["F2"] = "予定工数合計"
ws["I5"] = "根拠/備考"
ws["E5"] = "MAX工数"

for sheet in wb.worksheets:
    sheet.sheet_view.tabSelected = False
wb["本番工数"].sheet_view.tabSelected = True
wb.active = wb.sheetnames.index("本番工数")

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUTPUT)

ns_uri = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ns = {"main": ns_uri}
ET.register_namespace("", ns_uri)

with TemporaryDirectory() as tmp:
    tmp_root = Path(tmp)
    with ZipFile(OUTPUT, "r") as zin:
        zin.extractall(tmp_root)

    xml_path = sheet_xml_path(tmp_root, SHEET)
    tree = ET.parse(xml_path)
    root = tree.getroot()
    cells = {cell.attrib["r"]: cell for cell in root.findall(".//main:c", ns) if "r" in cell.attrib}
    for ref, value in CACHES.items():
        cell = cells.get(ref)
        if cell is None:
            continue
        for value_el in cell.findall("main:v", ns):
            cell.remove(value_el)
        v = ET.SubElement(cell, f"{{{ns_uri}}}v")
        cell.attrib.pop("t", None)
        v.text = str(value)
    tree.write(xml_path, encoding="utf-8", xml_declaration=True)

    patched = OUTPUT.with_suffix(".patched.xlsx")
    if patched.exists():
        patched.unlink()
    with ZipFile(patched, "w", ZIP_DEFLATED) as zout:
        for file in tmp_root.rglob("*"):
            if file.is_file():
                zout.write(file, file.relative_to(tmp_root).as_posix())
    shutil.move(patched, OUTPUT)

print(OUTPUT)
