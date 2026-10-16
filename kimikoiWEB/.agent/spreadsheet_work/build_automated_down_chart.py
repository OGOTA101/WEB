from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB")
SOURCE = ROOT / "outputs" / "kimikoi_kousu" / "きみこい・工数表_ダウンチャート.xlsx"
OUTPUT = ROOT / "outputs" / "kimikoi_kousu" / "きみこい・工数表_ダウンチャート_自動化版.xlsx"

BASE_SHEET = "本番工数"
CHART_SHEET = "ダウンチャート"
HOLIDAY_SHEET = "祝日設定"

SOURCE_FIRST_ROW = 22
SOURCE_LAST_ROW = 1019
TASK_ROWS = 520
DATE_DAYS = 730
FIRST_TASK_ROW = 6
FIRST_DATE_COL = 14  # N
HELPER_COLS = range(10, 14)  # J:M

HOLIDAYS = [
    date(2026, 1, 1),
    date(2026, 1, 12),
    date(2026, 2, 11),
    date(2026, 2, 23),
    date(2026, 3, 20),
    date(2026, 4, 29),
    date(2026, 5, 3),
    date(2026, 5, 4),
    date(2026, 5, 5),
    date(2026, 5, 6),
    date(2026, 7, 20),
    date(2026, 8, 11),
    date(2026, 9, 21),
    date(2026, 9, 22),
    date(2026, 9, 23),
    date(2026, 10, 12),
    date(2026, 11, 3),
    date(2026, 11, 23),
    date(2027, 1, 1),
    date(2027, 1, 11),
    date(2027, 2, 11),
    date(2027, 2, 23),
    date(2027, 3, 21),
    date(2027, 3, 22),
    date(2027, 4, 29),
    date(2027, 5, 3),
    date(2027, 5, 4),
    date(2027, 5, 5),
    date(2027, 7, 19),
    date(2027, 8, 11),
    date(2027, 9, 20),
    date(2027, 9, 23),
    date(2027, 10, 11),
    date(2027, 11, 3),
    date(2027, 11, 23),
    date(2028, 1, 1),
    date(2028, 1, 10),
    date(2028, 2, 11),
    date(2028, 2, 23),
    date(2028, 3, 20),
    date(2028, 4, 29),
    date(2028, 5, 3),
    date(2028, 5, 4),
    date(2028, 5, 5),
]


def f(ref: str) -> str:
    return f"'{BASE_SHEET}'!{ref}"


def date_from_work_index(index_ref: str) -> str:
    holidays = f"'{HOLIDAY_SHEET}'!$A$2:$A$80"
    return (
        f'IF($H$3="休みなし",$A$3+{index_ref}-1,'
        f'IF($H$3="土日休み",WORKDAY.INTL($A$3-1,{index_ref},1),'
        f"WORKDAY.INTL($A$3-1,{index_ref},1,{holidays})))"
    )


wb = load_workbook(SOURCE)
ws_base = wb[BASE_SHEET]

ws_base["AF21"] = "自動用_大項目"
ws_base["AG21"] = "自動用_中項目"
for source_row in range(SOURCE_FIRST_ROW, SOURCE_LAST_ROW + 1):
    previous = source_row - 1
    ws_base[f"AF{source_row}"] = (
        f'=IF($C{source_row}<>"",$C{source_row},IF(ROW()={SOURCE_FIRST_ROW},"",AF{previous}))'
    )
    ws_base[f"AG{source_row}"] = (
        f'=IF($D{source_row}<>"",$D{source_row},IF(ROW()={SOURCE_FIRST_ROW},"",AG{previous}))'
    )
for helper_col in ("AF", "AG"):
    ws_base.column_dimensions[helper_col].hidden = True

if CHART_SHEET in wb.sheetnames:
    del wb[CHART_SHEET]
if HOLIDAY_SHEET in wb.sheetnames:
    del wb[HOLIDAY_SHEET]

ws = wb.create_sheet(CHART_SHEET, index=1)
ws_holidays = wb.create_sheet(HOLIDAY_SHEET)
ws_holidays.sheet_state = "hidden"

ws_holidays["A1"] = "祝日"
for row, holiday in enumerate(HOLIDAYS, start=2):
    ws_holidays.cell(row, 1, holiday)
    ws_holidays.cell(row, 1).number_format = "yyyy/m/d"

dark = "1F4E5F"
dark2 = "173A47"
white = "FFFFFF"
input_fill = PatternFill("solid", fgColor="FFF2CC")
linked_fill = PatternFill("solid", fgColor="DDEBF7")
header_fill = PatternFill("solid", fgColor=dark)
subheader_fill = PatternFill("solid", fgColor=dark2)
weekend_fill = PatternFill("solid", fgColor="E7E6E6")
holiday_fill = PatternFill("solid", fgColor="D9D2E9")
bar_fill = PatternFill("solid", fgColor="9DC3E6")
green_fill = PatternFill("solid", fgColor="E2F0D9")
thin = Side(style="thin", color="B7B7B7")
hair = Side(style="hair", color="D9D9D9")
medium = Side(style="medium", color="666666")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
date_border = Border(left=hair, right=hair, top=thin, bottom=thin)

ws.sheet_view.showGridLines = False
ws.freeze_panes = "N6"

ws.merge_cells("A1:I1")
ws["A1"] = "きみこい 本番工数 ダウンチャート（自動化版）"
ws["A1"].fill = header_fill
ws["A1"].font = Font(bold=True, color=white, size=16)
ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 28

source_h_range = f"$H${SOURCE_FIRST_ROW}:$H${SOURCE_LAST_ROW}"
source_h_ref = f(source_h_range)

summary = [
    ("開始日", date(2026, 6, 1)),
    ("終了目安", f"=IFERROR(MAX($H${FIRST_TASK_ROW}:$H${FIRST_TASK_ROW + TASK_ROWS - 1}),\"\")"),
    ("タスク数", f'=COUNTIF({source_h_ref},">0")'),
    ("MAX工数合計", f'=SUMIF({source_h_ref},">0",{source_h_ref})'),
    ("バッファ係数", 1.25),
    ("予定工数合計", "=D3*E3"),
    ("日次消化力", 1.0),
    ("休日前提", "土日祝休み"),
    ("表示日数", "=MIN(730,MAX(365,IFERROR(B3-A3+30,365)))"),
]

for col, (label, value) in enumerate(summary, start=1):
    header = ws.cell(2, col, label)
    header.fill = subheader_fill
    header.font = Font(bold=True, color=white, size=9)
    header.alignment = Alignment(horizontal="center", vertical="center")
    header.border = border
    cell = ws.cell(3, col, value)
    cell.fill = input_fill if col in {1, 5, 7, 8} else linked_fill
    cell.font = Font(bold=True, size=9)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border

ws["A3"].number_format = "yyyy/m/d"
ws["B3"].number_format = "yyyy/m/d"
for ref in ("D3", "E3", "F3", "G3"):
    ws[ref].number_format = "0.0"

ws.merge_cells("A4:I4")
ws["A4"] = "黄色セルを変更すると、タスク数・終了目安・日付列・バー表示がExcel再計算で更新されます。タスク内容と工数は「本番工数」シートのMAX工数が0より大きい行から取得します。"
ws["A4"].fill = input_fill
ws["A4"].font = Font(size=9)
ws["A4"].alignment = Alignment(wrap_text=True, vertical="center")
ws["A4"].border = border
ws.row_dimensions[4].height = 28

dv = DataValidation(type="list", formula1='"土日祝休み,土日休み,休みなし"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(ws["H3"])

left_headers = ["No", "大項目", "中項目", "タスク", "MAX工数", "予定工数", "開始日", "終了日", "根拠/備考"]
for col, header in enumerate(left_headers, start=1):
    cell = ws.cell(5, col, header)
    cell.fill = header_fill
    cell.font = Font(bold=True, color=white)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(left=thin, right=thin, top=medium, bottom=medium)

helper_headers = ["元行", "累計前", "開始idx", "終了idx"]
for col, header in zip(HELPER_COLS, helper_headers):
    ws.cell(5, col, header)
    ws.column_dimensions[get_column_letter(col)].hidden = True

widths = {"A": 6, "B": 18, "C": 20, "D": 44, "E": 10, "F": 10, "G": 12, "H": 12, "I": 44}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

for offset in range(DATE_DAYS):
    col = FIRST_DATE_COL + offset
    letter = get_column_letter(col)
    ws.column_dimensions[letter].width = 3
    date_formula = f'=IF(COLUMN()-COLUMN($N$4)+1>$I$3,"",$A$3+COLUMN()-COLUMN($N$4))'
    month_formula = f'=IF({letter}$4="","",IF(OR(COLUMN()=COLUMN($N$4),DAY({letter}$4)=1),TEXT({letter}$4,"yyyy/m"),""))'
    weekday_formula = f'=IF({letter}$4="","",CHOOSE(WEEKDAY({letter}$4,2),"月","火","水","木","金","土","日"))'

    for row, formula in ((3, month_formula), (4, date_formula), (5, weekday_formula)):
        cell = ws.cell(row, col, formula)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = date_border
        cell.font = Font(size=8)
    ws.cell(3, col).fill = header_fill
    ws.cell(3, col).font = Font(size=8, bold=True, color=white)
    ws.cell(4, col).number_format = "d"

for idx in range(TASK_ROWS):
    row = FIRST_TASK_ROW + idx
    n = f"ROWS($A${FIRST_TASK_ROW}:A{row})"
    source_row_formula = (
        f'=IF({n}>$C$3,"",AGGREGATE(15,6,ROW({f(source_h_range)})/({f(source_h_range)}>0),{n}))'
    )
    major_lookup = f'=IF($J{row}="","",INDEX({f("$AF:$AF")},$J{row}))'
    mid_lookup = f'=IF($J{row}="","",INDEX({f("$AG:$AG")},$J{row}))'
    task_formula = f'=IF($J{row}="","",IF(INDEX({f("$E:$E")},$J{row})="",C{row},INDEX({f("$E:$E")},$J{row})))'
    max_formula = f'=IF($J{row}="","",INDEX({f("$H:$H")},$J{row}))'
    scheduled_formula = f'=IF($J{row}="","",E{row}*$E$3)'
    prev_formula = f'=IF($J{row}="","",SUM($F${FIRST_TASK_ROW}:F{row - 1}))' if row > FIRST_TASK_ROW else f'=IF($J{row}="","",0)'
    start_idx_formula = f'=IF($J{row}="","",INT(K{row}/$G$3)+1)'
    end_idx_formula = f'=IF($J{row}="","",ROUNDUP((K{row}+F{row})/$G$3,0))'
    start_date_formula = f'=IF($J{row}="","",{date_from_work_index(f"L{row}")})'
    end_date_formula = f'=IF($J{row}="","",{date_from_work_index(f"M{row}")})'
    note_formula = (
        f'=IF($J{row}="","",TEXTJOIN(" / ",TRUE,INDEX({f("$I:$I")},$J{row}),'
        f'INDEX({f("$J:$J")},$J{row}),INDEX({f("$L:$L")},$J{row})))'
    )

    formulas = [
        f'=IF($J{row}="","",{idx + 1})',
        major_lookup,
        mid_lookup,
        task_formula,
        max_formula,
        scheduled_formula,
        start_date_formula,
        end_date_formula,
        note_formula,
        source_row_formula,
        prev_formula,
        start_idx_formula,
        end_idx_formula,
    ]
    for col, formula in enumerate(formulas, start=1):
        cell = ws.cell(row, col, formula)
        cell.border = border if col <= 9 else Border()
        cell.font = Font(size=8)
        cell.alignment = Alignment(
            horizontal="left" if col in {2, 3, 4, 9} else "center",
            vertical="center",
            wrap_text=col in {4, 9},
        )
        if col in {5, 6}:
            cell.fill = linked_fill
            cell.number_format = "0.0"
        elif col in {7, 8}:
            cell.number_format = "yyyy/m/d"
        elif col == 2:
            cell.fill = green_fill
    ws.row_dimensions[row].height = 18

    for col in range(FIRST_DATE_COL, FIRST_DATE_COL + DATE_DAYS):
        cell = ws.cell(row, col)
        cell.border = Border(left=hair, right=hair, top=hair, bottom=hair)

last_task_row = FIRST_TASK_ROW + TASK_ROWS - 1
last_date_col = FIRST_DATE_COL + DATE_DAYS - 1
last_date_letter = get_column_letter(last_date_col)
date_range = f"{get_column_letter(FIRST_DATE_COL)}4:{last_date_letter}{last_task_row}"
task_bar_range = f"{get_column_letter(FIRST_DATE_COL)}{FIRST_TASK_ROW}:{last_date_letter}{last_task_row}"

first_date_letter = get_column_letter(FIRST_DATE_COL)
weekend_rule = FormulaRule(
    formula=[f'=AND({first_date_letter}$4<>"",$H$3<>"休みなし",WEEKDAY({first_date_letter}$4,2)>5)'],
    fill=weekend_fill,
)
holiday_rule = FormulaRule(
    formula=[
        f'=AND({first_date_letter}$4<>"",$H$3="土日祝休み",COUNTIF(\'{HOLIDAY_SHEET}\'!$A:$A,{first_date_letter}$4)>0)'
    ],
    fill=holiday_fill,
)
bar_rule = FormulaRule(
    formula=[
        f'=AND({first_date_letter}$4<>"",$J{FIRST_TASK_ROW}<>"",{first_date_letter}$4>=$G{FIRST_TASK_ROW},'
        f'{first_date_letter}$4<=$H{FIRST_TASK_ROW},OR($H$3="休みなし",AND($H$3="土日休み",WEEKDAY({first_date_letter}$4,2)<=5),'
        f'AND($H$3="土日祝休み",WEEKDAY({first_date_letter}$4,2)<=5,COUNTIF(\'{HOLIDAY_SHEET}\'!$A:$A,{first_date_letter}$4)=0)))'
    ],
    fill=bar_fill,
)
ws.conditional_formatting.add(date_range, weekend_rule)
ws.conditional_formatting.add(date_range, holiday_rule)
ws.conditional_formatting.add(task_bar_range, bar_rule)

legend_row = last_task_row + 3
ws.cell(legend_row, 1, "凡例")
ws.cell(legend_row, 1).font = Font(bold=True)
legend_items = [("予定作業", "9DC3E6"), ("土日", "E7E6E6"), ("祝日", "D9D2E9"), ("編集セル", "FFF2CC")]
for i, (label, color) in enumerate(legend_items, start=2):
    cell = ws.cell(legend_row, i, label)
    cell.fill = PatternFill("solid", fgColor=color)
    cell.border = border
    cell.alignment = Alignment(horizontal="center")

ws.auto_filter.ref = f"A5:I{last_task_row}"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0

for row in range(1, 6):
    ws.row_dimensions[row].height = 20
ws.row_dimensions[1].height = 28
ws.row_dimensions[4].height = 30
ws.row_dimensions[5].height = 22

for sheet in wb.worksheets:
    sheet.sheet_view.tabSelected = False
ws.sheet_view.tabSelected = True
wb.active = wb.sheetnames.index(CHART_SHEET)

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)
print(OUTPUT)
print("task_rows", TASK_ROWS)
print("date_days", DATE_DAYS)
