from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB")
SOURCE = ROOT / "outputs" / "kimikoi_kousu" / "きみこい・工数表_ダウンチャート.xlsx"
OUTPUT = ROOT / "outputs" / "kimikoi_kousu" / "きみこい・工数単価算出表_660万円_項目別.xlsx"

TARGET_AMOUNT = 6_600_000
HOURS_PER_DAY = 8.0
BUFFER_RATE = 1.25


def as_number(value) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value)
        except ValueError:
            return None
    return None


def clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


source_wb = load_workbook(SOURCE, data_only=True)
source_ws = source_wb["本番工数"]

major_totals = defaultdict(lambda: {"tasks": 0, "min": 0.0, "max": 0.0, "notes": []})
mid_totals = defaultdict(lambda: {"tasks": 0, "min": 0.0, "max": 0.0, "details": []})
major = ""
mid = ""

for row in range(22, source_ws.max_row + 1):
    major_value = clean_text(source_ws.cell(row, 3).value)
    mid_value = clean_text(source_ws.cell(row, 4).value)
    detail_value = clean_text(source_ws.cell(row, 5).value)
    max_effort = as_number(source_ws.cell(row, 8).value)
    min_effort = as_number(source_ws.cell(row, 7).value) or 0.0
    note = clean_text(source_ws.cell(row, 12).value)

    if major_value:
        major = major_value
    if mid_value:
        mid = mid_value
    if max_effort is None or max_effort <= 0:
        continue

    title = detail_value or mid or major
    major_totals[major]["tasks"] += 1
    major_totals[major]["min"] += min_effort
    major_totals[major]["max"] += max_effort
    if note and len(major_totals[major]["notes"]) < 3:
        major_totals[major]["notes"].append(note)

    key = (major, mid or "その他")
    mid_totals[key]["tasks"] += 1
    mid_totals[key]["min"] += min_effort
    mid_totals[key]["max"] += max_effort
    if title and len(mid_totals[key]["details"]) < 4:
        mid_totals[key]["details"].append(title)

total_raw_max = round(sum(v["max"] for v in major_totals.values()), 2)
total_raw_min = round(sum(v["min"] for v in major_totals.values()), 2)
total_days = round(total_raw_max * BUFFER_RATE, 2)
total_hours = round(total_days * HOURS_PER_DAY, 2)
hourly_rate = TARGET_AMOUNT / total_hours
daily_rate = hourly_rate * HOURS_PER_DAY
hourly_rate_display = round(hourly_rate, 1)
daily_rate_display = round(daily_rate, 1)

wb = Workbook()
ws = wb.active
ws.title = "工数単価算出表"
ws.sheet_view.showGridLines = False

dark = "1F4E5F"
dark2 = "173A47"
blue = "D9EAF7"
green = "E2F0D9"
yellow = "FFF2CC"
white = "FFFFFF"
gray = "F2F2F2"
line = Side(style="thin", color="B7B7B7")
medium = Side(style="medium", color="666666")
border = Border(left=line, right=line, top=line, bottom=line)

ws.merge_cells("A1:J1")
ws["A1"] = "きみこい 本番実装 工数単価算出表"
ws["A1"].fill = PatternFill("solid", fgColor=dark)
ws["A1"].font = Font(bold=True, color=white, size=16)
ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 30

ws.merge_cells("A2:J2")
ws["A2"] = "本番工数表の積上げ工数を時間換算し、税抜660万円を時間単価として逆算した概算データ"
ws["A2"].fill = PatternFill("solid", fgColor=dark2)
ws["A2"].font = Font(color=white, size=10)
ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

assumptions = [
    ("対象金額", TARGET_AMOUNT, "円", "税抜想定"),
    ("採用工数", total_days, "人日", "MAX工数に調整係数を加味"),
    ("合計時間", total_hours, "時間", f"1人日={HOURS_PER_DAY:g}時間換算"),
    ("算出時間単価", hourly_rate_display, "円/h", "対象金額 ÷ 合計時間"),
    ("算出人日単価", daily_rate_display, "円/人日", "算出時間単価 × 8時間"),
    ("生工数 MIN", total_raw_min, "人日", "参考値"),
    ("生工数 MAX", total_raw_max, "人日", "見積採用前の積上げ値"),
    ("調整係数", BUFFER_RATE, "倍", "デバッグ・統合QA・演出確認の予備"),
    ("開始時期", "6月/9月いずれも同額", "", "同一スコープ前提では総工数で算出"),
]

for col, header in enumerate(["前提項目", "値", "単位", "算出メモ"], start=1):
    cell = ws.cell(4, col, header)
    cell.fill = PatternFill("solid", fgColor=dark)
    cell.font = Font(bold=True, color=white)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border

for row_offset, (label, value, unit, note) in enumerate(assumptions, start=5):
    values = [label, value, unit, note]
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row_offset, col, value)
        cell.border = border
        cell.alignment = Alignment(horizontal="left" if col in {1, 4} else "right", vertical="center", wrap_text=col == 4)
        if col == 2:
            cell.fill = PatternFill("solid", fgColor=yellow if row_offset <= 9 else blue)
            cell.font = Font(bold=True)
        if label == "対象金額" and col == 2:
            cell.number_format = '¥#,##0'
        elif label in {"算出時間単価", "算出人日単価"} and col == 2:
            cell.number_format = '¥#,##0.0'
        elif isinstance(value, float):
            cell.number_format = "0.0"

section_row = 16
ws.merge_cells(start_row=section_row, start_column=1, end_row=section_row, end_column=10)
ws.cell(section_row, 1, "大項目別 工数・時間単価算出")
ws.cell(section_row, 1).fill = PatternFill("solid", fgColor=dark)
ws.cell(section_row, 1).font = Font(bold=True, color=white, size=12)

headers = ["No", "大項目", "タスク数", "MIN人日", "MAX人日", "採用人日", "採用時間", "時間単価", "算出金額", "補足"]
header_row = section_row + 1
for col, header in enumerate(headers, start=1):
    cell = ws.cell(header_row, col, header)
    cell.fill = PatternFill("solid", fgColor=dark2)
    cell.font = Font(bold=True, color=white)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(left=line, right=line, top=medium, bottom=medium)

major_start = header_row + 1
major_items = sorted(major_totals.items(), key=lambda item: item[0])
running_cost = 0
for index, (name, data) in enumerate(major_items, start=1):
    row = major_start + index - 1
    adopted_days = data["max"] * BUFFER_RATE
    hours = adopted_days * HOURS_PER_DAY
    if index == len(major_items):
        amount = TARGET_AMOUNT - running_cost
    else:
        amount = round(hours * hourly_rate)
        running_cost += amount
    values = [
        index,
        name,
        data["tasks"],
        data["min"],
        data["max"],
        adopted_days,
        hours,
        hourly_rate_display,
        amount,
        " / ".join(data["notes"]) or "本番実装工数として集計",
    ]
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row, col, value)
        cell.border = border
        cell.alignment = Alignment(horizontal="left" if col in {2, 10} else "right", vertical="center", wrap_text=col == 10)
        if col in {4, 5, 6, 7, 8, 9}:
            cell.fill = PatternFill("solid", fgColor=blue if col != 9 else green)
        if col in {4, 5, 6, 7}:
            cell.number_format = "0.0"
        elif col in {8, 9}:
            cell.number_format = '¥#,##0'
    ws.row_dimensions[row].height = 42

total_row = major_start + len(major_items)
total_values = [
    "",
    "合計",
    sum(v["tasks"] for v in major_totals.values()),
    total_raw_min,
    total_raw_max,
    total_days,
    total_hours,
    hourly_rate_display,
    TARGET_AMOUNT,
    "合計金額が660万円になるよう、総時間から単価を逆算",
]
for col, value in enumerate(total_values, start=1):
    cell = ws.cell(total_row, col, value)
    cell.fill = PatternFill("solid", fgColor=dark)
    cell.font = Font(bold=True, color=white)
    cell.border = Border(left=line, right=line, top=medium, bottom=medium)
    cell.alignment = Alignment(horizontal="left" if col in {2, 10} else "right", vertical="center", wrap_text=col == 10)
    if col in {4, 5, 6, 7}:
        cell.number_format = "0.0"
    elif col in {8, 9}:
        cell.number_format = '¥#,##0'

detail_row = total_row + 3
ws.merge_cells(start_row=detail_row, start_column=1, end_row=detail_row, end_column=10)
ws.cell(detail_row, 1, "中項目別 内訳")
ws.cell(detail_row, 1).fill = PatternFill("solid", fgColor=dark)
ws.cell(detail_row, 1).font = Font(bold=True, color=white, size=12)

detail_headers = ["No", "大項目", "中項目", "タスク数", "MAX人日", "採用人日", "採用時間", "算出金額", "主な作業", "備考"]
for col, header in enumerate(detail_headers, start=1):
    cell = ws.cell(detail_row + 1, col, header)
    cell.fill = PatternFill("solid", fgColor=dark2)
    cell.font = Font(bold=True, color=white)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(left=line, right=line, top=medium, bottom=medium)

detail_start = detail_row + 2
mid_items = sorted(mid_totals.items(), key=lambda item: (item[0][0], item[0][1]))
for index, ((major_name, mid_name), data) in enumerate(mid_items, start=1):
    row = detail_start + index - 1
    adopted_days = data["max"] * BUFFER_RATE
    hours = adopted_days * HOURS_PER_DAY
    amount = round(hours * hourly_rate)
    values = [
        index,
        major_name,
        mid_name,
        data["tasks"],
        data["max"],
        adopted_days,
        hours,
        amount,
        " / ".join(data["details"]),
        "工数ベース",
    ]
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row, col, value)
        cell.border = border
        cell.alignment = Alignment(horizontal="left" if col in {2, 3, 9, 10} else "right", vertical="center", wrap_text=col in {9, 10})
        if col in {5, 6, 7, 8}:
            cell.fill = PatternFill("solid", fgColor=blue if col != 8 else green)
        if col in {5, 6, 7}:
            cell.number_format = "0.0"
        elif col == 8:
            cell.number_format = '¥#,##0'
    ws.row_dimensions[row].height = 24

note_row = detail_start + len(mid_items) + 2
ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)
ws.cell(note_row, 1, "備考: 本表は、本番工数表の積上げタスクを時間換算し、対象金額660万円を総時間で割った単価から算出した概算です。追加仕様・素材制作・シナリオ制作・大幅な演出追加は別途見積前提です。")
ws.cell(note_row, 1).fill = PatternFill("solid", fgColor=gray)
ws.cell(note_row, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws.cell(note_row, 1).border = border
ws.row_dimensions[note_row].height = 40

chart = BarChart()
chart.title = "大項目別 算出金額"
chart.y_axis.title = "円"
chart.x_axis.title = "大項目"
data = Reference(ws, min_col=9, min_row=header_row, max_row=total_row - 1)
cats = Reference(ws, min_col=2, min_row=major_start, max_row=total_row - 1)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.height = 7
chart.width = 16
chart.legend = None
ws.add_chart(chart, "F4")

widths = {
    "A": 16,
    "B": 26,
    "C": 14,
    "D": 42,
    "E": 12,
    "F": 12,
    "G": 12,
    "H": 14,
    "I": 16,
    "J": 56,
}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

for row in range(1, ws.max_row + 1):
    if ws.row_dimensions[row].height is None:
        ws.row_dimensions[row].height = 20

ws.freeze_panes = "A18"
ws.auto_filter.ref = f"A{header_row}:J{total_row}"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)

print(OUTPUT)
print("raw_min", total_raw_min)
print("raw_max", total_raw_max)
print("days", total_days)
print("hours", total_hours)
print("hourly_rate", round(hourly_rate, 2))
