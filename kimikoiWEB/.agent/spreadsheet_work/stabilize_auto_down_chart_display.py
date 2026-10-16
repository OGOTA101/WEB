from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


WORKBOOK = Path(r"C:\Users\ihogo\Desktop\WEB\kimikoiWEB\outputs\kimikoi_kousu\きみこい・工数表_ダウンチャート_自動化完成版.xlsx")
SHEET = "ダウンチャート"
TASK_START_ROW = 6

VISIBLE_COLS_TO_FIX = range(1, 10)  # A:I
HELPER_COLS_TO_HIDE = range(10, 15)  # J:N


data_wb = load_workbook(WORKBOOK, data_only=True)
data_ws = data_wb[SHEET]

last_task_row = TASK_START_ROW
for row in range(TASK_START_ROW, data_ws.max_row + 1):
    if data_ws.cell(row, 1).value:
        last_task_row = row

visible_values = {
    (row, col): data_ws.cell(row, col).value
    for row in range(TASK_START_ROW, last_task_row + 1)
    for col in VISIBLE_COLS_TO_FIX
}

wb = load_workbook(WORKBOOK)
ws = wb[SHEET]

thin = Side(style="thin", color="B7B7B7")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
linked_fill = PatternFill("solid", fgColor="DDEBF7")

for row in range(TASK_START_ROW, last_task_row + 1):
    for col in VISIBLE_COLS_TO_FIX:
        cell = ws.cell(row, col)
        cell.value = visible_values[(row, col)]
        cell.border = border
        cell.font = Font(size=8)
        cell.alignment = Alignment(
            horizontal="left" if col in {2, 3, 4, 9} else "center",
            vertical="center",
            wrap_text=col in {4, 9},
        )
        if col in {5, 6}:
            cell.fill = linked_fill
            cell.number_format = "0.0"
        if col in {7, 8}:
            cell.number_format = "yyyy/m/d"

for col in HELPER_COLS_TO_HIDE:
    ws.column_dimensions[ws.cell(1, col).column_letter].hidden = True

ws.column_dimensions["A"].width = 7
ws.column_dimensions["B"].width = 22
ws.column_dimensions["C"].width = 24
ws.column_dimensions["D"].width = 48
ws.column_dimensions["E"].width = 10
ws.column_dimensions["F"].width = 12
ws.column_dimensions["G"].width = 12
ws.column_dimensions["H"].width = 12
ws.column_dimensions["I"].width = 40

for row in range(TASK_START_ROW, last_task_row + 1):
    ws.row_dimensions[row].height = 20

wb.save(WORKBOOK)
print(WORKBOOK)
print("fixed_visible_rows", last_task_row - TASK_START_ROW + 1)
