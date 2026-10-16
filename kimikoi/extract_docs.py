import pandas as pd
import docx
import sys
import os

with open("output_docs_utf8.txt", "w", encoding="utf-8") as f:
    def print_to_file(text):
        f.write(text + "\n")

    def read_xlsx(file_path):
        print_to_file(f"\n=================================\n--- FILE: {os.path.basename(file_path)} ---\n=================================")
        try:
            xls = pd.ExcelFile(file_path)
            for sheet_name in xls.sheet_names:
                print_to_file(f"\n=== Sheet: {sheet_name} ===")
                df = pd.read_excel(xls, sheet_name=sheet_name)
                print_to_file(df.to_string(index=False))
        except Exception as e:
            print_to_file(f"Error reading {file_path}: {e}")

    def read_docx(file_path):
        print_to_file(f"\n=================================\n--- FILE: {os.path.basename(file_path)} ---\n=================================")
        try:
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                print_to_file(para.text)
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.replace("\n", " ").strip() for cell in row.cells]
                    print_to_file("\t".join(row_data))
        except Exception as e:
            print_to_file(f"Error reading {file_path}: {e}")

    read_xlsx(r"c:\Users\ihogo\Desktop\WEB\kimikoi\きみこい・工数表.xlsx")
    read_docx(r"c:\Users\ihogo\Desktop\WEB\kimikoi\キミコイチート.docx")
    read_xlsx(r"c:\Users\ihogo\Desktop\WEB\kimikoi\キミコイ本制作要件_1201 1.xlsx")
